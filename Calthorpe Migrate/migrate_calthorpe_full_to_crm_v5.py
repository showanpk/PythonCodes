#!/usr/bin/env python3
"""
Saheli CRM - Calthorpe historical register migration
=====================================================

PURPOSE
-------
Migrates Calthorpe staff attendance registers into the current Saheli CRM schema.
Designed from the uploaded backend models/ApplicationDbContext and the three uploaded
Calthorpe Excel formats.

SOURCE PRIORITY (avoids overlapping workbooks)
----------------------------------------------
* May 2025 - October 2025 : Register for Exercise - Calthorpe.xlsx (wide monthly format)
* November - December 2025: Calthorpe Register NEW 2025 (2).xlsx (row-per-attendance format)
* January - March 2026    : Calthorpe Register 2026 (4).xlsx (row-per-attendance format)

IMPORTANT BEHAVIOUR
-------------------
* Existing sessions are reused; they are not duplicated.
* Existing FULL participants are matched by normalized Saheli Card Number.
* Missing FULL participants with a real Saheli Card Number can be created.
* No-card people are treated as LITE members.
* LITE members are matched first by normalized FirstName + LastName. Extra fields
  (DOB/postcode/phone) are used to resolve duplicate-name matches where possible.
* If no LITE name match exists, a new LiteMembers row is generated with a LITE-n ID.
* Existing attendance for the same member/session is not duplicated.
* Source rows that cannot be safely resolved are logged for review instead of being guessed.
* DRY RUN IS THE DEFAULT. Use --commit only after reviewing the preview report.\n* V5 removes three verified non-person staff-note rows and cleans 'helen new member' to Helen / Unknown.

The script intentionally does NOT overwrite existing participant/profile data.
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys
import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Iterable, Optional

# User-facing migration script: openpyxl is used only to read the source Excel registers.
from openpyxl import load_workbook

# -----------------------------------------------------------------------------
# CONFIGURATION
# -----------------------------------------------------------------------------

# Paste your REAL Azure SQL / SQL Server ODBC connection string here, OR set the
# environment variable SAHELI_SQL_CONNECTION_STRING before running the script.
# The backend ZIP contains Program.cs expecting ConnectionStrings:DefaultConnection,
# but the ZIP supplied here does not contain the secret appsettings connection value.
CONNECTION_STRING = os.getenv(
    "SAHELI_SQL_CONNECTION_STRING",
    r"Driver={ODBC Driver 18 for SQL Server};"
    r"Server=tcp:sahelihub.database.windows.net,1433;"
    r"Database=SaheliHubCRM;"
    r"Uid=sahelihubadmin;"
    r"Pwd=W7WZ7ZaG1YbMZ71gh%2xSFuR;"
    r"Encrypt=yes;"
    r"TrustServerCertificate=no;"
    r"Connection Timeout=30;"
)

BASE_DIR = Path(__file__).resolve().parent
OLD_WIDE_FILE = BASE_DIR / "Register for Exercise - Calthorpe.xlsx"
NEW_2025_FILE = BASE_DIR / "Calthorpe Register NEW 2025 (2).xlsx"
REGISTER_2026_FILE = BASE_DIR / "Calthorpe Register 2026 (4).xlsx"

VENUE_NAME = "Calthorpe Wellbeing Hub"
VENUE_ALIASES = {"calthorpe", "calthorpe wellbeing hub"}

MIGRATION_START_DATE = date(2025, 4, 1)
MIGRATION_END_DATE = date(2026, 3, 31)

CREATE_MISSING_FULL_PARTICIPANTS = True
CREATE_MISSING_LITE_MEMBERS = True
UPDATE_EXISTING_ATTENDANCE_TO_ATTENDED = True

# The wide May-Oct workbook contains activity/date but no session time.
# The script first learns activity/day times from the detailed Nov-Mar workbooks.
# If no usable time can be learned and no existing CRM session can be reused,
# deterministic placeholder times are used and explicitly marked in Session.Notes.
ALLOW_PLACEHOLDER_TIMES_FOR_WIDE_SOURCE = True
PLACEHOLDER_BASE_HOUR = 8
PLACEHOLDER_DURATION_MINUTES = 45
PLACEHOLDER_GAP_MINUTES = 15

DEFAULT_SESSION_FREQUENCY = "Historical"
DEFAULT_CATEGORY = "Physical Activity"  # <= 30 chars, satisfies current Sessions schema
MIGRATION_NOTE_PREFIX = "Historical Calthorpe register import"
NO_SURNAME_LABEL = "Unknown"

# -----------------------------------------------------------------------------
# CURRENT BACKEND SCHEMA USED BY THIS SCRIPT
# -----------------------------------------------------------------------------
# dbo.Participants: ParticipantID identity, SaheliCardNumber unique, FullName, DOB,
# Postcode, MobileNumber, Site, Notes, CreatedAt, etc.
# dbo.LiteMembers: Id guid, MembershipId unique, FirstName, LastName, DOB, Phone,
# Email, Address, Postcode, EmergencyName/Phone/Relation, HealthConditions,
# Gender, Ethnicity, CreatedAtUtc, CreatedByUserId.
# dbo.Sessions: SessionId identity, Frequency, Category, ActivityCategory, VenueName,
# ActivityName, IsRecurringWeekly, DayOfWeek, SessionDate, StartTime, EndTime,
# IsBookingRequired, IsCancelled, CreatedAtUtc, etc.
# dbo.SessionAttendance: AttendanceId identity, SessionId, AttendanceMemberKind,
# ParticipantId/LiteMemberId, MemberDisplayId, SaheliCardNumber, MemberName,
# session snapshot fields, Attended, Notes, CreatedAtUtc.
# Current backend has filtered unique indexes on (SessionId, ParticipantId) and
# (SessionId, LiteMemberId), plus FULL/LITE check constraints.

# -----------------------------------------------------------------------------
# NORMALIZATION / ALIASES
# -----------------------------------------------------------------------------

INVALID_TEXT = {"", "#n/a", "#ref!", "#value!", "#name?", "none", "null", "nan", "0"}

# These values appear in the vertical staff registers in the participant/name/card
# column but describe a non-delivered/cancelled session rather than a person. They
# must NEVER create LiteMembers or attendance rows. The session itself is preserved
# and marked cancelled where safe.
NON_PERSON_SESSION_STATUS = {
    "no session",
    "holiday",
    "cancelled",
    "canceled",
    "none attended",
    "no one attended",
    "eid",
    "party",
    "closed",
    "closure",
    "bank holiday",
    "no class",
    "no attendance",
    # Calthorpe 2026 source-specific staff notes that were entered in the
    # Saheli Card Number column and are not participant names.
    "term time only",
    "kate a l salsa was on instead",
    "maisie party",
}

# Source-specific participant-name annotations that should be cleaned before
# LiteMember matching/creation. "helen new member" is evidence of a real attendee
# whose surname was not captured; keep the attendance as Helen / Unknown rather
# than creating a bogus surname "new member".
SOURCE_PERSON_NAME_CLEANUP = {
    "helen new member": "Helen",
}

# Verified historic identity alias from the supplied Calthorpe registers.
# Evidence in the source files:
#   - July/August: SAH433 + Wellbeing Card 22664084 + Dega Ali
#   - September-November: 433/727 + the same person / same Wellbeing Card
#   - later detailed register: 727 + Dega Ali
# For attendance migration, 727 is treated as the canonical current FULL identity.
# This does NOT delete or update Participant 433 in dbo.Participants; it only prevents
# the same person being split across two participant identities during this import.
CARD_CANONICAL_OVERRIDES: dict[str, str] = {
    "433": "727",
}

# Composite-card override is kept for explicit audit readability. The canonical
# alias above already collapses 433/727 to 727 before lookup.
COMPOSITE_CARD_OVERRIDES: dict[str, str] = {
    "433/727": "727",
}

# Historic Saheli card numbers in these files are short numeric/SAH-prefixed values.
# Two 8-digit values were entered in the "Saheli Card number" column but do not
# match any CRM FULL participant and are consistent with Wellbeing Card format.
# When a 7+ digit numeric value has NO existing FULL match and a name is present,
# V5 treats it as a Lite-member source identifier rather than creating a bogus FULL.
SUSPICIOUS_LONG_NUMERIC_CARD_MIN_DIGITS = 7

ACTIVITY_RULES: list[tuple[list[str], str]] = [
    (["aerobics", "hiit"], "Aerobics/HIIT"),
    (["strength & stretch", "strength and stretch", "strength and strecth", "strength & strecth"], "Strength & Stretch"),
    (["chair based exercise", "chair exercise", "chair based"], "Chair Based Exercise"),
    (["social knit and crochet", "social knit", "crochet", "socail knit and crochet", "knit and crochet"], "Crochet"),
    (["body conditioning", "body conditining"], "Body Conditioning"),
    (["pilate floor base", "pilates floor based work", "pilates floor based", "pilates", "pilate"], "Pilates"),
    (["esol", "speak english"], "ESOL"),
    (["salsa", "belly dancing"], "Salsa/Belly Dancing"),
    (["circuits class", "circuit training", "circuits", "circuit"], "Circuit Training"),
    (["zumba"], "Zumba"),
    (["arts and craft", "arts and crafts", "a&c"], "Arts"),
    (["mens multisport", "men's multisport", "mens multi sports", "men's multi sports", "men's class", "mens class", "men's session", "mens session"], "Men's Multi Sports"),
    (["yoga"], "Yoga"),
    (["gardening"], "Gardening"),
    (["dance class", "dance fitness", "dance"], "Dance Fitness"),
    (["bhangra"], "Bhangra"),
    (["workshop", "menopause workshop", "salvation army workshop"], "Workshops"),
    (["salvation army"], "Salvation Army"),
    (["self defence", "self defense"], "Self Defence"),
    (["tennis"], "Tennis"),
    (["walk & talk", "walk and talk", "walk and wellbeing", "wlak and wellbeing", "thursday walk"], "Walk & Talk"),
    (["community police visit"], "Community Police Visit"),
    (["diamond art and mindfulness", "mindfulness", "creative"], "Mindfulness / Creative"),
    (["innerva"], "Innerva"),
]

# Display names written when a new historical session has to be created.
ACTIVITY_DISPLAY = {
    "Aerobics/HIIT": "Aerobics",
    "Strength & Stretch": "Strength & Stretch",
    "Chair Based Exercise": "Chair Based Exercise",
    "Crochet": "Social Knit & Crochet",
    "Body Conditioning": "Body Conditioning",
    "Pilates": "Pilates",
    "ESOL": "ESOL",
    "Salsa/Belly Dancing": "Salsa/Belly Dancing",
    "Circuit Training": "Circuit Training",
    "Zumba": "Zumba",
    "Arts": "Arts",
    "Men's Multi Sports": "Men's Multi Sports",
    "Yoga": "Yoga",
    "Gardening": "Gardening",
    "Dance Fitness": "Dance Fitness",
    "Bhangra": "Bhangra",
    "Workshops": "Workshops",
    "Salvation Army": "Salvation Army",
    "Self Defence": "Self Defence",
    "Tennis": "Tennis",
    "Walk & Talk": "Walk & Talk",
    "Community Police Visit": "Community Police Visit",
    "Mindfulness / Creative": "Mindfulness / Creative",
    "Innerva": "Innerva",
}


def clean_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"\s+", " ", str(value).strip())
    if text.lower() in INVALID_TEXT:
        return None
    return text or None


def normalize_name_piece(value: Any) -> str:
    text = clean_text(value) or ""
    text = text.casefold()
    text = re.sub(r"[^\w\s'-]", "", text, flags=re.UNICODE)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_name_key(first_name: Any, last_name: Any) -> str:
    return f"{normalize_name_piece(first_name)}|{normalize_name_piece(last_name)}"


def normalize_full_name_key(full_name: Any) -> str:
    text = normalize_name_piece(full_name)
    return text


def split_name(full_name: Any) -> tuple[Optional[str], Optional[str]]:
    text = clean_text(full_name)
    if not text:
        return None, None
    text = re.sub(r"\s+", " ", text).strip()
    parts = text.split(" ")
    if len(parts) == 1:
        return parts[0][:100], NO_SURNAME_LABEL
    return parts[0][:100], " ".join(parts[1:])[:100]


def normalize_card(value: Any) -> Optional[str]:
    """Return a simple Saheli card candidate, or None if the cell looks like a name."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            return None
        return str(int(value))
    text = clean_text(value)
    if not text:
        return None
    compact = re.sub(r"\s+", "", text).upper()
    # Real source cards are usually numeric or SAH-prefixed. Composite values are
    # handled separately by card_candidates().
    if re.fullmatch(r"\d+", compact):
        return str(int(compact))
    if re.fullmatch(r"SAH[-_]?\d+", compact):
        return compact.replace("_", "-")
    return None


def card_candidates(value: Any) -> list[str]:
    direct = normalize_card(value)
    if direct:
        return [direct]
    text = clean_text(value)
    if not text:
        return []
    # Handle historic cells such as 433/727 without silently picking a card.
    if "/" in text or "\\" in text:
        found = []
        for token in re.findall(r"(?:SAH[-_]?\d+|\d+)", text.upper()):
            c = normalize_card(token)
            if c and c not in found:
                found.append(c)
        return found
    return []


def normalize_card_key(value: Any) -> str:
    c = normalize_card(value)
    if not c:
        return ""
    key = re.sub(r"[^A-Z0-9]", "", c.upper())
    # Historical registers sometimes write the same Saheli number as SAH449 while
    # the CRM stores 449. Treat those as the same identity key.
    if key.startswith("SAH") and key[3:].isdigit():
        key = key[3:]
    if key.isdigit():
        key = str(int(key))
    return key


def normalize_postcode(value: Any) -> str:
    text = clean_text(value) or ""
    return re.sub(r"\s+", "", text).upper()


def normalize_phone(value: Any) -> str:
    text = clean_text(value) or ""
    digits = re.sub(r"\D", "", text)
    if digits.startswith("44"):
        digits = "0" + digits[2:]
    return digits


def normalize_status_label(value: Any) -> str:
    text = clean_text(value) or ""
    text = text.casefold().strip()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def is_non_person_session_status(value: Any) -> bool:
    return normalize_status_label(value) in NON_PERSON_SESSION_STATUS


def sanitize_source_person_name(value: Any) -> Optional[str]:
    """Clean known staff annotations without guessing a participant identity."""
    text = clean_text(value)
    if not text:
        return None
    key = normalize_status_label(text)
    cleaned = SOURCE_PERSON_NAME_CLEANUP.get(key)
    return cleaned if cleaned is not None else text


def composite_card_key(value: Any) -> str:
    candidates = card_candidates(value)
    if len(candidates) < 2:
        return ""
    return "/".join(sorted(normalize_card_key(c) for c in candidates))


def apply_card_canonical_override(value: Any) -> Optional[str]:
    """Map a verified historic Saheli card alias to the canonical current card."""
    card = normalize_card(value)
    if not card:
        return None
    key = normalize_card_key(card)
    target = CARD_CANONICAL_OVERRIDES.get(key)
    return target or card


def looks_like_misfiled_wellbeing_number(value: Any) -> bool:
    """True only for long numeric values that are very unlikely to be historic Saheli cards."""
    key = normalize_card_key(value)
    return bool(key.isdigit() and len(key) >= SUSPICIOUS_LONG_NUMERIC_CARD_MIN_DIGITS)


def canonical_activity(value: Any) -> Optional[str]:
    text = clean_text(value)
    if not text:
        return None
    key = re.sub(r"\s+", " ", text.casefold()).strip()
    for needles, canonical in ACTIVITY_RULES:
        if any(needle in key for needle in needles):
            return canonical
    return None


def activity_display_name(canonical: str, raw: Any = None) -> str:
    return ACTIVITY_DISPLAY.get(canonical, clean_text(raw) or canonical)[:150]


def normalize_venue(value: Any) -> str:
    return re.sub(r"\s+", " ", (clean_text(value) or "").casefold()).strip()


def is_yes_mark(value: Any) -> bool:
    if value is True:
        return True
    if isinstance(value, (int, float)):
        return float(value) == 1.0
    text = (clean_text(value) or "").casefold().strip()
    return text in {"yes", "y", "x", "✓", "✔", "present", "attended", "1"}


def safe_date(value: Any, sheet_year: int | None = None, sheet_month: int | None = None) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, datetime):
        result = value.date()
    elif isinstance(value, date):
        result = value
    elif isinstance(value, (int, float)):
        try:
            result = (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
        except Exception:
            return None
    else:
        text = clean_text(value)
        if not text:
            return None
        result = None
        for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%y", "%d/%m/%y"):
            try:
                result = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                pass
        if result is None:
            # Repair source typos such as 09.05.205 / 14.05.225 / 17.07.20254.
            m = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,5})", text)
            if m:
                dd, mm, yyyy = m.groups()
                if sheet_year is not None:
                    yyyy = str(sheet_year)
                elif len(yyyy) == 2:
                    yyyy = "20" + yyyy
                elif len(yyyy) > 4:
                    yyyy = yyyy[:4]
                try:
                    result = date(int(yyyy), int(mm), int(dd))
                except ValueError:
                    return None
            else:
                return None

    # Monthly wide tabs are authoritative for year/month. Correct obvious entry typos.
    if sheet_year is not None and sheet_month is not None:
        if result.year != sheet_year or result.month != sheet_month:
            try:
                result = date(sheet_year, sheet_month, result.day)
            except ValueError:
                return None
    return result


def safe_dob(value: Any) -> Optional[date]:
    d = safe_date(value)
    if not d:
        return None
    today = date.today()
    if d > today or d.year < 1900:
        return None
    return d


def parse_risk(value: Any) -> Optional[str]:
    text = clean_text(value)
    if not text:
        return None
    low = text.casefold()
    if low in {"low", "medium", "high"}:
        return low.title()
    # Keep other short source values, but not phone numbers accidentally typed into risk.
    if len(text) <= 100 and not re.fullmatch(r"\d{7,}", re.sub(r"\D", "", text)):
        return text[:100]
    return None


def normalize_header(value: Any) -> str:
    text = clean_text(value) or ""
    return re.sub(r"[^a-z0-9]", "", text.casefold())

# -----------------------------------------------------------------------------
# TIME PARSING / INFERENCE
# -----------------------------------------------------------------------------

EVENING_CANONICAL = {"Yoga", "Salsa/Belly Dancing", "Zumba"}


def _parse_clock(piece: str) -> tuple[int, int, Optional[str]]:
    p = piece.strip().casefold().replace(".", ":")
    p = re.sub(r"\s+", "", p)
    suffix = None
    if p.endswith("am"):
        suffix = "am"; p = p[:-2]
    elif p.endswith("pm"):
        suffix = "pm"; p = p[:-2]
    if ":" in p:
        h, m = p.split(":", 1)
    else:
        h, m = p, "0"
    return int(h), int(m[:2] or "0"), suffix


def parse_time_range(value: Any, canonical: Optional[str] = None) -> Optional[tuple[time, time]]:
    text = clean_text(value)
    if not text:
        return None
    s = text.casefold().replace("–", "-").replace("—", "-")
    s = re.sub(r"\s+", " ", s).strip()
    # Repair source typos e.g. 12-30-2pm / 12-3--2pm -> 12:30-2pm
    s = re.sub(r"^(\d{1,2})-(\d{2})-(\d{1,2})(am|pm)$", r"\1:\2-\3\4", s)
    s = re.sub(r"^(\d{1,2})-3--(\d{1,2})(am|pm)$", r"\1:30-\2\3", s)
    parts = re.split(r"\s*-\s*", s, maxsplit=1)
    if len(parts) != 2:
        return None
    try:
        sh, sm, ssuf = _parse_clock(parts[0])
        eh, em, esuf = _parse_clock(parts[1])
    except Exception:
        return None
    if sh > 12 or eh > 12 or sm > 59 or em > 59:
        return None

    def to24(h: int, suffix: Optional[str]) -> int:
        if suffix == "am":
            return 0 if h == 12 else h
        if suffix == "pm":
            return 12 if h == 12 else h + 12
        return h

    # If only the end has AM/PM, infer start from the same period except around noon.
    if ssuf is None and esuf is not None:
        if esuf == "pm":
            if sh == 12:
                ssuf = "pm"
            elif sh <= 7:
                ssuf = "pm"
            elif eh == 12 and sh < 12:
                # 11:00-12pm / 11:30-12:30pm crosses into noon.
                ssuf = "am"
            elif sh > eh:
                # 11:30-1:30pm means 11:30am to 1:30pm.
                ssuf = "am"
            else:
                ssuf = "pm"
        else:
            ssuf = "am"
    if esuf is None and ssuf is not None:
        esuf = ssuf

    # No suffix at all: infer known evening activities when the clock is <= 7.
    if ssuf is None and esuf is None and canonical in EVENING_CANONICAL and sh <= 7:
        ssuf = esuf = "pm"

    start_h = to24(sh, ssuf)
    end_h = to24(eh, esuf)

    start_minutes = start_h * 60 + sm
    end_minutes = end_h * 60 + em
    if end_minutes <= start_minutes:
        # Typical source shorthand 12:30-2:30 or 11-12; roll end forward 12h when sensible.
        if esuf is None or ssuf is None:
            end_minutes += 12 * 60
        elif end_minutes == start_minutes:
            return None
    if end_minutes <= start_minutes or end_minutes >= 24 * 60:
        return None
    return time(start_minutes // 60, start_minutes % 60), time(end_minutes // 60, end_minutes % 60)

# -----------------------------------------------------------------------------
# DATA CLASSES
# -----------------------------------------------------------------------------

@dataclass
class PersonSource:
    raw_card: Any = None
    wellbeing_card: Optional[str] = None
    full_name: Optional[str] = None
    dob: Optional[date] = None
    postcode: Optional[str] = None
    phone: Optional[str] = None
    emergency_name: Optional[str] = None
    emergency_phone: Optional[str] = None
    risk: Optional[str] = None

@dataclass
class SourceSession:
    source_file: str
    source_sheet: str
    source_ref: str
    session_date: date
    canonical_activity: str
    activity_name: str
    raw_activity: str
    start_time: Optional[time]
    end_time: Optional[time]
    time_quality: str  # SOURCE / INFERRED / PLACEHOLDER / UNKNOWN
    occurrence_index: int = 1
    occurrence_count: int = 1
    people: list[tuple[PersonSource, str]] = field(default_factory=list)  # person, source row ref
    is_cancelled: bool = False
    cancel_reason: Optional[str] = None

    @property
    def source_key(self) -> str:
        return f"{self.source_file}|{self.source_sheet}|{self.source_ref}"

@dataclass
class DbSession:
    session_id: int
    session_date: date
    venue_name: str
    activity_name: str
    canonical_activity: Optional[str]
    start_time: time
    end_time: time
    is_cancelled: bool
    attendance_count: int

@dataclass
class DbParticipant:
    participant_id: int
    card: str
    full_name: Optional[str]

@dataclass
class DbLite:
    lite_id: str
    membership_id: str
    first_name: str
    last_name: str
    dob: Optional[date]
    postcode: Optional[str]
    phone: Optional[str]

# -----------------------------------------------------------------------------
# SOURCE PARSERS
# -----------------------------------------------------------------------------

WIDE_SHEETS = {
    "May 2025": (2025, 5),
    "June 2025": (2025, 6),
    "July 2025": (2025, 7),
    "August 2025": (2025, 8),
    "September 2025": (2025, 9),
    "October 2025": (2025, 10),
}

HEADER_ALIASES = {
    "card": {"sahelicardnumber", "sahelicardno", "sahelicard"},
    "wellbeing": {"wellbeingcardno", "wellbeingcardnumber", "wellbeingcard"},
    "name": {"name", "fullname"},
    "dob": {"dob", "dateofbirth"},
    "postcode": {"postcode", "postcode"},
    "phone": {"phonenumber", "phone", "mobilenumber"},
    "emergency_name": {"emergencycontact", "emergencycontactname"},
    "emergency_phone": {"emergencynumber", "emergencycontactnumber", "emergencyphone"},
    "risk": {"riskassesment", "riskassessment", "riskstratification"},
}


def locate_metadata_columns(header_row: Iterable[Any]) -> dict[str, int]:
    result: dict[str, int] = {}
    for idx, value in enumerate(header_row, start=1):
        normalized = normalize_header(value)
        if not normalized:
            continue
        for field_name, aliases in HEADER_ALIASES.items():
            if normalized in aliases and field_name not in result:
                result[field_name] = idx
    return result


def cell(row: tuple[Any, ...] | list[Any], one_based_col: Optional[int]) -> Any:
    if not one_based_col:
        return None
    idx = one_based_col - 1
    return row[idx] if 0 <= idx < len(row) else None


def build_person_from_wide_row(row: tuple[Any, ...], cols: dict[str, int]) -> PersonSource:
    return PersonSource(
        raw_card=cell(row, cols.get("card")),
        wellbeing_card=clean_text(cell(row, cols.get("wellbeing"))),
        full_name=clean_text(cell(row, cols.get("name"))),
        dob=safe_dob(cell(row, cols.get("dob"))),
        postcode=clean_text(cell(row, cols.get("postcode"))),
        phone=clean_text(cell(row, cols.get("phone"))),
        emergency_name=clean_text(cell(row, cols.get("emergency_name"))),
        emergency_phone=clean_text(cell(row, cols.get("emergency_phone"))),
        risk=parse_risk(cell(row, cols.get("risk"))),
    )


def learn_time_templates() -> dict[tuple[str, int], tuple[time, time]]:
    """Learn the most common detailed-register time for canonical activity + weekday."""
    counts: dict[tuple[str, int], Counter] = defaultdict(Counter)
    for path, min_date, max_date in [
        (NEW_2025_FILE, date(2025, 11, 1), date(2025, 12, 31)),
        (REGISTER_2026_FILE, date(2026, 1, 1), date(2026, 3, 31)),
    ]:
        if not path.exists():
            continue
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            if normalize_header(ws.title) == "template":
                continue
            sheet_can = canonical_activity(ws.title)
            for row in ws.iter_rows(min_row=2, max_col=10, values_only=True):
                raw_activity = clean_text(row[0])
                d = safe_date(row[2])
                can = canonical_activity(raw_activity) or sheet_can
                if not d or not can or not (min_date <= d <= max_date):
                    continue
                parsed = parse_time_range(row[4], can)
                if parsed:
                    counts[(can, d.weekday())][parsed] += 1
        wb.close()
    return {key: counter.most_common(1)[0][0] for key, counter in counts.items() if counter}


def parse_wide_workbook(time_templates: dict[tuple[str, int], tuple[time, time]]) -> list[SourceSession]:
    sessions: list[SourceSession] = []
    if not OLD_WIDE_FILE.exists():
        raise FileNotFoundError(OLD_WIDE_FILE)
    wb = load_workbook(OLD_WIDE_FILE, read_only=True, data_only=True)
    for sheet_name, (sheet_year, sheet_month) in WIDE_SHEETS.items():
        sheet_lookup = {name.strip(): name for name in wb.sheetnames}
        actual_sheet_name = sheet_lookup.get(sheet_name.strip())
        if not actual_sheet_name:
            print(f"WARNING: wide source sheet missing: {sheet_name}")
            continue
        ws = wb[actual_sheet_name]
        header1 = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        date_row = [c for c in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
        header_row = [c for c in next(ws.iter_rows(min_row=3, max_row=3, values_only=True))]
        metadata = locate_metadata_columns(header_row)

        # Session columns are identified by a valid date in row 2 and an activity in row 3.
        event_columns: list[tuple[int, date, str, str]] = []
        for idx, raw_activity in enumerate(header_row, start=1):
            if idx - 1 >= len(date_row):
                continue
            d = safe_date(date_row[idx - 1], sheet_year, sheet_month)
            can = canonical_activity(raw_activity)
            if not d or not can or not (MIGRATION_START_DATE <= d <= MIGRATION_END_DATE):
                continue
            event_columns.append((idx, d, can, clean_text(raw_activity) or can))

        # Count same canonical/date occurrences so repeated classes are kept separate.
        group_counts = Counter((d, can) for _, d, can, _ in event_columns)
        group_seen: Counter = Counter()

        source_sessions: list[tuple[int, SourceSession]] = []
        placeholder_slot_by_date: Counter = Counter()
        for col_idx, d, can, raw_activity in event_columns:
            group_seen[(d, can)] += 1
            occurrence_index = group_seen[(d, can)]
            occurrence_count = group_counts[(d, can)]
            learned = time_templates.get((can, d.weekday()))
            if learned:
                start_t, end_t = learned
                quality = "INFERRED"
                # For repeated same activity/date without source times, offset subsequent
                # occurrences so separate source columns remain separate sessions.
                if occurrence_count > 1 and occurrence_index > 1:
                    delta = (occurrence_index - 1) * (PLACEHOLDER_DURATION_MINUTES + PLACEHOLDER_GAP_MINUTES)
                    start_dt = datetime.combine(d, start_t) + timedelta(minutes=delta)
                    end_dt = datetime.combine(d, end_t) + timedelta(minutes=delta)
                    if end_dt.date() == d:
                        start_t, end_t = start_dt.time().replace(second=0, microsecond=0), end_dt.time().replace(second=0, microsecond=0)
                    else:
                        learned = None
            if not learned:
                if ALLOW_PLACEHOLDER_TIMES_FOR_WIDE_SOURCE:
                    slot = placeholder_slot_by_date[d]
                    placeholder_slot_by_date[d] += 1
                    start_minutes = PLACEHOLDER_BASE_HOUR * 60 + slot * (PLACEHOLDER_DURATION_MINUTES + PLACEHOLDER_GAP_MINUTES)
                    # Keep deterministic placeholders inside the day.
                    start_minutes = start_minutes % (22 * 60)
                    end_minutes = start_minutes + PLACEHOLDER_DURATION_MINUTES
                    start_t = time(start_minutes // 60, start_minutes % 60)
                    end_t = time(end_minutes // 60, end_minutes % 60)
                    quality = "PLACEHOLDER"
                else:
                    start_t = end_t = None
                    quality = "UNKNOWN"

            session = SourceSession(
                source_file=OLD_WIDE_FILE.name,
                source_sheet=sheet_name,
                source_ref=f"COL-{col_idx}",
                session_date=d,
                canonical_activity=can,
                activity_name=activity_display_name(can, raw_activity),
                raw_activity=raw_activity,
                start_time=start_t,
                end_time=end_t,
                time_quality=quality,
                occurrence_index=occurrence_index,
                occurrence_count=occurrence_count,
            )
            source_sessions.append((col_idx, session))
            sessions.append(session)

        # Map attendance marks into the corresponding SourceSession.
        for excel_row_num, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
            person = build_person_from_wide_row(row, metadata)
            if not person.full_name and not card_candidates(person.raw_card):
                continue
            for col_idx, session in source_sessions:
                value = row[col_idx - 1] if col_idx - 1 < len(row) else None
                if is_yes_mark(value):
                    session.people.append((person, f"ROW-{excel_row_num}"))
        
    wb.close()
    return sessions


def parse_vertical_workbook(path: Path, min_date: date, max_date: date) -> list[SourceSession]:
    if not path.exists():
        raise FileNotFoundError(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    grouped: dict[tuple, SourceSession] = {}
    for ws in wb.worksheets:
        if normalize_header(ws.title) == "template":
            continue
        sheet_can = canonical_activity(ws.title)
        for excel_row_num, row in enumerate(ws.iter_rows(min_row=2, max_col=10, values_only=True), start=2):
            raw_activity = clean_text(row[0])
            d = safe_date(row[2])
            if not d or not (min_date <= d <= max_date) or not (MIGRATION_START_DATE <= d <= MIGRATION_END_DATE):
                continue
            can = canonical_activity(raw_activity) or sheet_can
            if not can:
                continue
            parsed_time = parse_time_range(row[4], can)
            if parsed_time:
                start_t, end_t = parsed_time
                time_quality = "SOURCE"
            else:
                start_t = end_t = None
                time_quality = "UNKNOWN"

            raw_card = row[5]
            raw_name = clean_text(row[6])

            # Several 2026 source rows place either a cardless participant name OR
            # a staff/session note in the Saheli Card Number column while the
            # formula-based Name cell is #N/A. Check exact known non-person notes
            # before treating text as a participant.
            if not card_candidates(raw_card) and not raw_name:
                possible_name = clean_text(raw_card)
                if possible_name and is_non_person_session_status(possible_name):
                    raw_name = possible_name
                    raw_card = None
                elif possible_name and re.search(r"[A-Za-z]", possible_name):
                    raw_name = sanitize_source_person_name(possible_name)
                    raw_card = None
            else:
                raw_name = sanitize_source_person_name(raw_name)

            key = (d, can, start_t, end_t, clean_text(raw_activity) or can)
            if key not in grouped:
                grouped[key] = SourceSession(
                    source_file=path.name,
                    source_sheet=ws.title,
                    source_ref=f"{d.isoformat()}|{clean_text(raw_activity) or can}|{clean_text(row[4]) or 'NO-TIME'}",
                    session_date=d,
                    canonical_activity=can,
                    activity_name=activity_display_name(can, raw_activity),
                    raw_activity=clean_text(raw_activity) or can,
                    start_time=start_t,
                    end_time=end_t,
                    time_quality=time_quality,
                )
            session = grouped[key]

            # Staff sometimes typed session status into the participant/name column.
            # Preserve the session as cancelled/no-delivery, but never create a Lite
            # member or attendance row for words such as Cancelled/Holiday/Eid.
            status_value = raw_name if raw_name else (clean_text(raw_card) if not card_candidates(raw_card) else None)
            if is_non_person_session_status(status_value):
                session.is_cancelled = True
                session.cancel_reason = clean_text(status_value)
                continue

            person = PersonSource(
                raw_card=raw_card,
                full_name=raw_name,
                emergency_name=clean_text(row[7]),
                emergency_phone=clean_text(row[8]),
                risk=parse_risk(row[9]),
            )
            if not person.full_name and not card_candidates(person.raw_card):
                continue
            session.people.append((person, f"ROW-{excel_row_num}"))

    # If a grouped source session contains real attendees, it was delivered even if
    # a stray status row also exists in that group. Keep the attendee evidence and
    # clear the cancellation flag rather than silently discarding people.
    for session in grouped.values():
        if session.people and session.is_cancelled:
            session.is_cancelled = False
            session.cancel_reason = None

    wb.close()
    return list(grouped.values())


def deduplicate_source_attendance(sessions: list[SourceSession]) -> tuple[int, int]:
    """Remove duplicate person marks within the exact same parsed source session."""
    removed = 0
    kept = 0
    for session in sessions:
        seen: set[str] = set()
        new_people = []
        for person, row_ref in session.people:
            cards = card_candidates(person.raw_card)
            if cards:
                identity = "CARD:" + "/".join(sorted(normalize_card_key(c) for c in cards))
            else:
                identity = "NAME:" + normalize_full_name_key(person.full_name)
            if not identity or identity in {"NAME:", "CARD:"}:
                new_people.append((person, row_ref)); kept += 1
                continue
            if identity in seen:
                removed += 1
                continue
            seen.add(identity)
            new_people.append((person, row_ref)); kept += 1
        session.people = new_people
    return kept, removed


def parse_all_sources() -> list[SourceSession]:
    time_templates = learn_time_templates()
    sessions = []
    sessions.extend(parse_wide_workbook(time_templates))
    sessions.extend(parse_vertical_workbook(NEW_2025_FILE, date(2025, 11, 1), date(2025, 12, 31)))
    sessions.extend(parse_vertical_workbook(REGISTER_2026_FILE, date(2026, 1, 1), date(2026, 3, 31)))
    # Stable sort is helpful for deterministic preview and ID generation.
    sessions.sort(key=lambda s: (s.session_date, s.start_time or time(23, 59), s.canonical_activity, s.source_key))
    deduplicate_source_attendance(sessions)
    return sessions

# -----------------------------------------------------------------------------
# DATABASE HELPERS
# -----------------------------------------------------------------------------

REQUIRED_COLUMNS = {
    "Participants": {"ParticipantID", "SaheliCardNumber", "FullName", "DateOfBirth", "Postcode", "MobileNumber", "Site", "Notes", "CreatedAt"},
    "LiteMembers": {"Id", "MembershipId", "FirstName", "LastName", "DateOfBirth", "Phone", "Email", "Address", "Postcode", "EmergencyName", "EmergencyPhone", "EmergencyRelation", "HealthConditions", "Gender", "Ethnicity", "CreatedAtUtc", "CreatedByUserId"},
    "Sessions": {"SessionId", "Frequency", "Category", "ActivityCategory", "VenueName", "ActivityName", "Notes", "IsRecurringWeekly", "DayOfWeek", "SessionDate", "StartTime", "EndTime", "IsBookingRequired", "IsCancelled", "CreatedAtUtc"},
    "SessionAttendance": {"AttendanceId", "SessionId", "AttendanceMemberKind", "ParticipantId", "LiteMemberId", "MemberDisplayId", "SaheliCardNumber", "MemberName", "Phone", "EmergencyName", "EmergencyPhone", "SessionName", "SessionDay", "SessionDate", "SessionMonth", "SessionStartTime", "SessionEndTime", "RiskStratification", "Attended", "Notes", "CreatedAtUtc", "UpdatedAtUtc"},
}


def validate_connection_string() -> None:
    if "YOUR_SQL_" in CONNECTION_STRING or "YOUR_SQL_SERVER" in CONNECTION_STRING:
        raise RuntimeError(
            "Database connection is not configured. Edit CONNECTION_STRING in the script "
            "or set SAHELI_SQL_CONNECTION_STRING."
        )


def preflight_schema(cur) -> None:
    for table_name, expected in REQUIRED_COLUMNS.items():
        rows = cur.execute(
            """
            SELECT COLUMN_NAME
            FROM INFORMATION_SCHEMA.COLUMNS
            WHERE TABLE_SCHEMA = 'dbo' AND TABLE_NAME = ?
            """,
            table_name,
        ).fetchall()
        actual = {r[0] for r in rows}
        missing = sorted(expected - actual)
        if missing:
            raise RuntimeError(f"dbo.{table_name} is missing required columns: {missing}")


def load_db_participants(cur) -> tuple[dict[str, DbParticipant], dict[str, list[DbParticipant]]]:
    by_card: dict[str, DbParticipant] = {}
    by_name: dict[str, list[DbParticipant]] = defaultdict(list)
    for row in cur.execute("SELECT ParticipantID, SaheliCardNumber, FullName FROM dbo.Participants"):
        p = DbParticipant(int(row[0]), clean_text(row[1]) or "", clean_text(row[2]))
        card_key = normalize_card_key(p.card)
        if card_key:
            by_card[card_key] = p
        name_key = normalize_full_name_key(p.full_name)
        if name_key:
            by_name[name_key].append(p)
    return by_card, by_name


def load_db_lites(cur) -> tuple[dict[str, list[DbLite]], dict[str, DbLite]]:
    by_name: dict[str, list[DbLite]] = defaultdict(list)
    by_id: dict[str, DbLite] = {}
    rows = cur.execute(
        """
        SELECT Id, MembershipId, FirstName, LastName, DateOfBirth, Postcode, Phone
        FROM dbo.LiteMembers
        """
    ).fetchall()
    for row in rows:
        dob = row[4]
        if isinstance(dob, datetime):
            dob = dob.date()
        lite = DbLite(str(row[0]), str(row[1]), str(row[2]), str(row[3]), dob, clean_text(row[5]), clean_text(row[6]))
        by_name[normalize_name_key(lite.first_name, lite.last_name)].append(lite)
        by_id[lite.lite_id.lower()] = lite
    return by_name, by_id


def load_db_sessions(cur) -> list[DbSession]:
    placeholders = ",".join("?" for _ in VENUE_ALIASES)
    sql = f"""
        SELECT s.SessionId, s.SessionDate, s.VenueName, s.ActivityName,
               s.StartTime, s.EndTime, s.IsCancelled,
               COUNT(a.AttendanceId) AS AttendanceCount
        FROM dbo.Sessions s
        LEFT JOIN dbo.SessionAttendance a ON a.SessionId = s.SessionId AND a.Attended = 1
        WHERE s.SessionDate >= ? AND s.SessionDate <= ?
          AND LOWER(LTRIM(RTRIM(s.VenueName))) IN ({placeholders})
        GROUP BY s.SessionId, s.SessionDate, s.VenueName, s.ActivityName,
                 s.StartTime, s.EndTime, s.IsCancelled
    """
    params = [MIGRATION_START_DATE, MIGRATION_END_DATE] + sorted(VENUE_ALIASES)
    result: list[DbSession] = []
    for row in cur.execute(sql, params).fetchall():
        d = row[1].date() if isinstance(row[1], datetime) else row[1]
        st = row[4]
        et = row[5]
        if isinstance(st, datetime): st = st.time()
        if isinstance(et, datetime): et = et.time()
        result.append(DbSession(int(row[0]), d, str(row[2]), str(row[3]), canonical_activity(row[3]), st, et, bool(row[6]), int(row[7] or 0)))
    return result


def load_existing_attendance_keys(cur) -> set[tuple[int, str, str]]:
    """Keys: (SessionId, FULL/LITE, member database ID as string)."""
    keys: set[tuple[int, str, str]] = set()
    sql = """
        SELECT a.SessionId, a.AttendanceMemberKind, a.ParticipantId, a.LiteMemberId
        FROM dbo.SessionAttendance a
        JOIN dbo.Sessions s ON s.SessionId = a.SessionId
        WHERE s.SessionDate >= ? AND s.SessionDate <= ?
          AND LOWER(LTRIM(RTRIM(s.VenueName))) IN ('calthorpe','calthorpe wellbeing hub')
    """
    for row in cur.execute(sql, MIGRATION_START_DATE, MIGRATION_END_DATE).fetchall():
        kind = (row[1] or "").upper()
        member_id = str(row[2]) if kind == "FULL" else str(row[3]).lower()
        if member_id:
            keys.add((int(row[0]), kind, member_id))
    return keys


def next_lite_membership_number(cur) -> int:
    # Current C# backend also generates LITE-n. We take a serializable/table lock during
    # the one-off transaction to prevent another importer from reserving the same value.
    rows = cur.execute("SELECT MembershipId FROM dbo.LiteMembers WITH (UPDLOCK, HOLDLOCK)").fetchall()
    max_num = 0
    for (membership_id,) in rows:
        text = clean_text(membership_id)
        if not text:
            continue
        m = re.fullmatch(r"LITE-(\d+)", text, re.IGNORECASE)
        if m:
            max_num = max(max_num, int(m.group(1)))
    return max_num + 1


def choose_lite_match(candidates: list[DbLite], person: PersonSource) -> tuple[Optional[DbLite], str]:
    if not candidates:
        return None, "NO_MATCH"
    if len(candidates) == 1:
        return candidates[0], "EXACT_NAME"

    # Resolve duplicate names using additional source fields where available.
    scored = []
    for c in candidates:
        score = 0
        comparable = 0
        if person.dob and c.dob:
            comparable += 1
            if person.dob == c.dob: score += 4
        if normalize_postcode(person.postcode) and normalize_postcode(c.postcode):
            comparable += 1
            if normalize_postcode(person.postcode) == normalize_postcode(c.postcode): score += 3
        if normalize_phone(person.phone) and normalize_phone(c.phone):
            comparable += 1
            if normalize_phone(person.phone) == normalize_phone(c.phone): score += 3
        scored.append((score, comparable, c))
    scored.sort(key=lambda x: (-x[0], -x[1], x[2].membership_id))
    if scored and scored[0][0] > 0 and (len(scored) == 1 or scored[0][0] > scored[1][0]):
        return scored[0][2], "EXACT_NAME_PLUS_DETAILS"
    return None, "AMBIGUOUS_NAME"


def insert_full_participant(cur, card: str, person: PersonSource) -> DbParticipant:
    full_name = clean_text(person.full_name)
    cur.execute(
        """
        INSERT INTO dbo.Participants
            (SaheliCardNumber, FullName, DateOfBirth, Postcode, MobileNumber,
             Site, Notes, CreatedAt)
        OUTPUT INSERTED.ParticipantID
        VALUES (?, ?, ?, ?, ?, ?, ?, SYSDATETIME())
        """,
        card[:50],
        full_name[:255] if full_name else None,
        person.dob,
        clean_text(person.postcode)[:20] if clean_text(person.postcode) else None,
        clean_text(person.phone),
        VENUE_NAME,
        f"{MIGRATION_NOTE_PREFIX}; created from historical source"[:4000],
    )
    participant_id = int(cur.fetchone()[0])
    return DbParticipant(participant_id, card, full_name)


def insert_lite_member(cur, membership_id: str, person: PersonSource) -> DbLite:
    first_name, last_name = split_name(person.full_name)
    if not first_name:
        raise ValueError("Cannot create LiteMember without a usable name")
    last_name = last_name or NO_SURNAME_LABEL
    lite_id = str(uuid.uuid4())
    cur.execute(
        """
        INSERT INTO dbo.LiteMembers
            (Id, MembershipId, FirstName, LastName, DateOfBirth, Phone, Email,
             Address, Postcode, EmergencyName, EmergencyPhone, EmergencyRelation,
             HealthConditions, Gender, Ethnicity, CreatedAtUtc, CreatedByUserId)
        VALUES
            (?, ?, ?, ?, ?, ?, NULL, NULL, ?, ?, ?, NULL,
             NULL, NULL, NULL, SYSUTCDATETIME(), NULL)
        """,
        lite_id,
        membership_id[:50],
        first_name[:100],
        last_name[:100],
        person.dob,
        clean_text(person.phone)[:30] if clean_text(person.phone) else None,
        clean_text(person.postcode)[:30] if clean_text(person.postcode) else None,
        clean_text(person.emergency_name)[:200] if clean_text(person.emergency_name) else None,
        clean_text(person.emergency_phone)[:30] if clean_text(person.emergency_phone) else None,
    )
    return DbLite(lite_id, membership_id, first_name, last_name, person.dob, person.postcode, person.phone)


def time_seconds(t: time) -> int:
    return t.hour * 3600 + t.minute * 60 + t.second


def session_candidates_for(source: SourceSession, db_sessions: list[DbSession]) -> list[DbSession]:
    return [
        s for s in db_sessions
        if s.session_date == source.session_date
        and s.canonical_activity == source.canonical_activity
        and normalize_venue(s.venue_name) in VENUE_ALIASES
    ]


def choose_existing_session(source: SourceSession, db_sessions: list[DbSession]) -> Optional[DbSession]:
    candidates = session_candidates_for(source, db_sessions)
    if not candidates:
        return None

    # With a trustworthy/source-derived time, prefer exact time match.
    if source.start_time and source.end_time and source.time_quality in {"SOURCE", "INFERRED"}:
        exact = [s for s in candidates if s.start_time == source.start_time and s.end_time == source.end_time]
        if exact:
            return sorted(exact, key=lambda s: (-s.attendance_count, s.is_cancelled, s.session_id))[0]
        same_start = [s for s in candidates if s.start_time == source.start_time]
        if len(same_start) == 1:
            return same_start[0]
        # Allow a modest source-vs-CRM time discrepancy before creating a duplicate.
        close = sorted(candidates, key=lambda s: abs(time_seconds(s.start_time) - time_seconds(source.start_time)))
        if close and abs(time_seconds(close[0].start_time) - time_seconds(source.start_time)) <= 30 * 60:
            return close[0]

    # Wide source has no actual time. If exactly one equivalent CRM session exists,
    # reuse it. If multiple exist, map repeated source columns by chronological ordinal.
    if len(candidates) == 1:
        return candidates[0]
    candidates = sorted(candidates, key=lambda s: (s.start_time, -s.attendance_count, s.is_cancelled, s.session_id))
    if 1 <= source.occurrence_index <= len(candidates):
        return candidates[source.occurrence_index - 1]

    # Otherwise use the established canonical ranking rather than blindly creating
    # another copy when a duplicate session already exists.
    return sorted(candidates, key=lambda s: (-s.attendance_count, s.is_cancelled, s.session_id))[0]


def create_session(cur, source: SourceSession) -> DbSession:
    if not source.start_time or not source.end_time:
        raise ValueError("Session time unresolved")
    if source.end_time <= source.start_time:
        raise ValueError("EndTime must be after StartTime")
    display = activity_display_name(source.canonical_activity, source.raw_activity)
    note = (
        f"{MIGRATION_NOTE_PREFIX}; source={source.source_file}/{source.source_sheet}/{source.source_ref}; "
        f"source_activity={source.raw_activity}; time_quality={source.time_quality}; "
        f"cancelled={int(source.is_cancelled)}; cancel_reason={source.cancel_reason or ''}"
    )[:500]
    category = (source.canonical_activity or DEFAULT_CATEGORY)[:30]
    cur.execute(
        """
        INSERT INTO dbo.Sessions
            (Frequency, Category, SubCategory, ActivityCategory, VenueName,
             AssignedStaffId, SessionProviderId, ActivityName, Notes,
             IsRecurringWeekly, DayOfWeek, SessionDate, ArrivalTime, StartTime,
             EndTime, Capacity, IsBookingRequired, IsCancelled, CreatedAtUtc)
        OUTPUT INSERTED.SessionId
        VALUES
            (?, ?, NULL, ?, ?, NULL, NULL, ?, ?,
             0, NULL, ?, NULL, ?, ?, NULL, 0, ?, SYSUTCDATETIME())
        """,
        DEFAULT_SESSION_FREQUENCY[:30], category, category, VENUE_NAME, display, note,
        source.session_date, source.start_time, source.end_time, 1 if source.is_cancelled else 0,
    )
    session_id = int(cur.fetchone()[0])
    return DbSession(session_id, source.session_date, VENUE_NAME, display, source.canonical_activity,
                     source.start_time, source.end_time, source.is_cancelled, 0)


def insert_attendance(cur, db_session: DbSession, member_kind: str, member_id: str,
                      display_id: str, member_name: Optional[str], person: PersonSource,
                      source: SourceSession, source_row_ref: str) -> None:
    is_full = member_kind == "FULL"
    participant_id = int(member_id) if is_full else None
    lite_id = None if is_full else member_id
    card = display_id if is_full else None
    wellbeing_note = f"; wellbeing_card={person.wellbeing_card}" if person.wellbeing_card else ""
    source_card_note = f"; source_card={clean_text(person.raw_card)}" if clean_text(person.raw_card) else ""
    note = (
        f"{MIGRATION_NOTE_PREFIX}; source={source.source_file}/{source.source_sheet}/{source_row_ref}; "
        f"source_session={source.source_ref}; time_quality={source.time_quality}{source_card_note}{wellbeing_note}"
    )[:500]
    cur.execute(
        """
        INSERT INTO dbo.SessionAttendance
            (SessionId, AttendanceMemberKind, ParticipantId, LiteMemberId,
             MemberDisplayId, SaheliCardNumber, MemberName, Phone,
             EmergencyName, EmergencyPhone, SessionName, SessionDay,
             SessionDate, SessionMonth, SessionStartTime, SessionEndTime,
             RiskStratification, Attended, Notes, CreatedAtUtc, UpdatedAtUtc)
        VALUES
            (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, SYSUTCDATETIME(), NULL)
        """,
        db_session.session_id,
        member_kind,
        participant_id,
        lite_id,
        display_id[:50],
        card[:50] if card else None,
        clean_text(member_name)[:200] if clean_text(member_name) else None,
        clean_text(person.phone)[:30] if clean_text(person.phone) else None,
        clean_text(person.emergency_name)[:200] if clean_text(person.emergency_name) else None,
        clean_text(person.emergency_phone)[:30] if clean_text(person.emergency_phone) else None,
        db_session.activity_name[:150],
        source.session_date.strftime("%A")[:20],
        source.session_date,
        source.session_date.strftime("%B")[:20],
        db_session.start_time,
        db_session.end_time,
        person.risk[:100] if person.risk else None,
        note,
    )

# -----------------------------------------------------------------------------
# MIGRATION
# -----------------------------------------------------------------------------

class MigrationLog:
    def __init__(self):
        self.rows: list[dict[str, Any]] = []
        self.counts: Counter = Counter()

    def add(self, action: str, source: Optional[SourceSession] = None, person: Optional[PersonSource] = None,
            detail: str = "", session_id: Optional[int] = None, member_ref: str = ""):
        self.counts[action] += 1
        self.rows.append({
            "Action": action,
            "Date": source.session_date.isoformat() if source else "",
            "Activity": source.activity_name if source else "",
            "SourceFile": source.source_file if source else "",
            "SourceSheet": source.source_sheet if source else "",
            "SourceSessionRef": source.source_ref if source else "",
            "SessionId": session_id or "",
            "SourceCard": clean_text(person.raw_card) if person else "",
            "SourceName": clean_text(person.full_name) if person else "",
            "MemberRef": member_ref,
            "Detail": detail,
        })

    def write_csv(self, path: Path):
        fieldnames = ["Action", "Date", "Activity", "SourceFile", "SourceSheet", "SourceSessionRef",
                      "SessionId", "SourceCard", "SourceName", "MemberRef", "Detail"]
        with path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(self.rows)


def resolve_member(cur, person: PersonSource, participants_by_card: dict[str, DbParticipant],
                   full_by_name: dict[str, list[DbParticipant]], lites_by_name: dict[str, list[DbLite]],
                   lite_number_state: list[int], log: MigrationLog, source: SourceSession,
                   composite_resolution_cache: dict[str, DbParticipant]):
    raw_candidates = card_candidates(person.raw_card)

    # Apply only verified identity aliases and de-duplicate the resulting candidate list.
    candidates: list[str] = []
    alias_applied = False
    for raw_candidate in raw_candidates:
        canonical = apply_card_canonical_override(raw_candidate)
        if not canonical:
            continue
        if normalize_card_key(canonical) != normalize_card_key(raw_candidate):
            alias_applied = True
        if normalize_card_key(canonical) not in {normalize_card_key(c) for c in candidates}:
            candidates.append(canonical)

    if candidates:
        matches = []
        for c in candidates:
            p = participants_by_card.get(normalize_card_key(c))
            if p and p not in matches:
                matches.append(p)

        if len(matches) == 1:
            p = matches[0]
            resolution = "MATCHED_FULL_CARD_ALIAS" if alias_applied else "MATCHED_FULL"
            return "FULL", str(p.participant_id), p.card, person.full_name or p.full_name, resolution

        if len(matches) > 1:
            comp_key = composite_card_key(person.raw_card)

            # 1) Explicit verified override, if configured.
            override = COMPOSITE_CARD_OVERRIDES.get(comp_key) or COMPOSITE_CARD_OVERRIDES.get(clean_text(person.raw_card) or "")
            if override:
                chosen = participants_by_card.get(normalize_card_key(override))
                if chosen and chosen in matches:
                    composite_resolution_cache[comp_key] = chosen
                    return "FULL", str(chosen.participant_id), chosen.card, person.full_name or chosen.full_name, "MATCHED_FULL_COMPOSITE_OVERRIDE"

            # 2) Reuse a decision already proven earlier in this same preview run.
            cached = composite_resolution_cache.get(comp_key)
            if cached and cached in matches:
                return "FULL", str(cached.participant_id), cached.card, person.full_name or cached.full_name, "MATCHED_FULL_COMPOSITE_CACHED"

            # 3) Resolve by exact normalized name only if exactly one candidate matches.
            source_name_key = normalize_full_name_key(person.full_name)
            if source_name_key:
                name_matches = [p for p in matches if normalize_full_name_key(p.full_name) == source_name_key]
                if len(name_matches) == 1:
                    chosen = name_matches[0]
                    composite_resolution_cache[comp_key] = chosen
                    return "FULL", str(chosen.participant_id), chosen.card, person.full_name or chosen.full_name, "MATCHED_FULL_COMPOSITE_BY_NAME"

            log.add(
                "REVIEW_AMBIGUOUS_COMPOSITE_CARD",
                source,
                person,
                detail=f"Multiple existing cards matched: {[p.card for p in matches]}; source_name={person.full_name or ''}",
            )
            return None

        # No existing FULL match. Long 7+ digit numeric values in these historic
        # files are treated as misfiled Wellbeing-style identifiers rather than
        # automatically creating a new FULL participant.
        if len(candidates) == 1 and looks_like_misfiled_wellbeing_number(candidates[0]) and clean_text(person.full_name):
            raw_long = normalize_card_key(candidates[0])
            if not person.wellbeing_card:
                person.wellbeing_card = raw_long
            log.add(
                "SOURCE_LONG_CARD_RECLASSIFIED_TO_LITE",
                source,
                person,
                detail=f"Unmatched {raw_long} treated as non-Saheli/Wellbeing-style identifier; resolving by Lite first+last name",
            )
            candidates = []
        elif len(candidates) > 1:
            log.add("REVIEW_UNMATCHED_COMPOSITE_CARD", source, person, detail=f"Composite card did not match CRM: {candidates}")
            return None
        else:
            card = candidates[0]
            if CREATE_MISSING_FULL_PARTICIPANTS:
                p = insert_full_participant(cur, card, person)
                participants_by_card[normalize_card_key(card)] = p
                if p.full_name:
                    full_by_name[normalize_full_name_key(p.full_name)].append(p)
                log.add("CREATED_FULL", source, person, detail=f"Created ParticipantID={p.participant_id}", member_ref=f"FULL:{p.participant_id}")
                return "FULL", str(p.participant_id), p.card, person.full_name or p.full_name, "CREATED_FULL"
            log.add("REVIEW_UNMATCHED_CARD", source, person, detail=f"Card {card} not found and CREATE_MISSING_FULL_PARTICIPANTS=False")
            return None

    # No Saheli Card: per migration rule, resolve/create in LiteMembers by first+last name.
    first_name, last_name = split_name(person.full_name)
    if not first_name:
        log.add("REVIEW_CARDLESS_NO_NAME", source, person, detail="Cannot create/match LiteMember without a name")
        return None
    last_name = last_name or NO_SURNAME_LABEL
    key = normalize_name_key(first_name, last_name)
    lite, match_reason = choose_lite_match(lites_by_name.get(key, []), person)
    if lite:
        return "LITE", lite.lite_id, lite.membership_id, f"{lite.first_name} {lite.last_name}".strip(), f"MATCHED_LITE_{match_reason}"
    if match_reason == "AMBIGUOUS_NAME":
        # Do not guess among two people with the exact same first/last name.
        log.add("REVIEW_AMBIGUOUS_LITE_NAME", source, person, detail=f"Multiple LiteMembers already have {first_name} {last_name}")
        return None
    if not CREATE_MISSING_LITE_MEMBERS:
        log.add("REVIEW_LITE_NOT_FOUND", source, person, detail="No LiteMember match and creation disabled")
        return None

    membership_id = f"LITE-{lite_number_state[0]}"
    lite_number_state[0] += 1
    lite = insert_lite_member(cur, membership_id, person)
    lites_by_name[key].append(lite)
    log.add("CREATED_LITE", source, person, detail=f"Created {membership_id}", member_ref=f"LITE:{lite.lite_id}")
    return "LITE", lite.lite_id, lite.membership_id, f"{lite.first_name} {lite.last_name}".strip(), "CREATED_LITE"


def run_migration(commit: bool) -> int:
    sessions = parse_all_sources()
    source_attendance = sum(len(s.people) for s in sessions)
    placeholder_sessions = sum(1 for s in sessions if s.time_quality == "PLACEHOLDER")
    unknown_time_sessions = sum(1 for s in sessions if not s.start_time or not s.end_time)
    cancelled_sessions = sum(1 for s in sessions if s.is_cancelled)
    zero_attendance_non_cancelled = sum(1 for s in sessions if not s.people and not s.is_cancelled)

    print("\n=== SOURCE AUDIT ===")
    print(f"Parsed source sessions        : {len(sessions):,}")
    print(f"Parsed attendance rows        : {source_attendance:,}")
    print(f"First source date             : {min((s.session_date for s in sessions), default=None)}")
    print(f"Last source date              : {max((s.session_date for s in sessions), default=None)}")
    print(f"Wide sessions using placeholder time: {placeholder_sessions:,}")
    print(f"Unresolved-time sessions      : {unknown_time_sessions:,}")
    print(f"Explicit cancelled/no-session : {cancelled_sessions:,}")
    print(f"Zero-attendance non-cancelled : {zero_attendance_non_cancelled:,}")
    by_month = Counter((s.session_date.year, s.session_date.month) for s in sessions for _ in s.people)
    for (year, month), count in sorted(by_month.items()):
        print(f"  {year}-{month:02d}: {count:,} attendance rows")

    if unknown_time_sessions:
        print("ERROR: Some source sessions have no usable/inferred time and placeholder mode is disabled.")
        return 2

    validate_connection_string()
    try:
        import pyodbc
    except ImportError as exc:
        raise RuntimeError("pyodbc is required. Run: pip install pyodbc openpyxl") from exc

    cn = pyodbc.connect(CONNECTION_STRING, autocommit=False)
    cur = cn.cursor()
    # One transaction and serializable isolation keep the one-off import deterministic.
    cur.execute("SET TRANSACTION ISOLATION LEVEL SERIALIZABLE;")
    log = MigrationLog()

    try:
        preflight_schema(cur)
        participants_by_card, full_by_name = load_db_participants(cur)
        lites_by_name, _ = load_db_lites(cur)
        db_sessions = load_db_sessions(cur)
        existing_attendance_keys = load_existing_attendance_keys(cur)
        lite_number_state = [next_lite_membership_number(cur)]
        composite_resolution_cache: dict[str, DbParticipant] = {}

        print("\n=== DATABASE PREFLIGHT ===")
        print(f"Existing Calthorpe sessions in period : {len(db_sessions):,}")
        print(f"Existing attendance member/session keys: {len(existing_attendance_keys):,}")
        print(f"Existing FULL participants loaded     : {len(participants_by_card):,}")
        print(f"Existing LITE name keys loaded        : {len(lites_by_name):,}")
        print(f"Next reserved Lite membership number  : {lite_number_state[0]}")

        for source in sessions:
            db_session = choose_existing_session(source, db_sessions)
            if db_session:
                log.add("EXISTING_SESSION", source, detail=f"Reused SessionId={db_session.session_id}; CRM time={db_session.start_time}-{db_session.end_time}", session_id=db_session.session_id)
                if source.is_cancelled:
                    if db_session.attendance_count > 0:
                        log.add("REVIEW_CANCELLED_SESSION_HAS_CRM_ATTENDANCE", source, detail=f"Source says cancelled ({source.cancel_reason}) but SessionId={db_session.session_id} has {db_session.attendance_count} CRM attendance rows", session_id=db_session.session_id)
                    elif not db_session.is_cancelled:
                        cur.execute("UPDATE dbo.Sessions SET IsCancelled = 1 WHERE SessionId = ?", db_session.session_id)
                        db_session.is_cancelled = True
                        log.add("MARKED_EXISTING_SESSION_CANCELLED", source, detail=f"Source status={source.cancel_reason}", session_id=db_session.session_id)
            else:
                try:
                    db_session = create_session(cur, source)
                except Exception as exc:
                    log.add("REVIEW_SESSION_NOT_CREATED", source, detail=str(exc))
                    continue
                db_sessions.append(db_session)
                log.add("NEW_SESSION", source, detail=f"Created SessionId={db_session.session_id}; time={db_session.start_time}-{db_session.end_time}; quality={source.time_quality}; cancelled={source.is_cancelled}", session_id=db_session.session_id)

            if source.is_cancelled:
                log.add("SOURCE_CANCELLED_SESSION", source, detail=f"No attendance migrated; source status={source.cancel_reason}", session_id=db_session.session_id)
                continue
            if not source.people:
                log.add("SOURCE_ZERO_ATTENDANCE_SESSION", source, detail="Source session has no participant attendance marks", session_id=db_session.session_id)

            for person, source_row_ref in source.people:
                resolved = resolve_member(
                    cur, person, participants_by_card, full_by_name, lites_by_name,
                    lite_number_state, log, source, composite_resolution_cache,
                )
                if not resolved:
                    continue
                member_kind, member_id, display_id, member_name, resolution = resolved
                member_key = (db_session.session_id, member_kind, str(member_id).lower() if member_kind == "LITE" else str(member_id))
                if member_key in existing_attendance_keys:
                    if UPDATE_EXISTING_ATTENDANCE_TO_ATTENDED:
                        if member_kind == "FULL":
                            cur.execute(
                                """UPDATE dbo.SessionAttendance
                                   SET Attended = 1,
                                       UpdatedAtUtc = CASE WHEN Attended = 0 THEN SYSUTCDATETIME() ELSE UpdatedAtUtc END
                                   WHERE SessionId = ? AND AttendanceMemberKind = 'FULL' AND ParticipantId = ?""",
                                db_session.session_id, int(member_id),
                            )
                        else:
                            cur.execute(
                                """UPDATE dbo.SessionAttendance
                                   SET Attended = 1,
                                       UpdatedAtUtc = CASE WHEN Attended = 0 THEN SYSUTCDATETIME() ELSE UpdatedAtUtc END
                                   WHERE SessionId = ? AND AttendanceMemberKind = 'LITE' AND LiteMemberId = ?""",
                                db_session.session_id, member_id,
                            )
                    log.add("ALREADY_IN_CRM", source, person, detail=f"{resolution}; attendance already exists", session_id=db_session.session_id, member_ref=f"{member_kind}:{member_id}")
                    continue

                try:
                    insert_attendance(cur, db_session, member_kind, member_id, display_id, member_name, person, source, source_row_ref)
                    existing_attendance_keys.add(member_key)
                    log.add("NEW_ATTENDANCE", source, person, detail=resolution, session_id=db_session.session_id, member_ref=f"{member_kind}:{member_id}")
                except Exception as exc:
                    # Unique constraints can still catch a race or an equivalent row we did not preload.
                    log.add("REVIEW_ATTENDANCE_INSERT_FAILED", source, person, detail=str(exc), session_id=db_session.session_id, member_ref=f"{member_kind}:{member_id}")
                    raise

        report_name = f"calthorpe_migration_{'commit' if commit else 'preview'}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        report_path = BASE_DIR / report_name
        log.write_csv(report_path)

        print("\n=== ACTION SUMMARY ===")
        for action, count in sorted(log.counts.items()):
            print(f"{action:36s} {count:,}")
        review_count = sum(v for k, v in log.counts.items() if k.startswith("REVIEW_"))
        print(f"\nReview-required rows/actions: {review_count:,}")
        print(f"Detailed migration report     : {report_path}")

        if commit:
            if review_count:
                print("\nCOMMIT BLOCKED: review-required items exist. No changes saved.")
                cn.rollback()
                return 3
            cn.commit()
            print("\nMode: COMMITTED")
        else:
            cn.rollback()
            print("\nMode: PREVIEW ONLY - transaction rolled back; database unchanged.")
            print("Run again with --commit only after the preview summary/report is correct.")
        return 0
    except Exception:
        cn.rollback()
        raise
    finally:
        cur.close()
        cn.close()

# -----------------------------------------------------------------------------
# AUDIT-ONLY MODE (NO DATABASE CONNECTION)
# -----------------------------------------------------------------------------

def run_source_audit_only() -> int:
    sessions = parse_all_sources()
    total_att = sum(len(s.people) for s in sessions)
    print("=== CALTHORPE SOURCE AUDIT (NO DATABASE) ===")
    print(f"Reporting window: {MIGRATION_START_DATE} to {MIGRATION_END_DATE}")
    print(f"Source sessions parsed : {len(sessions):,}")
    print(f"Attendance rows parsed : {total_att:,}")
    print(f"Cancelled/no-session   : {sum(1 for s in sessions if s.is_cancelled):,}")
    print(f"Zero-attendance sessions: {sum(1 for s in sessions if not s.people and not s.is_cancelled):,}")
    print(f"Date range             : {min((s.session_date for s in sessions), default=None)} to {max((s.session_date for s in sessions), default=None)}")
    by_month = Counter()
    by_file = Counter()
    by_activity = Counter()
    cardful = 0
    cardless = 0
    for s in sessions:
        for person, _ in s.people:
            by_month[(s.session_date.year, s.session_date.month)] += 1
            by_file[s.source_file] += 1
            by_activity[s.canonical_activity] += 1
            if card_candidates(person.raw_card): cardful += 1
            else: cardless += 1
    print("\nAttendance by month:")
    for ym, count in sorted(by_month.items()):
        print(f"  {ym[0]}-{ym[1]:02d}: {count:,}")
    print("\nAttendance by source file:")
    for k,v in by_file.items(): print(f"  {k}: {v:,}")
    print(f"\nAttendance rows with Saheli-card candidate: {cardful:,}")
    print(f"Attendance rows without Saheli card        : {cardless:,}")
    print("\nTop activities:")
    for k,v in by_activity.most_common(30): print(f"  {k:28s} {v:,}")
    print("\nSession time quality:")
    for k,v in Counter(s.time_quality for s in sessions).items(): print(f"  {k:12s} {v:,}")
    return 0

# -----------------------------------------------------------------------------
# CLI
# -----------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(description="Migrate Calthorpe historical attendance into Saheli CRM")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--audit-only", action="store_true", help="Parse/check Excel only; do not connect to database")
    mode.add_argument("--commit", action="store_true", help="Commit database changes (default is preview/rollback)")
    args = parser.parse_args()

    print("Saheli CRM - Calthorpe Historical Migration V5")
    print("Source directory:", BASE_DIR)
    for p in (OLD_WIDE_FILE, NEW_2025_FILE, REGISTER_2026_FILE):
        print(" ", "OK" if p.exists() else "MISSING", p.name)
    if any(not p.exists() for p in (OLD_WIDE_FILE, NEW_2025_FILE, REGISTER_2026_FILE)):
        print("ERROR: Put this script in the same folder as the three Excel files, using the exact filenames above.")
        return 2

    if args.audit_only:
        return run_source_audit_only()
    return run_migration(commit=args.commit)


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\nCancelled by user.")
        raise SystemExit(130)
    except Exception as exc:
        print(f"\nFATAL: {exc}", file=sys.stderr)
        raise
