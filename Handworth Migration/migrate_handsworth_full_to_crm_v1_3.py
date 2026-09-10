#!/usr/bin/env python3
"""
Saheli CRM - Handsworth historical migration V1.3

Safety model
------------
* --audit-only: parses Excel only; no DB connection.
* default: DB preview inside one transaction, then ROLLBACK.
* --commit: commits only when no REVIEW_* blockers remain.
* Existing Sessions / Participants / LiteMembers / attendance are reused.
* Existing CRM profile data is never overwritten.

Expected source files (same folder as script by default):
  Register for Exercise - Handsworth (1).xlsx
  Handsworth Register 2026 (2).xlsx
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys
import uuid
import warnings
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable, Optional

try:
    from openpyxl import load_workbook
except ImportError:
    print("ERROR: openpyxl is required. Run: py -m pip install openpyxl pyodbc", file=sys.stderr)
    raise

warnings.filterwarnings("ignore", message=r"Cell .* is marked as a date but the serial value.*")

SOURCE_2025 = "Register for Exercise - Handsworth (1).xlsx"
SOURCE_2026 = "Handsworth Register 2026 (2).xlsx"
DEFAULT_VENUE = "Handsworth"

# Canonical CRM ActivityName values already used by Saheli CRM.
ACTIVITY_ALIASES = {
    "pilates": "Pilates",
    "crochet_beginner": "Crochet for Beginners",
    "crochet_advanced": "Crochet for Advanced",
}

# Historical Handsworth start times learned from the later labels in the SAME 2025 register.
HISTORICAL_START_TIMES = {
    "Pilates": time(12, 0),
    "Crochet for Beginners": time(10, 0),
    "Crochet for Advanced": time(11, 0),
}

# End times are not stored in the historical source. Keep deterministic and documented.
DURATION_MINUTES = {
    "Pilates": 45,
    "Crochet for Beginners": 60,
    "Crochet for Advanced": 60,
}

# Explicit source corrections confirmed from paired 2026 rows.
DATE_CORRECTIONS = {
    # Row 121 follows the 24-Feb Beginner block; the repeated 17-Feb date is a source typo.
    ("Crochet and Knit Social", 121, date(2026, 2, 17)): date(2026, 2, 24),
    ("Crochet and Knit Social", 325, date(2005, 5, 12)): date(2026, 5, 12),
    ("Crochet and Knit Social", 371, date(2025, 5, 26)): date(2026, 5, 26),
}

INVALID_TEXT = {"", "#n/a", "#ref!", "#value!", "none", "null", "nan", "0"}

# Handsworth source identity corrections supported by repeated evidence inside the
# source workbooks. The 2026 register transitions from the 3-digit values below
# to the matching 15xx values for the same people. The old 3-digit numbers now
# belong to unrelated CRM participants, so they must never be used as CRM IDs.
SOURCE_CARD_CORRECTIONS = {
    "512": "1512",  # Madhubala Mahendra
    "514": "1514",  # Maureen Bennet
    "515": "1515",  # Shindo Kaur
    "516": "1516",  # Petrona Mason
}

SOURCE_CARD_CANONICAL_NAMES = {
    "1512": "Madhubala Mahendra",
    "1514": "Maureen Bennet",
    "1515": "Shindo Kaur",
    "1516": "Petrona Mason",
}

SOURCE_NAME_CANONICALISATION = {
    "maureen bennett": "Maureen Bennet",
    "madhubala mahandra": "Madhubala Mahendra",
    "madneep kaur": "Mandeep Kaur",
    "greta gritten": "Greta Gittens",
}

# Grace and Greta are distinct Handsworth CRM participants. The 2025 Grace row
# incorrectly carries Greta's card 1123, while later source rows identify Grace
# by name only. Exact unique CRM FullName is therefore the safe resolver for Grace.
SAFE_EXACT_FULL_NAME_REUSE = {"grace gittens"}


@dataclass
class PersonObservation:
    source_file: str
    sheet: str
    row: int
    name: str
    card: Optional[str] = None
    dob: Optional[date] = None
    postcode: Optional[str] = None
    emergency_name: Optional[str] = None
    emergency_phone: Optional[str] = None
    risk: Optional[str] = None

    @property
    def name_norm(self) -> str:
        return norm_name(self.name)


@dataclass
class AttendanceObservation:
    session: "SourceSession"
    person: PersonObservation
    source_ref: str


@dataclass
class SourceSession:
    source_file: str
    sheet: str
    source_ref: str
    source_date: date
    session_date: date
    source_activity: str
    activity_name: str
    start_time: time
    end_time: time
    time_quality: str
    date_correction: str = ""
    attendance: list[AttendanceObservation] = field(default_factory=list)

    @property
    def key(self) -> tuple:
        return (self.session_date, self.activity_name, self.start_time)


@dataclass
class DbParticipant:
    participant_id: int
    card: str
    full_name: str
    dob: Optional[date]
    postcode: Optional[str]
    mobile: Optional[str]


@dataclass
class DbLite:
    lite_id: str
    membership_id: str
    first_name: str
    last_name: str
    dob: Optional[date]
    phone: Optional[str]
    postcode: Optional[str]

    @property
    def full_name(self) -> str:
        return f"{self.first_name} {self.last_name}".strip()


@dataclass
class ResolvedMember:
    kind: str  # FULL / LITE
    participant_id: Optional[int]
    lite_id: Optional[str]
    display_id: str
    name: str
    card: Optional[str]
    phone: Optional[str]
    action: str
    detail: str


@dataclass
class Action:
    action: str
    severity: str = "INFO"
    source_file: str = ""
    sheet: str = ""
    source_ref: str = ""
    session_date: str = ""
    activity: str = ""
    start_time: str = ""
    source_name: str = ""
    source_card: str = ""
    member_kind: str = ""
    member_display_id: str = ""
    session_id: str = ""
    detail: str = ""


# --------------------------- normalisation helpers ---------------------------

def clean_text(v: Any) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    if s.lower() in INVALID_TEXT:
        return ""
    return re.sub(r"\s+", " ", s).strip()


def norm_name(v: Any) -> str:
    s = clean_text(v).lower()
    s = s.replace("’", "'")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def norm_simple(v: Any) -> str:
    s = clean_text(v).lower()
    return re.sub(r"[^a-z0-9]+", "", s)


def valid_card(v: Any) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        try:
            if float(v) <= 0:
                return None
            if float(v).is_integer():
                return str(int(v))
        except Exception:
            return None
    s = clean_text(v)
    if re.fullmatch(r"\d+(?:\.0+)?", s):
        n = int(float(s))
        return str(n) if n > 0 else None
    return None


def as_date(v: Any) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(v))).date()
        except Exception:
            return None
    s = clean_text(v)
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            pass
    return None


def as_time(v: Any) -> Optional[time]:
    if isinstance(v, datetime):
        return v.time().replace(second=0, microsecond=0)
    if isinstance(v, time):
        return v.replace(second=0, microsecond=0)
    s = clean_text(v).lower().replace(".", "")
    if not s:
        return None
    s = s.replace(" ", "")
    for fmt in ("%I%p", "%I:%M%p", "%I%M%p", "%H:%M", "%H%M"):
        try:
            return datetime.strptime(s, fmt).time()
        except Exception:
            pass
    return None


def plus_minutes(t: time, minutes: int) -> time:
    dt = datetime.combine(date(2000, 1, 1), t) + timedelta(minutes=minutes)
    return dt.time().replace(second=0, microsecond=0)


def safe_phone(v: Any) -> Optional[str]:
    # Emergency numbers only. Do not turn Excel date-formatted garbage into fake participant phones.
    if v is None or isinstance(v, (date, datetime)):
        return None
    s = clean_text(v)
    if not s:
        return None
    digits = re.sub(r"\D", "", s)
    if 9 <= len(digits) <= 15:
        if s.startswith("+"):
            return "+" + digits
        if len(digits) == 10 and not digits.startswith("0"):
            # Source frequently lost a leading zero because Excel stored the number numerically.
            return "0" + digits
        return digits
    return None


def normalize_postcode(v: Any) -> Optional[str]:
    s = clean_text(v).upper().replace(" ", "")
    return s or None


def is_attended(v: Any) -> bool:
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v) == 1.0
    return clean_text(v).lower() in {"yes", "y", "1", "x", "attended", "present", "✓", "✔"}


def map_activity(raw: Any) -> Optional[str]:
    s = norm_name(raw)
    if "pilates" in s:
        return ACTIVITY_ALIASES["pilates"]
    if "advanced" in s and "crochet" in s:
        return ACTIVITY_ALIASES["crochet_advanced"]
    if "crochet" in s:
        return ACTIVITY_ALIASES["crochet_beginner"]
    return None


def name_similarity(a: str, b: str) -> float:
    a, b = norm_name(a), norm_name(b)
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    return SequenceMatcher(None, a, b).ratio()


def split_name(full_name: str) -> tuple[str, str]:
    parts = clean_text(full_name).split()
    if len(parts) < 2:
        return (parts[0] if parts else "Unknown", "Unknown")
    return (" ".join(parts[:-1]), parts[-1])


# ------------------------------- source parser ------------------------------

def parse_2025(path: Path) -> tuple[list[SourceSession], list[PersonObservation]]:
    wb = load_workbook(path, data_only=True)
    ws = wb["2025 REGISTER"]
    people_by_row: dict[int, PersonObservation] = {}
    all_people: list[PersonObservation] = []

    for row in range(4, ws.max_row + 1):
        card_cell = ws.cell(row, 2).value
        name_cell = ws.cell(row, 4).value
        card = valid_card(card_cell)
        name = clean_text(name_cell)
        # Some source rows put a person's name into the card column.
        if not name and not card:
            possible_name = clean_text(card_cell)
            if possible_name and not valid_card(possible_name):
                name = possible_name
        if not name and not card:
            continue
        p = PersonObservation(
            source_file=path.name,
            sheet=ws.title,
            row=row,
            name=name,
            card=card,
            dob=as_date(ws.cell(row, 5).value),
            postcode=normalize_postcode(ws.cell(row, 6).value),
            emergency_name=clean_text(ws.cell(row, 7).value) or None,
            emergency_phone=safe_phone(ws.cell(row, 8).value),
            risk=clean_text(ws.cell(row, 9).value) or None,
        )
        people_by_row[row] = p
        all_people.append(p)

    sessions: list[SourceSession] = []
    for col in range(10, ws.max_column + 1):
        d = as_date(ws.cell(2, col).value)
        raw_activity = clean_text(ws.cell(3, col).value)
        if not d or not raw_activity:
            continue
        activity = map_activity(raw_activity)
        if not activity:
            continue
        source_time = None
        m = re.search(r"\b(\d{1,2}(?::\d{2})?\s*(?:am|pm))\b", raw_activity, re.I)
        if m:
            source_time = as_time(m.group(1))
        start = source_time or HISTORICAL_START_TIMES[activity]
        quality = "SOURCE" if source_time else "INFERRED_FROM_LATER_2025_LABELS"
        end = plus_minutes(start, DURATION_MINUTES[activity])
        session = SourceSession(
            source_file=path.name,
            sheet=ws.title,
            source_ref=f"Column {col}",
            source_date=d,
            session_date=d,
            source_activity=raw_activity,
            activity_name=activity,
            start_time=start,
            end_time=end,
            time_quality=quality,
        )
        for row, p in people_by_row.items():
            if is_attended(ws.cell(row, col).value):
                session.attendance.append(
                    AttendanceObservation(session=session, person=p, source_ref=f"{ws.title}!{ws.cell(row, col).coordinate}")
                )
        sessions.append(session)
    return sessions, all_people


def parse_2026_sheet(path: Path, ws) -> tuple[list[SourceSession], list[PersonObservation]]:
    sessions: list[SourceSession] = []
    people: list[PersonObservation] = []
    current: Optional[SourceSession] = None

    for row in range(2, ws.max_row + 1):
        raw_session = clean_text(ws.cell(row, 2).value)
        raw_date = as_date(ws.cell(row, 4).value)
        if raw_session and raw_date:
            activity = map_activity(raw_session)
            start = as_time(ws.cell(row, 6).value)
            if activity and start:
                corrected = DATE_CORRECTIONS.get((ws.title, row, raw_date), raw_date)
                correction = ""
                if corrected != raw_date:
                    correction = f"{raw_date.isoformat()} -> {corrected.isoformat()}"
                current = SourceSession(
                    source_file=path.name,
                    sheet=ws.title,
                    source_ref=f"Row {row}",
                    source_date=raw_date,
                    session_date=corrected,
                    source_activity=raw_session,
                    activity_name=activity,
                    start_time=start,
                    end_time=plus_minutes(start, DURATION_MINUTES[activity]),
                    time_quality="SOURCE_START_END_INFERRED",
                    date_correction=correction,
                )
                sessions.append(current)
            else:
                current = None

        if current is None:
            continue

        card_cell = ws.cell(row, 7).value
        name_cell = ws.cell(row, 8).value
        card = valid_card(card_cell)
        name = clean_text(name_cell)
        if not name and not card:
            possible_name = clean_text(card_cell)
            if possible_name and not valid_card(possible_name):
                name = possible_name
        if not name and not card:
            continue

        p = PersonObservation(
            source_file=path.name,
            sheet=ws.title,
            row=row,
            name=name,
            card=card,
            emergency_name=clean_text(ws.cell(row, 9).value) or None,
            emergency_phone=safe_phone(ws.cell(row, 10).value),
            risk=clean_text(ws.cell(row, 11).value) or None,
        )
        people.append(p)
        current.attendance.append(
            AttendanceObservation(session=current, person=p, source_ref=f"{ws.title}!A{row}:K{row}")
        )
    return sessions, people


def parse_2026(path: Path) -> tuple[list[SourceSession], list[PersonObservation]]:
    wb = load_workbook(path, data_only=True)
    sessions: list[SourceSession] = []
    people: list[PersonObservation] = []
    for sheet_name in ("Crochet and Knit Social", "Chair Based Pilates"):
        ws = wb[sheet_name]
        s, p = parse_2026_sheet(path, ws)
        sessions.extend(s)
        people.extend(p)
    return sessions, people


def dedupe_source_attendance(sessions: list[SourceSession]) -> list[Action]:
    actions: list[Action] = []
    for s in sessions:
        seen: set[tuple[str, str]] = set()
        kept: list[AttendanceObservation] = []
        for a in s.attendance:
            p = a.person
            identity = ("CARD", p.card, p.name_norm) if p.card else ("NAME", p.name_norm, "")
            if not identity[1]:
                kept.append(a)
                continue
            if identity in seen:
                actions.append(Action(
                    action="SKIP_SOURCE_DUPLICATE_ATTENDANCE",
                    severity="INFO",
                    source_file=s.source_file,
                    sheet=s.sheet,
                    source_ref=a.source_ref,
                    session_date=s.session_date.isoformat(),
                    activity=s.activity_name,
                    start_time=s.start_time.strftime("%H:%M"),
                    source_name=p.name,
                    source_card=p.card or "",
                    detail=f"Duplicate source attendance identity {identity[0]}={identity[1]} name={identity[2]}",
                ))
                continue
            seen.add(identity)
            kept.append(a)
        s.attendance = kept
    return actions


def apply_source_identity_corrections(people: Iterable[PersonObservation]) -> list[Action]:
    """Apply only source-supported Handsworth identity corrections before CRM matching."""
    actions: list[Action] = []
    for p in people:
        original_name = p.name
        original_card = p.card

        canonical = SOURCE_NAME_CANONICALISATION.get(p.name_norm)
        if canonical and clean_text(p.name) != canonical:
            p.name = canonical
            actions.append(Action(
                action="SOURCE_NAME_CANONICALISED", severity="INFO",
                source_file=p.source_file, sheet=p.sheet, source_ref=f"Row {p.row}",
                source_name=p.name, source_card=p.card or "",
                detail=f"Source name '{original_name}' canonicalised to '{canonical}' for cross-year identity matching.",
            ))

        # The 2025 Grace Gittens row incorrectly uses Greta Gittens' card 1123.
        # Clear only that one name/card combination so Grace cannot be attached to Greta.
        if p.card == "1123" and p.name_norm == "grace gittens":
            p.card = None
            actions.append(Action(
                action="SOURCE_WRONG_CARD_CLEARED", severity="INFO",
                source_file=p.source_file, sheet=p.sheet, source_ref=f"Row {p.row}",
                source_name=p.name, source_card="1123",
                detail="Grace Gittens is distinct from CRM card 1123 (Greta Gittens); source card cleared before CRM matching.",
            ))

        if p.card in SOURCE_CARD_CORRECTIONS:
            old = p.card
            new = SOURCE_CARD_CORRECTIONS[old]
            expected = SOURCE_CARD_CANONICAL_NAMES[new]
            # Only apply when the row is nameless or its name agrees with the
            # source-supported identity for this card transition.
            if (not p.name_norm) or p.name_norm == norm_name(expected):
                p.card = new
                if not p.name_norm:
                    p.name = expected
                actions.append(Action(
                    action="SOURCE_CARD_CORRECTED_TO_15XX", severity="INFO",
                    source_file=p.source_file, sheet=p.sheet, source_ref=f"Row {p.row}",
                    source_name=p.name, source_card=p.card,
                    detail=f"Handsworth source continuity resolves legacy/mistyped card {old} to {new} for {expected}.",
                ))

        # Later rows sometimes contain only the corrected 15xx card because the
        # workbook lookup formula returned #N/A. Restore the already-established
        # source identity so the migration does not create a nameless participant.
        if p.card in SOURCE_CARD_CANONICAL_NAMES and not p.name_norm:
            p.name = SOURCE_CARD_CANONICAL_NAMES[p.card]
            actions.append(Action(
                action="SOURCE_NAME_INFERRED_FROM_CORRECTED_CARD", severity="INFO",
                source_file=p.source_file, sheet=p.sheet, source_ref=f"Row {p.row}",
                source_name=p.name, source_card=p.card,
                detail=f"Name restored from repeated Handsworth source identity for corrected card {p.card}.",
            ))

    return actions


def build_source_identity_maps(people: Iterable[PersonObservation]):
    name_cards: dict[str, set[str]] = defaultdict(set)
    card_names: dict[str, set[str]] = defaultdict(set)
    for p in people:
        if p.card and p.name_norm:
            name_cards[p.name_norm].add(p.card)
            card_names[p.card].add(p.name_norm)
    return name_cards, card_names


def enrich_missing_cards(people: Iterable[PersonObservation], name_cards: dict[str, set[str]]) -> list[Action]:
    actions: list[Action] = []
    for p in people:
        if p.card or not p.name_norm:
            continue
        cards = name_cards.get(p.name_norm, set())
        if len(cards) == 1:
            p.card = next(iter(cards))
            actions.append(Action(
                action="SOURCE_CARD_RESOLVED_FROM_OTHER_YEAR",
                severity="INFO",
                source_file=p.source_file,
                sheet=p.sheet,
                source_ref=f"Row {p.row}",
                source_name=p.name,
                source_card=p.card,
                detail="Exact normalized name appears elsewhere in Handsworth source with one unique Saheli Card.",
            ))
    return actions


def source_card_conflicts(card_names: dict[str, set[str]]) -> dict[str, set[str]]:
    conflicts: dict[str, set[str]] = {}
    for card, names in card_names.items():
        ns = sorted(n for n in names if n)
        if len(ns) <= 1:
            continue
        # Treat small spelling variations as the same person; flag materially different names.
        min_sim = min(name_similarity(a, b) for i, a in enumerate(ns) for b in ns[i + 1:])
        if min_sim < 0.80:
            conflicts[card] = set(ns)
    return conflicts


def parse_sources(source_dir: Path):
    p2025 = source_dir / SOURCE_2025
    p2026 = source_dir / SOURCE_2026
    missing = [str(p) for p in (p2025, p2026) if not p.exists()]
    if missing:
        raise FileNotFoundError("Missing source file(s): " + ", ".join(missing))

    s25, p25 = parse_2025(p2025)
    s26, p26 = parse_2026(p2026)
    sessions = s25 + s26
    all_people = p25 + p26

    identity_correction_actions = apply_source_identity_corrections(all_people)
    name_cards, _ = build_source_identity_maps(all_people)
    resolution_actions = enrich_missing_cards(all_people, name_cards)
    # rebuild after enrichment
    name_cards, card_names = build_source_identity_maps(all_people)
    conflicts = source_card_conflicts(card_names)
    duplicate_actions = dedupe_source_attendance(sessions)

    # Only delivered sessions are migration candidates.
    delivered = [s for s in sessions if s.attendance]
    empty = [s for s in sessions if not s.attendance]

    # Safety guard: two delivered source blocks must never silently collapse into
    # the same CRM session key. If a future workbook contains an unrecognised
    # date/time duplication, block commit and force source review.
    key_groups: dict[tuple, list[SourceSession]] = defaultdict(list)
    for s in delivered:
        key_groups[s.key].append(s)
    duplicate_session_actions: list[Action] = []
    for key, group in key_groups.items():
        if len(group) > 1:
            refs = "; ".join(f"{x.sheet} {x.source_ref}" for x in group)
            duplicate_session_actions.append(Action(
                action="REVIEW_SOURCE_SESSION_DUPLICATE", severity="BLOCKER",
                source_file=group[0].source_file, sheet=group[0].sheet,
                source_ref=refs, session_date=group[0].session_date.isoformat(),
                activity=group[0].activity_name, start_time=group[0].start_time.strftime("%H:%M"),
                detail=f"Multiple delivered source blocks share CRM session key {key}: {refs}. Correct the source mapping before commit."
            ))

    return delivered, empty, all_people, conflicts, identity_correction_actions + resolution_actions + duplicate_actions + duplicate_session_actions


# ----------------------------- database helpers -----------------------------

def get_connection(connection_string: str):
    try:
        import pyodbc
    except ImportError:
        print("ERROR: pyodbc is required for DB preview/commit. Run: py -m pip install pyodbc", file=sys.stderr)
        raise
    return pyodbc.connect(connection_string, autocommit=False)


def fetch_dicts(cursor, sql: str, params: tuple = ()) -> list[dict[str, Any]]:
    cursor.execute(sql, params)
    cols = [c[0] for c in cursor.description]
    return [dict(zip(cols, row)) for row in cursor.fetchall()]


def verify_schema(cursor):
    required = {
        "Participants": {"ParticipantID", "SaheliCardNumber", "FullName", "DateOfBirth", "Postcode", "MobileNumber", "Site", "Notes"},
        "LiteMembers": {"Id", "MembershipId", "FirstName", "LastName", "DateOfBirth", "Phone", "Postcode"},
        "Sessions": {"SessionId", "Frequency", "Category", "VenueName", "ActivityName", "Notes", "IsRecurringWeekly", "DayOfWeek", "SessionDate", "StartTime", "EndTime", "IsBookingRequired", "IsCancelled"},
        "SessionAttendance": {"AttendanceId", "SessionId", "ParticipantId", "SessionName", "SessionDate", "SessionStartTime", "SessionEndTime", "SaheliCardNumber", "Attended", "AttendanceMemberKind", "LiteMemberId", "MemberDisplayId", "MemberName"},
    }
    rows = fetch_dicts(cursor, """
        SELECT t.name AS TableName, c.name AS ColumnName
        FROM sys.tables t
        JOIN sys.columns c ON c.object_id=t.object_id
        WHERE t.name IN ('Participants','LiteMembers','Sessions','SessionAttendance')
    """)
    found: dict[str, set[str]] = defaultdict(set)
    for r in rows:
        found[r["TableName"]].add(r["ColumnName"])
    errors = []
    for table, cols in required.items():
        missing = cols - found.get(table, set())
        if missing:
            errors.append(f"{table}: missing {sorted(missing)}")
    if errors:
        raise RuntimeError("CRM schema does not match expected migration schema: " + "; ".join(errors))


def load_db_people(cursor):
    participants: list[DbParticipant] = []
    for r in fetch_dicts(cursor, """
        SELECT ParticipantID, SaheliCardNumber, FullName, DateOfBirth, Postcode, MobileNumber
        FROM Participants
    """):
        participants.append(DbParticipant(
            participant_id=int(r["ParticipantID"]),
            card=clean_text(r["SaheliCardNumber"]),
            full_name=clean_text(r["FullName"]),
            dob=as_date(r["DateOfBirth"]),
            postcode=normalize_postcode(r["Postcode"]),
            mobile=clean_text(r["MobileNumber"]) or None,
        ))

    lites: list[DbLite] = []
    for r in fetch_dicts(cursor, """
        SELECT Id, MembershipId, FirstName, LastName, DateOfBirth, Phone, Postcode
        FROM LiteMembers
    """):
        lites.append(DbLite(
            lite_id=str(r["Id"]), membership_id=clean_text(r["MembershipId"]),
            first_name=clean_text(r["FirstName"]), last_name=clean_text(r["LastName"]),
            dob=as_date(r["DateOfBirth"]), phone=clean_text(r["Phone"]) or None,
            postcode=normalize_postcode(r["Postcode"]),
        ))
    return participants, lites


def next_lite_number(lites: list[DbLite]) -> int:
    nums = []
    for l in lites:
        m = re.fullmatch(r"LITE-(\d+)", l.membership_id.strip(), re.I)
        if m:
            nums.append(int(m.group(1)))
    return (max(nums) if nums else 0) + 1


def load_session_templates(cursor, venue: str) -> dict[str, dict[str, Any]]:
    templates: dict[str, dict[str, Any]] = {}
    # Prefer nearest same-venue template for each activity. If none, use same activity globally.
    for activity in ACTIVITY_ALIASES.values():
        same_venue = fetch_dicts(cursor, """
            SELECT TOP 1 Frequency, Category, SubCategory, ActivityCategory, Capacity, IsBookingRequired,
                         ActivityName, VenueName, SessionProviderId
            FROM Sessions
            WHERE LOWER(LTRIM(RTRIM(VenueName))) = LOWER(LTRIM(RTRIM(?)))
              AND LOWER(LTRIM(RTRIM(ActivityName))) = LOWER(LTRIM(RTRIM(?)))
            ORDER BY CASE WHEN SessionDate IS NULL THEN 1 ELSE 0 END, ABS(DATEDIFF(day, ISNULL(SessionDate, GETDATE()), GETDATE()))
        """, (venue, activity))
        rows = same_venue
        if not rows:
            rows = fetch_dicts(cursor, """
                SELECT TOP 1 Frequency, Category, SubCategory, ActivityCategory, Capacity, IsBookingRequired,
                             ActivityName, VenueName, SessionProviderId
                FROM Sessions
                WHERE LOWER(LTRIM(RTRIM(ActivityName))) = LOWER(LTRIM(RTRIM(?)))
                ORDER BY CASE WHEN SessionDate IS NULL THEN 1 ELSE 0 END, ABS(DATEDIFF(day, ISNULL(SessionDate, GETDATE()), GETDATE()))
            """, (activity,))
        if rows:
            templates[activity] = rows[0]
    return templates


def existing_session_candidates(cursor, venue: str, s: SourceSession) -> list[dict[str, Any]]:
    return fetch_dicts(cursor, """
        SELECT SessionId, ActivityName, SessionDate, StartTime, EndTime, VenueName, IsCancelled
        FROM Sessions
        WHERE SessionDate = ?
          AND LOWER(LTRIM(RTRIM(VenueName))) = LOWER(LTRIM(RTRIM(?)))
          AND LOWER(LTRIM(RTRIM(ActivityName))) = LOWER(LTRIM(RTRIM(?)))
        ORDER BY StartTime
    """, (s.session_date, venue, s.activity_name))


def choose_existing_session(candidates: list[dict[str, Any]], s: SourceSession) -> tuple[Optional[dict[str, Any]], Optional[str]]:
    if not candidates:
        return None, None
    exact = []
    for c in candidates:
        ct = as_time(c["StartTime"])
        if ct == s.start_time:
            exact.append(c)
    if len(exact) == 1:
        return exact[0], None
    if len(exact) > 1:
        return None, "Multiple CRM sessions match same date/activity/start time."
    # Old 2025 rows have inferred times. One same-day/activity CRM session is safer to reuse than creating a duplicate.
    if s.time_quality.startswith("INFERRED") and len(candidates) == 1:
        return candidates[0], None
    if len(candidates) == 1:
        ct = as_time(candidates[0]["StartTime"])
        if ct and abs((datetime.combine(date.min, ct) - datetime.combine(date.min, s.start_time)).total_seconds()) <= 30 * 60:
            return candidates[0], None
    return None, "CRM has same date/activity at a different or ambiguous time. Review before creating another session."


def insert_session(cursor, venue: str, s: SourceSession, template: dict[str, Any]) -> int:
    notes = (
        f"Historical Handsworth migration. Source={s.source_file}; {s.sheet}; {s.source_ref}; "
        f"time_quality={s.time_quality}; end_time_inferred={DURATION_MINUTES[s.activity_name]}min"
    )
    if s.date_correction:
        notes += f"; date_correction={s.date_correction}"
    cursor.execute("""
        INSERT INTO Sessions
            (Frequency, Category, SubCategory, ActivityCategory, VenueName, ActivityName, Notes,
             IsRecurringWeekly, DayOfWeek, SessionDate, ArrivalTime, StartTime, EndTime, Capacity,
             IsBookingRequired, IsCancelled, AssignedStaffId, RecurringSeriesId, SessionProviderId)
        OUTPUT INSERTED.SessionId
        VALUES (?, ?, ?, ?, ?, ?, ?, 0, NULL, ?, NULL, ?, ?, ?, ?, 0, NULL, NULL, NULL)
    """, (
        clean_text(template.get("Frequency")) or "Weekly",
        clean_text(template.get("Category")),
        clean_text(template.get("SubCategory")) or None,
        clean_text(template.get("ActivityCategory")) or None,
        venue, s.activity_name, notes,
        s.session_date, s.start_time, s.end_time,
        template.get("Capacity"), bool(template.get("IsBookingRequired") or False),
    ))
    return int(cursor.fetchone()[0])


def find_existing_attendance(cursor, session_id: int, member: ResolvedMember) -> Optional[dict[str, Any]]:
    if member.kind == "FULL":
        rows = fetch_dicts(cursor, """
            SELECT TOP 1 AttendanceId, Attended FROM SessionAttendance
            WHERE SessionId=? AND ParticipantId=?
            ORDER BY AttendanceId
        """, (session_id, member.participant_id))
    else:
        rows = fetch_dicts(cursor, """
            SELECT TOP 1 AttendanceId, Attended FROM SessionAttendance
            WHERE SessionId=? AND LiteMemberId=?
            ORDER BY AttendanceId
        """, (session_id, member.lite_id))
    return rows[0] if rows else None


def insert_attendance(cursor, session_id: int, s: SourceSession, p: PersonObservation, m: ResolvedMember):
    cursor.execute("""
        INSERT INTO SessionAttendance
            (SessionId, ParticipantId, SessionName, SessionDay, SessionDate, SessionMonth,
             SessionStartTime, SessionEndTime, SaheliCardNumber, RiskStratification, Attended,
             Notes, AttendanceMemberKind, LiteMemberId, MemberDisplayId, MemberName, Phone,
             EmergencyName, EmergencyPhone)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        session_id,
        m.participant_id if m.kind == "FULL" else None,
        s.activity_name,
        s.session_date.strftime("%A"),
        s.session_date,
        s.session_date.strftime("%B"),
        s.start_time,
        s.end_time,
        m.card if m.kind == "FULL" else None,
        p.risk,
        f"Historical Handsworth migration; source={s.source_file}; {p.sheet} row {p.row}",
        m.kind,
        m.lite_id if m.kind == "LITE" else None,
        m.display_id,
        m.name or p.name,
        m.phone,
        p.emergency_name,
        p.emergency_phone,
    ))


# ----------------------------- member resolution ----------------------------

def resolve_member(
    cursor,
    p: PersonObservation,
    participants_by_card: dict[str, DbParticipant],
    participants_by_name: dict[str, list[DbParticipant]],
    lites_by_name: dict[str, list[DbLite]],
    source_conflicting_cards: dict[str, set[str]],
    next_lite: list[int],
    venue: str,
    cache: dict[tuple, ResolvedMember],
) -> tuple[Optional[ResolvedMember], Optional[str]]:
    cache_key = ("CARD", p.card) if p.card else ("NAME", p.name_norm, p.dob, p.postcode)
    if cache_key in cache:
        return cache[cache_key], None

    if p.card:
        existing = participants_by_card.get(p.card)
        if existing:
            if p.name_norm and existing.full_name:
                sim = name_similarity(p.name, existing.full_name)
                if sim < 0.80:
                    return None, (
                        f"Saheli Card {p.card} belongs to CRM participant '{existing.full_name}', "
                        f"but source row says '{p.name}' (name similarity {sim:.2f})."
                    )
            m = ResolvedMember(
                kind="FULL", participant_id=existing.participant_id, lite_id=None,
                display_id=existing.card, name=existing.full_name or p.name,
                card=existing.card, phone=existing.mobile,
                action="REUSE_FULL_BY_CARD", detail=f"ParticipantID={existing.participant_id}",
            )
            cache[cache_key] = m
            return m, None

        if p.card in source_conflicting_cards:
            return None, f"Source uses card {p.card} against materially different names: {sorted(source_conflicting_cards[p.card])}"

        if not clean_text(p.name):
            return None, f"Saheli Card {p.card} is not in CRM and this source row has no usable name; refusing to create a nameless FULL participant."

        cursor.execute("""
            INSERT INTO Participants (SaheliCardNumber, FullName, DateOfBirth, Postcode, Site, Notes)
            OUTPUT INSERTED.ParticipantID
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            p.card, p.name or None, p.dob, p.postcode, venue,
            f"Created by historical Handsworth migration from {p.source_file}; source row {p.row}. Existing CRM data is never overwritten."
        ))
        pid = int(cursor.fetchone()[0])
        dbp = DbParticipant(pid, p.card, p.name, p.dob, p.postcode, None)
        participants_by_card[p.card] = dbp
        participants_by_name[norm_name(p.name)].append(dbp)
        m = ResolvedMember(
            kind="FULL", participant_id=pid, lite_id=None, display_id=p.card,
            name=p.name, card=p.card, phone=None,
            action="CREATE_FULL_FROM_VALID_CARD", detail=f"New ParticipantID={pid}",
        )
        cache[cache_key] = m
        return m, None

    # No source card: Grace Gittens is an explicitly reviewed Handsworth exception.
    # Her 2025 row carried Greta's card 1123 in error, while CRM has one exact
    # Handsworth FULL participant named Grace Gittens. Reuse only when the exact
    # normalized FullName is unique in CRM.
    full_candidates = participants_by_name.get(p.name_norm, [])
    if p.name_norm in SAFE_EXACT_FULL_NAME_REUSE:
        unique_exact = {c.participant_id: c for c in full_candidates}
        if len(unique_exact) == 1:
            c = next(iter(unique_exact.values()))
            m = ResolvedMember(
                kind="FULL", participant_id=c.participant_id, lite_id=None,
                display_id=c.card, name=c.full_name or p.name, card=c.card, phone=c.mobile,
                action="REUSE_FULL_BY_REVIEWED_EXACT_NAME",
                detail=f"ParticipantID={c.participant_id}; source wrong-card correction reviewed",
            )
            cache[cache_key] = m
            return m, None
        if len(unique_exact) > 1:
            return None, f"Multiple FULL CRM participants exactly match reviewed name '{p.name}'."

    # Otherwise only reuse FULL when exact name + DOB/postcode makes it safe.
    safe_full = []
    for c in full_candidates:
        if p.dob and c.dob and p.dob == c.dob:
            safe_full.append(c)
        elif p.postcode and c.postcode and p.postcode == c.postcode:
            safe_full.append(c)
    if len({c.participant_id for c in safe_full}) == 1:
        c = safe_full[0]
        m = ResolvedMember(
            kind="FULL", participant_id=c.participant_id, lite_id=None,
            display_id=c.card, name=c.full_name or p.name, card=c.card, phone=c.mobile,
            action="REUSE_FULL_BY_NAME_AND_PROFILE", detail=f"ParticipantID={c.participant_id}",
        )
        cache[cache_key] = m
        return m, None
    if len({c.participant_id for c in safe_full}) > 1:
        return None, f"Multiple FULL CRM participants match name/profile for '{p.name}'."

    # Lite matching follows Calthorpe approach: exact normalized name; resolve duplicates using DOB/postcode.
    lite_candidates = lites_by_name.get(p.name_norm, [])
    if len(lite_candidates) == 1:
        c = lite_candidates[0]
        if p.dob and c.dob and p.dob != c.dob:
            return None, f"Exact Lite name match '{p.name}' has conflicting DOB ({p.dob} vs {c.dob})."
        if p.postcode and c.postcode and p.postcode != c.postcode:
            return None, f"Exact Lite name match '{p.name}' has conflicting postcode ({p.postcode} vs {c.postcode})."
        m = ResolvedMember(
            kind="LITE", participant_id=None, lite_id=c.lite_id,
            display_id=c.membership_id, name=c.full_name, card=None, phone=c.phone,
            action="REUSE_LITE_BY_EXACT_NAME", detail=f"LiteMemberId={c.lite_id}",
        )
        cache[cache_key] = m
        return m, None

    if len(lite_candidates) > 1:
        narrowed = [c for c in lite_candidates if (p.dob and c.dob and p.dob == c.dob) or (p.postcode and c.postcode and p.postcode == c.postcode)]
        unique = {c.lite_id: c for c in narrowed}
        if len(unique) == 1:
            c = next(iter(unique.values()))
            m = ResolvedMember(
                kind="LITE", participant_id=None, lite_id=c.lite_id,
                display_id=c.membership_id, name=c.full_name, card=None, phone=c.phone,
                action="REUSE_LITE_BY_NAME_AND_PROFILE", detail=f"LiteMemberId={c.lite_id}",
            )
            cache[cache_key] = m
            return m, None
        return None, f"Multiple LiteMembers share the exact name '{p.name}' and source profile cannot resolve uniquely."

    first, last = split_name(p.name)
    lite_id = str(uuid.uuid4())
    membership_id = f"LITE-{next_lite[0]}"
    next_lite[0] += 1
    cursor.execute("""
        INSERT INTO LiteMembers
            (Id, MembershipId, FirstName, LastName, DateOfBirth, Phone, Email, Address, Postcode,
             EmergencyName, EmergencyPhone, EmergencyRelation, HealthConditions, Gender, Ethnicity, CreatedByUserId)
        VALUES (?, ?, ?, ?, ?, NULL, NULL, NULL, ?, ?, ?, NULL, NULL, NULL, NULL, NULL)
    """, (
        lite_id, membership_id, first, last, p.dob, p.postcode,
        p.emergency_name, p.emergency_phone,
    ))
    c = DbLite(lite_id, membership_id, first, last, p.dob, None, p.postcode)
    lites_by_name[p.name_norm].append(c)
    m = ResolvedMember(
        kind="LITE", participant_id=None, lite_id=lite_id, display_id=membership_id,
        name=c.full_name, card=None, phone=None,
        action="CREATE_LITE", detail=f"New LiteMemberId={lite_id}",
    )
    cache[cache_key] = m
    return m, None


# --------------------------------- reports ----------------------------------

def action_from(s: SourceSession, p: Optional[PersonObservation], action: str, detail: str, severity="INFO", **kwargs) -> Action:
    return Action(
        action=action,
        severity=severity,
        source_file=s.source_file,
        sheet=s.sheet,
        source_ref=(kwargs.get("source_ref") or (f"Row {p.row}" if p else s.source_ref)),
        session_date=s.session_date.isoformat(),
        activity=s.activity_name,
        start_time=s.start_time.strftime("%H:%M"),
        source_name=p.name if p else "",
        source_card=(p.card or "") if p else "",
        member_kind=kwargs.get("member_kind", ""),
        member_display_id=kwargs.get("member_display_id", ""),
        session_id=str(kwargs.get("session_id", "") or ""),
        detail=detail,
    )


def write_actions(actions: list[Action], out_dir: Path, prefix: str) -> Path:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = out_dir / f"{prefix}_{ts}.csv"
    fields = list(Action.__dataclass_fields__.keys())
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for a in actions:
            w.writerow({k: getattr(a, k) for k in fields})
    return path


def write_review_people(actions: list[Action], out_dir: Path) -> Optional[Path]:
    rows = [a for a in actions if a.action == "REVIEW_PARTICIPANT"]
    if not rows:
        return None
    grouped: dict[tuple[str, str, str], list[Action]] = defaultdict(list)
    for a in rows:
        grouped[((a.source_card or "").strip(), norm_name(a.source_name), a.detail)].append(a)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = out_dir / f"handsworth_participant_review_{ts}.csv"
    fields = ["attendance_rows", "source_card", "source_name", "reason", "example_date", "example_activity", "example_source_ref"]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for (card, _name_norm, reason), items in sorted(grouped.items(), key=lambda kv: (-len(kv[1]), kv[0][0], kv[0][1])):
            ex = items[0]
            w.writerow({
                "attendance_rows": len(items),
                "source_card": card,
                "source_name": ex.source_name,
                "reason": reason,
                "example_date": ex.session_date,
                "example_activity": ex.activity,
                "example_source_ref": ex.source_ref,
            })
    return path


def write_summary(actions: list[Action], sessions: list[SourceSession], empty: list[SourceSession], out_dir: Path, mode: str) -> Path:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = out_dir / f"handsworth_migration_summary_{ts}.txt"
    counts = Counter(a.action for a in actions)
    total_att = sum(len(s.attendance) for s in sessions)
    dates = [s.session_date for s in sessions]
    blockers = [a for a in actions if a.action.startswith("REVIEW_") or a.severity == "BLOCKER"]
    with path.open("w", encoding="utf-8") as f:
        f.write("Saheli CRM - Handsworth historical migration\n")
        f.write("===========================================\n\n")
        f.write(f"Mode: {mode}\n")
        f.write(f"Delivered source sessions: {len(sessions)}\n")
        f.write(f"Empty source placeholders skipped: {len(empty)}\n")
        f.write(f"Clean source attendance rows: {total_att}\n")
        if dates:
            f.write(f"Source period: {min(dates)} to {max(dates)}\n")
        f.write(f"Review blockers: {len(blockers)}\n")
        unique_review_people = {((a.source_card or '').strip(), norm_name(a.source_name)) for a in blockers if a.action == "REVIEW_PARTICIPANT"}
        unique_review_people.discard(("", ""))
        unique_new_lites = {a.member_display_id for a in actions if a.action == "CREATE_LITE" and a.member_display_id}
        unique_new_full = {a.member_display_id for a in actions if a.action == "CREATE_FULL_FROM_VALID_CARD" and a.member_display_id}
        f.write(f"Unique participant identities needing REVIEW_PARTICIPANT: {len(unique_review_people)}\n")
        f.write(f"Unique new Lite members proposed: {len(unique_new_lites)}\n")
        f.write(f"Unique new FULL participants proposed: {len(unique_new_full)}\n\n")
        f.write("Action counts (trace rows; member-resolution actions can repeat per attendance)\n--------------------------------------------------------------------------\n")
        for k, v in sorted(counts.items()):
            f.write(f"{k}: {v}\n")
        if blockers:
            f.write("\nBLOCKERS - do not commit until resolved\n---------------------------------------\n")
            for a in blockers:
                f.write(f"{a.session_date} | {a.activity} | {a.source_name} | {a.source_card} | {a.action} | {a.detail}\n")
    return path


def print_source_summary(sessions, empty, source_actions, conflicts):
    total = sum(len(s.attendance) for s in sessions)
    dates = [s.session_date for s in sessions]
    by_activity = defaultdict(lambda: [0, 0])
    for s in sessions:
        by_activity[s.activity_name][0] += 1
        by_activity[s.activity_name][1] += len(s.attendance)
    print("\nHandsworth source audit")
    print("-----------------------")
    print(f"Delivered sessions: {len(sessions)}")
    print(f"Empty placeholders skipped: {len(empty)}")
    print(f"Clean attendance rows: {total}")
    if dates:
        print(f"Period: {min(dates)} to {max(dates)}")
    for activity, (sc, ac) in sorted(by_activity.items()):
        print(f"  {activity}: {sc} sessions, {ac} attendance")
    print(f"Source card conflicts requiring CRM validation: {len(conflicts)}")
    print(f"Source resolution/dedup actions: {len(source_actions)}")


# -------------------------------- migration ---------------------------------

def run_db_preview_or_commit(args, sessions, empty, source_conflicts, source_actions):
    conn_str = args.connection_string or os.getenv("SAHELI_SQL_CONNECTION_STRING", "")
    if not conn_str:
        raise RuntimeError(
            "No SQL connection string. Set SAHELI_SQL_CONNECTION_STRING or pass --connection-string."
        )
    conn = get_connection(conn_str)
    cursor = conn.cursor()
    actions = list(source_actions)
    try:
        verify_schema(cursor)

        # Verify exact Handsworth venue exists before any creation.
        venues = fetch_dicts(cursor, """
            SELECT DISTINCT VenueName FROM Sessions
            WHERE LOWER(LTRIM(RTRIM(VenueName))) = LOWER(LTRIM(RTRIM(?)))
        """, (args.venue,))
        if not venues:
            actions.append(Action(
                action="REVIEW_VENUE_NOT_FOUND", severity="BLOCKER",
                detail=f"No existing Sessions.VenueName exactly matches '{args.venue}'. Use --venue with the exact CRM venue."
            ))

        participants, lites = load_db_people(cursor)
        participants_by_card = {p.card: p for p in participants if p.card}
        participants_by_name: dict[str, list[DbParticipant]] = defaultdict(list)
        for p in participants:
            if norm_name(p.full_name):
                participants_by_name[norm_name(p.full_name)].append(p)
        lites_by_name: dict[str, list[DbLite]] = defaultdict(list)
        for l in lites:
            if norm_name(l.full_name):
                lites_by_name[norm_name(l.full_name)].append(l)
        next_lite = [next_lite_number(lites)]
        member_cache: dict[tuple, ResolvedMember] = {}
        templates = load_session_templates(cursor, args.venue)

        for activity in sorted({s.activity_name for s in sessions}):
            if activity not in templates or not clean_text(templates[activity].get("Category")):
                actions.append(Action(
                    action="REVIEW_MISSING_SESSION_TEMPLATE", severity="BLOCKER",
                    activity=activity,
                    detail=f"Cannot safely create new '{activity}' sessions because no existing CRM session provides a non-null Category template. Existing matching sessions can still be reused."
                ))

        session_cache: dict[tuple, Optional[int]] = {}
        # Tracks attendance inserted earlier in THIS transaction so it is not mislabeled as pre-existing CRM data.
        created_attendance_keys: set[tuple] = set()
        for s in sorted(sessions, key=lambda x: (x.session_date, x.start_time, x.activity_name)):
            session_id: Optional[int] = None
            if s.key in session_cache:
                session_id = session_cache[s.key]
            else:
                candidates = existing_session_candidates(cursor, args.venue, s)
                chosen, problem = choose_existing_session(candidates, s)
                if problem:
                    actions.append(action_from(s, None, "REVIEW_SESSION_AMBIGUOUS", problem, severity="BLOCKER"))
                    session_cache[s.key] = None
                elif chosen:
                    session_id = int(chosen["SessionId"])
                    if bool(chosen.get("IsCancelled")):
                        actions.append(action_from(s, None, "REVIEW_EXISTING_SESSION_CANCELLED", f"SessionId={session_id} is cancelled in CRM.", severity="BLOCKER", session_id=session_id))
                        session_id = None
                    else:
                        actions.append(action_from(s, None, "REUSE_SESSION", f"Existing SessionId={session_id}", session_id=session_id))
                        session_cache[s.key] = session_id
                else:
                    template = templates.get(s.activity_name)
                    if not template or not clean_text(template.get("Category")):
                        actions.append(action_from(s, None, "REVIEW_CANNOT_CREATE_SESSION", "No safe session Category template is available.", severity="BLOCKER"))
                        session_cache[s.key] = None
                    else:
                        session_id = insert_session(cursor, args.venue, s, template)
                        actions.append(action_from(s, None, "CREATE_SESSION", f"New SessionId={session_id}; Category copied from existing '{s.activity_name}' template; source time preserved; end time inferred.", session_id=session_id))
                        session_cache[s.key] = session_id

            if not session_id:
                for a in s.attendance:
                    actions.append(action_from(s, a.person, "REVIEW_ATTENDANCE_BLOCKED_BY_SESSION", "Attendance cannot be processed until its session is resolved.", severity="BLOCKER", source_ref=a.source_ref))
                continue

            for a in s.attendance:
                p = a.person
                member, problem = resolve_member(
                    cursor, p, participants_by_card, participants_by_name, lites_by_name,
                    source_conflicts, next_lite, args.venue, member_cache
                )
                if problem or not member:
                    actions.append(action_from(s, p, "REVIEW_PARTICIPANT", problem or "Unable to resolve participant", severity="BLOCKER", session_id=session_id, source_ref=a.source_ref))
                    continue

                # Log member resolution once per source attendance for traceability.
                actions.append(action_from(
                    s, p, member.action, member.detail,
                    member_kind=member.kind, member_display_id=member.display_id,
                    session_id=session_id, source_ref=a.source_ref
                ))

                resolved_key = (
                    session_id,
                    member.kind,
                    member.participant_id if member.kind == "FULL" else member.lite_id,
                )
                existing_att = find_existing_attendance(cursor, session_id, member)
                if existing_att:
                    if bool(existing_att["Attended"]):
                        if resolved_key in created_attendance_keys:
                            actions.append(action_from(
                                s, p, "SKIP_RESOLVED_DUPLICATE_ATTENDANCE",
                                f"This source row resolves to the same member/session as an attendance inserted earlier in this migration transaction (AttendanceId={existing_att['AttendanceId']}).",
                                member_kind=member.kind, member_display_id=member.display_id,
                                session_id=session_id, source_ref=a.source_ref
                            ))
                        else:
                            actions.append(action_from(
                                s, p, "SKIP_EXISTING_ATTENDANCE", f"AttendanceId={existing_att['AttendanceId']} already existed in CRM and is attended.",
                                member_kind=member.kind, member_display_id=member.display_id,
                                session_id=session_id, source_ref=a.source_ref
                            ))
                    else:
                        actions.append(action_from(
                            s, p, "REVIEW_EXISTING_ATTENDANCE_NOT_ATTENDED",
                            f"AttendanceId={existing_att['AttendanceId']} already exists with Attended=0. Script will not overwrite it.",
                            severity="BLOCKER", member_kind=member.kind, member_display_id=member.display_id,
                            session_id=session_id, source_ref=a.source_ref
                        ))
                    continue

                insert_attendance(cursor, session_id, s, p, member)
                created_attendance_keys.add(resolved_key)
                actions.append(action_from(
                    s, p, "CREATE_ATTENDANCE", "New attended SessionAttendance row.",
                    member_kind=member.kind, member_display_id=member.display_id,
                    session_id=session_id, source_ref=a.source_ref
                ))

        blockers = [a for a in actions if a.action.startswith("REVIEW_") or a.severity == "BLOCKER"]
        mode = "COMMIT" if args.commit else "PREVIEW_ROLLBACK"
        if args.commit and not blockers:
            conn.commit()
            print("\nCOMMIT COMPLETE: Handsworth migration committed successfully.")
        else:
            conn.rollback()
            if args.commit and blockers:
                print(f"\nCOMMIT REFUSED: {len(blockers)} REVIEW/BLOCKER item(s) remain. Transaction rolled back.")
            else:
                print("\nPREVIEW COMPLETE: transaction rolled back; CRM was not changed.")

        out_dir = Path(args.output_dir).resolve()
        out_dir.mkdir(parents=True, exist_ok=True)
        csv_path = write_actions(actions, out_dir, "handsworth_migration_preview" if not args.commit else "handsworth_migration_commit")
        sum_path = write_summary(actions, sessions, empty, out_dir, mode)
        review_path = write_review_people(actions, out_dir)
        print(f"Actions CSV: {csv_path}")
        print(f"Summary:     {sum_path}")
        if review_path:
            print(f"Participant review CSV: {review_path}")
        print(f"Review blockers: {len(blockers)}")
        review_rows = [a for a in blockers if a.action == "REVIEW_PARTICIPANT"]
        review_grouped: dict[tuple[str, str, str], list[Action]] = defaultdict(list)
        for a in review_rows:
            review_grouped[((a.source_card or '').strip(), norm_name(a.source_name), a.detail)].append(a)
        unique_review_people = {(card, name_norm) for (card, name_norm, _reason) in review_grouped}
        unique_review_people.discard(("", ""))
        unique_new_lites = {a.member_display_id for a in actions if a.action == "CREATE_LITE" and a.member_display_id}
        unique_new_full = {a.member_display_id for a in actions if a.action == "CREATE_FULL_FROM_VALID_CARD" and a.member_display_id}
        resolved_dupes = [a for a in actions if a.action == "SKIP_RESOLVED_DUPLICATE_ATTENDANCE"]
        print(f"Unique participant identities needing review: {len(unique_review_people)}")
        print(f"Unique new Lite members proposed: {len(unique_new_lites)}")
        print(f"Unique new FULL participants proposed: {len(unique_new_full)}")
        print(f"Resolved duplicate attendance rows skipped inside this migration: {len(resolved_dupes)}")
        if review_grouped:
            print("\nParticipants needing review (unique reason groups):")
            for (card, _name_norm, reason), items in sorted(review_grouped.items(), key=lambda kv: (-len(kv[1]), kv[0][0], kv[0][1])):
                ex = items[0]
                card_text = card or "NO CARD"
                print(f"  {len(items):>3} row(s) | {card_text} | {ex.source_name} | {reason}")
        print("\nAction summary (trace rows; member actions may repeat per attendance):")
        for k, v in sorted(Counter(a.action for a in actions).items()):
            print(f"  {k}: {v}")
        return 2 if args.commit and blockers else 0
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def main():
    parser = argparse.ArgumentParser(description="Saheli CRM Handsworth historical migration")
    parser.add_argument("--source-dir", default=str(Path(__file__).resolve().parent), help="Folder containing the two Handsworth Excel files")
    parser.add_argument("--output-dir", default=str(Path(__file__).resolve().parent), help="Folder for audit/preview CSV and summary")
    parser.add_argument("--venue", default=DEFAULT_VENUE, help="Exact CRM Sessions.VenueName (default: Handsworth)")
    parser.add_argument("--connection-string", default="", help="Optional ODBC connection string; safer to use SAHELI_SQL_CONNECTION_STRING env variable")
    parser.add_argument("--audit-only", action="store_true", help="Parse Excel only; no DB connection")
    parser.add_argument("--commit", action="store_true", help="Commit only if zero REVIEW/BLOCKER items remain")
    args = parser.parse_args()

    if args.audit_only and args.commit:
        parser.error("--audit-only and --commit cannot be used together")

    source_dir = Path(args.source_dir).resolve()
    out_dir = Path(args.output_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    sessions, empty, all_people, conflicts, source_actions = parse_sources(source_dir)
    print_source_summary(sessions, empty, source_actions, conflicts)

    # Explicitly report materially conflicting source cards during audit.
    for card, names in sorted(conflicts.items()):
        print(f"  REVIEW source card {card}: {sorted(names)}")

    if args.audit_only:
        audit_actions = list(source_actions)
        for card, names in sorted(conflicts.items()):
            audit_actions.append(Action(
                action="REVIEW_SOURCE_CARD_CONFLICT", severity="BLOCKER",
                source_card=card, detail=f"Card is used against materially different source names: {sorted(names)}"
            ))
        csv_path = write_actions(audit_actions, out_dir, "handsworth_source_audit")
        sum_path = write_summary(audit_actions, sessions, empty, out_dir, "AUDIT_ONLY")
        print(f"\nAudit CSV: {csv_path}")
        print(f"Summary:   {sum_path}")
        print("No database connection was made.")
        return 0

    return run_db_preview_or_commit(args, sessions, empty, conflicts, source_actions)


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("Cancelled.", file=sys.stderr)
        sys.exit(130)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
