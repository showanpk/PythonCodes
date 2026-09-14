#!/usr/bin/env python3
"""
OCF CRM - Monday Attendance / Activity Migration PREVIEW V1

READ-ONLY preview for:
    Activity Type -> Activity -> Participant allocation -> Session -> Attendance

This script DOES NOT write to MySQL.

What it does:
- Reads all known OCF Monday attendance .xlsx exports in a folder.
- Treats each Monday attendance board as a historical Activity/run.
- Maps programme folders into OCF Activity Types:
      Girls Cricket / Boys Cricket -> Cricket
      Girls Football / Boys Football -> FootBall
      Girls Youth Club / Youth Club (Mixed) -> Youth Club
      HAF -> HAF
      XLR8 -> XLR8
- Adds a desired activity-only row for Boys Cricket (no historical sessions/attendance).
- Treats each participant row on a board as evidence of programme/activity enrolment.
- Matches source participants to the existing OCF participants table using
  normalized Full Name + Date of Birth only.
- Treats each dated Yes/No column as a session candidate.
- Flags explicit "Cancelled" columns as cancelled session candidates.
- Does NOT assume an all-blank date column was delivered.
- Detects possible duplicate source sessions across boards under the same programme.
- Inspects MySQL schema/table names for Activity Types, Activities, allocations,
  Sessions and attendance so the commit script can be built against the real schema.

Dependencies:
    py -m pip install openpyxl PyMySQL

Connection:
    $rawPassword = "YOUR_PASSWORD"
    $encodedPassword = [uri]::EscapeDataString($rawPassword)
    $env:OCF_MYSQL_CONNECTION_STRING="mysql://USER:$encodedPassword@IP:3306/DB?charset=utf8mb4"

Run:
    py -3.14 .\preview_ocf_monday_activity_sessions.py `
      --source-dir ".\Session Excels"

If Girls Cricket is outside that folder:
    py -3.14 .\preview_ocf_monday_activity_sessions.py `
      --source-dir ".\Session Excels" `
      --extra-file ".\Girls Cricket Attendance.xlsx"

Outputs are written beside the script unless --output-dir is supplied.

IMPORTANT:
This is a PREVIEW ONLY script. There is intentionally no --commit flag.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import parse_qs, unquote, urlparse

import pymysql
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


DEFAULT_CONNECTION_STRING = os.getenv(
    "OCF_MYSQL_CONNECTION_STRING",
    "mysql://YOUR_DB_USER:YOUR_DB_PASSWORD@YOUR_DB_HOST:3306/YOUR_DB_NAME?charset=utf8mb4",
)

NON_SESSION_HEADERS = {
    "name",
    "participant list",
    "girls participant list",
    "status",
    "date of birth",
    "age",
    "total sessions attended",
    "attendance %",
    "attendance%",
}

YES_VALUES = {"yes", "y", "present", "attended", "1"}
NO_VALUES = {"no", "n", "absent", "0"}

MONTHS = {
    "jan": 1, "january": 1,
    "feb": 2, "february": 2,
    "mar": 3, "march": 3,
    "apr": 4, "april": 4,
    "may": 5,
    "jun": 6, "june": 6,
    "jul": 7, "july": 7,
    "aug": 8, "august": 8,
    "sep": 9, "sept": 9, "september": 9,
    "oct": 10, "october": 10,
    "nov": 11, "november": 11,
    "dec": 12, "december": 12,
}

DAY_WORDS = {
    "mon", "monday", "tue", "tues", "tuesday", "wed", "weds", "wednesday",
    "thu", "thur", "thurs", "thursday", "fri", "friday", "sat", "saturday",
    "sun", "sunday",
}


@dataclass(frozen=True)
class ProgrammeConfig:
    source_programme: str
    activity_type: str


@dataclass
class SourceParticipant:
    file: str
    sheet: str
    source_programme: str
    activity_type: str
    activity_name: str
    row_number: int
    source_name: str
    source_dob: str | None
    source_status: str | None
    match_status: str = ""
    matched_participant_id: str | None = None
    matched_ocf_id: str | None = None
    matched_name: str | None = None
    match_detail: str = ""


@dataclass
class SessionCandidate:
    file: str
    sheet: str
    source_programme: str
    activity_type: str
    activity_name: str
    header: str
    session_date: str
    cancelled: bool
    yes_count: int
    no_count: int
    blank_count: int
    marked_count: int
    source_participant_count: int
    status: str
    possible_duplicate_group: str = ""


@dataclass
class FileAudit:
    file: str
    sheet: str
    source_programme: str
    activity_type: str
    activity_name: str
    inferred_year: int | None
    participant_rows: int
    matched_rows: int
    unmatched_rows: int
    ambiguous_rows: int
    date_columns: int
    delivered_session_candidates: int
    cancelled_session_candidates: int
    empty_unconfirmed_date_columns: int
    yes_attendance_marks: int
    no_attendance_marks: int
    warning: str = ""


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def normalize_name(value: Any) -> str:
    text = clean_text(value) or ""
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.casefold()
    text = text.replace("-", " ")
    text = re.sub(r"[^a-z0-9\s]", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def parse_dob(value: Any, epoch) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        try:
            parsed = from_excel(value, epoch)
            if isinstance(parsed, datetime):
                return parsed.date().isoformat()
            if isinstance(parsed, date):
                return parsed.isoformat()
        except Exception:
            return None

    text = clean_text(value)
    if not text:
        return None

    for fmt in (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%d.%m.%Y",
        "%m/%d/%Y",
        "%b %d, %Y",
        "%B %d, %Y",
    ):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def classify_file(filename: str) -> ProgrammeConfig | None:
    n = filename.casefold()

    if "girls_cricket" in n or "girls cricket" in n:
        return ProgrammeConfig("Girls Cricket", "Cricket")

    if n.startswith("girls_football") or "girls football" in n:
        return ProgrammeConfig("Girls Football", "FootBall")

    if n.startswith("girls_attendance") or "girls youth club" in n:
        return ProgrammeConfig("Girls Youth Club", "Youth Club")

    if n.startswith("haf_") or n.startswith("haf "):
        return ProgrammeConfig("HAF", "HAF")

    if n.startswith("ocf_football") or "ocf football" in n:
        return ProgrammeConfig("Boys Football", "FootBall")

    if n.startswith("ocf_youth_club") or "ocf youth club" in n:
        return ProgrammeConfig("Youth Club (Mixed)", "Youth Club")

    if n.startswith("xlr8") or "xlr8" in n:
        return ProgrammeConfig("XLR8", "XLR8")

    return None


def infer_year(filename: str, title1: str, title2: str) -> int | None:
    # Ignore the long Monday export ID at end of filename.
    basename = re.sub(r"_\d{9,}\.xlsx$", "", filename, flags=re.I)
    text = f"{basename} {title1} {title2}"

    full = re.findall(r"\b(20(?:24|25|26))\b", text)
    if full:
        counts = Counter(map(int, full))
        return counts.most_common(1)[0][0]

    short = re.findall(r"(?<!\d)(24|25|26)(?!\d)", text)
    if short:
        counts = Counter(2000 + int(v) for v in short)
        return counts.most_common(1)[0][0]

    return None


def parse_session_header(header: Any, inferred_year: int | None) -> tuple[str | None, bool]:
    text = clean_text(header)
    if not text:
        return None, False

    key = text.casefold().strip()
    if key in NON_SESSION_HEADERS:
        return None, False

    cancelled = "cancelled" in key or "canceled" in key

    # Explicit dd/mm/yyyy.
    m = re.search(r"\b(\d{1,2})/(\d{1,2})/(20\d{2})\b", key)
    if m:
        d, mo, y = map(int, m.groups())
        try:
            return date(y, mo, d).isoformat(), cancelled
        except ValueError:
            return None, cancelled

    if inferred_year is None:
        return None, cancelled

    # Remove parenthetical descriptors such as (HAF), (F), (cancelled).
    simple = re.sub(r"\([^)]*\)", " ", key)
    simple = simple.replace(",", " ")
    simple = re.sub(r"\s+", " ", simple).strip()

    tokens = simple.split()
    tokens = [t for t in tokens if t not in DAY_WORDS]

    day_num = None
    month_num = None

    for token in tokens:
        dm = re.fullmatch(r"(\d{1,2})(?:st|nd|rd|th)?", token)
        if dm and day_num is None:
            day_num = int(dm.group(1))
            continue
        if token in MONTHS:
            month_num = MONTHS[token]

    if day_num is None or month_num is None:
        return None, cancelled

    try:
        return date(inferred_year, month_num, day_num).isoformat(), cancelled
    except ValueError:
        return None, cancelled


def normalize_attendance(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return "blank"

    key = text.casefold()
    if key in YES_VALUES:
        return "yes"
    if key in NO_VALUES:
        return "no"
    return "other"


def activity_name_from_titles(title1: str | None, title2: str | None, filename: str) -> str:
    # Monday board name is normally row 1. Preserve that source wording where available.
    if title1:
        return title1.strip()
    if title2:
        return title2.strip()
    return re.sub(r"_\d{9,}\.xlsx$", "", filename, flags=re.I).replace("_", " ")


def find_header_row(ws) -> tuple[int, list[Any]]:
    for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = list(row)
        normalized = {str(v).strip().casefold() for v in values if v is not None}
        if "date of birth" in normalized and (
            "participant list" in normalized or "girls participant list" in normalized
        ):
            return row_number, values
        if row_number >= 10:
            break
    raise RuntimeError("Could not find attendance header row.")


def connect_mysql(connection_string: str):
    parsed = urlparse(connection_string)
    if parsed.scheme not in {"mysql", "mysql+pymysql"}:
        raise ValueError("Connection string must begin mysql://")
    if not parsed.hostname or not parsed.username or not parsed.path.strip("/"):
        raise ValueError("Connection string is incomplete.")

    query = parse_qs(parsed.query)
    kwargs = {
        "host": parsed.hostname,
        "port": parsed.port or 3306,
        "user": unquote(parsed.username),
        "password": unquote(parsed.password or ""),
        "database": parsed.path.lstrip("/"),
        "charset": query.get("charset", ["utf8mb4"])[0],
        "autocommit": False,
        "cursorclass": pymysql.cursors.DictCursor,
        "connect_timeout": 15,
        "read_timeout": 60,
        "write_timeout": 60,
    }
    if query.get("ssl", ["false"])[0].casefold() in {"1", "true", "yes"}:
        kwargs["ssl"] = {}
    conn = pymysql.connect(**kwargs)
    with conn.cursor() as cur:
        cur.execute("SET time_zone = '+00:00'")
    return conn


def get_columns(conn, table: str) -> list[dict[str, Any]]:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT
                COLUMN_NAME,
                DATA_TYPE,
                IS_NULLABLE,
                COLUMN_DEFAULT,
                COLUMN_KEY,
                EXTRA
            FROM information_schema.COLUMNS
            WHERE TABLE_SCHEMA = DATABASE()
              AND TABLE_NAME = %s
            ORDER BY ORDINAL_POSITION
            """,
            (table,),
        )
        return cur.fetchall()


def all_tables(conn) -> list[str]:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT TABLE_NAME
            FROM information_schema.TABLES
            WHERE TABLE_SCHEMA = DATABASE()
              AND TABLE_TYPE = 'BASE TABLE'
            ORDER BY TABLE_NAME
            """
        )
        return [row["TABLE_NAME"] for row in cur.fetchall()]


def related_schema_tables(conn) -> dict[str, list[dict[str, Any]]]:
    tables = all_tables(conn)
    terms = ("activity", "session", "attendance", "participant")
    related = {}
    for table in tables:
        name = table.casefold()
        if any(term in name for term in terms):
            related[table] = get_columns(conn, table)
    return related


def load_db_participants(conn) -> tuple[
    dict[tuple[str, str], list[dict[str, Any]]],
    dict[str, list[dict[str, Any]]],
]:
    cols = {c["COLUMN_NAME"].casefold() for c in get_columns(conn, "participants")}
    required = {"id", "full_name", "date_of_birth"}
    if not required.issubset(cols):
        raise RuntimeError(
            "participants table does not contain required columns id, full_name, date_of_birth."
        )

    select_cols = ["id", "full_name", "date_of_birth"]
    if "ocf_id" in cols:
        select_cols.append("ocf_id")
    else:
        select_cols.append("NULL AS ocf_id")

    with conn.cursor() as cur:
        cur.execute(f"SELECT {', '.join(select_cols)} FROM participants")
        rows = cur.fetchall()

    by_name_dob: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)

    for row in rows:
        name_key = normalize_name(row.get("full_name"))
        dob_raw = row.get("date_of_birth")
        dob = str(dob_raw) if dob_raw is not None else None
        if name_key:
            by_name[name_key].append(row)
        if name_key and dob:
            by_name_dob[(name_key, dob)].append(row)

    return by_name_dob, by_name


def resolve_participant(
    source_name: str,
    source_dob: str | None,
    by_name_dob,
    by_name,
) -> tuple[str, dict[str, Any] | None, str]:
    name_key = normalize_name(source_name)
    if not name_key:
        return "UNMATCHED", None, "Source name is blank/unusable."

    if source_dob:
        candidates = by_name_dob.get((name_key, source_dob), [])
        if len(candidates) == 1:
            return "MATCHED_NAME_DOB", candidates[0], "Exact normalized name + DOB."
        if len(candidates) > 1:
            return "AMBIGUOUS", None, f"{len(candidates)} CRM rows share exact normalized name + DOB."

        same_name = by_name.get(name_key, [])
        if len(same_name) == 1:
            row = same_name[0]
            return (
                "REVIEW_DOB_MISMATCH",
                None,
                f"Exact normalized name exists in CRM but DOB differs: CRM={row.get('date_of_birth')}, source={source_dob}.",
            )
        if len(same_name) > 1:
            return "AMBIGUOUS", None, f"{len(same_name)} CRM rows share this normalized name; no DOB match."

        return "UNMATCHED", None, "No CRM participant matched normalized name + DOB."

    same_name = by_name.get(name_key, [])
    if len(same_name) == 1:
        return "REVIEW_NAME_ONLY", None, "DOB missing in source; one same-name CRM participant exists but is not auto-matched."
    if len(same_name) > 1:
        return "AMBIGUOUS", None, f"DOB missing; {len(same_name)} CRM rows share normalized name."
    return "UNMATCHED", None, "DOB missing and no same-name CRM participant found."


def load_excel(path: Path, programme: ProgrammeConfig, by_name_dob, by_name):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active

    first_rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        first_rows.append(list(row))
        if i >= 3:
            break

    title1 = clean_text(first_rows[0][0]) if len(first_rows) >= 1 and first_rows[0] else None
    title2 = clean_text(first_rows[1][0]) if len(first_rows) >= 2 and first_rows[1] else None
    year = infer_year(path.name, title1 or "", title2 or "")
    activity_name = activity_name_from_titles(title1, title2, path.name)

    header_row_num, headers = find_header_row(ws)
    header_index = {}
    for i, h in enumerate(headers):
        if h is None:
            continue
        key = str(h).strip().casefold()
        header_index.setdefault(key, i)

    participant_name_col = header_index.get("girls participant list")
    if participant_name_col is None:
        participant_name_col = header_index.get("participant list")

    dob_col = header_index.get("date of birth")
    status_col = header_index.get("status")

    if participant_name_col is None or dob_col is None:
        raise RuntimeError("Participant List and/or Date of Birth column not found.")

    date_columns = []
    for index, h in enumerate(headers):
        session_date, cancelled = parse_session_header(h, year)
        if session_date:
            date_columns.append((index, clean_text(h) or "", session_date, cancelled))

    participants: list[SourceParticipant] = []
    attendance_by_col: dict[int, Counter] = {idx: Counter() for idx, *_ in date_columns}

    for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_number <= header_row_num:
            continue
        values = list(row)

        source_name = clean_text(values[participant_name_col] if participant_name_col < len(values) else None)
        source_dob = parse_dob(values[dob_col] if dob_col < len(values) else None, wb.epoch)

        # Skip summary / blank rows. Source participant name is the authoritative row identity.
        if not source_name:
            continue

        source_status = (
            clean_text(values[status_col])
            if status_col is not None and status_col < len(values)
            else None
        )

        match_status, match, detail = resolve_participant(
            source_name, source_dob, by_name_dob, by_name
        )

        participants.append(
            SourceParticipant(
                file=path.name,
                sheet=ws.title,
                source_programme=programme.source_programme,
                activity_type=programme.activity_type,
                activity_name=activity_name,
                row_number=row_number,
                source_name=source_name,
                source_dob=source_dob,
                source_status=source_status,
                match_status=match_status,
                matched_participant_id=str(match["id"]) if match else None,
                matched_ocf_id=match.get("ocf_id") if match else None,
                matched_name=match.get("full_name") if match else None,
                match_detail=detail,
            )
        )

        for col_index, _, _, _ in date_columns:
            value = values[col_index] if col_index < len(values) else None
            attendance_by_col[col_index][normalize_attendance(value)] += 1

    sessions: list[SessionCandidate] = []
    for col_index, header, session_date, cancelled in date_columns:
        counts = attendance_by_col[col_index]
        yes_count = counts["yes"]
        no_count = counts["no"]
        blank_count = counts["blank"]
        other_count = counts["other"]
        marked = yes_count + no_count

        if cancelled:
            status = "CANCELLED_SOURCE_SESSION"
        elif marked > 0:
            status = "DELIVERED_SOURCE_SESSION"
        else:
            status = "REVIEW_UNCONFIRMED_EMPTY_SESSION"

        if other_count:
            status += f"|REVIEW_{other_count}_UNKNOWN_ATTENDANCE_VALUES"

        sessions.append(
            SessionCandidate(
                file=path.name,
                sheet=ws.title,
                source_programme=programme.source_programme,
                activity_type=programme.activity_type,
                activity_name=activity_name,
                header=header,
                session_date=session_date,
                cancelled=cancelled,
                yes_count=yes_count,
                no_count=no_count,
                blank_count=blank_count,
                marked_count=marked,
                source_participant_count=len(participants),
                status=status,
            )
        )

    warning_parts = []
    if programme.source_programme == "Girls Cricket":
        if title2 and "youth club" in title2.casefold():
            warning_parts.append(
                "Girls Cricket file contains copied Girls Youth Club subtitle; classified as Girls Cricket from board/file title."
            )

    matched = sum(1 for p in participants if p.match_status == "MATCHED_NAME_DOB")
    ambiguous = sum(1 for p in participants if p.match_status == "AMBIGUOUS")
    unmatched = len(participants) - matched - ambiguous

    audit = FileAudit(
        file=path.name,
        sheet=ws.title,
        source_programme=programme.source_programme,
        activity_type=programme.activity_type,
        activity_name=activity_name,
        inferred_year=year,
        participant_rows=len(participants),
        matched_rows=matched,
        unmatched_rows=unmatched,
        ambiguous_rows=ambiguous,
        date_columns=len(sessions),
        delivered_session_candidates=sum("DELIVERED_SOURCE_SESSION" in s.status for s in sessions),
        cancelled_session_candidates=sum(s.cancelled for s in sessions),
        empty_unconfirmed_date_columns=sum("REVIEW_UNCONFIRMED_EMPTY_SESSION" in s.status for s in sessions),
        yes_attendance_marks=sum(s.yes_count for s in sessions),
        no_attendance_marks=sum(s.no_count for s in sessions),
        warning=" | ".join(warning_parts),
    )

    return audit, participants, sessions


def detect_duplicate_sessions(sessions: list[SessionCandidate]):
    groups: dict[tuple[str, str], list[SessionCandidate]] = defaultdict(list)

    for session in sessions:
        if "DELIVERED_SOURCE_SESSION" not in session.status:
            continue
        key = (session.source_programme, session.session_date)
        groups[key].append(session)

    counter = 1
    for (programme, session_date), rows in groups.items():
        # Multiple activities/runs in same programme on same date are possible.
        # Flag rather than silently dedupe.
        unique_activities = {r.activity_name for r in rows}
        if len(rows) > 1 and len(unique_activities) > 1:
            group_id = f"DUP-{counter:03d}"
            counter += 1
            for row in rows:
                row.possible_duplicate_group = group_id


def write_csv(path: Path, rows: list[dict[str, Any]]):
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fieldnames = list(rows[0].keys())
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def participant_dict(p: SourceParticipant) -> dict[str, Any]:
    return {
        "file": p.file,
        "sheet": p.sheet,
        "source_programme": p.source_programme,
        "activity_type": p.activity_type,
        "activity_name": p.activity_name,
        "source_row": p.row_number,
        "source_name": p.source_name,
        "source_dob": p.source_dob or "",
        "source_status": p.source_status or "",
        "match_status": p.match_status,
        "matched_participant_id": p.matched_participant_id or "",
        "matched_ocf_id": p.matched_ocf_id or "",
        "matched_name": p.matched_name or "",
        "match_detail": p.match_detail,
    }


def session_dict(s: SessionCandidate) -> dict[str, Any]:
    return {
        "file": s.file,
        "sheet": s.sheet,
        "source_programme": s.source_programme,
        "activity_type": s.activity_type,
        "activity_name": s.activity_name,
        "source_header": s.header,
        "session_date": s.session_date,
        "cancelled": "YES" if s.cancelled else "NO",
        "yes_attendance": s.yes_count,
        "no_attendance": s.no_count,
        "blank": s.blank_count,
        "marked": s.marked_count,
        "source_participant_rows": s.source_participant_count,
        "status": s.status,
        "possible_duplicate_group": s.possible_duplicate_group,
    }


def file_audit_dict(a: FileAudit) -> dict[str, Any]:
    return {
        "file": a.file,
        "sheet": a.sheet,
        "source_programme": a.source_programme,
        "activity_type": a.activity_type,
        "activity_name": a.activity_name,
        "inferred_year": a.inferred_year or "",
        "participant_rows": a.participant_rows,
        "matched_rows": a.matched_rows,
        "unmatched_or_review_rows": a.unmatched_rows,
        "ambiguous_rows": a.ambiguous_rows,
        "date_columns": a.date_columns,
        "delivered_session_candidates": a.delivered_session_candidates,
        "cancelled_session_candidates": a.cancelled_session_candidates,
        "empty_unconfirmed_date_columns": a.empty_unconfirmed_date_columns,
        "yes_attendance_marks": a.yes_attendance_marks,
        "no_attendance_marks": a.no_attendance_marks,
        "warning": a.warning,
    }


def write_schema_report(path: Path, related_schema: dict[str, list[dict[str, Any]]]):
    lines = []
    lines.append("OCF CRM RELATED MYSQL SCHEMA - READ ONLY")
    lines.append("=" * 72)
    lines.append("")
    for table, columns in related_schema.items():
        lines.append(f"[{table}]")
        for c in columns:
            lines.append(
                f"  {c['COLUMN_NAME']} | {c['DATA_TYPE']} | nullable={c['IS_NULLABLE']} "
                f"| default={c['COLUMN_DEFAULT']} | key={c['COLUMN_KEY']} | extra={c['EXTRA']}"
            )
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def parse_args():
    p = argparse.ArgumentParser(description="Preview OCF Monday activity/session migration.")
    p.add_argument("--source-dir", required=True, help="Folder containing attendance .xlsx files.")
    p.add_argument("--extra-file", action="append", default=[], help="Additional .xlsx file; may be repeated.")
    p.add_argument("--output-dir", default=".", help="Folder for preview CSVs/reports.")
    p.add_argument("--connection-string", default=DEFAULT_CONNECTION_STRING)
    return p.parse_args()


def main() -> int:
    args = parse_args()

    source_dir = Path(args.source_dir).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    if not source_dir.exists() or not source_dir.is_dir():
        print(f"ERROR: source directory not found: {source_dir}", file=sys.stderr)
        return 2

    xlsx_files = list(source_dir.glob("*.xlsx"))
    for extra in args.extra_file:
        path = Path(extra).expanduser().resolve()
        if path.exists() and path.suffix.casefold() == ".xlsx":
            xlsx_files.append(path)
        else:
            print(f"WARNING: extra file not found/skipped: {path}")

    # Deduplicate paths.
    xlsx_files = sorted({p.resolve() for p in xlsx_files})

    known_files = []
    skipped_files = []
    for path in xlsx_files:
        config = classify_file(path.name)
        if config:
            known_files.append((path, config))
        else:
            skipped_files.append(path)

    if "YOUR_DB_" in args.connection_string:
        print("ERROR: set OCF_MYSQL_CONNECTION_STRING before running.", file=sys.stderr)
        return 3

    try:
        conn = connect_mysql(args.connection_string)
    except Exception as exc:
        print(f"ERROR connecting to MySQL: {exc}", file=sys.stderr)
        return 3

    try:
        # Explicit read-only transaction.
        with conn.cursor() as cur:
            cur.execute("START TRANSACTION READ ONLY")

        by_name_dob, by_name = load_db_participants(conn)
        schema = related_schema_tables(conn)

        audits: list[FileAudit] = []
        participants: list[SourceParticipant] = []
        sessions: list[SessionCandidate] = []
        errors = []

        for path, config in known_files:
            try:
                audit, file_participants, file_sessions = load_excel(
                    path, config, by_name_dob, by_name
                )
                audits.append(audit)
                participants.extend(file_participants)
                sessions.extend(file_sessions)
            except Exception as exc:
                errors.append(
                    {
                        "file": path.name,
                        "error": str(exc),
                    }
                )

        detect_duplicate_sessions(sessions)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        prefix = output_dir / f"ocf_monday_activity_sessions_preview_{timestamp}"

        summary_path = Path(str(prefix) + "_files.csv")
        participants_path = Path(str(prefix) + "_participants.csv")
        sessions_path = Path(str(prefix) + "_sessions.csv")
        schema_path = Path(str(prefix) + "_schema.txt")
        errors_path = Path(str(prefix) + "_errors.csv")
        skipped_path = Path(str(prefix) + "_skipped_files.csv")

        write_csv(summary_path, [file_audit_dict(a) for a in audits])
        write_csv(participants_path, [participant_dict(p) for p in participants])
        write_csv(sessions_path, [session_dict(s) for s in sessions])
        write_schema_report(schema_path, schema)
        write_csv(errors_path, errors)
        write_csv(
            skipped_path,
            [{"file": p.name, "reason": "Filename not recognized as a configured OCF attendance export."}
             for p in skipped_files],
        )

        # Desired activity catalogue for the next commit stage.
        activities_path = Path(str(prefix) + "_desired_activities.csv")
        desired_activities = []
        seen = set()
        for a in audits:
            key = (a.activity_type, a.activity_name)
            if key not in seen:
                seen.add(key)
                desired_activities.append(
                    {
                        "activity_type": a.activity_type,
                        "activity_name": a.activity_name,
                        "source_programme": a.source_programme,
                        "source": "Monday attendance board",
                        "historical_sessions": "YES",
                    }
                )

        # User explicitly wants Boys Cricket available for future use with no historical attendance.
        desired_activities.append(
            {
                "activity_type": "Cricket",
                "activity_name": "Boys Cricket",
                "source_programme": "Boys Cricket",
                "source": "Activity-only requirement; no historical attendance source",
                "historical_sessions": "NO",
            }
        )
        write_csv(activities_path, desired_activities)

        match_counts = Counter(p.match_status for p in participants)
        delivered = [s for s in sessions if "DELIVERED_SOURCE_SESSION" in s.status]
        cancelled = [s for s in sessions if s.cancelled]
        empty = [s for s in sessions if "REVIEW_UNCONFIRMED_EMPTY_SESSION" in s.status]
        dup_groups = sorted({s.possible_duplicate_group for s in sessions if s.possible_duplicate_group})

        # Unique participant-to-activity-type allocations supported by source roster.
        allocation_keys = {
            (p.matched_participant_id, p.activity_type)
            for p in participants
            if p.match_status == "MATCHED_NAME_DOB" and p.matched_participant_id
        }

        # Unique participant-to-board/activity enrolments, for information.
        enrolment_keys = {
            (p.matched_participant_id, p.activity_name)
            for p in participants
            if p.match_status == "MATCHED_NAME_DOB" and p.matched_participant_id
        }

        print()
        print("OCF Monday Activity / Session Migration PREVIEW")
        print("------------------------------------------------")
        print(f"Known attendance Excel files:       {len(known_files)}")
        print(f"Skipped/unrecognised Excel files:   {len(skipped_files)}")
        print(f"Workbook parse errors:              {len(errors)}")
        print()
        print("Participant roster rows")
        print(f"  Total source rows:                {len(participants)}")
        print(f"  Exact Name + DOB matched:         {match_counts['MATCHED_NAME_DOB']}")
        print(f"  Unmatched:                        {match_counts['UNMATCHED']}")
        print(f"  Review name-only:                 {match_counts['REVIEW_NAME_ONLY']}")
        print(f"  Review DOB mismatch:              {match_counts['REVIEW_DOB_MISMATCH']}")
        print(f"  Ambiguous:                        {match_counts['AMBIGUOUS']}")
        print()
        print("Allocation / enrolment candidates")
        print(f"  Unique participant Activity Types:{len(allocation_keys):>8}")
        print(f"  Unique participant Activities:    {len(enrolment_keys):>8}")
        print()
        print("Session source columns")
        print(f"  Total dated columns:              {len(sessions)}")
        print(f"  Delivered candidates:             {len(delivered)}")
        print(f"  Explicit cancelled candidates:    {len(cancelled)}")
        print(f"  Empty/unconfirmed date columns:   {len(empty)}")
        print(f"  Possible cross-board duplicates:  {len(dup_groups)} groups")
        print()
        print("Attendance source marks")
        print(f"  Yes / present marks:              {sum(s.yes_count for s in sessions)}")
        print(f"  No / absent marks:                {sum(s.no_count for s in sessions)}")
        print()
        print("Desired Activity Types")
        for value in sorted({row['activity_type'] for row in desired_activities}):
            print(f"  - {value}")
        print()
        print("Boys Cricket")
        print("  Desired Activity only:            YES")
        print("  Historical sessions:              0")
        print("  Historical attendance:            0")
        print()
        print("Files written:")
        print(f"  {summary_path}")
        print(f"  {participants_path}")
        print(f"  {sessions_path}")
        print(f"  {activities_path}")
        print(f"  {schema_path}")
        if errors:
            print(f"  {errors_path}")
        if skipped_files:
            print(f"  {skipped_path}")
        print()
        print("PREVIEW COMPLETE: MySQL database was NOT changed.")
        print("Do NOT create a commit script until the schema report and REVIEW rows are checked.")

        conn.rollback()
        return 0

    except Exception as exc:
        try:
            conn.rollback()
        except Exception:
            pass
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    finally:
        conn.close()


if __name__ == "__main__":
    raise SystemExit(main())
