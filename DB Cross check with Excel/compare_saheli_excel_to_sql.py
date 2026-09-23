from pathlib import Path
from datetime import datetime, date
from getpass import getpass
import re

import pandas as pd
import pyodbc
import openpyxl


# ============================================================
# SAHELI CRM - LOCAL PARTICIPANT DATA AUDIT
#
# READ ONLY:
#   - Reads the Excel file
#   - SELECTs dbo.Participants
#   - Does NOT INSERT / UPDATE / DELETE anything
#
# Output stays on this PC.
# ============================================================


# ------------------------------------------------------------
# NORMALISATION
# ------------------------------------------------------------

def blank(value):
    if value is None:
        return True

    try:
        return pd.isna(value)
    except Exception:
        return False


def text(value):
    if blank(value):
        return ""

    value = str(value).strip()

    # Collapse repeated spaces
    value = re.sub(r"\s+", " ", value)

    return value


def norm_text(value):
    return text(value).upper()


def norm_name(value):
    return norm_text(value)


def norm_card(value):
    if blank(value):
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    s = text(value)

    # Handle Excel values such as 1060.0
    if re.fullmatch(r"\d+\.0", s):
        return s[:-2]

    return s


def norm_date(value):
    if blank(value):
        return ""

    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    try:
        parsed = pd.to_datetime(value, dayfirst=True, errors="coerce")

        if pd.isna(parsed):
            return text(value)

        return parsed.strftime("%Y-%m-%d")

    except Exception:
        return text(value)


def norm_mobile(value):
    if blank(value):
        return ""

    s = str(value).strip()

    # Remove spaces, brackets, dashes etc.
    digits = re.sub(r"\D", "", s)

    if not digits:
        return ""

    # 0044xxxxxxxxxx -> 0xxxxxxxxxx
    if digits.startswith("0044"):
        digits = "0" + digits[4:]

    # 44xxxxxxxxxx -> 0xxxxxxxxxx
    elif digits.startswith("44") and len(digits) >= 11:
        digits = "0" + digits[2:]

    return digits


def norm_postcode(value):
    if blank(value):
        return ""

    return re.sub(r"\s+", "", str(value)).upper()


def norm_email(value):
    if blank(value):
        return ""

    return str(value).strip().lower()


def norm_yes_no(value):
    if blank(value):
        return ""

    if isinstance(value, bool):
        return "YES" if value else "NO"

    if isinstance(value, (int, float)):
        if value == 1:
            return "YES"
        if value == 0:
            return "NO"

    s = norm_text(value).replace(";", "")

    if s in {"YES", "Y", "TRUE", "1"}:
        return "YES"

    if s in {"NO", "N", "FALSE", "0"}:
        return "NO"

    return s


def norm_multi(value):
    """
    Useful for fields such as referral reason where order of
    semi-colon separated options may differ.
    """

    if blank(value):
        return ""

    s = str(value)

    values = [
        re.sub(r"\s+", " ", x.strip()).upper()
        for x in s.split(";")
        if x.strip()
    ]

    return ";".join(sorted(values))


def original_value(value):
    if blank(value):
        return ""

    if isinstance(value, (datetime, date, pd.Timestamp)):
        return norm_date(value)

    return str(value).strip()


# ------------------------------------------------------------
# EXCEL LABEL HANDLING
# ------------------------------------------------------------

def clean_excel_label(value):
    if blank(value):
        return ""

    s = str(value).strip().lower()

    s = s.replace(":", "")
    s = re.sub(r"\s+", " ", s)

    return s


EXCEL_FIELD_MAP = {
    "registration date": "RegistrationDate",
    "saheli card number": "SaheliCardNumber",
    "full name": "FullName",
    "date of birth": "DateOfBirth",
    "age": "Age",
    "address": "Address",
    "postcode": "Postcode",
    "email": "Email",
    "mobile/home no": "MobileNumber",
    "gender": "Gender",
    "is your gender the same as assigned at birth?": "GenderSameAsBirth",
    "ethnicity": "Ethnicity",
    "preferred spoken language": "PreferredLanguage",
    "religion": "Religion",
    "sexuality": "Sexuality",
    "occupation": "Occupation",
    "living alone": "LivingAlone",
    "caring responsibilities": "CaringResponsibilities",
    "referral reason": "ReferralReason",
    "how heard about saheli hub?": "HeardAboutSaheli",
    "gp surgery name": "GPSurgeryName",
    "staff member": "StaffMember",
    "site": "Site",
    "notes": "Notes",
}


# ------------------------------------------------------------
# FIELD COMPARISON RULES
# ------------------------------------------------------------

FIELD_NORMALISERS = {
    "SaheliCardNumber": norm_card,
    "FullName": norm_name,
    "DateOfBirth": norm_date,
    "RegistrationDate": norm_date,
    "Address": norm_text,
    "Postcode": norm_postcode,
    "Email": norm_email,
    "MobileNumber": norm_mobile,
    "Gender": norm_text,
    "GenderSameAsBirth": norm_yes_no,
    "Ethnicity": norm_text,
    "PreferredLanguage": norm_text,
    "Religion": norm_text,
    "Sexuality": norm_text,
    "Occupation": norm_text,
    "LivingAlone": norm_yes_no,
    "CaringResponsibilities": norm_yes_no,
    "ReferralReason": norm_multi,
    "HeardAboutSaheli": norm_text,
    "GPSurgeryName": norm_text,
    "StaffMember": norm_text,
    "Site": norm_text,
    "Notes": norm_text,
}


CORE_IDENTITY_FIELDS = {
    "FullName",
    "DateOfBirth",
    "MobileNumber",
    "Postcode",
}


# ------------------------------------------------------------
# READ THE VERTICAL SAHELI EXCEL FORMAT
# ------------------------------------------------------------

def read_saheli_excel(excel_path):
    wb = openpyxl.load_workbook(
        excel_path,
        data_only=True,
        read_only=True
    )

    records = []

    for ws in wb.worksheets:

        labels = {}

        for row in range(1, ws.max_row + 1):
            raw_label = ws.cell(row=row, column=1).value
            clean_label = clean_excel_label(raw_label)

            if clean_label:
                labels[clean_label] = row

        card_row = labels.get("saheli card number")

        if not card_row:
            print(
                f"Skipping sheet '{ws.title}' "
                "- Saheli Card Number row not found."
            )
            continue

        # Each participant is stored in a separate column
        # starting at column B.
        for column in range(2, ws.max_column + 1):

            card = ws.cell(
                row=card_row,
                column=column
            ).value

            if blank(card):
                continue

            record = {
                "SourceSheet": ws.title,
                "SourceColumn": column,
            }

            for excel_label, sql_field in EXCEL_FIELD_MAP.items():

                row = labels.get(excel_label)

                if row:
                    record[sql_field] = ws.cell(
                        row=row,
                        column=column
                    ).value
                else:
                    record[sql_field] = None

            record["SaheliCardNumber"] = norm_card(
                record["SaheliCardNumber"]
            )

            records.append(record)

    wb.close()

    df = pd.DataFrame(records)

    if df.empty:
        raise RuntimeError(
            "No participant records were found in the Excel workbook."
        )

    return df


# ------------------------------------------------------------
# READ SQL - SELECT ONLY
# ------------------------------------------------------------

SQL_QUERY = """
SELECT
    ParticipantID,
    SaheliCardNumber,
    FullName,
    DateOfBirth,
    Age,
    Address,
    Postcode,
    Email,
    MobileNumber,
    Gender,
    GenderSameAsBirth,
    Ethnicity,
    PreferredLanguage,
    Religion,
    Sexuality,
    Occupation,
    LivingAlone,
    CaringResponsibilities,
    ReferralReason,
    HeardAboutSaheli,
    GPSurgeryName,
    CreatedAt,
    HasHealthConditionOrDisability,
    HealthConditionDetails,
    StaffMember,
    Site,
    Notes,
    RegistrationDate,
    GPSurgeryId
FROM dbo.Participants;
"""


def read_sql_participants(server, database, username, password):

    connection_string = (
        "DRIVER={ODBC Driver 18 for SQL Server};"
        f"SERVER=tcp:{server},1433;"
        f"DATABASE={database};"
        f"UID={username};"
        f"PWD={password};"
        "Encrypt=yes;"
        "TrustServerCertificate=no;"
        "Connection Timeout=30;"
        "ApplicationIntent=ReadOnly;"
    )

    connection = pyodbc.connect(connection_string)

    try:
        cursor = connection.cursor()

        # Safety: SELECT only
        cursor.execute(SQL_QUERY)

        columns = [
            column[0]
            for column in cursor.description
        ]

        rows = cursor.fetchall()

        df = pd.DataFrame.from_records(
            rows,
            columns=columns
        )

    finally:
        connection.close()

    df["SaheliCardNumber"] = (
        df["SaheliCardNumber"]
        .apply(norm_card)
    )

    return df


# ------------------------------------------------------------
# DUPLICATE CHECK
# ------------------------------------------------------------

def find_sql_duplicates(sql_df):

    df = sql_df.copy()

    df["_Name"] = df["FullName"].apply(norm_name)
    df["_DOB"] = df["DateOfBirth"].apply(norm_date)
    df["_Mobile"] = df["MobileNumber"].apply(norm_mobile)

    duplicate_mask = df.duplicated(
        subset=["_Name", "_DOB"],
        keep=False
    )

    duplicate_df = df[
        duplicate_mask
        & (df["_Name"] != "")
        & (df["_DOB"] != "")
    ].copy()

    if duplicate_df.empty:
        return pd.DataFrame()

    duplicate_df["SameMobile"] = duplicate_df.groupby(
        ["_Name", "_DOB"]
    )["_Mobile"].transform(
        lambda x: len({
            v for v in x
            if v
        }) == 1
    )

    return duplicate_df[
        [
            "ParticipantID",
            "SaheliCardNumber",
            "FullName",
            "DateOfBirth",
            "MobileNumber",
            "Postcode",
            "RegistrationDate",
            "CreatedAt",
            "SameMobile",
        ]
    ].sort_values(
        ["FullName", "DateOfBirth", "SaheliCardNumber"]
    )


# ------------------------------------------------------------
# COMPARISON
# ------------------------------------------------------------

def compare(excel_df, sql_df):

    excel_by_card = {
        norm_card(row["SaheliCardNumber"]): row
        for _, row in excel_df.iterrows()
    }

    sql_by_card = {
        norm_card(row["SaheliCardNumber"]): row
        for _, row in sql_df.iterrows()
    }

    excel_cards = set(excel_by_card)
    sql_cards = set(sql_by_card)

    both_cards = sorted(
        excel_cards & sql_cards,
        key=lambda x: (len(x), x)
    )

    excel_only_cards = sorted(
        excel_cards - sql_cards,
        key=lambda x: (len(x), x)
    )

    sql_only_cards = sorted(
        sql_cards - excel_cards,
        key=lambda x: (len(x), x)
    )

    profile_rows = []
    difference_rows = []
    wrong_identity_rows = []

    fields_to_compare = [
        field
        for field in EXCEL_FIELD_MAP.values()
        if field != "Age"
        and field != "SaheliCardNumber"
    ]

    for card in both_cards:

        excel_row = excel_by_card[card]
        sql_row = sql_by_card[card]

        mismatches = []
        core_mismatches = []

        for field in fields_to_compare:

            normaliser = FIELD_NORMALISERS[field]

            excel_value = excel_row.get(field)
            sql_value = sql_row.get(field)

            excel_normalised = normaliser(excel_value)
            sql_normalised = normaliser(sql_value)

            if excel_normalised != sql_normalised:

                mismatches.append(field)

                severity = (
                    "CORE IDENTITY"
                    if field in CORE_IDENTITY_FIELDS
                    else "PROFILE"
                )

                if field in CORE_IDENTITY_FIELDS:
                    core_mismatches.append(field)

                difference_rows.append(
                    {
                        "SaheliCardNumber": card,
                        "ParticipantID": sql_row.get(
                            "ParticipantID"
                        ),
                        "CurrentSQLFullName": sql_row.get(
                            "FullName"
                        ),
                        "Field": field,
                        "Severity": severity,
                        "ExcelValue": original_value(
                            excel_value
                        ),
                        "SQLValue": original_value(
                            sql_value
                        ),
                    }
                )

        profile_rows.append(
            {
                "SaheliCardNumber": card,
                "ParticipantID": sql_row.get(
                    "ParticipantID"
                ),
                "ExcelFullName": excel_row.get(
                    "FullName"
                ),
                "SQLFullName": sql_row.get(
                    "FullName"
                ),
                "ExcelDOB": norm_date(
                    excel_row.get("DateOfBirth")
                ),
                "SQLDOB": norm_date(
                    sql_row.get("DateOfBirth")
                ),
                "ExcelMobile": original_value(
                    excel_row.get("MobileNumber")
                ),
                "SQLMobile": original_value(
                    sql_row.get("MobileNumber")
                ),
                "ExcelPostcode": original_value(
                    excel_row.get("Postcode")
                ),
                "SQLPostcode": original_value(
                    sql_row.get("Postcode")
                ),
                "CoreIdentityStatus":
                    "MATCH"
                    if not core_mismatches
                    else "REVIEW",
                "ProfileStatus":
                    "MATCH"
                    if not mismatches
                    else "REVIEW",
                "MismatchCount": len(mismatches),
                "MismatchFields": ", ".join(
                    mismatches
                ),
            }
        )

        # ----------------------------------------------------
        # Detect:
        #
        # SQL card X contains the identity belonging to
        # another Excel card Y.
        # ----------------------------------------------------

        sql_name = norm_name(
            sql_row.get("FullName")
        )
        sql_dob = norm_date(
            sql_row.get("DateOfBirth")
        )
        sql_mobile = norm_mobile(
            sql_row.get("MobileNumber")
        )

        own_excel_name = norm_name(
            excel_row.get("FullName")
        )
        own_excel_dob = norm_date(
            excel_row.get("DateOfBirth")
        )
        own_excel_mobile = norm_mobile(
            excel_row.get("MobileNumber")
        )

        own_strong_match = (
            sql_name != ""
            and sql_dob != ""
            and sql_mobile != ""
            and sql_name == own_excel_name
            and sql_dob == own_excel_dob
            and sql_mobile == own_excel_mobile
        )

        # Only search for somebody else when this card does
        # not correctly match its own Excel identity.
        if not own_strong_match:

            for other_card, other_excel in excel_by_card.items():

                if other_card == card:
                    continue

                other_name = norm_name(
                    other_excel.get("FullName")
                )
                other_dob = norm_date(
                    other_excel.get("DateOfBirth")
                )
                other_mobile = norm_mobile(
                    other_excel.get("MobileNumber")
                )

                confidence = None

                if (
                    sql_name
                    and sql_dob
                    and sql_mobile
                    and sql_name == other_name
                    and sql_dob == other_dob
                    and sql_mobile == other_mobile
                ):
                    confidence = "STRONG"

                elif (
                    sql_name
                    and sql_dob
                    and sql_name == other_name
                    and sql_dob == other_dob
                ):
                    confidence = "MEDIUM"

                if confidence:

                    wrong_identity_rows.append(
                        {
                            "SQLSaheliCardNumber": card,
                            "ParticipantID": sql_row.get(
                                "ParticipantID"
                            ),
                            "SQLCurrentName": sql_row.get(
                                "FullName"
                            ),

                            "ExcelExpectedNameForThisCard":
                                excel_row.get("FullName"),

                            "SQLCurrentDOB":
                                norm_date(
                                    sql_row.get(
                                        "DateOfBirth"
                                    )
                                ),

                            "SQLCurrentMobile":
                                original_value(
                                    sql_row.get(
                                        "MobileNumber"
                                    )
                                ),

                            "AppearsToBelongToExcelCard":
                                other_card,

                            "MatchedExcelName":
                                other_excel.get(
                                    "FullName"
                                ),

                            "MatchedExcelDOB":
                                norm_date(
                                    other_excel.get(
                                        "DateOfBirth"
                                    )
                                ),

                            "MatchedExcelMobile":
                                original_value(
                                    other_excel.get(
                                        "MobileNumber"
                                    )
                                ),

                            "Confidence":
                                confidence,

                            "Explanation":
                                (
                                    f"SQL card {card} "
                                    f"currently contains identity "
                                    f"details matching Excel card "
                                    f"{other_card}."
                                ),
                        }
                    )

    # --------------------------------------------------------
    # Excel only
    # --------------------------------------------------------

    excel_only_rows = []

    for card in excel_only_cards:

        row = excel_by_card[card]

        excel_only_rows.append(
            {
                "SaheliCardNumber": card,
                "FullName": row.get("FullName"),
                "DateOfBirth": norm_date(
                    row.get("DateOfBirth")
                ),
                "MobileNumber": row.get(
                    "MobileNumber"
                ),
                "Postcode": row.get("Postcode"),
                "RegistrationDate": norm_date(
                    row.get("RegistrationDate")
                ),
                "SourceSheet": row.get(
                    "SourceSheet"
                ),
            }
        )

    # --------------------------------------------------------
    # SQL only
    # --------------------------------------------------------

    sql_only_rows = []

    for card in sql_only_cards:

        row = sql_by_card[card]

        sql_only_rows.append(
            {
                "ParticipantID": row.get(
                    "ParticipantID"
                ),
                "SaheliCardNumber": card,
                "FullName": row.get("FullName"),
                "DateOfBirth": norm_date(
                    row.get("DateOfBirth")
                ),
                "MobileNumber": row.get(
                    "MobileNumber"
                ),
                "Postcode": row.get("Postcode"),
                "RegistrationDate": norm_date(
                    row.get("RegistrationDate")
                ),
                "CreatedAt": original_value(
                    row.get("CreatedAt")
                ),
            }
        )

    return (
        pd.DataFrame(profile_rows),
        pd.DataFrame(difference_rows),
        pd.DataFrame(wrong_identity_rows),
        pd.DataFrame(excel_only_rows),
        pd.DataFrame(sql_only_rows),
    )


# ------------------------------------------------------------
# OUTPUT FORMATTING
# ------------------------------------------------------------

def format_workbook(path):

    wb = openpyxl.load_workbook(path)

    for ws in wb.worksheets:

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Bold first row
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(
                bold=True
            )

        # Sensible column widths
        for column_cells in ws.columns:

            max_length = 0

            column_letter = (
                column_cells[0].column_letter
            )

            for cell in column_cells[:1000]:

                value = (
                    ""
                    if cell.value is None
                    else str(cell.value)
                )

                max_length = max(
                    max_length,
                    len(value)
                )

            ws.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 10),
                55
            )

    wb.save(path)


# ------------------------------------------------------------
# MAIN
# ------------------------------------------------------------

def main():

    print()
    print("=" * 65)
    print("SAHELI CRM - LOCAL EXCEL vs SQL PARTICIPANT AUDIT")
    print("=" * 65)
    print()
    print("This script is READ ONLY.")
    print("It does not update or delete CRM data.")
    print()

    excel_path = Path(
        input(
            "Full path to master Excel file: "
        ).strip().strip('"')
    )

    if not excel_path.exists():
        raise FileNotFoundError(
            f"Excel file not found: {excel_path}"
        )

    print()
    print("Azure SQL connection")
    print("--------------------")

    server = input(
        "Server, e.g. xxx.database.windows.net: "
    ).strip()

    database = input(
        "Database [SaheliHubCRM]: "
    ).strip()

    if not database:
        database = "SaheliHubCRM"

    username = input(
        "SQL username: "
    ).strip()

    password = getpass(
        "SQL password: "
    )

    print()
    print("Reading master Excel locally...")

    excel_df = read_saheli_excel(
        excel_path
    )

    print(
        f"Excel participants found: "
        f"{len(excel_df)}"
    )

    print()
    print(
        "Reading dbo.Participants "
        "using SELECT only..."
    )

    sql_df = read_sql_participants(
        server,
        database,
        username,
        password
    )

    print(
        f"SQL participants found: "
        f"{len(sql_df)}"
    )

    print()
    print("Comparing locally...")

    (
        comparison_df,
        differences_df,
        wrong_identity_df,
        excel_only_df,
        sql_only_df,
    ) = compare(
        excel_df,
        sql_df
    )

    duplicate_sql_df = find_sql_duplicates(
        sql_df
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    output_dir = (
        Path.cwd()
        / "SaheliCRM_Audit"
    )

    output_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    report_path = output_dir / (
        f"Saheli_Participant_Audit_"
        f"{timestamp}.xlsx"
    )

    sql_snapshot_path = output_dir / (
        f"SQL_Participants_Snapshot_"
        f"{timestamp}.csv"
    )

    excel_snapshot_path = output_dir / (
        f"Excel_Participants_Snapshot_"
        f"{timestamp}.csv"
    )

    # Save local snapshots
    sql_df.to_csv(
        sql_snapshot_path,
        index=False,
        encoding="utf-8-sig"
    )

    excel_df.to_csv(
        excel_snapshot_path,
        index=False,
        encoding="utf-8-sig"
    )

    both_count = len(
        set(excel_df["SaheliCardNumber"])
        & set(sql_df["SaheliCardNumber"])
    )

    core_review_count = 0

    if not comparison_df.empty:
        core_review_count = (
            comparison_df[
                "CoreIdentityStatus"
            ]
            .eq("REVIEW")
            .sum()
        )

    strong_wrong = 0
    medium_wrong = 0

    if not wrong_identity_df.empty:

        strong_wrong = (
            wrong_identity_df[
                "Confidence"
            ]
            .eq("STRONG")
            .sum()
        )

        medium_wrong = (
            wrong_identity_df[
                "Confidence"
            ]
            .eq("MEDIUM")
            .sum()
        )

    summary_df = pd.DataFrame(
        [
            {
                "Metric":
                    "Participants in master Excel",
                "Count": len(excel_df),
            },
            {
                "Metric":
                    "Participants in SQL",
                "Count": len(sql_df),
            },
            {
                "Metric":
                    "Saheli cards found in both",
                "Count": both_count,
            },
            {
                "Metric":
                    "Excel cards missing from SQL",
                "Count": len(excel_only_df),
            },
            {
                "Metric":
                    "SQL cards missing from Excel",
                "Count": len(sql_only_df),
            },
            {
                "Metric":
                    "Cards with core identity differences",
                "Count": int(
                    core_review_count
                ),
            },
            {
                "Metric":
                    "Possible wrong identity - STRONG",
                "Count": int(strong_wrong),
            },
            {
                "Metric":
                    "Possible wrong identity - MEDIUM",
                "Count": int(medium_wrong),
            },
            {
                "Metric":
                    "SQL duplicate identity rows",
                "Count": len(
                    duplicate_sql_df
                ),
            },
        ]
    )

    print()
    print("Creating local audit workbook...")

    with pd.ExcelWriter(
        report_path,
        engine="openpyxl"
    ) as writer:

        summary_df.to_excel(
            writer,
            sheet_name="Summary",
            index=False
        )

        comparison_df.to_excel(
            writer,
            sheet_name="ProfileComparison",
            index=False
        )

        differences_df.to_excel(
            writer,
            sheet_name="FieldDifferences",
            index=False
        )

        wrong_identity_df.to_excel(
            writer,
            sheet_name="PossibleWrongIdentity",
            index=False
        )

        duplicate_sql_df.to_excel(
            writer,
            sheet_name="SQLDuplicateIdentity",
            index=False
        )

        excel_only_df.to_excel(
            writer,
            sheet_name="ExcelOnly",
            index=False
        )

        sql_only_df.to_excel(
            writer,
            sheet_name="SQLOnly",
            index=False
        )

    format_workbook(report_path)

    print()
    print("=" * 65)
    print("AUDIT COMPLETE")
    print("=" * 65)

    print(
        f"\nReport:\n{report_path}"
    )

    print(
        f"\nLocal SQL snapshot:\n"
        f"{sql_snapshot_path}"
    )

    print(
        f"\nLocal Excel snapshot:\n"
        f"{excel_snapshot_path}"
    )

    print()
    print(
        "No SQL records were changed."
    )


if __name__ == "__main__":
    main()