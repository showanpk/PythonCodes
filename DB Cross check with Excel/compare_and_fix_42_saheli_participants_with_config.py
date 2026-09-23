import argparse
import re
import sys
from copy import copy
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pyodbc


# ============================================================
# SAHELI CRM - FIX ONLY THE 42 IDENTIFIED PARTICIPANT RECORDS
#
# What this script does:
#   1. Reads the master Excel locally.
#   2. Reads ONLY the 42 identified dbo.Participants rows.
#   3. Matches Excel -> SQL by SaheliCardNumber.
#   4. Creates a local BEFORE backup + comparison workbook.
#   5. Preview mode is the default: NO SQL changes.
#   6. With --commit, updates ONLY mapped dbo.Participants fields
#      for those exact 42 cards, inside one SQL transaction.
#   7. Keeps ParticipantID and SaheliCardNumber unchanged.
#   8. Does NOT touch assessments, attendance, funding, forms,
#      emergency-contact child tables, or any unrelated participant.
#
# IMPORTANT:
#   Excel is treated as the master for the mapped participant fields.
#   For these 42 records, blank Excel cells become SQL NULL so that
#   wrong/stale details from the copied profile are not left behind.
# ============================================================


DEFAULT_EXCEL = r"C:\Users\shonk\Downloads\Full Registration for SAHELI (1).xlsx"

# ============================================================
# LOCAL SQL CONFIG
#
# Fill these four values once. The script will NOT ask for the
# SQL password in PowerShell/Terminal.
#
# SECURITY: keep this file only on your PC. Do NOT commit it to
# GitHub while SQL_PASSWORD contains the real password.
# ============================================================
SQL_SERVER = r"tcp:sahelihub.database.windows.net,1433"
SQL_DATABASE = r"SaheliHubCRM"
SQL_USERNAME = r"sahelihubadmin"
SQL_PASSWORD = r"W7WZ7ZaG1YbMZ71gh%2xSFuR;"

TARGET_CARDS = [
    "1209", "1212",
    "405", "716",
    "454", "861",
    "433", "727",
    "340", "585",
    "342", "467",
    "1060", "1062",
    "199", "1199",
    "345", "568",
    "1613", "1635",
    "98", "1626",
    "2039", "947",
    "404", "1601",
    "1380", "1360",
    "28", "86",
    "22664079", "489",
    "349", "341",
    "407", "423",
    "9", "1172",
    "643", "463",
    "2002", "951",
]

if len(TARGET_CARDS) != 42 or len(set(TARGET_CARDS)) != 42:
    raise RuntimeError("TARGET_CARDS must contain exactly 42 unique card numbers.")


# Excel header -> dbo.Participants column.
# Only fields that are safe to map directly are included.
HEADER_TO_SQL = {
    "registrationdate": "RegistrationDate",
    "sahelicardnumber": "SaheliCardNumber",
    "fullname": "FullName",
    "dateofbirth": "DateOfBirth",
    "age": "Age",
    "address": "Address",
    "postcode": "Postcode",
    "email": "Email",
    "mobilehomeno": "MobileNumber",
    "mobilenumber": "MobileNumber",
    "gender": "Gender",
    "isyourgendersameasassignedatbirth": "GenderSameAsBirth",
    "ethnicity": "Ethnicity",
    "preferredspokenlanguage": "PreferredLanguage",
    "preferredlanguage": "PreferredLanguage",
    "religion": "Religion",
    "caringresponsibilities": "CaringResponsibilities",
    "livingalone": "LivingAlone",
    "sexuality": "Sexuality",
    "occupation": "Occupation",
    "referralreason": "ReferralReason",
    "howheardaboutsahelihub": "HeardAboutSaheli",
    "gpsurgeryname": "GPSurgeryName",
    "notes": "Notes",
    "staffmember": "StaffMember",
    "site": "Site",
}

# These are intentionally NOT changed in this first pass:
# ParticipantID, SaheliCardNumber, CreatedAt, GPSurgeryId,
# HasHealthConditionOrDisability, HealthConditionDetails,
# emergency contacts, consents, and all assessment columns.
UPDATE_COLUMNS = [
    "FullName",
    "DateOfBirth",
    "Age",
    "Address",
    "Postcode",
    "Email",
    "MobileNumber",
    "Gender",
    "GenderSameAsBirth",
    "Ethnicity",
    "PreferredLanguage",
    "Religion",
    "Sexuality",
    "Occupation",
    "LivingAlone",
    "CaringResponsibilities",
    "ReferralReason",
    "HeardAboutSaheli",
    "GPSurgeryName",
    "StaffMember",
    "Site",
    "Notes",
    "RegistrationDate",
]

BIT_COLUMNS = {
    "GenderSameAsBirth",
    "LivingAlone",
    "CaringResponsibilities",
}

DATE_COLUMNS = {
    "DateOfBirth",
    "RegistrationDate",
}

INT_COLUMNS = {
    "Age",
}


def clean_header(value):
    if value is None:
        return ""
    s = str(value).strip().lower()
    # Removes spaces, punctuation, NBSP and line breaks.
    return re.sub(r"[^a-z0-9]+", "", s)


def normalise_card(value):
    if value is None:
        return ""

    if isinstance(value, bool):
        return str(value)

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()

    s = str(value).strip()
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]

    return s


def empty_to_none(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if value == "":
            return None
    return value


def parse_date(value, field_name):
    value = empty_to_none(value)
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    s = str(value).strip()

    formats = (
        "%d/%m/%Y",
        "%d/%m/%y",
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d-%m-%y",
        "%d.%m.%Y",
    )

    for fmt in formats:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass

    raise ValueError(f"Could not parse {field_name} date value: {value!r}")


def parse_int(value, field_name):
    value = empty_to_none(value)
    if value is None:
        return None

    if isinstance(value, bool):
        raise ValueError(f"Invalid integer for {field_name}: {value!r}")

    if isinstance(value, (int, float)):
        return int(value)

    s = str(value).strip()

    try:
        return int(float(s))
    except ValueError as exc:
        raise ValueError(f"Invalid integer for {field_name}: {value!r}") from exc


def parse_bit(value, field_name):
    value = empty_to_none(value)
    if value is None:
        return None

    if isinstance(value, bool):
        return value

    if isinstance(value, (int, float)):
        if value == 1:
            return True
        if value == 0:
            return False

    s = str(value).strip().lower()

    true_values = {
        "yes", "y", "true", "1",
        "same", "yes - same", "yes same",
    }
    false_values = {
        "no", "n", "false", "0",
        "different", "no - different", "no different",
    }

    if s in true_values:
        return True

    if s in false_values:
        return False

    # Do not guess with unusual text.
    raise ValueError(
        f"Cannot safely convert {field_name}={value!r} to SQL bit. "
        "Fix the Excel value or remove this field from UPDATE_COLUMNS."
    )


def sql_value(column, value):
    """Convert an Excel value into the appropriate SQL parameter value."""
    value = empty_to_none(value)

    if column in DATE_COLUMNS:
        return parse_date(value, column)

    if column in INT_COLUMNS:
        return parse_int(value, column)

    if column in BIT_COLUMNS:
        return parse_bit(value, column)

    if value is None:
        return None

    # Preserve phone numbers/card-like text as entered, but strip surrounding whitespace.
    return str(value).strip()


def display_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "1" if value else "0"
    return str(value)


def comparable(column, value):
    if value is None:
        return ""

    if column in DATE_COLUMNS:
        try:
            parsed = parse_date(value, column)
            return "" if parsed is None else parsed.isoformat()
        except Exception:
            return str(value).strip()

    if column in BIT_COLUMNS:
        try:
            parsed = parse_bit(value, column)
            if parsed is None:
                return ""
            return "1" if parsed else "0"
        except Exception:
            return str(value).strip().lower()

    if column == "MobileNumber":
        digits = re.sub(r"\D", "", str(value))
        if digits.startswith("0044"):
            digits = "0" + digits[4:]
        elif digits.startswith("44") and len(digits) >= 11:
            digits = "0" + digits[2:]
        return digits

    if column == "Postcode":
        return re.sub(r"\s+", "", str(value)).upper()

    return re.sub(r"\s+", " ", str(value).strip()).upper()


def detect_horizontal_table(ws):
    """
    Detects a normal Excel export:
        row N = headers across columns
        rows below = one participant per row
    """
    max_scan_rows = min(ws.max_row, 25)

    for row_num in range(1, max_scan_rows + 1):
        headers = {}
        for col_num in range(1, ws.max_column + 1):
            key = clean_header(ws.cell(row_num, col_num).value)
            if key:
                headers[key] = col_num

        if "sahelicardnumber" in headers and "fullname" in headers:
            return row_num, headers

    return None, None


def read_horizontal_sheet(ws, header_row, headers):
    records = {}

    mapped_columns = {}
    for header_key, col_num in headers.items():
        sql_col = HEADER_TO_SQL.get(header_key)
        if sql_col:
            mapped_columns[sql_col] = col_num

    card_col = mapped_columns.get("SaheliCardNumber")
    if not card_col:
        return records

    for row_num in range(header_row + 1, ws.max_row + 1):
        card = normalise_card(ws.cell(row_num, card_col).value)
        if card not in TARGET_CARDS:
            continue

        if card in records:
            raise RuntimeError(
                f"Excel contains duplicate Saheli Card Number {card}. "
                "The script will not guess which row is correct."
            )

        record = {
            "_sheet": ws.title,
            "_row": row_num,
            "SaheliCardNumber": card,
        }

        for sql_col, col_num in mapped_columns.items():
            record[sql_col] = ws.cell(row_num, col_num).value

        records[card] = record

    return records


def detect_vertical_sheet(ws):
    """
    Supports the older SAHELI format:
        column A = labels
        B/C/... = one participant per column
    """
    labels = {}

    for row_num in range(1, ws.max_row + 1):
        key = clean_header(ws.cell(row_num, 1).value)
        if key:
            labels[key] = row_num

    if "sahelicardnumber" in labels and "fullname" in labels:
        return labels

    return None


def read_vertical_sheet(ws, labels):
    records = {}

    mapped_rows = {}
    for header_key, row_num in labels.items():
        sql_col = HEADER_TO_SQL.get(header_key)
        if sql_col:
            mapped_rows[sql_col] = row_num

    card_row = mapped_rows.get("SaheliCardNumber")
    if not card_row:
        return records

    for col_num in range(2, ws.max_column + 1):
        card = normalise_card(ws.cell(card_row, col_num).value)
        if card not in TARGET_CARDS:
            continue

        if card in records:
            raise RuntimeError(
                f"Excel contains duplicate Saheli Card Number {card}. "
                "The script will not guess which column is correct."
            )

        record = {
            "_sheet": ws.title,
            "_column": col_num,
            "SaheliCardNumber": card,
        }

        for sql_col, row_num in mapped_rows.items():
            record[sql_col] = ws.cell(row_num, col_num).value

        records[card] = record

    return records


def read_master_excel(path):
    print(f"Reading master Excel: {path}")

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    all_records = {}

    try:
        for ws in wb.worksheets:
            header_row, horizontal_headers = detect_horizontal_table(ws)

            if header_row:
                found = read_horizontal_sheet(ws, header_row, horizontal_headers)
            else:
                vertical_labels = detect_vertical_sheet(ws)
                if vertical_labels:
                    found = read_vertical_sheet(ws, vertical_labels)
                else:
                    found = {}

            for card, record in found.items():
                if card in all_records:
                    raise RuntimeError(
                        f"Card {card} appears more than once across workbook sheets. "
                        "The script will not continue."
                    )
                all_records[card] = record
    finally:
        wb.close()

    return all_records


def connect_sql(server, database, username, password):
    connection_string = (
        "DRIVER={ODBC Driver 18 for SQL Server};"
        f"SERVER=tcp:{server},1433;"
        f"DATABASE={database};"
        f"UID={username};"
        f"PWD={password};"
        "Encrypt=yes;"
        "TrustServerCertificate=no;"
        "Connection Timeout=30;"
    )
    return pyodbc.connect(connection_string, autocommit=False)


def fetch_sql_targets(conn):
    placeholders = ",".join("?" for _ in TARGET_CARDS)

    sql = f"""
    SELECT
        ParticipantID,
        SaheliCardNumber,
        FullName,
        DateOfBirth,
        Age,
        Address,
        Postcode,
        Email,
        MobileNumber,
        Gender,
        GenderSameAsBirth,
        Ethnicity,
        PreferredLanguage,
        Religion,
        Sexuality,
        Occupation,
        LivingAlone,
        CaringResponsibilities,
        ReferralReason,
        HeardAboutSaheli,
        GPSurgeryName,
        CreatedAt,
        HasHealthConditionOrDisability,
        HealthConditionDetails,
        StaffMember,
        Site,
        Notes,
        RegistrationDate,
        GPSurgeryId
    FROM dbo.Participants
    WHERE CAST(SaheliCardNumber AS nvarchar(100)) IN ({placeholders});
    """

    cur = conn.cursor()
    cur.execute(sql, TARGET_CARDS)

    columns = [x[0] for x in cur.description]
    rows = cur.fetchall()

    result = {}

    for row in rows:
        item = dict(zip(columns, row))
        card = normalise_card(item["SaheliCardNumber"])

        if card in result:
            raise RuntimeError(
                f"SQL has more than one dbo.Participants row for card {card}. "
                "No changes will be made."
            )

        result[card] = item

    return result


def validate(excel_records, sql_records):
    errors = []

    missing_excel = [card for card in TARGET_CARDS if card not in excel_records]
    missing_sql = [card for card in TARGET_CARDS if card not in sql_records]

    if missing_excel:
        errors.append(
            "Missing from master Excel: " + ", ".join(missing_excel)
        )

    if missing_sql:
        errors.append(
            "Missing from dbo.Participants: " + ", ".join(missing_sql)
        )

    # Require a basic identity in Excel before we allow any update.
    for card in TARGET_CARDS:
        record = excel_records.get(card)
        if not record:
            continue

        full_name = empty_to_none(record.get("FullName"))
        dob = empty_to_none(record.get("DateOfBirth"))

        if not full_name:
            errors.append(f"Excel card {card}: Full Name is blank.")

        if not dob:
            errors.append(f"Excel card {card}: Date of Birth is blank.")
        else:
            try:
                parse_date(dob, "DateOfBirth")
            except Exception as exc:
                errors.append(f"Excel card {card}: {exc}")

        # Validate converted data now, before any SQL update starts.
        for column in UPDATE_COLUMNS:
            try:
                sql_value(column, record.get(column))
            except Exception as exc:
                errors.append(f"Excel card {card}, {column}: {exc}")

    return errors


def build_comparison(excel_records, sql_records):
    summary_rows = []
    diff_rows = []

    for card in TARGET_CARDS:
        x = excel_records[card]
        s = sql_records[card]

        differences = []

        for col in UPDATE_COLUMNS:
            excel_value = sql_value(col, x.get(col))
            sql_current = s.get(col)

            if comparable(col, excel_value) != comparable(col, sql_current):
                differences.append(col)
                diff_rows.append({
                    "SaheliCardNumber": card,
                    "ParticipantID": s["ParticipantID"],
                    "CurrentSQLName": s.get("FullName"),
                    "ExpectedExcelName": x.get("FullName"),
                    "Field": col,
                    "SQLBefore": display_value(sql_current),
                    "ExcelMaster": display_value(excel_value),
                })

        summary_rows.append({
            "SaheliCardNumber": card,
            "ParticipantID": s["ParticipantID"],
            "SQLCurrentName": s.get("FullName"),
            "ExcelCorrectName": x.get("FullName"),
            "SQLCurrentDOB": display_value(s.get("DateOfBirth")),
            "ExcelCorrectDOB": display_value(sql_value("DateOfBirth", x.get("DateOfBirth"))),
            "SQLCurrentMobile": display_value(s.get("MobileNumber")),
            "ExcelCorrectMobile": display_value(sql_value("MobileNumber", x.get("MobileNumber"))),
            "NeedsUpdate": "YES" if differences else "NO",
            "DifferenceCount": len(differences),
            "DifferenceFields": ", ".join(differences),
            "ExcelSource": (
                f"{x.get('_sheet')} row {x.get('_row')}"
                if x.get("_row")
                else f"{x.get('_sheet')} column {x.get('_column')}"
            ),
        })

    return summary_rows, diff_rows


def create_local_audit(output_path, sql_records, excel_records, summary_rows, diff_rows):
    wb = openpyxl.Workbook()

    # Remove default content later by reusing active sheet.
    ws = wb.active
    ws.title = "Comparison"

    comparison_headers = [
        "SaheliCardNumber",
        "ParticipantID",
        "SQLCurrentName",
        "ExcelCorrectName",
        "SQLCurrentDOB",
        "ExcelCorrectDOB",
        "SQLCurrentMobile",
        "ExcelCorrectMobile",
        "NeedsUpdate",
        "DifferenceCount",
        "DifferenceFields",
        "ExcelSource",
    ]
    ws.append(comparison_headers)

    for item in summary_rows:
        ws.append([item.get(h) for h in comparison_headers])

    ws2 = wb.create_sheet("FieldDifferences")
    diff_headers = [
        "SaheliCardNumber",
        "ParticipantID",
        "CurrentSQLName",
        "ExpectedExcelName",
        "Field",
        "SQLBefore",
        "ExcelMaster",
    ]
    ws2.append(diff_headers)

    for item in diff_rows:
        ws2.append([item.get(h) for h in diff_headers])

    ws3 = wb.create_sheet("SQL_BEFORE_Backup")
    sql_headers = [
        "ParticipantID",
        "SaheliCardNumber",
        "FullName",
        "DateOfBirth",
        "Age",
        "Address",
        "Postcode",
        "Email",
        "MobileNumber",
        "Gender",
        "GenderSameAsBirth",
        "Ethnicity",
        "PreferredLanguage",
        "Religion",
        "Sexuality",
        "Occupation",
        "LivingAlone",
        "CaringResponsibilities",
        "ReferralReason",
        "HeardAboutSaheli",
        "GPSurgeryName",
        "CreatedAt",
        "HasHealthConditionOrDisability",
        "HealthConditionDetails",
        "StaffMember",
        "Site",
        "Notes",
        "RegistrationDate",
        "GPSurgeryId",
    ]
    ws3.append(sql_headers)

    for card in TARGET_CARDS:
        item = sql_records[card]
        ws3.append([display_value(item.get(h)) for h in sql_headers])

    ws4 = wb.create_sheet("Excel_Master_42")
    excel_headers = ["SaheliCardNumber"] + UPDATE_COLUMNS
    ws4.append(excel_headers)

    for card in TARGET_CARDS:
        item = excel_records[card]
        row = []
        for h in excel_headers:
            if h == "SaheliCardNumber":
                value = card
            else:
                try:
                    value = sql_value(h, item.get(h))
                except Exception:
                    value = item.get(h)
            row.append(display_value(value))
        ws4.append(row)

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.font = openpyxl.styles.Font(bold=True)

        for column_cells in sheet.columns:
            width = 10
            for cell in list(column_cells)[:500]:
                if cell.value is not None:
                    width = max(width, min(len(str(cell.value)) + 2, 60))
            sheet.column_dimensions[column_cells[0].column_letter].width = width

    wb.save(output_path)


def perform_updates(conn, excel_records, sql_records):
    set_clause = ", ".join(f"[{col}] = ?" for col in UPDATE_COLUMNS)

    sql = f"""
    UPDATE dbo.Participants
    SET {set_clause}
    WHERE ParticipantID = ?
      AND CAST(SaheliCardNumber AS nvarchar(100)) = ?;
    """

    cur = conn.cursor()
    updated = []

    for card in TARGET_CARDS:
        excel_row = excel_records[card]
        sql_row = sql_records[card]

        params = [
            sql_value(col, excel_row.get(col))
            for col in UPDATE_COLUMNS
        ]

        params.extend([
            sql_row["ParticipantID"],
            card,
        ])

        cur.execute(sql, params)

        if cur.rowcount != 1:
            raise RuntimeError(
                f"Card {card}: expected to update exactly 1 row, "
                f"but SQL reported {cur.rowcount}. Entire transaction will roll back."
            )

        updated.append(card)

    return updated


def verify_after(conn, excel_records):
    after = fetch_sql_targets(conn)
    mismatches = []

    for card in TARGET_CARDS:
        expected = excel_records[card]
        actual = after[card]

        for col in UPDATE_COLUMNS:
            excel_value = sql_value(col, expected.get(col))
            sql_value_after = actual.get(col)

            if comparable(col, excel_value) != comparable(col, sql_value_after):
                mismatches.append(
                    f"{card} / {col}: SQL={sql_value_after!r}, Excel={excel_value!r}"
                )

    return after, mismatches


def main():
    parser = argparse.ArgumentParser(
        description="Preview or fix only the 42 identified Saheli participant records."
    )
    parser.add_argument(
        "--excel",
        default=DEFAULT_EXCEL,
        help=f"Master Excel path. Default: {DEFAULT_EXCEL}",
    )
    parser.add_argument(
        "--commit",
        action="store_true",
        help="Actually update the 42 dbo.Participants rows. Without this flag the script is preview-only.",
    )
    args = parser.parse_args()

    excel_path = Path(args.excel)

    print()
    print("=" * 72)
    print("SAHELI CRM - 42 PARTICIPANT REPAIR")
    print("=" * 72)
    print(f"Mode: {'COMMIT' if args.commit else 'PREVIEW ONLY'}")
    print(f"Excel: {excel_path}")
    print("Target records: exactly 42 fixed card numbers")
    print()

    if not excel_path.exists():
        print(f"ERROR: Excel file does not exist:\n{excel_path}")
        sys.exit(1)

    excel_records = read_master_excel(excel_path)

    print(f"Target Excel rows found: {len(excel_records)} / 42")

    server = SQL_SERVER.strip()
    database = SQL_DATABASE.strip() or "SaheliHubCRM"
    username = SQL_USERNAME.strip()
    password = SQL_PASSWORD

    config_errors = []
    if not server or "YOUR_SERVER" in server:
        config_errors.append("Set SQL_SERVER at the top of the script.")
    if not username or "YOUR_SQL_USERNAME" in username:
        config_errors.append("Set SQL_USERNAME at the top of the script.")
    if not password or "YOUR_SQL_PASSWORD" in password:
        config_errors.append("Set SQL_PASSWORD at the top of the script.")

    if config_errors:
        print()
        print("SQL CONFIG IS NOT FILLED IN")
        print("-" * 72)
        for error in config_errors:
            print(" - " + error)
        print()
        print("Edit SQL_SERVER, SQL_DATABASE, SQL_USERNAME and SQL_PASSWORD")
        print("near the top of this Python file, save it, then run again.")
        sys.exit(1)

    print(f"SQL server: {server}")
    print(f"Database: {database}")
    print(f"SQL user: {username}")
    print("SQL password: loaded from local script config")

    conn = None

    try:
        conn = connect_sql(server, database, username, password)
        sql_records = fetch_sql_targets(conn)

        print(f"Target SQL rows found: {len(sql_records)} / 42")

        errors = validate(excel_records, sql_records)

        if errors:
            print()
            print("VALIDATION FAILED - NOTHING WAS UPDATED")
            print("-" * 72)
            for error in errors:
                print(" - " + error)
            if conn:
                conn.rollback()
            sys.exit(2)

        summary_rows, diff_rows = build_comparison(
            excel_records,
            sql_records
        )

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_dir = Path.cwd() / "Saheli_42_Participant_Fix"
        output_dir.mkdir(parents=True, exist_ok=True)

        audit_path = output_dir / f"Saheli_42_Before_Comparison_{timestamp}.xlsx"

        create_local_audit(
            audit_path,
            sql_records,
            excel_records,
            summary_rows,
            diff_rows,
        )

        needs_update = [x for x in summary_rows if x["NeedsUpdate"] == "YES"]

        print()
        print(f"Local backup/comparison created:\n{audit_path}")
        print()
        print(f"Records needing changes: {len(needs_update)} / 42")
        print(f"Field differences found: {len(diff_rows)}")

        print()
        print("Preview of identities:")
        print("-" * 72)

        for item in summary_rows:
            marker = "*" if item["NeedsUpdate"] == "YES" else " "
            print(
                f"{marker} Card {item['SaheliCardNumber']}: "
                f"SQL '{item['SQLCurrentName']}' -> "
                f"Excel '{item['ExcelCorrectName']}' "
                f"({item['DifferenceCount']} field differences)"
            )

        if not args.commit:
            conn.rollback()
            print()
            print("=" * 72)
            print("PREVIEW COMPLETE - NO SQL DATA WAS CHANGED")
            print("=" * 72)
            print()
            print("Review the local Excel report first.")
            print("When satisfied, run the same script with --commit.")
            return

        print()
        print("WARNING:")
        print("This will overwrite the mapped dbo.Participants fields for")
        print("ONLY these 42 Saheli card numbers using the Excel as master.")
        print("ParticipantID and SaheliCardNumber will NOT be changed.")
        print("Assessments/attendance/funding/forms will NOT be changed.")
        print()

        confirmation = input(
            'Type exactly UPDATE 42 PARTICIPANTS to continue: '
        ).strip()

        if confirmation != "UPDATE 42 PARTICIPANTS":
            conn.rollback()
            print("Cancelled. No SQL data was changed.")
            return

        updated = perform_updates(
            conn,
            excel_records,
            sql_records
        )

        # Verify inside the same transaction before commit.
        after_records, verification_errors = verify_after(
            conn,
            excel_records
        )

        if verification_errors:
            conn.rollback()

            print()
            print("VERIFICATION FAILED - ALL CHANGES ROLLED BACK")
            print("-" * 72)

            for error in verification_errors[:100]:
                print(" - " + error)

            if len(verification_errors) > 100:
                print(f"... plus {len(verification_errors) - 100} more")

            sys.exit(3)

        conn.commit()

        after_path = output_dir / f"Saheli_42_AFTER_{timestamp}.xlsx"

        after_summary, after_diff = build_comparison(
            excel_records,
            after_records
        )

        create_local_audit(
            after_path,
            after_records,
            excel_records,
            after_summary,
            after_diff,
        )

        print()
        print("=" * 72)
        print("SUCCESS")
        print("=" * 72)
        print(f"42 rows updated and transaction committed.")
        print(f"Verification differences after update: {len(after_diff)}")
        print(f"After-check report:\n{after_path}")
        print()
        print("ParticipantID and SaheliCardNumber were preserved.")
        print("No assessment, attendance, funding, form or other child table was changed.")

    except Exception as exc:
        if conn is not None:
            try:
                conn.rollback()
            except Exception:
                pass

        print()
        print("=" * 72)
        print("ERROR - TRANSACTION ROLLED BACK / NOTHING COMMITTED")
        print("=" * 72)
        print(str(exc))
        raise

    finally:
        if conn is not None:
            conn.close()


if __name__ == "__main__":
    main()
