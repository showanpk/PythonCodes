"""Safely preview or repair exactly 42 Saheli dbo.Participants profiles.

Preview is the default and performs SELECT queries only.  Commit mode must be
requested explicitly with --commit, updates only records with available and
valid master data, runs the eligible batch in one transaction, re-reads and
verifies every intended field, and rolls the entire batch back on any error.

Assessments and all participant-linked child tables are intentionally outside
the scope of this program.
"""

from __future__ import annotations

import argparse
import re
import sys
import time
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import pyodbc

import saheli_full_readonly_audit as audit


DEFAULT_EXCEL = Path(r"C:\Users\shonk\Downloads\Full Registration for SAHELI (1).xlsx")
DEFAULT_AUDIT = Path(__file__).with_name("Saheli_Full_Excel_SQL_Audit_20260923_114145.xlsx")
DEFAULT_CONFIG = audit.DEFAULT_CONFIG

TARGET_CARDS = [
    "1209", "1212", "405", "716", "454", "861", "433", "727", "340", "585",
    "342", "467", "1060", "1062", "199", "1199", "345", "568", "1613", "1635",
    "98", "1626", "2039", "947", "404", "1601", "1380", "1360", "28", "86",
    "22664079", "489", "349", "341", "407", "423", "9", "1172", "643", "463",
    "2002", "951",
]

UPDATE_FIELDS = [
    "FullName", "DateOfBirth", "Age", "Address", "Postcode", "Email",
    "MobileNumber", "Gender", "GenderSameAsBirth", "Ethnicity",
    "PreferredLanguage", "Religion", "Sexuality", "Occupation", "LivingAlone",
    "CaringResponsibilities", "ReferralReason", "HeardAboutSaheli",
    "GPSurgeryName", "StaffMember", "Site", "Notes", "RegistrationDate",
]

EXCLUDED_FIELDS = [
    "ParticipantID", "SaheliCardNumber", "CreatedAt", "GPSurgeryId",
    "HasHealthConditionOrDisability", "HealthConditionDetails",
    "Emergency contact fields", "Consent/media/WhatsApp/declaration fields",
    "Relationship status", "All assessment and child-table fields",
]

DATE_FIELDS = {"DateOfBirth", "RegistrationDate"}
BIT_FIELDS = {"GenderSameAsBirth", "LivingAlone", "CaringResponsibilities"}
INT_FIELDS = {"Age"}
SQL_COLUMNS = ["ParticipantID", "SaheliCardNumber", *UPDATE_FIELDS,
               "CreatedAt", "GPSurgeryId", "HasHealthConditionOrDisability", "HealthConditionDetails"]

if len(TARGET_CARDS) != 42 or len(set(TARGET_CARDS)) != 42:
    raise RuntimeError("TARGET_CARDS must contain exactly 42 unique card numbers.")
if set(UPDATE_FIELDS) & {"ParticipantID", "SaheliCardNumber"}:
    raise RuntimeError("Identity keys must never appear in UPDATE_FIELDS.")


def empty_to_none(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def parse_date(value: Any, field: str) -> date | None:
    value = empty_to_none(value)
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    formats = ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d.%m.%Y", "%m/%d/%Y")
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"cannot convert {value!r} to a date")


def parse_int(value: Any, field: str) -> int | None:
    value = empty_to_none(value)
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError(f"cannot convert boolean {value!r} to an integer")
    if isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
        if number != number.to_integral_value():
            raise ValueError(f"{value!r} is not a whole number")
        return int(number)
    text = str(value).strip()
    if not re.fullmatch(r"[+-]?\d+", text):
        raise ValueError(f"cannot convert {value!r} to an integer")
    return int(text)


def parse_bit(value: Any, field: str) -> bool | None:
    value = empty_to_none(value)
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float, Decimal)) and value in (0, 1):
        return bool(value)
    text = re.sub(r"\s+", " ", str(value).strip()).casefold()
    if text in {"yes", "y", "true", "1", "same", "yes same", "yes - same"}:
        return True
    if text in {"no", "n", "false", "0", "different", "no different", "no - different"}:
        return False
    raise ValueError(f"cannot safely convert {value!r} to SQL bit")


def convert_excel_value(field: str, value: Any) -> Any:
    if field in DATE_FIELDS:
        return parse_date(value, field)
    if field in INT_FIELDS:
        return parse_int(value, field)
    if field in BIT_FIELDS:
        return parse_bit(value, field)
    value = empty_to_none(value)
    return None if value is None else str(value).strip()


def values_equal(field: str, expected: Any, actual: Any) -> bool:
    if field in DATE_FIELDS:
        try:
            return parse_date(expected, field) == parse_date(actual, field)
        except ValueError:
            return False
    if field in INT_FIELDS:
        try:
            return parse_int(expected, field) == parse_int(actual, field)
        except ValueError:
            return False
    if field in BIT_FIELDS:
        try:
            return parse_bit(expected, field) == parse_bit(actual, field)
        except ValueError:
            return False
    return empty_to_none(expected) == empty_to_none(actual)


def display(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    return value


def connection_string(config: dict[str, str], read_only: bool) -> str:
    server = config["SQL_SERVER"]
    if server.casefold().startswith("tcp:"):
        server = server[4:]
    if "," not in server:
        server += ",1433"
    intent = "ApplicationIntent=ReadOnly;" if read_only else ""
    return (
        "DRIVER={ODBC Driver 18 for SQL Server};"
        f"SERVER=tcp:{server};DATABASE={config['SQL_DATABASE']};"
        f"UID={config['SQL_USERNAME']};PWD={config['SQL_PASSWORD']};"
        "Encrypt=yes;TrustServerCertificate=no;Connection Timeout=30;"
        f"{intent}"
    )


class ParticipantConnection:
    """Narrow database surface: target SELECTs plus one fixed UPDATE."""

    def __init__(self, config: dict[str, str], commit_mode: bool):
        self.commit_mode = commit_mode
        self.connection = pyodbc.connect(connection_string(config, read_only=not commit_mode), autocommit=False)
        self.connection.cursor().execute("SET XACT_ABORT ON;")

    def fetch_targets(self) -> list[dict[str, Any]]:
        placeholders = ",".join("?" for _ in TARGET_CARDS)
        columns = ",".join(f"[{field}]" for field in SQL_COLUMNS)
        cursor = self.connection.cursor()
        cursor.execute(
            f"SELECT {columns} FROM dbo.Participants WHERE SaheliCardNumber IN ({placeholders})",
            tuple(TARGET_CARDS),
        )
        names = [item[0] for item in cursor.description]
        return [dict(zip(names, row)) for row in cursor.fetchall()]

    def fetch_metadata(self) -> list[dict[str, Any]]:
        placeholders = ",".join("?" for _ in UPDATE_FIELDS)
        cursor = self.connection.cursor()
        cursor.execute(
            "SELECT c.name AS ColumnName,t.name AS DataType,c.max_length AS MaxLength,c.is_nullable AS IsNullable "
            "FROM sys.columns c JOIN sys.types t ON t.user_type_id=c.user_type_id "
            "WHERE c.object_id=OBJECT_ID('dbo.Participants') AND c.name IN (" + placeholders + ")",
            tuple(UPDATE_FIELDS),
        )
        names = [item[0] for item in cursor.description]
        return [dict(zip(names, row)) for row in cursor.fetchall()]

    def update_participant(self, participant_id: int, card: str, values: dict[str, Any]) -> int:
        if not self.commit_mode:
            raise RuntimeError("Safety guard: UPDATE is unavailable outside --commit mode.")
        fields = [field for field in UPDATE_FIELDS if field in values]
        if not fields or set(values) != set(fields) or not set(fields).issubset(UPDATE_FIELDS):
            raise RuntimeError("Safety guard: update contains no fields or an unapproved field.")
        assignments = ",".join(f"[{field}]=?" for field in fields)
        parameters = [values[field] for field in fields] + [participant_id, card]
        cursor = self.connection.cursor()
        cursor.execute(
            f"UPDATE dbo.Participants SET {assignments} WHERE ParticipantID=? AND SaheliCardNumber=?",
            tuple(parameters),
        )
        return cursor.rowcount

    def commit(self) -> None:
        if not self.commit_mode:
            raise RuntimeError("Safety guard: commit unavailable in preview mode.")
        self.connection.commit()

    def rollback(self) -> None:
        self.connection.rollback()

    def close(self) -> None:
        self.connection.close()


def audit_report_records(path: Path) -> dict[str, dict[str, Any]]:
    if not path.is_file():
        return {}
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "ParticipantComparison" not in workbook.sheetnames:
            return {}
        rows = workbook["ParticipantComparison"].iter_rows(values_only=True)
        headers = list(next(rows))
        output = {}
        for row in rows:
            item = dict(zip(headers, row))
            card = audit.clean_card(item.get("SaheliCardNumber"))
            if card in TARGET_CARDS:
                output[card] = item
        return output
    finally:
        workbook.close()


def build_validation(excel: audit.ExcelData, sql_rows: list[dict[str, Any]],
                     metadata: list[dict[str, Any]], prior_audit: dict[str, dict[str, Any]]):
    issues: list[dict[str, Any]] = []
    converted: dict[str, dict[str, Any]] = {}
    skipped_fields: dict[str, dict[str, str]] = {}
    sql_groups: dict[str, list[dict[str, Any]]] = {}
    for row in sql_rows:
        sql_groups.setdefault(audit.clean_card(row.get("SaheliCardNumber")), []).append(row)

    mapped_headers = {
        str(row.get("SQLMapping", "")).rsplit(".", 1)[-1]
        for row in excel.structure if str(row.get("SQLMapping", "")).startswith("dbo.Participants.")
    }
    for field in UPDATE_FIELDS:
        if field not in mapped_headers:
            issues.append({"Severity": "BLOCKING", "Status": "VALIDATION_ERROR", "SaheliCardNumber": "ALL", "Field": field,
                           "Issue": "Required Excel header was not mapped; SQL NULL will not be inferred."})

    metadata_by_name = {row["ColumnName"]: row for row in metadata}
    for field in UPDATE_FIELDS:
        if field not in metadata_by_name:
            issues.append({"Severity": "BLOCKING", "Status": "VALIDATION_ERROR", "SaheliCardNumber": "ALL", "Field": field,
                           "Issue": "Approved field does not exist in live dbo.Participants metadata."})

    for card in TARGET_CARDS:
        excel_count = (1 if card in excel.participants else 0) + max(0, len(excel.duplicate_cards.get(card, [])) - 1)
        if excel_count == 0:
            issues.append({"Severity": "BLOCKED", "Status": "BLOCKED_MISSING_MASTER_DATA",
                           "SaheliCardNumber": card, "Field": "SaheliCardNumber",
                           "Issue": "Master Excel row is missing; no data will be derived from another card."})
        elif excel_count != 1:
            issues.append({"Severity": "BLOCKING", "Status": "VALIDATION_ERROR",
                           "SaheliCardNumber": card, "Field": "SaheliCardNumber",
                           "Issue": f"Expected exactly one Excel row; found {excel_count}."})
        sql_count = len(sql_groups.get(card, []))
        if sql_count != 1:
            issues.append({"Severity": "BLOCKING", "Status": "VALIDATION_ERROR", "SaheliCardNumber": card, "Field": "SaheliCardNumber",
                           "Issue": f"Expected exactly one dbo.Participants row; found {sql_count}."})
        elif sql_groups[card][0].get("ParticipantID") is None:
            issues.append({"Severity": "BLOCKING", "Status": "VALIDATION_ERROR", "SaheliCardNumber": card, "Field": "ParticipantID",
                           "Issue": "ParticipantID was not resolved."})

        if card not in prior_audit:
            issues.append({"Severity": "BLOCKING", "Status": "VALIDATION_ERROR", "SaheliCardNumber": card, "Field": "Prior audit",
                           "Issue": "Card is absent from the supplied full-audit ParticipantComparison sheet."})
        elif sql_count == 1 and prior_audit[card].get("ParticipantID") != sql_groups[card][0].get("ParticipantID"):
            issues.append({"Severity": "BLOCKING", "Status": "VALIDATION_ERROR", "SaheliCardNumber": card, "Field": "ParticipantID",
                           "Issue": "Live ParticipantID differs from the supplied full-audit report."})

        if excel_count == 1 and card in excel.participants:
            values: dict[str, Any] = {}
            for field in UPDATE_FIELDS:
                if field not in excel.participants[card]:
                    issues.append({"Severity": "BLOCKING", "Status": "VALIDATION_ERROR", "SaheliCardNumber": card, "Field": field,
                                   "Issue": "Field absent from parsed Excel row; value will not be treated as blank."})
                    continue
                raw = excel.participants[card].get(field)
                if field == "Age":
                    try:
                        converted_age = convert_excel_value(field, raw)
                        if converted_age is None:
                            raise ValueError("blank")
                        values[field] = converted_age
                    except ValueError:
                        raw_text = str(raw).strip()
                        reason = f"Excel formula error {raw_text}" if raw_text.startswith("#") else f"Excel Age is blank or non-numeric: {raw!r}"
                        skipped_fields.setdefault(card, {})[field] = reason
                        issues.append({"Severity": "WARNING", "Status": "WARNING_SKIPPED_AGE",
                                       "SaheliCardNumber": card, "Field": "Age",
                                       "Issue": f"AgeAction=PRESERVE_SQL_VALUE; Reason={reason}"})
                    continue
                try:
                    converted_value = convert_excel_value(field, raw)
                    column_meta = metadata_by_name.get(field)
                    if column_meta and converted_value is None and not bool(column_meta["IsNullable"]):
                        raise ValueError("Excel is blank but the SQL column is not nullable")
                    if column_meta and isinstance(converted_value, str) and int(column_meta["MaxLength"]) > 0:
                        max_chars = int(column_meta["MaxLength"])
                        if str(column_meta["DataType"]).casefold() in {"nvarchar", "nchar"}:
                            max_chars //= 2
                        if len(converted_value) > max_chars:
                            raise ValueError(f"text length {len(converted_value)} exceeds SQL limit {max_chars}")
                    values[field] = converted_value
                except ValueError as exc:
                    issues.append({"Severity": "BLOCKING", "Status": "VALIDATION_ERROR", "SaheliCardNumber": card, "Field": field,
                                   "Issue": f"Conversion failed: {exc}"})
            required_count = len(UPDATE_FIELDS) - (1 if "Age" in skipped_fields.get(card, {}) else 0)
            if len(values) == required_count:
                converted[card] = values

    return issues, converted, sql_groups, skipped_fields


def compare_targets(converted: dict[str, dict[str, Any]], sql_groups: dict[str, list[dict[str, Any]]],
                    skipped_fields: dict[str, dict[str, str]], issues: list[dict[str, Any]]):
    before_after = []
    changes = []
    for card in TARGET_CARDS:
        sql = sql_groups.get(card, [None])[0] if len(sql_groups.get(card, [])) == 1 else None
        expected = converted.get(card)
        changed_fields = []
        if sql and expected:
            for field in UPDATE_FIELDS:
                if field not in expected:
                    continue
                if not values_equal(field, expected[field], sql.get(field)):
                    changed_fields.append(field)
                    changes.append({
                        "SaheliCardNumber": card, "ParticipantID": sql.get("ParticipantID"),
                        "SQLCurrentName": sql.get("FullName"), "ExcelCorrectName": expected.get("FullName"),
                        "Field": field, "SQLBefore": display(sql.get(field)),
                        "ExcelMaster": display(expected.get(field)),
                        "WillSetSQLNull": "YES" if expected.get(field) is None else "NO",
                    })
        card_issues = [item for item in issues if item.get("SaheliCardNumber") == card]
        missing_master = any(item.get("Status") == "BLOCKED_MISSING_MASTER_DATA" for item in card_issues)
        validation_error = any(item.get("Status") == "VALIDATION_ERROR" for item in card_issues)
        skipped_age = "Age" in skipped_fields.get(card, {})
        if missing_master:
            status = "BLOCKED_MISSING_MASTER_DATA"
        elif validation_error or not sql or not expected:
            status = "VALIDATION_ERROR"
        elif skipped_age:
            status = "WARNING_SKIPPED_AGE"
        elif changed_fields:
            status = "READY_TO_UPDATE"
        else:
            status = "ALREADY_CORRECT"
        before_after.append({
            "SaheliCardNumber": card,
            "ParticipantID": "" if not sql else sql.get("ParticipantID"),
            "SQLCurrentName": "" if not sql else sql.get("FullName"),
            "ExcelCorrectName": "" if not expected else expected.get("FullName"),
            "SQLDOB": "" if not sql else display(sql.get("DateOfBirth")),
            "ExcelDOB": "" if not expected else display(expected.get("DateOfBirth")),
            "SQLMobile": "" if not sql else sql.get("MobileNumber"),
            "ExcelMobile": "" if not expected else expected.get("MobileNumber"),
            "ChangeCount": len(changed_fields), "FieldsThatWillChange": ", ".join(changed_fields),
            "Status": status,
            "AgeAction": "PRESERVE_SQL_VALUE" if skipped_age else ("USE_EXCEL_VALUE" if expected else "NOT_APPLICABLE"),
            "AgeReason": skipped_fields.get(card, {}).get("Age", ""),
        })
    return before_after, changes


def excel_master_rows(excel: audit.ExcelData) -> list[dict[str, Any]]:
    rows = []
    for card in TARGET_CARDS:
        item = excel.participants.get(card, {})
        row = {"SaheliCardNumber": card, "ExcelRowFound": "YES" if item else "NO",
               "SourceSheet": item.get("_sheet", ""), "SourceRow": item.get("_row", "")}
        row.update({field: display(item.get(field)) for field in UPDATE_FIELDS})
        rows.append(row)
    return rows


def sql_backup_rows(sql_groups: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    rows = []
    for card in TARGET_CARDS:
        members = sql_groups.get(card, [])
        if len(members) == 1:
            rows.append({key: display(value) for key, value in members[0].items()})
        else:
            rows.append({"SaheliCardNumber": card, "ValidationStatus": f"Expected one SQL row; found {len(members)}"})
    return rows


def report_summary(mode: str, issues, before_after, committed: bool = False,
                   actual_changed: int | None = None) -> list[dict[str, Any]]:
    eligible_statuses = {"READY_TO_UPDATE", "ALREADY_CORRECT", "WARNING_SKIPPED_AGE"}
    valid = [row for row in before_after if row["Status"] in eligible_statuses]
    changing = [row for row in valid if row["ChangeCount"] > 0]
    blocked = [row for row in before_after if row["Status"] not in eligible_statuses]
    rows = [
        {"Metric": "Mode", "Value": mode},
        {"Metric": "SQL committed", "Value": "YES" if committed else "NO"},
        {"Metric": "Target card count", "Value": len(TARGET_CARDS)},
        {"Metric": "Targets fully validated", "Value": len(valid)},
        {"Metric": "Targets blocked", "Value": len(blocked)},
        {"Metric": "Blocking validation issues", "Value": sum(i.get("Severity") != "WARNING" for i in issues)},
        {"Metric": "Warnings", "Value": sum(i.get("Severity") == "WARNING" for i in issues)},
        {"Metric": "Records requiring field changes", "Value": len(changing)},
        {"Metric": "Records already correct", "Value": len(valid) - len(changing)},
        {"Metric": "Mapped fields", "Value": len(UPDATE_FIELDS)},
        {"Metric": "Mapped-field changes", "Value": sum(row["ChangeCount"] for row in valid)},
        {"Metric": "ParticipantID changes", "Value": 0},
        {"Metric": "SaheliCardNumber changes", "Value": 0},
        {"Metric": "Other tables changed", "Value": 0},
    ]
    if actual_changed is not None:
        rows.extend([
            {"Metric": "Records actually changed", "Value": actual_changed},
            {"Metric": "Records already correct before repair", "Value": len(valid) - actual_changed},
            {"Metric": "Unresolved mapped-field differences", "Value": sum(row["ChangeCount"] for row in valid)},
            {"Metric": "Verification result", "Value": "PASSED" if committed else "NOT COMMITTED"},
        ])
    return rows


def create_report(path: Path, mode: str, excel: audit.ExcelData, sql_groups,
                  before_after, changes, issues, skipped_fields, committed: bool = False,
                  actual_changed: int | None = None) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    audit.write_sheet(workbook, "Summary", report_summary(mode, issues, before_after, committed, actual_changed))
    audit.write_sheet(workbook, "BeforeAfter", before_after)
    audit.write_sheet(workbook, "FieldChanges", changes)
    audit.write_sheet(workbook, "SQLBeforeBackup", sql_backup_rows(sql_groups))
    audit.write_sheet(workbook, "ExcelMaster42", excel_master_rows(excel))
    audit.write_sheet(workbook, "ValidationIssues", issues)
    audit.write_sheet(workbook, "BlockedRecords", [row for row in before_after if row["Status"] in {"BLOCKED_MISSING_MASTER_DATA", "VALIDATION_ERROR"}])
    skipped_rows = [{"SaheliCardNumber": card, "Field": field, "Action": "PRESERVE_SQL_VALUE", "Reason": reason}
                    for card, fields in skipped_fields.items() for field, reason in fields.items()]
    audit.write_sheet(workbook, "SkippedFields", skipped_rows)
    workbook.properties.title = f"Saheli 42 Participant Repair {mode}"
    workbook.properties.subject = "Controlled dbo.Participants-only profile repair"
    workbook.properties.description = "Assessments and linked child tables are excluded."
    workbook.save(path)
    validate_report(path)


def validate_report(path: Path) -> None:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        required = {"Summary", "BeforeAfter", "FieldChanges", "SQLBeforeBackup", "ExcelMaster42",
                    "ValidationIssues", "BlockedRecords", "SkippedFields"}
        missing = required - set(workbook.sheetnames)
        if missing:
            raise RuntimeError(f"Report is missing sheets: {sorted(missing)}")
        if workbook["BeforeAfter"].max_row != 43:
            raise RuntimeError("BeforeAfter must contain exactly 42 target rows plus its header.")
        if workbook["ExcelMaster42"].max_row != 43:
            raise RuntimeError("ExcelMaster42 must contain exactly 42 target rows plus its header.")
        for sheet in workbook.worksheets:
            if sheet.max_row < 2:
                raise RuntimeError(f"Report sheet is empty: {sheet.title}")
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f" and isinstance(cell.value, str) and any(
                        token in cell.value for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")
                    ):
                        raise RuntimeError(f"Formula error in {sheet.title}!{cell.coordinate}")
    finally:
        workbook.close()


def verify_after(before_sql: dict[str, list[dict[str, Any]]], after_rows: list[dict[str, Any]],
                 converted: dict[str, dict[str, Any]], skipped_fields: dict[str, dict[str, str]]) -> list[str]:
    errors = []
    after_groups: dict[str, list[dict[str, Any]]] = {}
    for row in after_rows:
        after_groups.setdefault(audit.clean_card(row.get("SaheliCardNumber")), []).append(row)
    for card in converted:
        before = before_sql[card][0]
        members = after_groups.get(card, [])
        if len(members) != 1:
            errors.append(f"Card {card}: expected exactly one post-update SQL row; found {len(members)}")
            continue
        after = members[0]
        if before["ParticipantID"] != after["ParticipantID"]:
            errors.append(f"Card {card}: ParticipantID changed")
        if audit.clean_card(after["SaheliCardNumber"]) != card:
            errors.append(f"Card {card}: SaheliCardNumber changed")
        for field in converted[card]:
            if not values_equal(field, converted[card][field], after.get(field)):
                errors.append(f"Card {card}, {field}: post-update value does not match Excel")
        for field in skipped_fields.get(card, {}):
            if not values_equal(field, before.get(field), after.get(field)):
                errors.append(f"Card {card}, {field}: skipped field did not preserve its SQL BEFORE value")
    return errors


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Preview or commit the controlled 42-card dbo.Participants repair")
    parser.add_argument("--commit", action="store_true", help="Perform the transaction after all validation succeeds")
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL, help="Authoritative master workbook")
    parser.add_argument("--audit-report", type=Path, default=DEFAULT_AUDIT, help="Previously generated full audit")
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG, help="Python file containing the existing SQL config constants")
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).parent, help="Local report directory")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    mode = "COMMIT" if args.commit else "PREVIEW ONLY"
    print("=" * 72)
    print(f"SAHELI 42 PARTICIPANT PROFILE REPAIR - {mode}")
    print("=" * 72)
    print("Scope: dbo.Participants profile fields only; assessments and child tables excluded.")
    if not args.excel.is_file():
        raise FileNotFoundError(args.excel)
    if not args.audit_report.is_file():
        raise FileNotFoundError(args.audit_report)
    if not args.config.is_file():
        raise FileNotFoundError(args.config)

    started = time.perf_counter()
    excel = audit.read_excel(args.excel)
    prior_audit = audit_report_records(args.audit_report)
    config = audit.load_sql_config(args.config)
    db = ParticipantConnection(config, commit_mode=args.commit)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    try:
        print("Reading the 42 target dbo.Participants rows...", flush=True)
        sql_rows = db.fetch_targets()
        metadata = db.fetch_metadata()
        issues, converted, sql_groups, skipped_fields = build_validation(excel, sql_rows, metadata, prior_audit)
        before_after, changes = compare_targets(converted, sql_groups, skipped_fields, issues)
        eligible_statuses = {"READY_TO_UPDATE", "ALREADY_CORRECT", "WARNING_SKIPPED_AGE"}
        eligible_cards = [row["SaheliCardNumber"] for row in before_after if row["Status"] in eligible_statuses]
        changing_cards = [row["SaheliCardNumber"] for row in before_after
                          if row["Status"] in eligible_statuses and row["ChangeCount"] > 0]
        already_correct = [row["SaheliCardNumber"] for row in before_after
                           if row["Status"] == "ALREADY_CORRECT"]
        blocked_cards = [row["SaheliCardNumber"] for row in before_after
                         if row["Status"] in {"BLOCKED_MISSING_MASTER_DATA", "VALIDATION_ERROR"}]
        eligible_converted = {card: converted[card] for card in eligible_cards}
        print("Eligible target cards: " + ", ".join(eligible_cards))
        print("Records requiring update: " + (", ".join(changing_cards) or "NONE"))
        print("Already correct: " + (", ".join(already_correct) or "NONE"))
        print("Blocked: " + (", ".join(blocked_cards) or "NONE"))

        if not args.commit:
            db.rollback()
            output = args.output_dir / f"Saheli_42_Participant_Repair_Preview_{timestamp}.xlsx"
            create_report(output, "PREVIEW ONLY", excel, sql_groups, before_after, changes, issues,
                          skipped_fields, committed=False)
            print(f"Preview report: {output}")
            print(f"Eligible records: {len(eligible_cards)}")
            print(f"Records requiring update: {len(changing_cards)}")
            print(f"Blocked records: {len(blocked_cards)}")
            print(f"Warnings: {sum(item.get('Severity') == 'WARNING' for item in issues)}")
            print("No SQL data was modified.")
            return 0

        fatal_issues = [item for item in issues
                        if item.get("Severity") == "BLOCKING" and item.get("SaheliCardNumber") == "ALL"]
        if fatal_issues:
            db.rollback()
            output = args.output_dir / f"Saheli_42_Participant_Repair_Preview_{timestamp}.xlsx"
            create_report(output, "COMMIT BLOCKED - PREVIEW", excel, sql_groups, before_after, changes, issues,
                          skipped_fields, committed=False)
            print(f"COMMIT ABORTED: {len(fatal_issues)} validation error(s).")
            print(f"Validation report: {output}")
            print("Transaction rolled back; no SQL data was modified.")
            return 2

        confirmation = input("Type UPDATE VALIDATED PARTICIPANTS to continue: ").strip()
        if confirmation != "UPDATE VALIDATED PARTICIPANTS":
            db.rollback()
            print("Confirmation did not match. Transaction rolled back; no SQL data was modified.")
            return 3

        before_path = args.output_dir / f"Saheli_42_Participant_Repair_BEFORE_{timestamp}.xlsx"
        create_report(before_path, "BEFORE COMMIT BACKUP", excel, sql_groups, before_after, changes, issues,
                      skipped_fields, committed=False)
        print(f"Local BEFORE backup created: {before_path}", flush=True)

        print(f"Updating {len(changing_cards)} changed participant profiles in one transaction...", flush=True)
        for card in changing_cards:
            participant_id = sql_groups[card][0]["ParticipantID"]
            affected = db.update_participant(participant_id, card, converted[card])
            if affected != 1:
                raise RuntimeError(f"Card {card}: UPDATE affected {affected} rows; expected exactly 1")

        print(f"Re-reading and verifying all {len(eligible_cards)} eligible rows before commit...", flush=True)
        after_rows = db.fetch_targets()
        verification_errors = verify_after(sql_groups, after_rows, eligible_converted, skipped_fields)
        if verification_errors:
            raise RuntimeError("Post-update verification failed: " + "; ".join(verification_errors[:20]))
        db.commit()
        print("Transaction committed after successful verification.", flush=True)

        after_groups: dict[str, list[dict[str, Any]]] = {}
        for row in after_rows:
            after_groups.setdefault(audit.clean_card(row["SaheliCardNumber"]), []).append(row)
        after_before_after, after_changes = compare_targets(converted, after_groups, skipped_fields, issues)
        after_path = args.output_dir / f"Saheli_42_Participant_Repair_AFTER_{timestamp}.xlsx"
        create_report(after_path, "AFTER COMMIT", excel, sql_groups, after_before_after, after_changes, issues,
                      skipped_fields, committed=True, actual_changed=len(changing_cards))
        print(f"AFTER report: {after_path}")
        print(f"Eligible records: {len(eligible_cards)}")
        print(f"Records changed: {len(changing_cards)}")
        print(f"Records already correct: {len(eligible_cards)-len(changing_cards)}")
        print(f"Blocked records: {len(blocked_cards)}")
        print(f"Skipped fields: {sum(len(fields) for fields in skipped_fields.values())}")
        print("Unresolved mapped-field differences: 0")
        return 0
    except Exception:
        db.rollback()
        print("ERROR: transaction rolled back; no partial SQL update was retained.", file=sys.stderr)
        raise
    finally:
        db.close()
        print(f"Runtime: {time.perf_counter()-started:.1f}s")


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("Cancelled; transaction rolled back.", file=sys.stderr)
        raise SystemExit(130)
