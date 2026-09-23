"""Complete, local, read-only Saheli Excel versus Azure SQL audit.

This program deliberately has no SQL write path.  Every database statement is
checked before execution, the connection requests ApplicationIntent=ReadOnly,
and the transaction is rolled back before the connection is closed.
"""

from __future__ import annotations

import argparse
import ast
import re
import sys
import time
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import pyodbc
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_EXCEL = Path(r"C:\Users\shonk\Downloads\Full Registration for SAHELI (1).xlsx")
DEFAULT_CONFIG = Path(__file__).with_name("compare_and_fix_42_saheli_participants_with_config.py")
OUTPUT_PREFIX = "Saheli_Full_Excel_SQL_Audit"

FORBIDDEN_SQL = re.compile(
    r"\b(UPDATE|DELETE|INSERT|MERGE|TRUNCATE|DROP|ALTER|CREATE|EXEC(?:UTE)?|GRANT|REVOKE|DENY)\b",
    re.IGNORECASE,
)

PARTICIPANT_COLUMNS = [
    "ParticipantID", "SaheliCardNumber", "FullName", "DateOfBirth", "Age",
    "Address", "Postcode", "Email", "MobileNumber", "Gender",
    "GenderSameAsBirth", "Ethnicity", "PreferredLanguage", "Religion",
    "Sexuality", "Occupation", "LivingAlone", "CaringResponsibilities",
    "ReferralReason", "HeardAboutSaheli", "GPSurgeryName", "CreatedAt",
    "HasHealthConditionOrDisability", "HealthConditionDetails", "StaffMember",
    "Site", "Notes", "RegistrationDate", "GPSurgeryId",
]

REGISTRATION_MAP = {
    "registrationdate": "RegistrationDate",
    "sahelicardnumber": "SaheliCardNumber",
    "fullname": "FullName",
    "dateofbirth": "DateOfBirth",
    "age": "Age",
    "address": "Address",
    "postcode": "Postcode",
    "email": "Email",
    "mobilehomeno": "MobileNumber",
    "mobilenumber": "MobileNumber",
    "gender": "Gender",
    "isyourgenderthesameasassignedatbirth": "GenderSameAsBirth",
    "healthconditionsdisability": "HealthConditionDetails",
    "ethnicity": "Ethnicity",
    "preferredspokenlanguage": "PreferredLanguage",
    "preferredlanguage": "PreferredLanguage",
    "religion": "Religion",
    "caringresponsibilities": "CaringResponsibilities",
    "livingalone": "LivingAlone",
    "sexuality": "Sexuality",
    "occupation": "Occupation",
    "referralreason": "ReferralReason",
    "howheardaboutsahelihub": "HeardAboutSaheli",
    "gpsurgeryname": "GPSurgeryName",
    "notes": "Notes",
    "staffmember": "StaffMember",
    "site": "Site",
}

CORE_FIELDS = ("FullName", "DateOfBirth", "MobileNumber", "Postcode")
DATE_FIELDS = {"DateOfBirth", "RegistrationDate", "AssessmentDate", "NextReviewDate", "LastBpmeasurementDate"}
BOOL_FIELDS = {
    "GenderSameAsBirth", "LivingAlone", "CaringResponsibilities",
    "HasHealthConditionOrDisability", "HasHealthCondition", "DoctorAdvisedNoExercise",
    "ChestPain", "SugaryDrinkIntake", "HighCholesterol", "TakesPrescribedMedication",
    "ReferredToDoctor", "ShortnessOfBreath", "BprecordedWithGp",
}

# Exact label-to-schema mappings based on the live database metadata. Ambiguous
# fields are intentionally omitted and listed in the report's UnmappedFields tab.
ASSESSMENT_MAP = {
    "assessmentdate": ("Assessment_Master", "AssessmentDate"),
    "weightkg": ("Assessment_BodyComposition", "WeightKg"),
    "heightcm": ("Assessment_BodyComposition", "HeightCm"),
    "bmi": ("Assessment_BodyComposition", "Bmivalue"),
    "bmiresults": ("Assessment_BodyComposition", "Bmicategory"),
    "waistcm": ("Assessment_BodyComposition", "WaistCm"),
    "hipcm": ("Assessment_BodyComposition", "HipCm"),
    "waisttohipratiocm": ("Assessment_BodyComposition", "WaistHipRatio"),
    "bodyfatpercentageresult": ("Assessment_BodyComposition", "BodyFatScore"),
    "bodyfatpercentagescore": ("Assessment_BodyComposition", "BodyFatCategory"),
    "visceralfatlevelresult": ("Assessment_BodyComposition", "VisceralFatScore"),
    "visceralfatlevelscore": ("Assessment_BodyComposition", "VisceralFatCategory"),
    "skeletalmusclepercentage": ("Assessment_BodyComposition", "SkeletalMuscleScore"),
    "skeletalmusclescore": ("Assessment_BodyComposition", "SkeletalMuscleCategory"),
    "restingmetabolism": ("Assessment_BodyComposition", "RestingMetabolism"),
    "doyouhaveanyhealthcondition": ("Assessment_HealthScreening", "HasHealthCondition"),
    "whendidyoulastmeasureyourbloodpressure": ("Assessment_HealthScreening", "LastBpmeasurementDate"),
    "haveyourecordedyourbloodpressuremeasurementandregistereditwithagporpharmacistyesnonotsure": ("Assessment_HealthScreening", "BprecordedWithGp"),
    "whatisahealthybloodpressureforanadult": ("Assessment_HealthScreening", "KnowledgeHealthyBp"),
    "whyisahighbloodpressuredangerous": ("Assessment_HealthScreening", "KnowledgeBprisk"),
    "howcanyouhelpreduceyourbloodpressure": ("Assessment_HealthScreening", "KnowledgeBpreduction"),
    "bloodpressurelevel": ("Assessment_HealthScreening", "Bplevel"),
    "doyouhaveaheartcondition": ("Assessment_HealthScreening", "HeartConditionTypes"),
    "heartratebpm": ("Assessment_HealthScreening", "HeartRateBpm"),
    "atrialfibrillationresult": ("Assessment_HealthScreening", "AtrialFibrillationResult"),
    "heartage": ("Assessment_HealthScreening", "HeartAge"),
    "didyourdoctoradviseyounottoexercise": ("Assessment_HealthScreening", "DoctorAdvisedNoExercise"),
    "doyoufeelpaininchestatrestduringactivity": ("Assessment_HealthScreening", "ChestPain"),
    "doyouhaveshortnessofbreath": ("Assessment_HealthScreening", "ShortnessOfBreath"),
    "doyouhavediabetes": ("Assessment_HealthScreening", "DiabetesType"),
    "diabetesrisk": ("Assessment_HealthScreening", "DiabetesRisk"),
    "glucoselevelmgdl": ("Assessment_HealthScreening", "GlucoseLevel"),
    "hba1c": ("Assessment_HealthScreening", "HbA1c"),
    "doyoutakesugarydrinksincludingchai": ("Assessment_HealthScreening", "SugaryDrinkIntake"),
    "doyouexperiencethefollowinghealthissues": ("Assessment_HealthScreening", "OtherHealthIssues"),
    "doyouhaveabonejointcondition": ("Assessment_HealthScreening", "BoneJointConditions"),
    "doyoutakeanyprescribedmedication": ("Assessment_HealthScreening", "TakesPrescribedMedication"),
    "referredtodoctorforanyconcerningresults": ("Assessment_HealthScreening", "ReferredToDoctor"),
    "riskstratificationscore": ("Assessment_HealthScreening", "RiskStratification"),
    "comments": ("Assessment_HealthScreening", "HealthComments"),
    "howwelldoyoumanageyourhealthconditionsratingoutof10": ("Assessment_HealthScreening", "SelfManagementScore"),
    "inthepastweekonhowmanydayshaveyoudoneatotalof30minsormoreofphysicalactivitywhichwasenoughtoraiseyourbreathingrate": ("Assessment_PhysicalActivity", "ActiveDaysPerWeek"),
    "physicalactivitylevel": ("Assessment_PhysicalActivity", "ActivityLevel"),
    "commentspa": ("Assessment_PhysicalActivity", "ActivityComments"),
    "ivebeenfeelingoptimisticaboutthefuture": ("Assessment_WEMWBS", "FeelingOptimistic"),
    "ivebeenfeelinguseful": ("Assessment_WEMWBS", "FeelingUseful"),
    "ivebeenfeelingrelaxed": ("Assessment_WEMWBS", "FeelingRelaxed"),
    "ivebeenfeelinginterestedinotherpeople": ("Assessment_WEMWBS", "FeelingInterestedInPeople"),
    "ivehadenergytospare": ("Assessment_WEMWBS", "EnergyToSpare"),
    "ivebeendealingwithproblemswell": ("Assessment_WEMWBS", "DealingWithProblems"),
    "ivebeenthinkingclearly": ("Assessment_WEMWBS", "ThinkingClearly"),
    "ivebeenfeelinggoodaboutmyself": ("Assessment_WEMWBS", "FeelingGoodAboutSelf"),
    "ivebeenfeelingclosetootherpeople": ("Assessment_WEMWBS", "FeelingCloseToOthers"),
    "ivebeenfeelingconfident": ("Assessment_WEMWBS", "FeelingConfident"),
    "ivebeenabletomakeupmyownmindaboutthings": ("Assessment_WEMWBS", "MakingOwnMindUp"),
    "ivebeenfeelingloved": ("Assessment_WEMWBS", "FeelingLoved"),
    "ivebeeninterestedinnewthings": ("Assessment_WEMWBS", "InterestedInNewThings"),
    "ivebeenfeelingcheerful": ("Assessment_WEMWBS", "FeelingCheerful"),
    "wemwbs": ("Assessment_WEMWBS", "Wemwbscomments"),
    "nourishmentratethequalityofthefoodyouputintoyourbodyonadailybasis": ("Assessment_Lifestyle", "Nourishment"),
    "movementratehowoftenandforhowlongyoumoveyourbodyonadailybasis": ("Assessment_Lifestyle", "Movement"),
    "connectednessratehowwellyoustayconnectedwithfamilyfriendsandyourhigherpower": ("Assessment_Lifestyle", "Connectedness"),
    "sleepratethequalityofyoursleep": ("Assessment_Lifestyle", "SleepQuality"),
    "happyselfratehowoftenandforhowlongyouperformpositivepracticesgratitudevirtueawarenessmeditationprayeretc": ("Assessment_Lifestyle", "HappySelf"),
    "resilienceratehowwellyouareabletomanagestressinyourlife": ("Assessment_Lifestyle", "Resilience"),
    "greenandblueratehowoftenandhowlongyouspendinnatureoroutdoors": ("Assessment_Lifestyle", "GreenBlueSpace"),
    "screentimeratehowhappyyouarewithyourcurrentamountofscreentime": ("Assessment_Lifestyle", "ScreenTime"),
    "substanceuseratehowcomfortableyouarewithanycurrentsubstanceusesmokingalcoholdrugs": ("Assessment_Lifestyle", "SubstanceUse"),
    "purposeratehowwellyoufeelyouarefulfillingyourpassionpurposeorvocationinlife": ("Assessment_Lifestyle", "Purpose"),
    "comments3": ("Assessment_Lifestyle", "LifestyleComments"),
    "howoftendoyoufeelthatyoulackcompanionship": ("Assessment_SocialIsolation", "LackCompanionship"),
    "howoftendoyoufeelleftout": ("Assessment_SocialIsolation", "FeelLeftOut"),
    "howoftendoyoufeelisolatedfromothers": ("Assessment_SocialIsolation", "FeelIsolated"),
    "socialisolation": ("Assessment_SocialIsolation", "SocialIsolationComments"),
    "howconfidentareyoutojoinactivities": ("Assessment_CommunityConfidence", "ConfidenceToJoin"),
    "howmanyhobbiesandpassionsdoyouhave": ("Assessment_CommunityConfidence", "NumberOfHobbies"),
    "howinvolvedyoufeelinyourcommunity": ("Assessment_CommunityConfidence", "CommunityInvolvement"),
    "howmuchyouknowaboutlocalsupportservices": ("Assessment_CommunityConfidence", "ServiceAwareness"),
    "whatareyouraimsgoals": ("Assessment_AimsGoals", "AimsGoals"),
    "comments5": ("Assessment_AimsGoals", "AimsDescription"),
    "whatreasonsstopyoufromjoiningactivities": ("Assessment_Barriers", "Barriers"),
    "comments6": ("Assessment_Barriers", "BarrierComments"),
    "whatareyourpreferredactivities": ("Assessment_PreferredActivities", "PreferredActivities"),
    "comments7": ("Assessment_PreferredActivities", "PreferredActivityComments"),
    "dateofnextreviewappointment": ("Assessment_PreferredActivities", "PreferredNextReviewDate"),
}

ASSESSMENT_TABLES = [
    "Assessment_Master", "Assessments", "Assessment_BodyComposition",
    "Assessment_HealthScreening", "Assessment_PhysicalActivity", "Assessment_WEMWBS",
    "Assessment_Lifestyle", "Assessment_SocialIsolation", "Assessment_CommunityConfidence",
    "Assessment_AimsGoals", "Assessment_Barriers", "Assessment_PreferredActivities",
]


def clean_key(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).strip().lower().replace("�", "'").replace("’", "'")
    return re.sub(r"[^a-z0-9]+", "", s)


def clean_card(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, (int, float, Decimal)) and float(value).is_integer():
        return str(int(value))
    s = str(value).strip()
    return re.sub(r"\.0$", "", s)


def norm_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def norm_name(value: Any) -> str:
    return norm_text(value)


def norm_date(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s[:10], fmt).date().isoformat()
        except ValueError:
            pass
    return norm_text(value)


def norm_phone(value: Any) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    if digits.startswith("0044"):
        return "0" + digits[4:]
    if digits.startswith("44") and len(digits) >= 11:
        return "0" + digits[2:]
    return digits


def norm_postcode(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).upper()


def norm_email(value: Any) -> str:
    return str(value or "").strip().casefold()


def norm_bool(value: Any) -> str:
    if value is None or str(value).strip() == "":
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, (int, float, Decimal)) and value in (0, 1):
        return "yes" if value == 1 else "no"
    s = norm_text(value)
    if s in {"yes", "y", "true", "1", "same", "yes same", "yes - same"}:
        return "yes"
    if s in {"no", "n", "false", "0", "different", "no different", "no - different"}:
        return "no"
    return s


def norm_number(value: Any) -> str:
    if value is None or str(value).strip() == "":
        return ""
    try:
        n = Decimal(str(value).strip().replace(",", ""))
        return format(n.normalize(), "f")
    except (InvalidOperation, ValueError):
        return norm_text(value)


def comparable(field: str, value: Any) -> str:
    if field in DATE_FIELDS or "date" in field.casefold():
        return norm_date(value)
    if field == "FullName":
        return norm_name(value)
    if field == "MobileNumber":
        return norm_phone(value)
    if field == "Postcode":
        return norm_postcode(value)
    if field == "Email":
        return norm_email(value)
    if field in BOOL_FIELDS:
        return norm_bool(value)
    if isinstance(value, (int, float, Decimal)):
        return norm_number(value)
    return norm_text(value)


def display(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    return value


def ordinal_pattern() -> re.Pattern[str]:
    return re.compile(r"^(\d+)(?:st|nd|rd|th)\s+Assessment\s+(.*)$", re.IGNORECASE)


@dataclass
class ExcelData:
    participants: dict[str, dict[str, Any]]
    duplicate_cards: dict[str, list[dict[str, Any]]]
    assessments: dict[tuple[str, int], dict[str, Any]]
    structure: list[dict[str, Any]]
    unmapped: list[dict[str, Any]]
    sheet_name: str
    reported_rows: int
    reported_cols: int
    meaningful_last_col: int
    max_populated_assessment: int


def progress(current: int, total: int, seen: set[int]) -> None:
    pct = int(current * 100 / max(total, 1))
    milestone = min(100, (pct // 25) * 25)
    if milestone and milestone not in seen:
        seen.add(milestone)
        print(f"  {milestone}%", flush=True)


def read_excel(path: Path) -> ExcelData:
    print("Loading workbook...", flush=True)
    started = time.perf_counter()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        candidates = []
        for ws in wb.worksheets:
            first = next(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True), ())
            keys = {clean_key(v) for v in first}
            if "sahelicardnumber" in keys and "fullname" in keys:
                candidates.append(ws)
        if not candidates:
            raise RuntimeError("No horizontal participant sheet with Saheli Card Number and Full Name was found.")
        ws = candidates[0]
        print(f"Sheet: {ws.title}", flush=True)
        print("Detected participant layout: horizontal (one participant per row)", flush=True)
        print(f"Reported range: {ws.calculate_dimension()} ({ws.max_row} rows x {ws.max_column} columns)", flush=True)

        rows = ws.iter_rows(values_only=True)
        headers = list(next(rows))
        # One pass over rows. Values beyond the last genuinely populated column
        # are never revisited; this avoids the previous repeated 952-column scans.
        raw_rows: list[tuple[int, tuple[Any, ...]]] = []
        col_counts = [0] * len(headers)
        seen_progress: set[int] = set()
        for row_no, row in enumerate(rows, 2):
            raw_rows.append((row_no, row))
            for i, value in enumerate(row):
                if value is not None and (not isinstance(value, str) or value.strip()):
                    col_counts[i] += 1
            progress(row_no - 1, ws.max_row, seen_progress)

        populated_indices = [i for i, count in enumerate(col_counts) if count]
        meaningful_last = max(populated_indices) + 1 if populated_indices else len(headers)
        print(f"Meaningful data columns: 1-{meaningful_last}; ignoring {len(headers)-meaningful_last} header/formatted-only columns.", flush=True)

        registration_cols: dict[str, int] = {}
        assessment_cols: dict[int, list[tuple[int, str, str, str]]] = defaultdict(list)
        structure: list[dict[str, Any]] = []
        unmapped: list[dict[str, Any]] = []
        ap = ordinal_pattern()
        max_populated_assessment = 0
        for i, header in enumerate(headers[:meaningful_last]):
            h = str(header or "").strip()
            key = clean_key(h)
            match = ap.match(h)
            mapping = ""
            assessment_number = ""
            data_count = col_counts[i]
            if match:
                assessment_number = int(match.group(1))
                label = match.group(2).strip()
                label_key = clean_key(label)
                target = ASSESSMENT_MAP.get(label_key)
                if target:
                    assessment_cols[assessment_number].append((i, label, target[0], target[1]))
                    mapping = f"dbo.{target[0]}.{target[1]}"
                elif label_key == "bloodpressuresystolicdiastolic":
                    assessment_cols[assessment_number].append((i, label, "SPECIAL", "BloodPressure"))
                    mapping = "dbo.Assessment_HealthScreening.SystolicBp + DiastolicBp"
                else:
                    unmapped.append({"Area": "Assessment", "ExcelHeader": h, "Reason": "No reliable SQL mapping"})
                if data_count:
                    max_populated_assessment = max(max_populated_assessment, int(assessment_number))
            elif key in REGISTRATION_MAP:
                registration_cols[REGISTRATION_MAP[key]] = i
                mapping = f"dbo.Participants.{REGISTRATION_MAP[key]}"
            elif i < 34:
                unmapped.append({"Area": "Registration", "ExcelHeader": h, "Reason": "No direct dbo.Participants column or intentionally excluded consent/emergency field"})
            structure.append({
                "Sheet": ws.title, "ColumnNumber": i + 1, "ExcelHeader": h,
                "NonBlankDataCells": data_count, "AssessmentNumber": assessment_number,
                "SQLMapping": mapping,
            })

        card_col = registration_cols.get("SaheliCardNumber")
        if card_col is None:
            raise RuntimeError("Saheli Card Number column could not be mapped.")
        print("Reading registration fields...", flush=True)
        participants: dict[str, dict[str, Any]] = {}
        duplicate_cards: dict[str, list[dict[str, Any]]] = defaultdict(list)
        assessments: dict[tuple[str, int], dict[str, Any]] = {}
        for row_no, row in raw_rows:
            card = clean_card(row[card_col] if card_col < len(row) else None)
            if not card:
                continue
            participant = {field: row[idx] if idx < len(row) else None for field, idx in registration_cols.items()}
            participant.update({"SaheliCardNumber": card, "_sheet": ws.title, "_row": row_no})
            if card in participants:
                if not duplicate_cards[card]:
                    duplicate_cards[card].append(participants[card])
                duplicate_cards[card].append(participant)
            else:
                participants[card] = participant
            for number, columns in assessment_cols.items():
                record = {"SaheliCardNumber": card, "AssessmentNumber": number, "_row": row_no}
                has_value = False
                for idx, _label, table, column in columns:
                    value = row[idx] if idx < len(row) else None
                    if value is not None and (not isinstance(value, str) or value.strip()):
                        has_value = True
                    if table == "SPECIAL" and column == "BloodPressure":
                        text = str(value or "")
                        parts = re.findall(r"\d+(?:\.\d+)?", text)
                        record["SystolicBp"] = parts[0] if parts else None
                        record["DiastolicBp"] = parts[1] if len(parts) > 1 else None
                    else:
                        record[column] = value
                if has_value:
                    assessments[(card, number)] = record

        print(f"Participants found: {len(participants):,}", flush=True)
        print(f"Excel assessments found: {len(assessments):,}; maximum populated assessment number: {max_populated_assessment}", flush=True)
        print(f"Workbook read in {time.perf_counter()-started:.1f}s", flush=True)
        return ExcelData(participants, dict(duplicate_cards), assessments, structure, unmapped,
                         ws.title, ws.max_row, ws.max_column, meaningful_last, max_populated_assessment)
    finally:
        wb.close()


def load_sql_config(path: Path) -> dict[str, str]:
    tree = ast.parse(path.read_text(encoding="utf-8-sig"), filename=str(path))
    wanted = {"SQL_SERVER", "SQL_DATABASE", "SQL_USERNAME", "SQL_PASSWORD"}
    values: dict[str, str] = {}
    for node in tree.body:
        if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name):
            name = node.targets[0].id
            if name in wanted:
                values[name] = str(ast.literal_eval(node.value))
    missing = sorted(wanted - values.keys())
    if missing:
        raise RuntimeError(f"Missing SQL config values in {path.name}: {', '.join(missing)}")
    return values


class ReadOnlySql:
    def __init__(self, config: dict[str, str]):
        server = config["SQL_SERVER"]
        if server.casefold().startswith("tcp:"):
            server = server[4:]
        if "," not in server:
            server += ",1433"
        cs = (
            "DRIVER={ODBC Driver 18 for SQL Server};"
            f"SERVER=tcp:{server};DATABASE={config['SQL_DATABASE']};"
            f"UID={config['SQL_USERNAME']};PWD={config['SQL_PASSWORD']};"
            "Encrypt=yes;TrustServerCertificate=no;Connection Timeout=30;"
            "ApplicationIntent=ReadOnly;"
        )
        self.connection = pyodbc.connect(cs, autocommit=False)

    def execute(self, statement: str, params: Iterable[Any] = ()):  # pyodbc cursor
        stripped = re.sub(r"/\*.*?\*/|--[^\r\n]*", " ", statement, flags=re.DOTALL).strip()
        if FORBIDDEN_SQL.search(stripped) or not (stripped.upper().startswith("SELECT") or stripped.upper().startswith("WITH")):
            raise RuntimeError("Safety guard rejected a non-read-only SQL statement.")
        return self.connection.cursor().execute(statement, tuple(params))

    def close(self) -> None:
        self.connection.rollback()
        self.connection.close()


def rows_as_dicts(cursor) -> list[dict[str, Any]]:
    columns = [d[0] for d in cursor.description]
    return [dict(zip(columns, row)) for row in cursor.fetchall()]


def fetch_schema(db: ReadOnlySql) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    columns = rows_as_dicts(db.execute("""
        SELECT s.name AS SchemaName,t.name AS TableName,c.column_id AS ColumnOrder,
               c.name AS ColumnName,ty.name AS DataType,c.max_length AS MaxLength,
               c.precision AS PrecisionValue,c.scale AS ScaleValue,c.is_nullable AS IsNullable
        FROM sys.tables t JOIN sys.schemas s ON s.schema_id=t.schema_id
        JOIN sys.columns c ON c.object_id=t.object_id
        JOIN sys.types ty ON ty.user_type_id=c.user_type_id
        WHERE c.name LIKE '%Participant%' OR c.name LIKE '%SaheliCard%'
           OR t.name LIKE 'Assessment[_]%' OR t.name='Assessments' OR t.name='Participants'
        ORDER BY s.name,t.name,c.column_id
    """))
    fks = rows_as_dicts(db.execute("""
        SELECT sch.name AS SchemaName,t.name AS TableName,c.name AS ColumnName,
               schp.name AS ReferencedSchema,tp.name AS ReferencedTable,cp.name AS ReferencedColumn,
               fk.name AS ConstraintName
        FROM sys.foreign_key_columns fkc JOIN sys.foreign_keys fk ON fk.object_id=fkc.constraint_object_id
        JOIN sys.tables t ON t.object_id=fkc.parent_object_id JOIN sys.schemas sch ON sch.schema_id=t.schema_id
        JOIN sys.columns c ON c.object_id=t.object_id AND c.column_id=fkc.parent_column_id
        JOIN sys.tables tp ON tp.object_id=fkc.referenced_object_id JOIN sys.schemas schp ON schp.schema_id=tp.schema_id
        JOIN sys.columns cp ON cp.object_id=tp.object_id AND cp.column_id=fkc.referenced_column_id
        WHERE tp.name='Participants' OR t.name LIKE 'Assessment[_]%' OR t.name='Assessments'
        ORDER BY sch.name,t.name
    """))
    fk_lookup = {(x["SchemaName"], x["TableName"], x["ColumnName"]): x for x in fks}
    for row in columns:
        fk = fk_lookup.get((row["SchemaName"], row["TableName"], row["ColumnName"]))
        row["ForeignKeyTarget"] = "" if not fk else f"{fk['ReferencedSchema']}.{fk['ReferencedTable']}.{fk['ReferencedColumn']}"
        row["ForeignKeyName"] = "" if not fk else fk["ConstraintName"]
    return columns, fks


def fetch_participants(db: ReadOnlySql) -> list[dict[str, Any]]:
    quoted = ",".join(f"[{x}]" for x in PARTICIPANT_COLUMNS)
    return rows_as_dicts(db.execute(f"SELECT {quoted} FROM dbo.Participants"))


def fetch_assessments(db: ReadOnlySql) -> list[dict[str, Any]]:
    return rows_as_dicts(db.execute("""
        SELECT m.AssessmentID,m.SaheliCardNumber,m.AssessmentNumber,m.AssessmentDate,
          a.StaffMember,a.Site,a.RiskStratificationScore,a.NextReviewDate,
          b.WeightKg,b.HeightCm,b.Bmicategory,b.Bmivalue,b.WaistCm,b.HipCm,b.WaistHipRatio,
          b.BodyFatCategory,b.BodyFatScore,b.VisceralFatCategory,b.VisceralFatScore,
          b.SkeletalMuscleCategory,b.SkeletalMuscleScore,b.RestingMetabolism,
          h.HasHealthCondition,h.LastBpmeasurementDate,h.BprecordedWithGp,h.KnowledgeHealthyBp,
          h.KnowledgeBprisk,h.KnowledgeBpreduction,h.SystolicBp,h.DiastolicBp,h.Bplevel,
          h.HeartConditionTypes,h.HeartRateBpm,h.AtrialFibrillationResult,h.HeartAge,
          h.DoctorAdvisedNoExercise,h.ChestPain,h.ShortnessOfBreath,h.DiabetesType,h.DiabetesRisk,
          h.GlucoseLevel,h.HbA1c,h.SugaryDrinkIntake,h.HighCholesterol,h.OtherHealthIssues,
          h.BoneJointConditions,h.TakesPrescribedMedication,h.ReferredToDoctor,h.RiskStratification,
          h.HealthComments,h.SelfManagementScore,
          p.ActiveDaysPerWeek,p.ActivityLevel,p.ActivityComments,
          w.FeelingOptimistic,w.FeelingUseful,w.FeelingRelaxed,w.FeelingInterestedInPeople,
          w.EnergyToSpare,w.DealingWithProblems,w.ThinkingClearly,w.FeelingGoodAboutSelf,
          w.FeelingCloseToOthers,w.FeelingConfident,w.MakingOwnMindUp,w.FeelingLoved,
          w.InterestedInNewThings,w.FeelingCheerful,w.Wemwbscomments,
          l.Nourishment,l.Movement,l.Connectedness,l.SleepQuality,l.HappySelf,l.Resilience,
          l.GreenBlueSpace,l.ScreenTime,l.SubstanceUse,l.Purpose,l.LifestyleComments,
          si.LackCompanionship,si.FeelLeftOut,si.FeelIsolated,si.SocialIsolationComments,
          cc.ConfidenceToJoin,cc.NumberOfHobbies,cc.CommunityInvolvement,cc.ServiceAwareness,
          ag.AimsGoals,ag.AimsDescription,br.Barriers,br.BarrierComments,
          pa.PreferredActivities,pa.ActivityComments AS PreferredActivityComments,pa.NextReviewDate AS PreferredNextReviewDate
        FROM dbo.Assessment_Master m
        LEFT JOIN dbo.Assessments a ON a.SaheliCardNumber=m.SaheliCardNumber AND a.AssessmentNumber=m.AssessmentNumber
        LEFT JOIN dbo.Assessment_BodyComposition b ON b.SaheliCardNumber=m.SaheliCardNumber AND b.AssessmentNumber=m.AssessmentNumber
        LEFT JOIN dbo.Assessment_HealthScreening h ON h.SaheliCardNumber=m.SaheliCardNumber AND h.AssessmentNumber=m.AssessmentNumber
        LEFT JOIN dbo.Assessment_PhysicalActivity p ON p.SaheliCardNumber=m.SaheliCardNumber AND p.AssessmentNumber=m.AssessmentNumber
        LEFT JOIN dbo.Assessment_WEMWBS w ON w.SaheliCardNumber=m.SaheliCardNumber AND w.AssessmentNumber=m.AssessmentNumber
        LEFT JOIN dbo.Assessment_Lifestyle l ON l.SaheliCardNumber=m.SaheliCardNumber AND l.AssessmentNumber=m.AssessmentNumber
        LEFT JOIN dbo.Assessment_SocialIsolation si ON si.SaheliCardNumber=m.SaheliCardNumber AND si.AssessmentNumber=m.AssessmentNumber
        LEFT JOIN dbo.Assessment_CommunityConfidence cc ON cc.SaheliCardNumber=m.SaheliCardNumber AND cc.AssessmentNumber=m.AssessmentNumber
        LEFT JOIN dbo.Assessment_AimsGoals ag ON ag.SaheliCardNumber=m.SaheliCardNumber AND ag.AssessmentNumber=m.AssessmentNumber
        LEFT JOIN dbo.Assessment_Barriers br ON br.SaheliCardNumber=m.SaheliCardNumber AND br.AssessmentNumber=m.AssessmentNumber
        LEFT JOIN dbo.Assessment_PreferredActivities pa ON pa.SaheliCardNumber=m.SaheliCardNumber AND pa.AssessmentNumber=m.AssessmentNumber
    """))


def group_duplicates(records: Iterable[dict[str, Any]], source: str) -> list[dict[str, Any]]:
    definitions = [
        ("NAME+DOB", lambda r: (norm_name(r.get("FullName")), norm_date(r.get("DateOfBirth")))),
        ("NAME+DOB+MOBILE", lambda r: (norm_name(r.get("FullName")), norm_date(r.get("DateOfBirth")), norm_phone(r.get("MobileNumber")))),
        ("MOBILE", lambda r: (norm_phone(r.get("MobileNumber")),)),
        ("POSTCODE+DOB", lambda r: (norm_postcode(r.get("Postcode")), norm_date(r.get("DateOfBirth")))),
    ]
    output = []
    for rule, key_func in definitions:
        groups: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
        for record in records:
            key = key_func(record)
            if all(key):
                groups[key].append(record)
        number = 0
        for key, members in sorted(groups.items()):
            if len(members) < 2:
                continue
            number += 1
            cards = ", ".join(sorted((clean_card(m.get("SaheliCardNumber")) for m in members), key=lambda x: (len(x), x)))
            for member in members:
                output.append({
                    "Source": source, "Rule": rule, "GroupNumber": number,
                    "GroupKey": " | ".join(key), "GroupSize": len(members), "AllCardNumbers": cards,
                    "ParticipantID": member.get("ParticipantID", ""),
                    "SaheliCardNumber": clean_card(member.get("SaheliCardNumber")),
                    "FullName": member.get("FullName", ""), "DateOfBirth": display(member.get("DateOfBirth")),
                    "MobileNumber": member.get("MobileNumber", ""), "Postcode": member.get("Postcode", ""),
                })
    return output


def excel_duplicate_card_rows(excel: ExcelData) -> list[dict[str, Any]]:
    output = []
    group_number = 0
    for card, members in sorted(excel.duplicate_cards.items(), key=lambda x: (len(x[0]), x[0])):
        group_number += 1
        for member in members:
            output.append({
                "Source": "Excel", "Rule": "CARD", "GroupNumber": group_number,
                "GroupKey": card, "GroupSize": len(members), "AllCardNumbers": card,
                "ParticipantID": "", "SaheliCardNumber": card,
                "FullName": member.get("FullName", ""), "DateOfBirth": display(member.get("DateOfBirth")),
                "MobileNumber": member.get("MobileNumber", ""), "Postcode": member.get("Postcode", ""),
                "ExcelSourceRow": member.get("_row", ""),
            })
    return output


def identity_candidates(sql_record: dict[str, Any], excel_by_card: dict[str, dict[str, Any]], own_card: str) -> list[dict[str, Any]]:
    sn, sd, sm, sp = (norm_name(sql_record.get("FullName")), norm_date(sql_record.get("DateOfBirth")),
                      norm_phone(sql_record.get("MobileNumber")), norm_postcode(sql_record.get("Postcode")))
    candidates = []
    for card, e in excel_by_card.items():
        if card == own_card:
            continue
        en, ed, em, ep = (norm_name(e.get("FullName")), norm_date(e.get("DateOfBirth")),
                          norm_phone(e.get("MobileNumber")), norm_postcode(e.get("Postcode")))
        matched = [label for label, a, b in (("Name", sn, en), ("DOB", sd, ed), ("Mobile", sm, em), ("Postcode", sp, ep)) if a and a == b]
        confidence = ""
        if sn and sd and sm and (sn, sd, sm) == (en, ed, em):
            confidence = "STRONG"
        elif sn and sd and (sn, sd) == (en, ed):
            confidence = "MEDIUM"
        elif len(matched) >= 3:
            confidence = "MEDIUM"
        if confidence:
            candidates.append({"card": card, "record": e, "confidence": confidence, "matched": ", ".join(matched)})
    return candidates


def compare_participants(excel: ExcelData, sql_rows: list[dict[str, Any]]):
    sql_by_card = {clean_card(r["SaheliCardNumber"]): r for r in sql_rows}
    all_cards = sorted(set(excel.participants) | set(sql_by_card), key=lambda x: (len(x), x))
    comparison, differences, conflicts, wrong = [], [], [], []
    suspicious_cards: set[str] = set()
    fields = [x for x in PARTICIPANT_COLUMNS if x not in {"ParticipantID", "SaheliCardNumber", "CreatedAt", "GPSurgeryId"} and x in set(REGISTRATION_MAP.values()) | {"HasHealthConditionOrDisability", "HealthConditionDetails"}]
    for card in all_cards:
        e, s = excel.participants.get(card), sql_by_card.get(card)
        if e is None:
            classification, mismatches, core = "SQL ONLY", [], []
        elif s is None:
            classification, mismatches, core = "EXCEL ONLY", [], []
        else:
            mismatches = [f for f in fields if comparable(f, e.get(f)) != comparable(f, s.get(f))]
            core = [f for f in CORE_FIELDS if comparable(f, e.get(f)) != comparable(f, s.get(f))]
            candidates = identity_candidates(s, excel.participants, card) if core else []
            if candidates:
                classification = "POSSIBLE WRONG IDENTITY"
                suspicious_cards.add(card)
                for c in candidates:
                    wrong.append({
                        "SQLCard": card, "ParticipantID": s.get("ParticipantID"),
                        "ExpectedExcelName": e.get("FullName"), "CurrentSQLName": s.get("FullName"),
                        "CurrentSQLDOB": display(s.get("DateOfBirth")), "CurrentSQLMobile": s.get("MobileNumber"),
                        "MatchesExcelCard": c["card"], "MatchedExcelName": c["record"].get("FullName"),
                        "MatchedFields": c["matched"], "Confidence": c["confidence"],
                        "Explanation": f"SQL card {card} identity matches Excel card {c['card']} ({c['confidence']}). No repair performed.",
                    })
            elif core:
                classification = "CORE IDENTITY DIFFERENCE"
                suspicious_cards.add(card)
            elif mismatches:
                classification = "PROFILE DIFFERENCE"
            else:
                classification = "MATCH"
            for field in mismatches:
                differences.append({
                    "SaheliCardNumber": card, "ParticipantID": s.get("ParticipantID"),
                    "FullName": s.get("FullName"), "Field": field,
                    "Severity": "CORE IDENTITY" if field in CORE_FIELDS else "PROFILE",
                    "ExcelValue": display(e.get(field)), "SQLValue": display(s.get(field)),
                    "ExcelNormalised": comparable(field, e.get(field)), "SQLNormalised": comparable(field, s.get(field)),
                })
            if core:
                conflicts.append({
                    "SaheliCardNumber": card, "ParticipantID": s.get("ParticipantID"),
                    "ExcelFullName": e.get("FullName"), "SQLFullName": s.get("FullName"),
                    "ExcelDOB": display(e.get("DateOfBirth")), "SQLDOB": display(s.get("DateOfBirth")),
                    "ExcelMobile": e.get("MobileNumber"), "SQLMobile": s.get("MobileNumber"),
                    "ExcelPostcode": e.get("Postcode"), "SQLPostcode": s.get("Postcode"),
                    "CoreDifferenceFields": ", ".join(core), "Classification": classification,
                })
        comparison.append({
            "SaheliCardNumber": card, "ParticipantID": "" if s is None else s.get("ParticipantID"),
            "Classification": classification, "ExcelFullName": "" if e is None else e.get("FullName"),
            "SQLFullName": "" if s is None else s.get("FullName"),
            "ExcelDOB": "" if e is None else display(e.get("DateOfBirth")),
            "SQLDOB": "" if s is None else display(s.get("DateOfBirth")),
            "ExcelMobile": "" if e is None else e.get("MobileNumber"),
            "SQLMobile": "" if s is None else s.get("MobileNumber"),
            "ExcelPostcode": "" if e is None else e.get("Postcode"),
            "SQLPostcode": "" if s is None else s.get("Postcode"),
            "MismatchCount": len(mismatches), "MismatchFields": ", ".join(mismatches),
            "ExcelSource": "" if e is None else f"{e['_sheet']} row {e['_row']}",
        })
    return comparison, differences, conflicts, wrong, suspicious_cards, sql_by_card


def assessment_fingerprint(record: dict[str, Any]) -> tuple[str, ...] | None:
    fields = ("AssessmentDate", "WeightKg", "HeightCm", "WaistCm", "HipCm", "HeartRateBpm", "HbA1c", "GlucoseLevel")
    values = tuple(comparable(field, record.get(field)) for field in fields)
    if sum(bool(x) for x in values) < 3:
        return None
    return values


def compare_assessments(excel: ExcelData, sql_rows: list[dict[str, Any]]):
    sql_by_key = {(clean_card(r["SaheliCardNumber"]), int(r["AssessmentNumber"])): r for r in sql_rows}
    all_keys = sorted(set(excel.assessments) | set(sql_by_key), key=lambda x: (len(x[0]), x[0], x[1]))
    fingerprint_index: dict[tuple[str, ...], list[tuple[str, int]]] = defaultdict(list)
    for key, record in excel.assessments.items():
        fp = assessment_fingerprint(record)
        if fp:
            fingerprint_index[fp].append(key)
    summary, differences, wrong = [], [], []
    suspicious_cards: set[str] = set()
    excluded = {"SaheliCardNumber", "AssessmentNumber", "_row", "AssessmentID", "StaffMember", "Site", "RiskStratificationScore", "NextReviewDate", "PreferredActivityComments", "PreferredNextReviewDate"}
    for key in all_keys:
        card, number = key
        e, s = excel.assessments.get(key), sql_by_key.get(key)
        if e is None:
            classification, mismatch_fields = "SQL ASSESSMENT MISSING FROM EXCEL", []
        elif s is None:
            classification, mismatch_fields = "EXCEL ASSESSMENT MISSING FROM SQL", []
        else:
            fields = sorted((set(e) & set(s)) - excluded)
            mismatch_fields = [f for f in fields if comparable(f, e.get(f)) != comparable(f, s.get(f))]
            classification = "FIELD DIFFERENCE" if mismatch_fields else "MATCH"
            for field in mismatch_fields:
                differences.append({
                    "SaheliCardNumber": card, "AssessmentNumber": number,
                    "ExcelAssessmentDate": display(e.get("AssessmentDate")), "SQLAssessmentDate": display(s.get("AssessmentDate")),
                    "Field": field, "ExcelValue": display(e.get(field)), "SQLValue": display(s.get(field)),
                    "ExcelNormalised": comparable(field, e.get(field)), "SQLNormalised": comparable(field, s.get(field)),
                })
            fp = assessment_fingerprint(s)
            if mismatch_fields and fp:
                alternatives = [k for k in fingerprint_index.get(fp, []) if k[0] != card]
                for other_card, other_number in alternatives:
                    suspicious_cards.update((card, other_card))
                    wrong.append({
                        "SQLCard": card, "SQLAssessmentNumber": number,
                        "SQLAssessmentDate": display(s.get("AssessmentDate")),
                        "MatchesExcelCard": other_card, "MatchesExcelAssessmentNumber": other_number,
                        "Confidence": "STRONG", "MatchedFingerprintFields": "Assessment date + at least 2 measured values",
                        "Explanation": "Exact multi-field assessment fingerprint matches another Excel participant. No reassociation performed.",
                    })
        summary.append({
            "SaheliCardNumber": card, "AssessmentNumber": number, "Classification": classification,
            "ExcelAssessmentDate": "" if e is None else display(e.get("AssessmentDate")),
            "SQLAssessmentDate": "" if s is None else display(s.get("AssessmentDate")),
            "FieldDifferenceCount": len(mismatch_fields), "DifferenceFields": ", ".join(mismatch_fields),
        })
    return summary, differences, wrong, suspicious_cards


def discover_linked_columns(schema_rows: list[dict[str, Any]]) -> list[tuple[str, str, str]]:
    found = set()
    for row in schema_rows:
        col = row["ColumnName"]
        if col.casefold() in {"participantid", "recipientparticipantid", "participant_id"}:
            found.add((row["SchemaName"], row["TableName"], col))
        elif col.casefold() in {"sahelicardnumber", "recipientsahelicardnumber"}:
            found.add((row["SchemaName"], row["TableName"], col))
    return sorted(found)


def quote_identifier(value: str) -> str:
    return "[" + value.replace("]", "]]" ) + "]"


def fetch_linked_counts(db: ReadOnlySql, schema_rows, suspicious_cards: set[str], sql_by_card) -> list[dict[str, Any]]:
    if not suspicious_cards:
        return []
    output = []
    participants_by_id = {int(r["ParticipantID"]): card for card, r in sql_by_card.items()}
    for schema, table, column in discover_linked_columns(schema_rows):
        if table == "Participants":
            continue
        is_id = "participantid" in column.casefold() and "sahelicard" not in column.casefold()
        if is_id:
            ids = [sql_by_card[c]["ParticipantID"] for c in suspicious_cards if c in sql_by_card]
            values = ids
        else:
            values = sorted(suspicious_cards)
        if not values:
            continue
        placeholders = ",".join("?" for _ in values)
        qtable = f"{quote_identifier(schema)}.{quote_identifier(table)}"
        qcolumn = quote_identifier(column)
        query = f"SELECT {qcolumn} AS LinkValue, COUNT_BIG(*) AS LinkedRowCount FROM {qtable} WHERE {qcolumn} IN ({placeholders}) GROUP BY {qcolumn}"
        for row in rows_as_dicts(db.execute(query, values)):
            link = row["LinkValue"]
            card = participants_by_id.get(int(link), "") if is_id and link is not None else clean_card(link)
            output.append({"SaheliCardNumber": card, "ParticipantID": sql_by_card.get(card, {}).get("ParticipantID", ""),
                           "Schema": schema, "Table": table, "LinkColumn": column, "LinkedRowCount": row["LinkedRowCount"]})
    return output


def assessment_counts(records: Iterable[dict[str, Any]]) -> dict[str, int]:
    result: dict[str, int] = defaultdict(int)
    for row in records:
        result[clean_card(row["SaheliCardNumber"])] += 1
    return dict(result)


def write_sheet(wb: Workbook, name: str, rows: list[dict[str, Any]], headers: list[str] | None = None) -> None:
    ws = wb.create_sheet(name)
    if not headers:
        headers = list(rows[0]) if rows else ["Message"]
    ws.append(headers)
    if rows:
        for row in rows:
            ws.append([display(row.get(h, "")) for h in headers])
    else:
        ws.append(["No rows"] + [""] * (len(headers) - 1))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34
    for col_idx, cells in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 500), max_col=ws.max_column), 1):
        width = min(55, max(10, max(len(str(c.value or "")) for c in cells) + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd"
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    if "Classification" in headers:
        col = get_column_letter(headers.index("Classification") + 1)
        end = max(2, ws.max_row)
        ws.conditional_formatting.add(f"{col}2:{col}{end}", FormulaRule(formula=[f'ISNUMBER(SEARCH("MATCH",{col}2))'], fill=PatternFill("solid", fgColor="E2F0D9")))
        ws.conditional_formatting.add(f"{col}2:{col}{end}", FormulaRule(formula=[f'ISNUMBER(SEARCH("DIFFERENCE",{col}2))'], fill=PatternFill("solid", fgColor="FFF2CC")))
        ws.conditional_formatting.add(f"{col}2:{col}{end}", FormulaRule(formula=[f'ISNUMBER(SEARCH("WRONG",{col}2))'], fill=PatternFill("solid", fgColor="F4CCCC")))


def create_report(path: Path, excel: ExcelData, sql_rows, schema_rows, fk_rows,
                  comparison, differences, conflicts, wrong_identity,
                  assessment_summary, assessment_differences, wrong_assessments,
                  linked_counts, sql_duplicates, excel_duplicates) -> None:
    counts = defaultdict(int)
    for r in comparison:
        counts[r["Classification"]] += 1
    assess_counts = defaultdict(int)
    for r in assessment_summary:
        assess_counts[r["Classification"]] += 1
    summary = [
        {"Metric": "Audit mode", "Count": "READ ONLY - no SQL changes"},
        {"Metric": "Excel participant rows", "Count": len(excel.participants) + sum(len(v) - 1 for v in excel.duplicate_cards.values())},
        {"Metric": "Excel unique card count", "Count": len(excel.participants)},
        {"Metric": "Excel repeated card groups", "Count": len(excel.duplicate_cards)},
        {"Metric": "SQL participant count", "Count": len(sql_rows)},
        {"Metric": "Cards in both", "Count": sum(1 for r in comparison if r["Classification"] not in {"EXCEL ONLY", "SQL ONLY"})},
        {"Metric": "Excel-only cards", "Count": counts["EXCEL ONLY"]},
        {"Metric": "SQL-only cards", "Count": counts["SQL ONLY"]},
        {"Metric": "Core identity conflicts", "Count": len(conflicts)},
        {"Metric": "Possible wrong identities", "Count": len(wrong_identity)},
        {"Metric": "SQL duplicate Name+DOB groups", "Count": len({r["GroupNumber"] for r in sql_duplicates if r["Rule"] == "NAME+DOB"})},
        {"Metric": "Excel duplicate Name+DOB groups", "Count": len({r["GroupNumber"] for r in excel_duplicates if r["Rule"] == "NAME+DOB"})},
        {"Metric": "Excel assessment count", "Count": len(excel.assessments)},
        {"Metric": "SQL assessment count", "Count": len(assessment_counts(sql_rows=[])) if False else sum(1 for r in assessment_summary if r["Classification"] != "EXCEL ASSESSMENT MISSING FROM SQL")},
        {"Metric": "Assessment field differences", "Count": len(assessment_differences)},
        {"Metric": "Excel assessments missing from SQL", "Count": assess_counts["EXCEL ASSESSMENT MISSING FROM SQL"]},
        {"Metric": "SQL assessments missing from Excel", "Count": assess_counts["SQL ASSESSMENT MISSING FROM EXCEL"]},
        {"Metric": "Possible wrong assessments", "Count": len(wrong_assessments)},
        {"Metric": "Excel reported range", "Count": f"{excel.reported_rows} rows x {excel.reported_cols} columns"},
        {"Metric": "Excel meaningful last data column", "Count": excel.meaningful_last_col},
        {"Metric": "Maximum populated Excel assessment number", "Count": excel.max_populated_assessment},
    ]
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "Summary", summary)
    write_sheet(wb, "ParticipantComparison", comparison)
    write_sheet(wb, "CoreIdentityConflicts", conflicts)
    write_sheet(wb, "PossibleWrongIdentity", wrong_identity)
    write_sheet(wb, "FieldDifferences", differences)
    write_sheet(wb, "ExcelOnlyParticipants", [r for r in comparison if r["Classification"] == "EXCEL ONLY"])
    write_sheet(wb, "SQLOnlyParticipants", [r for r in comparison if r["Classification"] == "SQL ONLY"])
    write_sheet(wb, "SQLDuplicateGroups", sql_duplicates)
    write_sheet(wb, "ExcelDuplicateGroups", excel_duplicates)
    write_sheet(wb, "AssessmentSummary", assessment_summary)
    write_sheet(wb, "AssessmentDifferences", assessment_differences)
    write_sheet(wb, "PossibleWrongAssessments", wrong_assessments)
    write_sheet(wb, "LinkedDataCounts", linked_counts)
    write_sheet(wb, "SQLSchemaMap", schema_rows)
    write_sheet(wb, "SQLForeignKeys", fk_rows)
    write_sheet(wb, "ExcelStructure", excel.structure)
    write_sheet(wb, "UnmappedFields", excel.unmapped)
    wb.properties.title = "Saheli Full Excel SQL Audit"
    wb.properties.subject = "Read-only local participant and health assessment comparison"
    wb.properties.description = "Generated locally. No SQL data was modified."
    wb.save(path)


def validate_report(path: Path) -> None:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        required = {"Summary", "ParticipantComparison", "CoreIdentityConflicts", "AssessmentSummary", "SQLSchemaMap", "ExcelStructure"}
        missing = required - set(wb.sheetnames)
        if missing:
            raise RuntimeError(f"Report validation failed; missing sheets: {sorted(missing)}")
        for ws in wb.worksheets:
            if ws.max_row < 2 or ws.max_column < 1:
                raise RuntimeError(f"Report validation failed; empty sheet: {ws.title}")
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "f" and isinstance(cell.value, str) and any(x in cell.value for x in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")):
                        raise RuntimeError(f"Formula error text found in {ws.title}!{cell.coordinate}")
    finally:
        wb.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Read-only Saheli Excel versus Azure SQL audit")
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL, help="Master Excel workbook")
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG, help="Existing Python file containing SQL config constants")
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).parent, help="Local report directory")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    print("=" * 72)
    print("SAHELI FULL EXCEL / SQL AUDIT - READ ONLY")
    print("=" * 72)
    if not args.excel.is_file():
        raise FileNotFoundError(args.excel)
    if not args.config.is_file():
        raise FileNotFoundError(args.config)
    started = time.perf_counter()
    excel = read_excel(args.excel)
    print("Connecting to Azure SQL with read-only application intent...", flush=True)
    db = ReadOnlySql(load_sql_config(args.config))
    try:
        print("Inspecting SQL metadata...", flush=True)
        schema_rows, fk_rows = fetch_schema(db)
        print("Reading dbo.Participants with SELECT...", flush=True)
        sql_participants = fetch_participants(db)
        print(f"SQL participants found: {len(sql_participants):,}", flush=True)
        print("Reading assessment parent/child tables with SELECT...", flush=True)
        sql_assessments = fetch_assessments(db)
        print(f"SQL assessments found: {len(sql_assessments):,}", flush=True)
        print("Comparing participant identities locally...", flush=True)
        comparison, differences, conflicts, wrong_identity, suspicious, sql_by_card = compare_participants(excel, sql_participants)
        print("Comparing health assessments locally...", flush=True)
        assessment_summary, assessment_differences, wrong_assessments, assessment_suspicious = compare_assessments(excel, sql_assessments)
        suspicious.update(assessment_suspicious)
        print("Calculating duplicate groups locally...", flush=True)
        sql_duplicates = group_duplicates(sql_participants, "SQL")
        excel_duplicates = excel_duplicate_card_rows(excel) + group_duplicates(excel.participants.values(), "Excel")
        print(f"Reading linked-data counts for {len(suspicious):,} suspicious cards...", flush=True)
        linked_counts = fetch_linked_counts(db, schema_rows, suspicious, sql_by_card)
    finally:
        db.close()
        print("SQL connection rolled back and closed; no SQL data modified.", flush=True)

    args.output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = args.output_dir / f"{OUTPUT_PREFIX}_{timestamp}.xlsx"
    print("Creating local audit workbook...", flush=True)
    create_report(output, excel, sql_participants, schema_rows, fk_rows,
                  comparison, differences, conflicts, wrong_identity,
                  assessment_summary, assessment_differences, wrong_assessments,
                  linked_counts, sql_duplicates, excel_duplicates)
    print("Validating report structure and error cells...", flush=True)
    validate_report(output)
    print("=" * 72)
    print("AUDIT COMPLETE")
    print(f"Report: {output}")
    print(f"Runtime: {time.perf_counter()-started:.1f} seconds")
    print(f"Core identity conflicts: {len(conflicts):,}")
    print(f"Possible wrong identities: {len(wrong_identity):,}")
    print(f"Assessment field differences: {len(assessment_differences):,}")
    print("READ ONLY confirmed: all SQL statements were SELECT/metadata queries.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("Audit cancelled; no SQL data was modified.", file=sys.stderr)
        raise SystemExit(130)
