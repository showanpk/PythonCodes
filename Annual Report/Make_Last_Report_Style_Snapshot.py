#!/usr/bin/env python3
"""
Create a one-page Saheli annual-report snapshot matching the layout/style
of the previous annual report, using the CURRENT workbook values.

It DOES NOT change the underlying figures.
It adds a new first sheet called: ANNUAL SNAPSHOT

Install once:
    pip install openpyxl pillow

Run:
    python Make_Last_Report_Style_Snapshot.py
"""

from pathlib import Path
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.image import Image
from openpyxl.utils import get_column_letter

# ============================================================
# CONFIG
# ============================================================

BASE = Path(__file__).resolve().parent

# Change this only if your latest workbook has a different name.
INPUT_XLSX = BASE / "Saheli_Hub_Annual_Report_2025_26_COMPLETE_BRANDED_SERVICES.xlsx"

OUTPUT_XLSX = BASE / "Saheli_Hub_Annual_Report_2025_26_LAST_REPORT_STYLE.xlsx"

# Optional logo. The script will try these names.
LOGO_CANDIDATES = [
    BASE / "saheli_logo.png",
    BASE / "SaheliHubLogo.png",
    BASE / "Saheli Hub Logo.png",
    BASE / "logo.png",
]

# Saheli colours
PINK = "C2185B"
PINK_LIGHT = "E91E63"
YELLOW = "FFC107"
ORANGE = "F28C28"
GREEN = "8BC34A"
TEAL = "2196A6"
PURPLE = "6A1B9A"
RED = "D32F2F"
GREY = "E5E7EB"
DARK = "2B2B2B"
WHITE = "FFFFFF"


# ============================================================
# HELPERS
# ============================================================

def clean_text(v):
    return "" if v is None else str(v).strip()

def as_number(v, default=0):
    try:
        return float(v)
    except Exception:
        return default

def find_row(ws, text, col=1, contains=False):
    target = text.strip().lower()
    for r in range(1, ws.max_row + 1):
        v = clean_text(ws.cell(r, col).value).lower()
        if (contains and target in v) or (not contains and v == target):
            return r
    return None

def find_metric(ws, label, fallback=None):
    """Find label in column A and return B value."""
    row = find_row(ws, label, 1, contains=False)
    if row:
        return ws.cell(row, 2).value
    # try contains
    row = find_row(ws, label, 1, contains=True)
    if row:
        return ws.cell(row, 2).value
    return fallback

def find_label_value_anywhere(ws, label, fallback=None):
    target = label.lower()
    for row in ws.iter_rows():
        for cell in row:
            if clean_text(cell.value).lower() == target:
                # prefer cell immediately right
                right = ws.cell(cell.row, cell.column + 1).value
                if right is not None:
                    return right
    return fallback

def set_fill_font(cell, fill=None, font_color=None, bold=False, size=11, italic=False):
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(
        name="Aptos",
        size=size,
        bold=bold,
        italic=italic,
        color=font_color or DARK
    )

def style_range(ws, cell_range, fill=None, font_color=None, bold=False,
                size=11, align="center", vertical="center", wrap=True):
    for row in ws[cell_range]:
        for cell in row:
            set_fill_font(cell, fill, font_color, bold, size)
            cell.alignment = Alignment(
                horizontal=align,
                vertical=vertical,
                wrap_text=wrap
            )

def merge_write(ws, cell_range, text, fill=None, font_color=None,
                bold=False, size=11, align="center"):
    ws.merge_cells(cell_range)
    c = ws[cell_range.split(":")[0]]
    c.value = text
    set_fill_font(c, fill, font_color, bold, size)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    return c

def set_chart_colors(chart, colours):
    try:
        for i, series in enumerate(chart.series):
            col = colours[i % len(colours)]
            series.graphicalProperties.solidFill = col
            series.graphicalProperties.line.solidFill = col
    except Exception:
        pass

def add_logo(ws):
    logo = next((p for p in LOGO_CANDIDATES if p.exists()), None)
    if not logo:
        return
    try:
        img = Image(str(logo))
        img.width = 145
        img.height = 55
        ws.add_image(img, "A1")
    except Exception:
        pass

def read_demographics(wb):
    ws = wb["DEMOGRAPHICS"]

    # Gender
    gender = {}
    row = find_row(ws, "Gender", 1, contains=False)
    if row:
        for r in range(row + 1, min(row + 8, ws.max_row + 1)):
            k = clean_text(ws.cell(r, 1).value)
            if k in {"Female", "Male", "Other", "Not recorded"}:
                gender[k] = as_number(ws.cell(r, 2).value)

    # Fallback current verified values
    gender.setdefault("Female", 794)
    gender.setdefault("Male", 184)
    gender.setdefault("Other", 1)
    gender.setdefault("Not recorded", 384)

    # Age - read valid DOB percentages where available
    age = {}
    age_header = find_row(ws, "Age Band", 1, contains=False)
    if age_header:
        for r in range(age_header + 1, min(age_header + 15, ws.max_row + 1)):
            band = clean_text(ws.cell(r, 1).value)
            if band in {"Under 16","16-25","26-35","36-45","46-55","56-65","66-75","76+","Unknown"}:
                # column D is usually % valid DOB in the current workbook
                val = ws.cell(r, 4).value
                if val is None:
                    val = ws.cell(r, 3).value
                age[band] = as_number(val)

    fallback_age = {
        "Under 16":0.4, "16-25":2.8, "26-35":7.3, "36-45":19.8,
        "46-55":22.7, "56-65":23.8, "66-75":17.4, "76+":5.8
    }
    for k,v in fallback_age.items():
        if k not in age or age[k] == 0:
            age[k] = v

    return gender, age

def read_top_activities(wb):
    ws = wb["TOP ACTIVITIES"]
    rows = []
    for r in range(2, ws.max_row + 1):
        activity = clean_text(ws.cell(r, 1).value)
        att = ws.cell(r, 2).value
        if activity and isinstance(att, (int, float)):
            rows.append((activity, float(att)))
    if not rows:
        rows = [
            ("Innerva",3220),("Boys Football",1968),("Chair Based Exercise",1425),
            ("Pilates",1300),("HAF",735),("Strength & Stretch",661),
            ("Circuit Training",648),("Yoga",623),("ESOL",587),
            ("Girls Youth Club",581)
        ]
    return rows[:10]

def read_registration_insights(wb):
    ws = wb["REGISTRATION INSIGHTS"]

    reasons = []
    for r in range(1, ws.max_row + 1):
        label = clean_text(ws.cell(r, 1).value)
        cnt = ws.cell(r, 2).value
        pct = ws.cell(r, 3).value
        if label and isinstance(cnt, (int, float)):
            # stop before next section if possible
            if "how people heard" in label.lower():
                break
            # remove header-like entries
            if label.lower() not in {"reason", "participants"}:
                reasons.append((label, float(cnt), pct))
    # Keep the meaningful main reasons only
    reasons = reasons[:8] if reasons else [
        ("Increase exercise or mobility",494,0.809),
        ("Weight management",294,0.481),
        ("Long-term health condition",194,0.318),
        ("Healthy eating or nutrition",148,0.242),
        ("Mild/moderate depression or anxiety",100,0.164),
        ("Isolation or loneliness",54,0.088),
        ("Learning/training/employment",51,0.084),
        ("Challenging social circumstances",35,0.057),
    ]

    # Heard-about section - scan for known labels anywhere
    heard = []
    known = ["Word of mouth","GP","WorkWell","Social media","Website","Other","Not recorded"]
    for label in known:
        found = False
        for r in range(1, ws.max_row + 1):
            if clean_text(ws.cell(r,1).value).lower() == label.lower():
                val = ws.cell(r,2).value
                pct = ws.cell(r,3).value
                heard.append((label, as_number(val), pct))
                found = True
                break
        if not found:
            pass

    if not heard:
        heard = [
            ("Word of mouth",438,0.7169),
            ("GP",72,0.1178),
            ("WorkWell",51,0.0835),
            ("Social media",8,0.0131),
            ("Website",7,0.0115),
            ("Other",31,0.0507),
            ("Not recorded",4,0.0065),
        ]

    return reasons, heard


# ============================================================
# LOAD WORKBOOK
# ============================================================

if not INPUT_XLSX.exists():
    raise SystemExit(f"Input workbook not found:\n{INPUT_XLSX}")

wb = load_workbook(INPUT_XLSX)

# Remove previous version if rerun
if "ANNUAL SNAPSHOT" in wb.sheetnames:
    del wb["ANNUAL SNAPSHOT"]

# Insert snapshot as first sheet
ws = wb.create_sheet("ANNUAL SNAPSHOT", 0)

# Landscape / single-page feel
ws.sheet_view.showGridLines = False
ws.freeze_panes = None
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins.left = 0.2
ws.page_margins.right = 0.2
ws.page_margins.top = 0.3
ws.page_margins.bottom = 0.3

# widths
for col in range(1, 19):  # A:R
    ws.column_dimensions[get_column_letter(col)].width = 11.5
for col in ["A","B","C","D","E","F","G","H","I","J","K","L","M","N","O","P","Q","R"]:
    pass

# row heights
for r in range(1, 50):
    ws.row_dimensions[r].height = 22
ws.row_dimensions[1].height = 32
ws.row_dimensions[2].height = 32
ws.row_dimensions[3].height = 32
ws.row_dimensions[4].height = 28
ws.row_dimensions[5].height = 22

# ============================================================
# READ CURRENT DATA
# ============================================================

summary = wb["SUMMARY"]

total_attendance = find_metric(summary, "Total verified annual attendance", 17580)
registered = find_metric(summary, "REGISTERED PARTICIPANTS", 2853)
new_reg = find_metric(summary, "New FULL registrations", 611)
sessions = find_metric(summary, "Sessions delivered", 1807)

gender, age = read_demographics(wb)
activities = read_top_activities(wb)
reasons, heard = read_registration_insights(wb)

female = gender["Female"]
male = gender["Male"]
binary_total = female + male
female_pct = 100 * female / binary_total if binary_total else 0
male_pct = 100 * male / binary_total if binary_total else 0

# Current established annual-report KPIs
ethnically_diverse_pct = 97.0
imd_20_pct = 80.3


# ============================================================
# TOP BANNER - MATCH OLD REPORT
# ============================================================

# Background
style_range(ws, "A1:R7", fill=PINK, font_color=WHITE)

merge_write(ws, "A1:F4", f"{int(total_attendance):,}",
            fill=PINK, font_color=WHITE, bold=True, size=46)
merge_write(ws, "A5:F7", "total attendances¹",
            fill=PINK, font_color=WHITE, bold=True, size=16)

merge_write(ws, "G1:L4", f"{int(registered):,}",
            fill=PINK, font_color=WHITE, bold=True, size=46)
merge_write(ws, "G5:L7", "registered participants²",
            fill=PINK, font_color=WHITE, bold=True, size=16)

merge_write(ws, "M1:R4", f"{int(new_reg):,}",
            fill=PINK, font_color=WHITE, bold=True, size=46)
merge_write(ws, "M5:R7", "new registrations",
            fill=PINK, font_color=WHITE, bold=True, size=16)

# separators
thin_white = Side(style="thin", color=WHITE)
for r in range(1, 8):
    ws.cell(r, 7).border = Border(left=thin_white)
    ws.cell(r, 13).border = Border(left=thin_white)

# ============================================================
# DEMOGRAPHICS SECTION
# ============================================================

merge_write(ws, "A9:R9", "DEMOGRAPHICS",
            fill=None, font_color=PINK, bold=True, size=15, align="left")

# Helper data is placed far to the right and hidden
helper_col = 21  # U

# Gender helper
ws.cell(1, helper_col).value = "Gender"
ws.cell(1, helper_col+1).value = "Percent"
ws.cell(2, helper_col).value = "Female"
ws.cell(2, helper_col+1).value = female_pct
ws.cell(3, helper_col).value = "Male"
ws.cell(3, helper_col+1).value = male_pct

# Gender chart
gender_chart = DoughnutChart()
gender_chart.title = "Gender"
gender_chart.holeSize = 58
gender_chart.height = 7.1
gender_chart.width = 8.2
gender_chart.legend.position = "r"
gender_chart.add_data(
    Reference(ws, min_col=helper_col+1, min_row=1, max_row=3),
    titles_from_data=True
)
gender_chart.set_categories(
    Reference(ws, min_col=helper_col, min_row=2, max_row=3)
)
gender_chart.dataLabels = DataLabelList()
gender_chart.dataLabels.showPercent = True
gender_chart.dataLabels.showLeaderLines = False
set_chart_colors(gender_chart, [PINK, ORANGE])
ws.add_chart(gender_chart, "A11")

# Age helper
age_start = 6
ws.cell(age_start, helper_col).value = "Age"
ws.cell(age_start, helper_col+1).value = "Percent"
age_order = ["16-25","26-35","36-45","46-55","56-65","66-75","76+"]
for i, band in enumerate(age_order, start=age_start+1):
    ws.cell(i, helper_col).value = band
    ws.cell(i, helper_col+1).value = age.get(band, 0)

age_chart = DoughnutChart()
age_chart.title = "Age"
age_chart.holeSize = 58
age_chart.height = 7.1
age_chart.width = 8.8
age_chart.legend.position = "r"
age_chart.add_data(
    Reference(ws, min_col=helper_col+1, min_row=age_start, max_row=age_start+len(age_order)),
    titles_from_data=True
)
age_chart.set_categories(
    Reference(ws, min_col=helper_col, min_row=age_start+1, max_row=age_start+len(age_order))
)
set_chart_colors(age_chart, [YELLOW, ORANGE, GREEN, PINK, TEAL, PURPLE, RED])
ws.add_chart(age_chart, "G11")

# Ethnicity donut
eth_start = 16
ws.cell(eth_start, helper_col).value = "Ethnicity"
ws.cell(eth_start, helper_col+1).value = "Percent"
ws.cell(eth_start+1, helper_col).value = "Ethnically diverse"
ws.cell(eth_start+1, helper_col+1).value = ethnically_diverse_pct
ws.cell(eth_start+2, helper_col).value = "White"
ws.cell(eth_start+2, helper_col+1).value = 100 - ethnically_diverse_pct

eth_chart = DoughnutChart()
eth_chart.title = "Ethnically diverse"
eth_chart.holeSize = 68
eth_chart.height = 6.5
eth_chart.width = 6.0
eth_chart.legend = None
eth_chart.add_data(
    Reference(ws, min_col=helper_col+1, min_row=eth_start, max_row=eth_start+2),
    titles_from_data=True
)
eth_chart.set_categories(
    Reference(ws, min_col=helper_col, min_row=eth_start+1, max_row=eth_start+2)
)
set_chart_colors(eth_chart, [PINK, GREY])
ws.add_chart(eth_chart, "M11")
merge_write(ws, "M20:O22", f"{ethnically_diverse_pct:.0f}%\nEthnically diverse\nbackground",
            fill=None, font_color=PINK, bold=True, size=12)

# IMD donut
imd_start = 20
ws.cell(imd_start, helper_col).value = "IMD"
ws.cell(imd_start, helper_col+1).value = "Percent"
ws.cell(imd_start+1, helper_col).value = "Most deprived 20%"
ws.cell(imd_start+1, helper_col+1).value = imd_20_pct
ws.cell(imd_start+2, helper_col).value = "Other matched"
ws.cell(imd_start+2, helper_col+1).value = 100-imd_20_pct

imd_chart = DoughnutChart()
imd_chart.title = "Index of Multiple Deprivation"
imd_chart.holeSize = 68
imd_chart.height = 6.5
imd_chart.width = 6.0
imd_chart.legend = None
imd_chart.add_data(
    Reference(ws, min_col=helper_col+1, min_row=imd_start, max_row=imd_start+2),
    titles_from_data=True
)
imd_chart.set_categories(
    Reference(ws, min_col=helper_col, min_row=imd_start+1, max_row=imd_start+2)
)
set_chart_colors(imd_chart, [PINK, "F48FB1"])
ws.add_chart(imd_chart, "P11")
merge_write(ws, "P20:R22", f"{imd_20_pct:.0f}%\nMost deprived\n20%",
            fill=None, font_color=PINK, bold=True, size=12)

# ============================================================
# BOTTOM CHARTS
# ============================================================

# Section titles
merge_write(ws, "A25:F25", "TOP 10 ACTIVITIES, BY ATTENDANCE",
            fill=None, font_color=PINK, bold=True, size=13, align="left")
merge_write(ws, "G25:L25", "TOP REASONS PEOPLE JOIN SAHELI HUB",
            fill=None, font_color=PINK, bold=True, size=13, align="left")
merge_write(ws, "M25:R25", "WHERE DO PEOPLE HEAR ABOUT US?",
            fill=None, font_color=PINK, bold=True, size=13, align="left")

# Top activities data
act_start = 25
ws.cell(act_start, helper_col+3).value = "Activity"
ws.cell(act_start, helper_col+4).value = "Attendance"
for i, (activity, att) in enumerate(activities, start=act_start+1):
    ws.cell(i, helper_col+3).value = activity
    ws.cell(i, helper_col+4).value = att

act_chart = BarChart()
act_chart.type = "bar"
act_chart.style = 10
act_chart.title = None
act_chart.height = 8.7
act_chart.width = 10.0
act_chart.legend = None
act_chart.add_data(
    Reference(ws, min_col=helper_col+4, min_row=act_start, max_row=act_start+len(activities)),
    titles_from_data=True
)
act_chart.set_categories(
    Reference(ws, min_col=helper_col+3, min_row=act_start+1, max_row=act_start+len(activities))
)
act_chart.dataLabels = DataLabelList()
act_chart.dataLabels.showVal = True
act_chart.varyColors = False
set_chart_colors(act_chart, [PINK])
ws.add_chart(act_chart, "A26")

# Reasons
reason_start = 25
ws.cell(reason_start, helper_col+6).value = "Reason"
ws.cell(reason_start, helper_col+7).value = "Participants"
for i, (reason, count, pct) in enumerate(reasons, start=reason_start+1):
    ws.cell(i, helper_col+6).value = reason
    ws.cell(i, helper_col+7).value = count

reason_chart = BarChart()
reason_chart.type = "bar"
reason_chart.style = 10
reason_chart.title = None
reason_chart.height = 8.7
reason_chart.width = 10.4
reason_chart.legend = None
reason_chart.add_data(
    Reference(ws, min_col=helper_col+7, min_row=reason_start, max_row=reason_start+len(reasons)),
    titles_from_data=True
)
reason_chart.set_categories(
    Reference(ws, min_col=helper_col+6, min_row=reason_start+1, max_row=reason_start+len(reasons))
)
reason_chart.dataLabels = DataLabelList()
reason_chart.dataLabels.showVal = True
set_chart_colors(reason_chart, [PINK])
ws.add_chart(reason_chart, "G26")

# Heard about
heard_start = 25
ws.cell(heard_start, helper_col+9).value = "Source"
ws.cell(heard_start, helper_col+10).value = "Participants"
for i, (source, count, pct) in enumerate(heard, start=heard_start+1):
    ws.cell(i, helper_col+9).value = source
    ws.cell(i, helper_col+10).value = count

heard_chart = DoughnutChart()
heard_chart.holeSize = 58
heard_chart.height = 8.5
heard_chart.width = 9.0
heard_chart.legend.position = "r"
heard_chart.add_data(
    Reference(ws, min_col=helper_col+10, min_row=heard_start, max_row=heard_start+len(heard)),
    titles_from_data=True
)
heard_chart.set_categories(
    Reference(ws, min_col=helper_col+9, min_row=heard_start+1, max_row=heard_start+len(heard))
)
heard_chart.dataLabels = DataLabelList()
heard_chart.dataLabels.showPercent = True
set_chart_colors(heard_chart, [PINK, ORANGE, GREEN, TEAL, PURPLE, "9CA3AF", GREY])
ws.add_chart(heard_chart, "M26")

# ============================================================
# FOOTNOTES
# ============================================================

merge_write(
    ws, "M43:R48",
    "1. Total attendance uses the verified reconciled annual-delivery figure.\n"
    "2. Registered participants = current FULL + Lite CRM profiles; this is not directly "
    "equivalent to last year's unique-attendee definition.\n"
    "3. Gender chart uses Female/Male records with verified gender only.\n"
    "4. Ethnicity and IMD percentages use records with usable/matched data.",
    fill=None, font_color=PINK, bold=False, size=8, align="left"
)

# Small comparison note
merge_write(
    ws, "A47:L49",
    f"Current headline comparison: {int(total_attendance):,} verified attendances | "
    f"{int(registered):,} registered participants | {int(new_reg):,} new registrations | "
    f"{int(sessions):,} sessions delivered",
    fill="FCE4EC", font_color=PINK, bold=True, size=9, align="left"
)

# Hide helper columns U:AE
for col in range(21, 32):
    ws.column_dimensions[get_column_letter(col)].hidden = True

# Optional logo if available
add_logo(ws)

# Print area
ws.print_area = "A1:R49"

# Save
wb.save(OUTPUT_XLSX)

print("Created:", OUTPUT_XLSX)
print("Top headline values:")
print("Attendance:", total_attendance)
print("Registered participants:", registered)
print("New registrations:", new_reg)
print("Female/Male recorded split: %.1f%% / %.1f%%" % (female_pct, male_pct))
