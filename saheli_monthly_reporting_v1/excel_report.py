from __future__ import annotations

from pathlib import Path
import re
import math
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MAGENTA = "D41474"
MAGENTA_LIGHT = "F8DCEB"
YELLOW = "FFC000"
NAVY = "0F4C81"
NAVY_DARK = "17365D"
WHITE = "FFFFFF"
LIGHT_GREY = "F3F4F6"
MID_GREY = "D9E1F2"
GREEN = "E2F0D9"
RED = "FCE4D6"
AMBER = "FFF2CC"
TEXT = "222222"


def safe_filename(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*]+', "_", str(value))
    value = re.sub(r"\s+", "_", value.strip())
    return value[:120] or "Unknown"


def _safe_sheet_name(value: str, existing: set[str]) -> str:
    raw = str(value or "Unknown").strip()
    low = raw.casefold()

    if "calthorpe" in low:
        base = "CALTHORPE"
    elif "alum rock" in low or low == "arcc":
        base = "ARCC"
    elif "handsworth" in low:
        base = "HANDSWORTH"
    elif "unknown" in low:
        base = "UNKNOWN"
    else:
        base = re.sub(r"[\[\]:*?/\\]", "", raw).upper()
        base = re.sub(r"\s+", " ", base).strip()
        if len(base) > 28:
            base = base[:28].rstrip()
        if not base:
            base = "LOCATION"

    candidate = base[:31]
    counter = 2
    while candidate in existing:
        suffix = f" {counter}"
        candidate = f"{base[:31-len(suffix)]}{suffix}"
        counter += 1
    existing.add(candidate)
    return candidate


def _write_df(writer, sheet_name: str, df: pd.DataFrame, startrow: int = 0, startcol: int = 0):
    if df is None:
        df = pd.DataFrame()
    df.to_excel(
        writer,
        sheet_name=sheet_name[:31],
        index=False,
        startrow=startrow,
        startcol=startcol,
    )


def _as_number(v):
    if pd.isna(v):
        return None
    if hasattr(v, "item"):
        v = v.item()
    return v


def _metric_value(df: pd.DataFrame, metric: str, column: str):
    if df is None or df.empty:
        return 0
    row = df[df["Metric"].eq(metric)]
    if row.empty:
        return 0
    return _as_number(row.iloc[0][column]) or 0


def _style_title(ws, title: str, subtitle: str, end_col: int = 8):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c = ws.cell(1, 1, title)
    c.fill = PatternFill("solid", fgColor=MAGENTA)
    c.font = Font(color=WHITE, bold=True, size=17)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    s = ws.cell(2, 1, subtitle)
    s.fill = PatternFill("solid", fgColor=MAGENTA_LIGHT)
    s.font = Font(color=TEXT, italic=True, size=10)
    s.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 22


def _section_header(ws, row: int, col_start: int, col_end: int, text: str):
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
    c = ws.cell(row, col_start, text)
    c.fill = PatternFill("solid", fgColor=YELLOW)
    c.font = Font(bold=True, color=TEXT, size=11)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 20


def _table_header(ws, row: int, col_start: int, headers: list[str]):
    for idx, header in enumerate(headers, start=col_start):
        c = ws.cell(row, idx, header)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(color=WHITE, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _apply_thin_borders(ws, min_row: int, max_row: int, min_col: int, max_col: int):
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = border


def _autowidth(ws, max_row_scan: int = 250):
    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        for row_idx in range(1, min(ws.max_row, max_row_scan) + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        width = min(max(max_len + 2, 11), 34)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _format_data_header_row(ws, row: int, start_col: int, end_col: int):
    for c in ws.iter_cols(min_col=start_col, max_col=end_col, min_row=row, max_row=row):
        cell = c[0]
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _trend_fill(cell, value):
    text = str(value or "").upper()
    if text == "UP":
        cell.fill = PatternFill("solid", fgColor=GREEN)
    elif text == "DOWN":
        cell.fill = PatternFill("solid", fgColor=RED)
    elif text == "STABLE":
        cell.fill = PatternFill("solid", fgColor=AMBER)


def _add_bar_chart(
    ws,
    title: str,
    category_col: int,
    series_start_col: int,
    series_end_col: int,
    header_row: int,
    data_start_row: int,
    data_end_row: int,
    position: str,
    *,
    horizontal: bool = False,
    width: float = 14.5,
    height: float = 8.0,
):
    """Add a compact management chart using an existing visible table."""
    if data_end_row < data_start_row:
        return

    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.style = 10
    chart.grouping = "clustered"
    chart.overlap = 0
    chart.title = title
    chart.height = height
    chart.width = width
    chart.legend.position = "b"

    data = Reference(
        ws,
        min_col=series_start_col,
        max_col=series_end_col,
        min_row=header_row,
        max_row=data_end_row,
    )
    categories = Reference(
        ws,
        min_col=category_col,
        min_row=data_start_row,
        max_row=data_end_row,
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    # Keep labels readable for longer activity/location names.
    if horizontal:
        chart.y_axis.title = ""
    else:
        chart.x_axis.title = ""

    ws.add_chart(chart, position)


def _build_summary_sheet(
    ws,
    overall: pd.DataFrame,
    locations: pd.DataFrame,
    activities: pd.DataFrame,
    registrations: pd.DataFrame,
    demographics,
    assessment_activity: pd.DataFrame,
    outcomes: pd.DataFrame,
    quality: pd.DataFrame,
    report_label: str,
    previous_label: str,
):
    """Build an SMT-friendly roll-up of the complete monthly workbook."""
    _style_title(
        ws,
        f"Saheli Hub Monthly Performance Report – {report_label}",
        f"Reporting period: {report_label} | Comparison: {previous_label}",
        end_col=8,
    )

    # ------------------------------------------------------------------
    # 1. Headline metrics + month-on-month comparison
    # ------------------------------------------------------------------
    _section_header(ws, 4, 1, 2, "Headline metrics")
    _section_header(ws, 4, 4, 8, "Month-on-month comparison")

    headline_metrics = [
        "Sessions Delivered",
        "Total Attendance",
        "Unique Participants",
        "New Registrations",
        "Health Assessments",
        "Follow-up Assessments",
    ]

    for i, metric in enumerate(headline_metrics, start=5):
        ws.cell(i, 1, metric)
        value = _metric_value(overall, metric, "Current Month")
        c = ws.cell(i, 2, value)
        c.font = Font(color=MAGENTA, bold=True)
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=MAGENTA_LIGHT)
        ws.cell(i, 2).fill = PatternFill("solid", fgColor=MAGENTA_LIGHT)

    headers = ["Measure", previous_label, report_label, "Difference", "% Change"]
    _table_header(ws, 5, 4, headers)

    for r, metric in enumerate(headline_metrics, start=6):
        prev = _metric_value(overall, metric, "Previous Month")
        curr = _metric_value(overall, metric, "Current Month")
        diff = curr - prev
        pct = None if not prev else diff / prev
        values = [metric, prev, curr, diff, pct]
        for j, val in enumerate(values, start=4):
            ws.cell(r, j, _as_number(val))
        if pct is not None:
            ws.cell(r, 8).number_format = "0.0%"
    _apply_thin_borders(ws, 5, 11, 4, 8)

    # ------------------------------------------------------------------
    # 2. Every location at a glance
    # ------------------------------------------------------------------
    start_row = 14
    _section_header(ws, start_row, 1, 8, "Location performance")
    location_cols = [
        "Location",
        "Previous Sessions",
        "Current Sessions",
        "Previous Attendance",
        "Current Attendance",
        "Previous Unique Participants",
        "Current Unique Participants",
        "Trend",
    ]
    _table_header(ws, start_row + 1, 1, location_cols)

    location_chart_start = None
    location_chart_end = None
    if locations is not None and not locations.empty:
        # Use to_dict instead of itertuples()._asdict(): pandas sanitises column
        # names with spaces in namedtuples, which caused blank summary cells.
        for i, record in enumerate(locations.to_dict("records"), start=start_row + 2):
            for j, col in enumerate(location_cols, start=1):
                ws.cell(i, j, _as_number(record.get(col)))
            _trend_fill(ws.cell(i, 8), record.get("Trend"))
        _apply_thin_borders(
            ws,
            start_row + 1,
            start_row + 1 + len(locations),
            1,
            8,
        )
        location_chart_start = start_row + 2
        location_chart_end = min(start_row + 1 + len(locations), location_chart_start + 9)

    row = start_row + 3 + (len(locations) if locations is not None else 0)

    # ------------------------------------------------------------------
    # 3. Top activities across Saheli
    # ------------------------------------------------------------------
    _section_header(ws, row, 1, 8, "Top activities")
    row += 1
    top_headers = [
        "Location", "Category", "Activity", "Sessions",
        "Attendance", "Unique participants", "Avg / session",
    ]
    _table_header(ws, row, 1, top_headers)
    top_header_row = row
    top = pd.DataFrame()
    if activities is not None and not activities.empty:
        top = activities.copy().sort_values(
            ["Attendance", "UniqueParticipants"], ascending=[False, False]
        ).head(10)
    top_data_start = None
    top_data_end = None
    if not top.empty:
        top_data_start = row + 1
        top_data_end = row + len(top)
        for ridx, rec in enumerate(top.to_dict("records"), start=row + 1):
            vals = [
                rec.get("Location"),
                rec.get("ActivityCategoryResolved"),
                rec.get("ActivityName"),
                rec.get("Sessions"),
                rec.get("Attendance"),
                rec.get("UniqueParticipants"),
                rec.get("Average Attendance / Session"),
            ]
            for cidx, val in enumerate(vals, start=1):
                ws.cell(ridx, cidx, _as_number(val))
        _apply_thin_borders(ws, row, row + len(top), 1, 7)
        row += len(top) + 2
    else:
        ws.cell(row + 1, 1, "No activity data")
        row += 3

    # ------------------------------------------------------------------
    # 4. Registrations by location
    # ------------------------------------------------------------------
    _section_header(ws, row, 1, 8, "Registration performance")
    row += 1
    reg_headers = ["Location", previous_label, report_label, "Difference"]
    _table_header(ws, row, 1, reg_headers)
    reg_header_row = row
    reg_view = pd.DataFrame()
    if registrations is not None and not registrations.empty:
        reg_view = registrations.copy().sort_values(
            ["Current Registrations", "Location"], ascending=[False, True]
        )
    reg_data_start = None
    reg_data_end = None
    if not reg_view.empty:
        reg_data_start = row + 1
        reg_data_end = min(row + len(reg_view), reg_data_start + 9)
        for ridx, rec in enumerate(reg_view.to_dict("records"), start=row + 1):
            vals = [
                rec.get("Location"),
                rec.get("Previous Registrations", 0),
                rec.get("Current Registrations", 0),
                rec.get("Change", 0),
            ]
            for cidx, val in enumerate(vals, start=1):
                ws.cell(ridx, cidx, _as_number(val))
        _apply_thin_borders(ws, row, row + len(reg_view), 1, 4)
        row += len(reg_view) + 2
    else:
        ws.cell(row + 1, 1, "No registration data")
        row += 3

    # ------------------------------------------------------------------
    # 5. Demographic snapshot (all new registrations)
    # ------------------------------------------------------------------
    _section_header(ws, row, 1, 8, "Demographic snapshot")
    row += 1
    _table_header(ws, row, 1, ["Dimension", "Category", "Participants", "% of dimension"])
    gender, ethnicity, age, disability = demographics
    demographic_blocks = [
        ("Gender", gender, "Gender"),
        ("Age", age, "Age Band"),
        ("Ethnicity", ethnicity, "Ethnicity"),
        ("Disability / health condition", disability, "Disability"),
    ]
    demo_rows = []
    for dimension, df, category_col in demographic_blocks:
        if df is None or df.empty or category_col not in df.columns:
            continue
        agg = df.groupby(category_col, dropna=False)["Participants"].sum().reset_index()
        total = float(agg["Participants"].sum())
        agg = agg.sort_values("Participants", ascending=False)
        # Keep the summary readable; complete detail remains in DEMOGRAPHICS.
        for rec in agg.head(5).to_dict("records"):
            count = int(rec.get("Participants", 0) or 0)
            demo_rows.append([
                dimension,
                rec.get(category_col),
                count,
                (count / total) if total else None,
            ])
    if demo_rows:
        for ridx, values in enumerate(demo_rows, start=row + 1):
            for cidx, val in enumerate(values, start=1):
                ws.cell(ridx, cidx, _as_number(val))
            if values[3] is not None:
                ws.cell(ridx, 4).number_format = "0.0%"
        _apply_thin_borders(ws, row, row + len(demo_rows), 1, 4)
        row += len(demo_rows) + 2
    else:
        ws.cell(row + 1, 1, "No demographic data")
        row += 3

    # ------------------------------------------------------------------
    # 6. Outcomes roll-up
    # ------------------------------------------------------------------
    _section_header(ws, row, 1, 8, "Outcome snapshot")
    row += 1
    _table_header(ws, row, 1, ["Outcome", "Paired", "Improved", "Improvement %"])
    outcome_header_row = row

    outcome_specs = [
        ("Confidence to join", "Confidence Paired", "Improved Confidence To Join"),
        ("Feeling confident", "Feeling Confident Paired", "Improved Feeling Confident"),
        ("Movement", "Movement Paired", "Improved Movement"),
        ("Less isolated", "Isolation Paired", "Less Isolated"),
        ("More active days", "Active Days Paired", "More Active Days"),
    ]
    outcome_rows = []
    if outcomes is not None and not outcomes.empty:
        for label, paired_col, improved_col in outcome_specs:
            paired = int(pd.to_numeric(outcomes.get(paired_col), errors="coerce").fillna(0).sum()) if paired_col in outcomes.columns else 0
            improved = int(pd.to_numeric(outcomes.get(improved_col), errors="coerce").fillna(0).sum()) if improved_col in outcomes.columns else 0
            pct = improved / paired if paired else None
            outcome_rows.append([label, paired, improved, pct])
    outcome_data_start = None
    outcome_data_end = None
    if outcome_rows:
        outcome_data_start = row + 1
        outcome_data_end = row + len(outcome_rows)
        for ridx, values in enumerate(outcome_rows, start=row + 1):
            for cidx, val in enumerate(values, start=1):
                ws.cell(ridx, cidx, _as_number(val))
            if values[3] is not None:
                ws.cell(ridx, 4).number_format = "0.0%"
        _apply_thin_borders(ws, row, row + len(outcome_rows), 1, 4)
        row += len(outcome_rows) + 2
    else:
        ws.cell(row + 1, 1, "No paired outcome data")
        row += 3

    # ------------------------------------------------------------------
    # 7. Data quality roll-up
    # ------------------------------------------------------------------
    _section_header(ws, row, 1, 8, "Data quality snapshot")
    row += 1
    dq_headers = [
        "New registrations", "Any core gap", "Gap %", "Missing DOB",
        "Missing postcode", "Missing gender", "Missing ethnicity", "Missing mobile",
    ]
    _table_header(ws, row, 1, dq_headers)
    if quality is not None and not quality.empty:
        def sum_col(name):
            return int(pd.to_numeric(quality[name], errors="coerce").fillna(0).sum()) if name in quality.columns else 0

        total_reg = sum_col("New Registrations")
        any_gap = sum_col("Participants With Any Core Gap")
        vals = [
            total_reg,
            any_gap,
            (any_gap / total_reg) if total_reg else None,
            sum_col("Missing DOB"),
            sum_col("Missing Postcode"),
            sum_col("Missing Gender"),
            sum_col("Missing Ethnicity"),
            sum_col("Missing Mobile"),
        ]
        for cidx, val in enumerate(vals, start=1):
            ws.cell(row + 1, cidx, _as_number(val))
        if vals[2] is not None:
            ws.cell(row + 1, 3).number_format = "0.0%"
        _apply_thin_borders(ws, row, row + 1, 1, 8)
        row += 3
    else:
        ws.cell(row + 1, 1, "No data-quality information")
        row += 3

    # ------------------------------------------------------------------
    # 8. Key changes / areas to look at
    # ------------------------------------------------------------------
    _section_header(ws, row, 1, 8, "Key location changes")
    row += 1
    _table_header(ws, row, 1, ["Location", "Attendance change", "% change", "Trend"])
    changes = pd.DataFrame()
    if locations is not None and not locations.empty:
        changes = locations[[
            "Location", "Attendance Change", "Attendance % Change", "Trend"
        ]].copy()
        changes = changes.sort_values("Attendance Change", ascending=False)
        if len(changes) > 8:
            top_up = changes.head(4)
            top_down = changes.tail(4)
            changes = pd.concat([top_up, top_down]).drop_duplicates("Location")
    if not changes.empty:
        for ridx, rec in enumerate(changes.to_dict("records"), start=row + 1):
            vals = [
                rec.get("Location"),
                rec.get("Attendance Change"),
                None if pd.isna(rec.get("Attendance % Change")) else rec.get("Attendance % Change") / 100,
                rec.get("Trend"),
            ]
            for cidx, val in enumerate(vals, start=1):
                ws.cell(ridx, cidx, _as_number(val))
            if vals[2] is not None:
                ws.cell(ridx, 3).number_format = "0.0%"
            _trend_fill(ws.cell(ridx, 4), rec.get("Trend"))
        _apply_thin_borders(ws, row, row + len(changes), 1, 4)

    # ------------------------------------------------------------------
    # 9. Visual dashboard charts (kept to the right of the tables)
    # ------------------------------------------------------------------
    if location_chart_start is not None and location_chart_end is not None:
        _add_bar_chart(
            ws,
            f"Attendance by location – {previous_label} vs {report_label}",
            category_col=1,
            series_start_col=4,
            series_end_col=5,
            header_row=start_row + 1,
            data_start_row=location_chart_start,
            data_end_row=location_chart_end,
            position="J4",
            horizontal=True,
            width=16.5,
            height=9.0,
        )

    if top_data_start is not None and top_data_end is not None:
        _add_bar_chart(
            ws,
            f"Top activities by attendance – {report_label}",
            category_col=3,
            series_start_col=5,
            series_end_col=5,
            header_row=top_header_row,
            data_start_row=top_data_start,
            data_end_row=top_data_end,
            position="J22",
            horizontal=True,
            width=16.5,
            height=9.0,
        )

    if outcome_data_start is not None and outcome_data_end is not None:
        _add_bar_chart(
            ws,
            f"Paired outcomes – {report_label}",
            category_col=1,
            series_start_col=2,
            series_end_col=3,
            header_row=outcome_header_row,
            data_start_row=outcome_data_start,
            data_end_row=outcome_data_end,
            position="J40",
            horizontal=True,
            width=16.5,
            height=8.5,
        )

    if reg_data_start is not None and reg_data_end is not None:
        _add_bar_chart(
            ws,
            f"Registrations by location – {previous_label} vs {report_label}",
            category_col=1,
            series_start_col=2,
            series_end_col=3,
            header_row=reg_header_row,
            data_start_row=reg_data_start,
            data_end_row=reg_data_end,
            position="J57",
            horizontal=True,
            width=16.5,
            height=9.0,
        )

    ws.freeze_panes = "A5"
    _autowidth(ws)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["D"].width = 24

def _build_demographics_sheet(ws, demographics, report_label: str):
    gender, ethnicity, age, disability = demographics
    _style_title(ws, f"Demographics – {report_label}", "New registrations during the reporting month", end_col=7)

    blocks = [
        ("Gender", gender),
        ("Ethnicity", ethnicity),
        ("Age profile", age),
        ("Disability / health condition", disability),
    ]

    row = 4
    for title, df in blocks:
        _section_header(ws, row, 1, 7, title)
        row += 1
        if df is None or df.empty:
            ws.cell(row, 1, "No data")
            row += 3
            continue
        headers = list(df.columns)
        _table_header(ws, row, 1, headers)
        for ridx, rec in enumerate(df.itertuples(index=False), start=row + 1):
            for cidx, val in enumerate(rec, start=1):
                ws.cell(ridx, cidx, _as_number(val))
        _apply_thin_borders(ws, row, row + len(df), 1, len(headers))
        row += len(df) + 3

    ws.freeze_panes = "A4"
    _autowidth(ws)


def _build_top_activities_sheet(ws, activities: pd.DataFrame, report_label: str):
    _style_title(ws, f"Top Activities – {report_label}", "Current-month delivery ranked by attendance", end_col=7)
    _section_header(ws, 4, 1, 7, "Activity performance across all locations")

    if activities is None or activities.empty:
        ws.cell(5, 1, "No activity data")
        return

    cols = [
        "Location",
        "ActivityCategoryResolved",
        "ActivityName",
        "Sessions",
        "Attendance",
        "UniqueParticipants",
        "Average Attendance / Session",
    ]
    view = activities[cols].copy().sort_values("Attendance", ascending=False)
    headers = [
        "Location",
        "Category",
        "Activity",
        "Sessions",
        "Attendance",
        "Unique participants",
        "Avg attendance / session",
    ]
    _table_header(ws, 5, 1, headers)
    for ridx, row in enumerate(view.itertuples(index=False), start=6):
        for cidx, val in enumerate(row, start=1):
            ws.cell(ridx, cidx, _as_number(val))
    _apply_thin_borders(ws, 5, 5 + len(view), 1, 7)
    top_end = min(5 + len(view), 15)
    _add_bar_chart(
        ws,
        f"Top 10 activities by attendance – {report_label}",
        category_col=3,
        series_start_col=5,
        series_end_col=5,
        header_row=5,
        data_start_row=6,
        data_end_row=top_end,
        position="I4",
        horizontal=True,
        width=16.0,
        height=9.0,
    )
    ws.freeze_panes = "A6"
    _autowidth(ws)


def _build_registration_sheet(ws, registrations, quality, report_label: str, previous_label: str):
    _style_title(ws, f"Registration Insights – {report_label}", f"New registrations and data quality compared with {previous_label}", end_col=8)

    _section_header(ws, 4, 1, 5, "Registrations by location")
    if registrations is not None and not registrations.empty:
        headers = list(registrations.columns)
        _table_header(ws, 5, 1, headers)
        for ridx, row in enumerate(registrations.itertuples(index=False), start=6):
            for cidx, val in enumerate(row, start=1):
                ws.cell(ridx, cidx, _as_number(val))
        _apply_thin_borders(ws, 5, 5 + len(registrations), 1, len(headers))

    qstart = 8 + (len(registrations) if registrations is not None else 0)
    _section_header(ws, qstart, 1, 8, "Data quality for new registrations")
    if quality is not None and not quality.empty:
        headers = list(quality.columns)
        _table_header(ws, qstart + 1, 1, headers)
        for ridx, row in enumerate(quality.itertuples(index=False), start=qstart + 2):
            for cidx, val in enumerate(row, start=1):
                ws.cell(ridx, cidx, _as_number(val))
        _apply_thin_borders(ws, qstart + 1, qstart + 1 + len(quality), 1, len(headers))

    if registrations is not None and not registrations.empty:
        reg_end = min(5 + len(registrations), 15)
        # Expected order: Location, Previous Registrations, Current Registrations, Change.
        _add_bar_chart(
            ws,
            f"Registrations by location – {previous_label} vs {report_label}",
            category_col=1,
            series_start_col=2,
            series_end_col=3,
            header_row=5,
            data_start_row=6,
            data_end_row=reg_end,
            position="J4",
            horizontal=True,
            width=16.0,
            height=9.0,
        )

    ws.freeze_panes = "A5"
    _autowidth(ws)



def _build_outcomes_sheet(ws, outcomes: pd.DataFrame, assessment_activity: pd.DataFrame, report_label: str):
    _style_title(
        ws,
        f"Outcomes – {report_label}",
        "Paired assessment outcomes and assessment activity by participant site",
        end_col=10,
    )

    _section_header(ws, 4, 1, 8, "Assessment activity by location")
    row = 5
    if assessment_activity is not None and not assessment_activity.empty:
        headers = list(assessment_activity.columns)
        _table_header(ws, row, 1, headers)
        for ridx, rec in enumerate(assessment_activity.to_dict("records"), start=row + 1):
            for cidx, header in enumerate(headers, start=1):
                ws.cell(ridx, cidx, _as_number(rec.get(header)))
        _apply_thin_borders(ws, row, row + len(assessment_activity), 1, len(headers))
        row += len(assessment_activity) + 3
    else:
        ws.cell(row, 1, "No assessment activity")
        row += 3

    _section_header(ws, row, 1, 10, "Paired outcomes by location")
    row += 1
    if outcomes is not None and not outcomes.empty:
        headers = list(outcomes.columns)
        _table_header(ws, row, 1, headers)
        for ridx, rec in enumerate(outcomes.to_dict("records"), start=row + 1):
            for cidx, header in enumerate(headers, start=1):
                val = rec.get(header)
                ws.cell(ridx, cidx, _as_number(val))
                if str(header).startswith("%") and val is not None and not pd.isna(val):
                    ws.cell(ridx, cidx, float(val) / 100)
                    ws.cell(ridx, cidx).number_format = "0.0%"
        _apply_thin_borders(ws, row, row + len(outcomes), 1, len(headers))
    else:
        ws.cell(row, 1, "No paired outcome data")

    if assessment_activity is not None and not assessment_activity.empty:
        assess_end = min(5 + len(assessment_activity), 15)
        # Expected first columns: Location, Previous Assessments, Current Assessments.
        _add_bar_chart(
            ws,
            f"Health assessments by location – {report_label}",
            category_col=1,
            series_start_col=2,
            series_end_col=3,
            header_row=5,
            data_start_row=6,
            data_end_row=assess_end,
            position="L4",
            horizontal=True,
            width=16.0,
            height=9.0,
        )

    ws.freeze_panes = "A5"
    _autowidth(ws)


def _build_data_quality_sheet(ws, quality: pd.DataFrame, report_label: str):
    _style_title(
        ws,
        f"Data Quality – {report_label}",
        "Core-field completeness for new registrations in the reporting month",
        end_col=10,
    )
    _section_header(ws, 4, 1, 10, "Data quality by location")
    if quality is None or quality.empty:
        ws.cell(5, 1, "No data-quality information")
        return

    headers = list(quality.columns)
    _table_header(ws, 5, 1, headers)
    for ridx, rec in enumerate(quality.to_dict("records"), start=6):
        for cidx, header in enumerate(headers, start=1):
            value = rec.get(header)
            if header == "Data Gap %" and value is not None and not pd.isna(value):
                ws.cell(ridx, cidx, float(value) / 100)
                ws.cell(ridx, cidx).number_format = "0.0%"
            else:
                ws.cell(ridx, cidx, _as_number(value))
    _apply_thin_borders(ws, 5, 5 + len(quality), 1, len(headers))
    if "Participants With Any Core Gap" in headers:
        gap_col = headers.index("Participants With Any Core Gap") + 1
        dq_end = min(5 + len(quality), 15)
        _add_bar_chart(
            ws,
            f"Registrations with core data gaps – {report_label}",
            category_col=1,
            series_start_col=gap_col,
            series_end_col=gap_col,
            header_row=5,
            data_start_row=6,
            data_end_row=dq_end,
            position="L4",
            horizontal=True,
            width=16.0,
            height=9.0,
        )
    ws.freeze_panes = "A6"
    _autowidth(ws)

def _write_small_table(ws, start_row: int, title: str, df: pd.DataFrame, max_cols: int = 8):
    _section_header(ws, start_row, 1, max_cols, title)
    header_row = start_row + 1
    if df is None or df.empty:
        ws.cell(header_row, 1, "No data")
        return header_row + 2

    headers = list(df.columns)[:max_cols]
    _table_header(ws, header_row, 1, headers)
    for ridx, row in enumerate(df[headers].itertuples(index=False), start=header_row + 1):
        for cidx, val in enumerate(row, start=1):
            ws.cell(ridx, cidx, _as_number(val))
    _apply_thin_borders(ws, header_row, header_row + len(df), 1, len(headers))
    return header_row + len(df) + 2


def _build_location_sheet(
    ws,
    location: str,
    report_label: str,
    previous_label: str,
    locations: pd.DataFrame,
    categories: pd.DataFrame,
    activities: pd.DataFrame,
    registrations: pd.DataFrame,
    demographics,
    assessment_activity: pd.DataFrame,
    outcomes: pd.DataFrame,
    quality: pd.DataFrame,
):
    _style_title(
        ws,
        f"{location} – Monthly Performance",
        f"{report_label} performance compared with {previous_label}",
        end_col=8,
    )

    loc = locations[locations["Location"].eq(location)].copy() if locations is not None and not locations.empty else pd.DataFrame()

    _section_header(ws, 4, 1, 2, "Headline metrics")
    _section_header(ws, 4, 4, 8, "Month-on-month comparison")

    if not loc.empty:
        r = loc.iloc[0]
        headline = [
            ("Sessions delivered", r.get("Current Sessions", 0)),
            ("Attendance", r.get("Current Attendance", 0)),
            ("Unique participants", r.get("Current Unique Participants", 0)),
        ]
    else:
        headline = [
            ("Sessions delivered", 0),
            ("Attendance", 0),
            ("Unique participants", 0),
        ]

    reg = registrations[registrations["Location"].eq(location)].copy() if registrations is not None and not registrations.empty else pd.DataFrame()
    cur_reg = int(reg.iloc[0].get("Current Registrations", 0)) if not reg.empty else 0
    headline.append(("New registrations", cur_reg))

    ass = assessment_activity[assessment_activity["Location"].eq(location)].copy() if assessment_activity is not None and not assessment_activity.empty else pd.DataFrame()
    cur_ass = int(ass.iloc[0].get("Current Assessments", 0)) if not ass.empty else 0
    cur_follow = int(ass.iloc[0].get("Current Follow-ups", 0)) if not ass.empty else 0
    headline.append(("Health assessments", cur_ass))
    headline.append(("Follow-up assessments", cur_follow))

    for i, (name, val) in enumerate(headline, start=5):
        ws.cell(i, 1, name)
        ws.cell(i, 2, _as_number(val))
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=MAGENTA_LIGHT)
        ws.cell(i, 2).fill = PatternFill("solid", fgColor=MAGENTA_LIGHT)
        ws.cell(i, 2).font = Font(color=MAGENTA, bold=True)

    _table_header(ws, 5, 4, ["Measure", previous_label, report_label, "Difference", "Trend"])

    compare_rows = []
    if not loc.empty:
        lr = loc.iloc[0]
        compare_rows.extend([
            ("Sessions", lr.get("Previous Sessions", 0), lr.get("Current Sessions", 0)),
            ("Attendance", lr.get("Previous Attendance", 0), lr.get("Current Attendance", 0)),
            ("Unique participants", lr.get("Previous Unique Participants", 0), lr.get("Current Unique Participants", 0)),
        ])
    else:
        compare_rows.extend([("Sessions", 0, 0), ("Attendance", 0, 0), ("Unique participants", 0, 0)])

    prev_reg = int(reg.iloc[0].get("Previous Registrations", 0)) if not reg.empty else 0
    compare_rows.append(("New registrations", prev_reg, cur_reg))
    prev_ass = int(ass.iloc[0].get("Previous Assessments", 0)) if not ass.empty else 0
    compare_rows.append(("Health assessments", prev_ass, cur_ass))
    prev_follow = int(ass.iloc[0].get("Previous Follow-ups", 0)) if not ass.empty else 0
    compare_rows.append(("Follow-up assessments", prev_follow, cur_follow))

    for ridx, (measure, prev, curr) in enumerate(compare_rows, start=6):
        diff = curr - prev
        trend = "UP" if diff > 0 else "DOWN" if diff < 0 else "STABLE"
        vals = [measure, prev, curr, diff, trend]
        for cidx, val in enumerate(vals, start=4):
            ws.cell(ridx, cidx, _as_number(val))
        _trend_fill(ws.cell(ridx, 8), trend)
    _apply_thin_borders(ws, 5, 11, 4, 8)

    next_row = 14

    cat = categories[categories["Location"].eq(location)].copy() if categories is not None and not categories.empty else pd.DataFrame()
    if not cat.empty:
        cat = cat.rename(columns={
            "ActivityCategoryResolved": "Activity Category",
            "CategoryResolved": "Category",
            "SubCategoryResolved": "Subcategory",
        })
        wanted = [
            "Activity Category", "Category", "Subcategory",
            "Previous Sessions", "Current Sessions",
            "Previous Attendance", "Current Attendance",
            "Previous UniqueParticipants", "Current UniqueParticipants",
            "Attendance Change", "Trend",
        ]
        wanted = [c for c in wanted if c in cat.columns]
        cat = cat[wanted]
    next_row = _write_small_table(ws, next_row, "Category performance", cat, max_cols=11)

    act = activities[activities["Location"].eq(location)].copy() if activities is not None and not activities.empty else pd.DataFrame()
    if not act.empty:
        act = act.rename(columns={
            "ActivityCategoryResolved": "Category",
            "UniqueParticipants": "Unique participants",
        })
        wanted = [
            "Category", "ActivityName", "Sessions", "Attendance",
            "Unique participants", "Average Attendance / Session",
        ]
        wanted = [c for c in wanted if c in act.columns]
        act = act[wanted].sort_values("Attendance", ascending=False)
    activity_table_start = next_row
    next_row = _write_small_table(ws, next_row, "Top activities", act, max_cols=8)
    if act is not None and not act.empty:
        activity_header_row = activity_table_start + 1
        activity_data_start = activity_table_start + 2
        activity_data_end = min(activity_header_row + len(act), activity_data_start + 7)
        # Columns after rename are: Category, ActivityName, Sessions, Attendance, ...
        _add_bar_chart(
            ws,
            f"Top activities by attendance – {report_label}",
            category_col=2,
            series_start_col=4,
            series_end_col=4,
            header_row=activity_header_row,
            data_start_row=activity_data_start,
            data_end_row=activity_data_end,
            position="J4",
            horizontal=True,
            width=16.0,
            height=9.0,
        )

    out = outcomes[outcomes["Location"].eq(location)].copy() if outcomes is not None and not outcomes.empty and "Location" in outcomes.columns else pd.DataFrame()
    next_row = _write_small_table(ws, next_row, "Outcomes", out, max_cols=10)

    q = quality[quality["Location"].eq(location)].copy() if quality is not None and not quality.empty else pd.DataFrame()
    next_row = _write_small_table(ws, next_row, "Data quality", q, max_cols=10)

    gender, ethnicity, age, disability = demographics
    demo_frames = []
    for label, df, col in [
        ("Gender", gender, "Gender"),
        ("Age band", age, "Age Band"),
        ("Ethnicity", ethnicity, "Ethnicity"),
        ("Disability / health condition", disability, "Disability"),
    ]:
        if df is None or df.empty:
            continue
        x = df[df["Location"].eq(location)].copy()
        if x.empty:
            continue
        category_col = [c for c in x.columns if c not in {"Location", "Participants"}]
        if not category_col:
            continue
        x = x[[category_col[0], "Participants"]].copy()
        x.insert(0, "Breakdown", label)
        x = x.rename(columns={category_col[0]: "Category"})
        demo_frames.append(x)
    demo = pd.concat(demo_frames, ignore_index=True) if demo_frames else pd.DataFrame()
    _write_small_table(ws, next_row, "Demographics", demo, max_cols=5)

    ws.freeze_panes = "A5"
    _autowidth(ws)
    ws.column_dimensions["A"].width = 28


def write_single_workbook(
    path: Path,
    report_label: str,
    previous_label: str,
    overall: pd.DataFrame,
    locations: pd.DataFrame,
    categories: pd.DataFrame,
    activities: pd.DataFrame,
    registrations: pd.DataFrame,
    demographics: tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame],
    assessment_activity: pd.DataFrame,
    outcomes: pd.DataFrame,
    quality: pd.DataFrame,
):
    """
    Create ONE Excel workbook containing:
      SUMMARY
      DEMOGRAPHICS
      TOP ACTIVITIES
      REGISTRATION INSIGHTS
      OUTCOMES
      DATA QUALITY
      one sheet for every location
    """
    path.parent.mkdir(parents=True, exist_ok=True)

    core_sheets = [
        "SUMMARY",
        "DEMOGRAPHICS",
        "TOP ACTIVITIES",
        "REGISTRATION INSIGHTS",
        "OUTCOMES",
        "DATA QUALITY",
    ]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name in core_sheets:
            pd.DataFrame().to_excel(writer, sheet_name=name, index=False)

        existing = set(core_sheets)
        location_sheet_map = {}
        all_locations = []
        if locations is not None and not locations.empty:
            all_locations = locations["Location"].dropna().astype(str).tolist()
        for location in all_locations:
            sheet_name = _safe_sheet_name(location, existing)
            location_sheet_map[location] = sheet_name
            pd.DataFrame().to_excel(writer, sheet_name=sheet_name, index=False)

    wb = load_workbook(path)

    _build_summary_sheet(
        wb["SUMMARY"],
        overall,
        locations,
        activities,
        registrations,
        demographics,
        assessment_activity,
        outcomes,
        quality,
        report_label,
        previous_label,
    )
    _build_demographics_sheet(wb["DEMOGRAPHICS"], demographics, report_label)
    _build_top_activities_sheet(wb["TOP ACTIVITIES"], activities, report_label)
    _build_registration_sheet(
        wb["REGISTRATION INSIGHTS"], registrations, quality, report_label, previous_label
    )
    _build_outcomes_sheet(wb["OUTCOMES"], outcomes, assessment_activity, report_label)
    _build_data_quality_sheet(wb["DATA QUALITY"], quality, report_label)

    for location, sheet_name in location_sheet_map.items():
        _build_location_sheet(
            wb[sheet_name],
            location,
            report_label,
            previous_label,
            locations,
            categories,
            activities,
            registrations,
            demographics,
            assessment_activity,
            outcomes,
            quality,
        )

    for ws in wb.worksheets:
        if ws["A1"].value is None and ws.max_row == 1 and ws.max_column == 1:
            ws.delete_rows(1, 1)

    wb.save(path)
    return location_sheet_map

