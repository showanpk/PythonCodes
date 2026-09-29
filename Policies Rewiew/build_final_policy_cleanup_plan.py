from pathlib import Path
from datetime import date
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# CONFIGURATION
# ============================================================

BASE_FOLDER = Path(
    r"C:\Users\shonk\source\PythonCodes"
)

DUPLICATE_REPORT = BASE_FOLDER / (
    "Saheli_Duplicate_Reconciliation_20260929.xlsx"
)

VERSION_REPORT = BASE_FOLDER / (
    "Saheli_Uncertain_Version_Comparison_20260929.xlsx"
)

TODAY = date.today()

OUTPUT_FILE = BASE_FOLDER / (
    f"Saheli_Final_Policy_Cleanup_Plan_"
    f"{TODAY.strftime('%Y%m%d')}.xlsx"
)


# ============================================================
# SAFETY
# ============================================================

print("=" * 78)
print("SAHELI FINAL POLICY CLEAN-UP PLAN")
print("=" * 78)
print()
print("READ-ONLY DECISION REPORT")
print()
print("This script WILL NOT:")
print("  - delete policies")
print("  - rename policies")
print("  - move policies")
print("  - modify policies")
print()
print("It only creates a final administrative decision workbook.")
print()


# ============================================================
# HELPERS
# ============================================================

def clean(value):
    if value is None:
        return ""

    value = str(value)
    value = value.replace("\xa0", " ")
    value = value.replace("\u200b", "")
    value = re.sub(r"\s+", " ", value)

    return value.strip()


def normalise_name(value):
    value = clean(value).lower()

    value = value.replace("_", " ")
    value = re.sub(r"\s+", " ", value)

    return value.strip()


def find_column(headers, *possible_names):
    for name in possible_names:
        if name in headers:
            return headers[name]

    return None


# ============================================================
# VERIFY INPUT FILES
# ============================================================

if not DUPLICATE_REPORT.exists():
    raise FileNotFoundError(
        f"Duplicate report not found:\n{DUPLICATE_REPORT}"
    )

if not VERSION_REPORT.exists():
    raise FileNotFoundError(
        f"Version comparison report not found:\n{VERSION_REPORT}"
    )


# ============================================================
# LOAD DUPLICATE RECONCILIATION
# ============================================================

duplicate_wb = load_workbook(
    DUPLICATE_REPORT,
    read_only=True,
    data_only=True
)

if "Duplicate Reconciliation" not in duplicate_wb.sheetnames:
    raise ValueError(
        "Duplicate Reconciliation sheet not found."
    )

duplicate_ws = duplicate_wb[
    "Duplicate Reconciliation"
]

dup_headers = {
    clean(cell.value): cell.column
    for cell in duplicate_ws[1]
}


required_dup_columns = [
    "File A",
    "File B",
    "Classification",
    "Content Similarity %",
    "Last Reviewed A",
    "Last Reviewed B",
    "Next Review A",
    "Next Review B",
]

for column in required_dup_columns:
    if column not in dup_headers:
        raise ValueError(
            f"Missing duplicate-report column: {column}"
        )


duplicate_pairs = []


for row in range(
    2,
    duplicate_ws.max_row + 1
):

    duplicate_pairs.append({
        "file_a": clean(
            duplicate_ws.cell(
                row,
                dup_headers["File A"]
            ).value
        ),

        "file_b": clean(
            duplicate_ws.cell(
                row,
                dup_headers["File B"]
            ).value
        ),

        "classification": clean(
            duplicate_ws.cell(
                row,
                dup_headers["Classification"]
            ).value
        ),

        "similarity": clean(
            duplicate_ws.cell(
                row,
                dup_headers["Content Similarity %"]
            ).value
        ),

        "last_a": clean(
            duplicate_ws.cell(
                row,
                dup_headers["Last Reviewed A"]
            ).value
        ),

        "last_b": clean(
            duplicate_ws.cell(
                row,
                dup_headers["Last Reviewed B"]
            ).value
        ),

        "next_a": clean(
            duplicate_ws.cell(
                row,
                dup_headers["Next Review A"]
            ).value
        ),

        "next_b": clean(
            duplicate_ws.cell(
                row,
                dup_headers["Next Review B"]
            ).value
        ),
    })


duplicate_wb.close()


# ============================================================
# LOAD VERSION COMPARISON
# ============================================================

version_wb = load_workbook(
    VERSION_REPORT,
    read_only=True,
    data_only=True
)

if "Version Comparison" not in version_wb.sheetnames:
    raise ValueError(
        "Version Comparison sheet not found."
    )

version_ws = version_wb[
    "Version Comparison"
]

ver_headers = {
    clean(cell.value): cell.column
    for cell in version_ws[1]
}


required_ver_columns = [
    "Policy / Topic",
    "File A",
    "File B",
    "Content Similarity %",
    "Substantive Change Lines",
    "Metadata Only?",
    "Master Assessment",
    "Likely Master",
    "Reason",
]

for column in required_ver_columns:
    if column not in ver_headers:
        raise ValueError(
            f"Missing version-report column: {column}"
        )


version_results = {}


for row in range(
    2,
    version_ws.max_row + 1
):

    topic = clean(
        version_ws.cell(
            row,
            ver_headers["Policy / Topic"]
        ).value
    )

    result = {
        "topic": topic,

        "file_a": clean(
            version_ws.cell(
                row,
                ver_headers["File A"]
            ).value
        ),

        "file_b": clean(
            version_ws.cell(
                row,
                ver_headers["File B"]
            ).value
        ),

        "similarity": clean(
            version_ws.cell(
                row,
                ver_headers["Content Similarity %"]
            ).value
        ),

        "substantive": clean(
            version_ws.cell(
                row,
                ver_headers["Substantive Change Lines"]
            ).value
        ),

        "metadata_only": clean(
            version_ws.cell(
                row,
                ver_headers["Metadata Only?"]
            ).value
        ),

        "master_assessment": clean(
            version_ws.cell(
                row,
                ver_headers["Master Assessment"]
            ).value
        ),

        "likely_master": clean(
            version_ws.cell(
                row,
                ver_headers["Likely Master"]
            ).value
        ),

        "reason": clean(
            version_ws.cell(
                row,
                ver_headers["Reason"]
            ).value
        ),
    }

    pair_key = frozenset([
        normalise_name(result["file_a"]),
        normalise_name(result["file_b"])
    ])

    version_results[
        pair_key
    ] = result


version_wb.close()


# ============================================================
# PROTECTED PAIRS
# ============================================================

# These must NOT receive an automatic delete recommendation.

PROTECTED_TOPICS = {
    "absence management",
    "computers email internet",
    "data protection",
}


# ============================================================
# BUILD FINAL DECISIONS
# ============================================================

decisions = []


for pair in duplicate_pairs:

    file_a = pair["file_a"]
    file_b = pair["file_b"]

    pair_key = frozenset([
        normalise_name(file_a),
        normalise_name(file_b)
    ])

    version_info = version_results.get(
        pair_key
    )

    topic = ""

    if version_info:
        topic = version_info[
            "topic"
        ]

    # --------------------------------------------------------
    # CASE 1:
    # IDENTICAL EXTRACTED TEXT
    # --------------------------------------------------------

    if pair["classification"] == "IDENTICAL TEXT CONTENT":

        decision = (
            "DUPLICATE CANDIDATE"
        )

        master_file = ""

        duplicate_file = ""

        action = (
            "Verify formatting, signatures, embedded objects "
            "and approval status. If confirmed equivalent, "
            "retain one master and archive/remove the redundant copy."
        )

        deletion_safe = (
            "VERIFY FIRST"
        )

        rationale = (
            "Extracted text content is identical."
        )


    # --------------------------------------------------------
    # CASE 2:
    # EXACT BINARY DUPLICATE
    # --------------------------------------------------------

    elif pair["classification"] == "EXACT FILE DUPLICATE":

        decision = (
            "DUPLICATE CANDIDATE"
        )

        master_file = ""

        duplicate_file = ""

        action = (
            "Files are binary-identical. Confirm preferred "
            "master filename/location before removing redundancy."
        )

        deletion_safe = (
            "VERIFY FIRST"
        )

        rationale = (
            "Files have identical binary content."
        )


    # --------------------------------------------------------
    # CASE 3:
    # VERSION CHECK PROVIDED A LIKELY MASTER
    # --------------------------------------------------------

    elif (
        version_info
        and
        version_info[
            "master_assessment"
        ]
        ==
        "LIKELY MASTER"
    ):

        master_file = (
            version_info[
                "likely_master"
            ]
        )

        if normalise_name(
            master_file
        ) == normalise_name(
            file_a
        ):

            duplicate_file = (
                file_b
            )

        else:

            duplicate_file = (
                file_a
            )

        decision = (
            "KEEP MASTER"
        )

        action = (
            "Retain the likely current master. "
            "Before removing the older version, verify "
            "document approval/signature status."
        )

        deletion_safe = (
            "VERIFY OLDER VERSION FIRST"
        )

        rationale = (
            version_info[
                "reason"
            ]
        )


    # --------------------------------------------------------
    # CASE 4:
    # MANUAL REVIEW REQUIRED
    # --------------------------------------------------------

    elif (
        version_info
        and
        version_info[
            "master_assessment"
        ]
        ==
        "MANUAL REVIEW REQUIRED"
    ):

        decision = (
            "MANUAL REVIEW - DO NOT DELETE"
        )

        master_file = ""
        duplicate_file = ""

        action = (
            "Compare the substantive policy changes and "
            "approval history manually. Both versions must "
            "remain until the current approved policy is confirmed."
        )

        deletion_safe = (
            "NO"
        )

        rationale = (
            version_info[
                "reason"
            ]
        )


    # --------------------------------------------------------
    # CASE 5:
    # UNRESOLVED VERSION
    # --------------------------------------------------------

    elif (
        version_info
        and
        version_info[
            "master_assessment"
        ]
        ==
        "UNRESOLVED"
    ):

        decision = (
            "MANUAL REVIEW - DO NOT DELETE"
        )

        master_file = ""
        duplicate_file = ""

        action = (
            "Keep both files. Review the changed content "
            "and approval/version history before choosing "
            "a master."
        )

        deletion_safe = (
            "NO"
        )

        rationale = (
            version_info[
                "reason"
            ]
        )


    # --------------------------------------------------------
    # FALLBACK
    # --------------------------------------------------------

    else:

        decision = (
            "MANUAL REVIEW - DO NOT DELETE"
        )

        master_file = ""
        duplicate_file = ""

        action = (
            "No sufficiently strong automated evidence "
            "supports removal."
        )

        deletion_safe = (
            "NO"
        )

        rationale = (
            "Automated classification is insufficient."
        )


    # ========================================================
    # EXTRA SAFETY FOR THE THREE PROTECTED TOPICS
    # ========================================================

    topic_key = normalise_name(
        topic
    )

    if topic_key in PROTECTED_TOPICS:

        if decision != (
            "KEEP MASTER"
        ):

            decision = (
                "MANUAL REVIEW - DO NOT DELETE"
            )

            master_file = ""
            duplicate_file = ""

            deletion_safe = (
                "NO"
            )


    decisions.append({
        "topic":
            topic,

        "file_a":
            file_a,

        "file_b":
            file_b,

        "classification":
            pair[
                "classification"
            ],

        "similarity":
            pair[
                "similarity"
            ],

        "last_a":
            pair[
                "last_a"
            ],

        "last_b":
            pair[
                "last_b"
            ],

        "next_a":
            pair[
                "next_a"
            ],

        "next_b":
            pair[
                "next_b"
            ],

        "decision":
            decision,

        "master_file":
            master_file,

        "duplicate_file":
            duplicate_file,

        "deletion_safe":
            deletion_safe,

        "action":
            action,

        "rationale":
            rationale,

        "substantive":
            (
                version_info[
                    "substantive"
                ]
                if version_info
                else ""
            ),

        "metadata_only":
            (
                version_info[
                    "metadata_only"
                ]
                if version_info
                else ""
            ),
    })


# ============================================================
# CREATE FINAL WORKBOOK
# ============================================================

wb = Workbook()


# ============================================================
# SUMMARY
# ============================================================

summary = wb.active
summary.title = "Summary"


summary.append([
    "SAHELI FINAL POLICY CLEAN-UP PLAN",
    ""
])

summary.append([
    "Generated",
    TODAY.strftime("%d/%m/%Y")
])

summary.append([
    "Duplicate/version pairs assessed",
    len(decisions)
])

summary.append([
    "Duplicate candidates",
    sum(
        1
        for item in decisions
        if item["decision"]
        ==
        "DUPLICATE CANDIDATE"
    )
])

summary.append([
    "Likely master identified",
    sum(
        1
        for item in decisions
        if item["decision"]
        ==
        "KEEP MASTER"
    )
])

summary.append([
    "Manual review / protected",
    sum(
        1
        for item in decisions
        if item["decision"]
        ==
        "MANUAL REVIEW - DO NOT DELETE"
    )
])

summary.append([
    "Automatic deletions",
    0
])

summary.append([
    "Important",
    (
        "This workbook is a clean-up decision plan only. "
        "No source policy has been deleted, renamed, "
        "moved or modified."
    )
])


# ============================================================
# CLEAN-UP DECISIONS
# ============================================================

ws = wb.create_sheet(
    "Clean-up Decisions"
)


headers = [
    "Policy / Topic",
    "File A",
    "File B",
    "Original Classification",
    "Content Similarity %",
    "Last Reviewed A",
    "Last Reviewed B",
    "Next Review A",
    "Next Review B",
    "Substantive Change Lines",
    "Metadata Only?",
    "Final Classification",
    "Master File",
    "Duplicate / Older File",
    "Safe to Delete Now?",
    "Required Action",
    "Decision Rationale",
    "Management Confirmation",
]


ws.append(
    headers
)


for item in decisions:

    ws.append([
        item[
            "topic"
        ],

        item[
            "file_a"
        ],

        item[
            "file_b"
        ],

        item[
            "classification"
        ],

        item[
            "similarity"
        ],

        item[
            "last_a"
        ],

        item[
            "last_b"
        ],

        item[
            "next_a"
        ],

        item[
            "next_b"
        ],

        item[
            "substantive"
        ],

        item[
            "metadata_only"
        ],

        item[
            "decision"
        ],

        item[
            "master_file"
        ],

        item[
            "duplicate_file"
        ],

        item[
            "deletion_safe"
        ],

        item[
            "action"
        ],

        item[
            "rationale"
        ],

        "",  # Management confirmation
    ])


# ============================================================
# MASTER FILES SHEET
# ============================================================

master_ws = wb.create_sheet(
    "Likely Masters"
)


master_ws.append([
    "Policy / Topic",
    "Likely Master File",
    "Older / Duplicate File",
    "Basis",
    "Status",
])


for item in decisions:

    if item[
        "decision"
    ] == "KEEP MASTER":

        master_ws.append([
            item[
                "topic"
            ],

            item[
                "master_file"
            ],

            item[
                "duplicate_file"
            ],

            item[
                "rationale"
            ],

            (
                "Verify approval/signature "
                "before removing older version"
            ),
        ])


# ============================================================
# DUPLICATE CANDIDATES SHEET
# ============================================================

duplicate_ws = wb.create_sheet(
    "Duplicate Candidates"
)


duplicate_ws.append([
    "File A",
    "File B",
    "Similarity %",
    "Reason",
    "Required Check",
    "Confirmed Master",
])


for item in decisions:

    if item[
        "decision"
    ] == "DUPLICATE CANDIDATE":

        duplicate_ws.append([
            item[
                "file_a"
            ],

            item[
                "file_b"
            ],

            item[
                "similarity"
            ],

            item[
                "rationale"
            ],

            (
                "Confirm formatting, signatures, "
                "embedded objects and approval status"
            ),

            "",
        ])


# ============================================================
# MANUAL REVIEW SHEET
# ============================================================

manual_ws = wb.create_sheet(
    "Manual Review"
)


manual_ws.append([
    "Policy / Topic",
    "File A",
    "File B",
    "Similarity %",
    "Substantive Change Lines",
    "Reason",
    "Required Action",
    "Status",
])


for item in decisions:

    if item[
        "decision"
    ] == (
        "MANUAL REVIEW - DO NOT DELETE"
    ):

        manual_ws.append([
            item[
                "topic"
            ],

            item[
                "file_a"
            ],

            item[
                "file_b"
            ],

            item[
                "similarity"
            ],

            item[
                "substantive"
            ],

            item[
                "rationale"
            ],

            item[
                "action"
            ],

            "PROTECTED - KEEP BOTH",
        ])


# ============================================================
# FORMATTING
# ============================================================

HEADER_FILL = PatternFill(
    "solid",
    fgColor="1F4E78"
)

HEADER_FONT = Font(
    bold=True,
    color="FFFFFF"
)

GREEN_FILL = PatternFill(
    "solid",
    fgColor="D9EAD3"
)

YELLOW_FILL = PatternFill(
    "solid",
    fgColor="FFF2CC"
)

RED_FILL = PatternFill(
    "solid",
    fgColor="F4CCCC"
)

BLUE_FILL = PatternFill(
    "solid",
    fgColor="D9EAF7"
)

thin_border = Border(
    left=Side(
        style="thin",
        color="D9E1F2"
    ),
    right=Side(
        style="thin",
        color="D9E1F2"
    ),
    top=Side(
        style="thin",
        color="D9E1F2"
    ),
    bottom=Side(
        style="thin",
        color="D9E1F2"
    ),
)


def format_sheet(sheet):

    sheet.freeze_panes = "A2"

    if sheet.max_row > 1:
        sheet.auto_filter.ref = (
            sheet.dimensions
        )

    for cell in sheet[1]:

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    for row in sheet.iter_rows():

        for cell in row:

            cell.border = thin_border

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    for column in sheet.columns:

        maximum = 0

        letter = get_column_letter(
            column[0].column
        )

        for cell in column:

            maximum = max(
                maximum,
                len(
                    clean(
                        cell.value
                    )
                )
            )

        sheet.column_dimensions[
            letter
        ].width = min(
            max(
                maximum + 2,
                12
            ),
            50
        )


for sheet in [
    summary,
    ws,
    master_ws,
    duplicate_ws,
    manual_ws,
]:

    format_sheet(
        sheet
    )


# ============================================================
# COLOUR FINAL CLASSIFICATIONS
# ============================================================

FINAL_CLASSIFICATION_COLUMN = 12


for row_number in range(
    2,
    ws.max_row + 1
):

    cell = ws.cell(
        row=row_number,
        column=FINAL_CLASSIFICATION_COLUMN
    )

    value = clean(
        cell.value
    )

    if value == "KEEP MASTER":

        cell.fill = GREEN_FILL

    elif value == (
        "DUPLICATE CANDIDATE"
    ):

        cell.fill = YELLOW_FILL

    elif value == (
        "MANUAL REVIEW - DO NOT DELETE"
    ):

        cell.fill = RED_FILL


# ============================================================
# ADD TABLE
# ============================================================

if ws.max_row >= 2:

    table_ref = (
        f"A1:"
        f"{get_column_letter(ws.max_column)}"
        f"{ws.max_row}"
    )

    table = Table(
        displayName=(
            "FinalPolicyCleanupTable"
        ),
        ref=table_ref
    )

    table.tableStyleInfo = (
        TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
    )

    ws.add_table(
        table
    )


# ============================================================
# SAVE
# ============================================================

OUTPUT_FOLDER = (
    OUTPUT_FILE.parent
)

OUTPUT_FOLDER.mkdir(
    parents=True,
    exist_ok=True
)

wb.save(
    OUTPUT_FILE
)


# ============================================================
# TERMINAL SUMMARY
# ============================================================

duplicate_count = sum(
    1
    for item in decisions
    if item["decision"]
    ==
    "DUPLICATE CANDIDATE"
)

master_count = sum(
    1
    for item in decisions
    if item["decision"]
    ==
    "KEEP MASTER"
)

manual_count = sum(
    1
    for item in decisions
    if item["decision"]
    ==
    "MANUAL REVIEW - DO NOT DELETE"
)


print()
print("=" * 78)
print("FINAL CLEAN-UP PLAN COMPLETE")
print("=" * 78)

print(
    f"Pairs assessed       : "
    f"{len(decisions)}"
)

print(
    f"Duplicate candidates : "
    f"{duplicate_count}"
)

print(
    f"Likely masters       : "
    f"{master_count}"
)

print(
    f"Manual review        : "
    f"{manual_count}"
)

print(
    "Automatic deletions : 0"
)

print()
print("FINAL DECISIONS")
print("-" * 78)


for item in decisions:

    print()
    print(
        f"{item['file_a']}"
    )

    print(
        f"  vs {item['file_b']}"
    )

    print(
        f"  -> {item['decision']}"
    )

    if item[
        "master_file"
    ]:

        print(
            f"  Master: "
            f"{item['master_file']}"
        )

    if item[
        "duplicate_file"
    ]:

        print(
            f"  Older/duplicate: "
            f"{item['duplicate_file']}"
        )


print()
print("=" * 78)

print(
    "Report created:"
)

print(
    OUTPUT_FILE
)

print()
print(
    "NO POLICY DOCUMENTS WERE "
    "MODIFIED OR DELETED."
)

print("=" * 78)