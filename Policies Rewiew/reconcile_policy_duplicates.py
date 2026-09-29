from pathlib import Path
from datetime import date
from difflib import SequenceMatcher
import hashlib
import re

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# CONFIGURATION
# ============================================================

POLICY_FOLDER = Path(
    r"C:\Users\shonk\Saheli Hub\Saheli Hub - Policies"
)

AUDIT_FILE = Path(
    r"C:\Users\shonk\source\PythonCodes\Saheli_Policy_Final_Audit_20260929.xlsx"
)

OUTPUT_FOLDER = Path(
    r"C:\Users\shonk\source\PythonCodes"
)

TODAY = date.today()

OUTPUT_FILE = OUTPUT_FOLDER / (
    f"Saheli_Duplicate_Reconciliation_{TODAY.strftime('%Y%m%d')}.xlsx"
)


# ============================================================
# SAFETY
# ============================================================

print("=" * 78)
print("SAHELI POLICY DUPLICATE RECONCILIATION - CORRECTED")
print("=" * 78)
print()
print("READ-ONLY MODE")
print("This script WILL NOT:")
print("  - delete documents")
print("  - rename documents")
print("  - move documents")
print("  - edit documents")
print()
print("It only reads documents and creates a new Excel report.")
print()


# ============================================================
# GENERAL HELPERS
# ============================================================

def clean(value):
    if value is None:
        return ""

    value = str(value)
    value = value.replace("\xa0", " ")
    value = value.replace("\u200b", "")
    value = re.sub(r"\s+", " ", value)

    return value.strip()


def filename_lookup_key(filename):
    """
    Used ONLY for finding the real file on disk.

    This fixes differences such as:

    Audit workbook:
        Computers Email Internet Policy1.docx

    Actual Windows filename:
        Computers Email  Internet Policy1.docx

    Multiple spaces, underscores and dash spacing are normalised.
    """

    value = clean(filename).lower()

    value = value.replace("_", " ")

    value = re.sub(
        r"\s*-\s*",
        " - ",
        value
    )

    value = re.sub(
        r"\s+",
        " ",
        value
    )

    return value.strip()


def canonical_policy_name(filename):
    """
    Used for duplicate comparison.

    Removes obvious Windows copy naming.

    DOES NOT remove years such as:
        Policy 23
        Policy 24
        Policy 26

    because those may indicate real versions.
    """

    value = filename_lookup_key(filename)

    value = re.sub(
        r"\.docx$",
        "",
        value,
        flags=re.IGNORECASE
    )

    value = re.sub(
        r"\s*-\s*copy(?:\s*\(\d+\))?$",
        "",
        value,
        flags=re.IGNORECASE
    )

    value = re.sub(
        r"\s+copy(?:\s*\(\d+\))?$",
        "",
        value,
        flags=re.IGNORECASE
    )

    value = re.sub(
        r"\s+",
        " ",
        value
    )

    return value.strip()


def sha256_file(path):
    sha = hashlib.sha256()

    with open(path, "rb") as f:
        while True:
            block = f.read(
                1024 * 1024
            )

            if not block:
                break

            sha.update(block)

    return sha.hexdigest()


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx(path):
    doc = Document(str(path))

    paragraphs = []

    for paragraph in doc.paragraphs:
        value = clean(
            paragraph.text
        )

        if value:
            paragraphs.append(value)

    table_rows = []

    for table in doc.tables:

        for row in table.rows:

            cells = [
                clean(cell.text)
                for cell in row.cells
            ]

            cells = [
                cell
                for cell in cells
                if cell
            ]

            if cells:
                table_rows.append(
                    " | ".join(cells)
                )

    full_text = "\n".join(
        paragraphs + table_rows
    )

    return full_text


def normalise_content(text):
    """
    Normalises text for comparison only.

    Policy body text is NOT written to the output workbook.
    """

    text = text.lower()

    text = text.replace(
        "\xa0",
        " "
    )

    text = text.replace(
        "\u200b",
        ""
    )

    text = re.sub(
        r"[“”]",
        '"',
        text
    )

    text = re.sub(
        r"[‘’]",
        "'",
        text
    )

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


def content_hash(text):
    normalised = normalise_content(
        text
    )

    return hashlib.sha256(
        normalised.encode(
            "utf-8",
            errors="ignore"
        )
    ).hexdigest()


# ============================================================
# METADATA EXTRACTION
# ============================================================

def get_table_rows(path):
    doc = Document(str(path))

    rows = []

    for table in doc.tables:

        for row in table.rows:

            cells = [
                clean(cell.text)
                for cell in row.cells
            ]

            rows.append(cells)

    return rows


def find_value_after_label(
    cells,
    labels
):

    for index, cell in enumerate(
        cells
    ):

        lower = cell.lower()

        for label in labels:

            if label.lower() in lower:

                # Same-cell value
                if ":" in cell:

                    after = clean(
                        cell.split(
                            ":",
                            1
                        )[1]
                    )

                    if after:
                        return after

                # Following-cell value
                for next_index in range(
                    index + 1,
                    len(cells)
                ):

                    candidate = clean(
                        cells[next_index]
                    )

                    if candidate:
                        return candidate

    return ""


def extract_metadata(path):

    result = {
        "version": "",
        "last_reviewed": "",
        "next_review": "",
        "approved_date": "",
        "author": "",
        "ceo_approval": "",
        "board_approval": "",
    }

    rows = get_table_rows(path)

    for cells in rows:

        if not result["version"]:
            result["version"] = (
                find_value_after_label(
                    cells,
                    [
                        "version no.",
                        "version number"
                    ]
                )
            )

        if not result["last_reviewed"]:
            result["last_reviewed"] = (
                find_value_after_label(
                    cells,
                    [
                        "last reviewed",
                        "last review"
                    ]
                )
            )

        if not result["next_review"]:
            result["next_review"] = (
                find_value_after_label(
                    cells,
                    [
                        "next review",
                        "review date",
                        "review due"
                    ]
                )
            )

        if not result["approved_date"]:
            result["approved_date"] = (
                find_value_after_label(
                    cells,
                    [
                        "approved by trustees",
                        "approved date",
                        "date approved"
                    ]
                )
            )

        if not result["author"]:
            result["author"] = (
                find_value_after_label(
                    cells,
                    ["author"]
                )
            )

        if not result["ceo_approval"]:
            result["ceo_approval"] = (
                find_value_after_label(
                    cells,
                    [
                        "authorised by ceo/coo",
                        "authorised by ceo",
                        "authorized by ceo"
                    ]
                )
            )

        if not result["board_approval"]:
            result["board_approval"] = (
                find_value_after_label(
                    cells,
                    [
                        "authorised by board chairperson",
                        "authorized by board chairperson",
                        "board chairperson or vice chair"
                    ]
                )
            )

    return result


# ============================================================
# LOAD DUPLICATE FLAGS FROM FIRST AUDIT
# ============================================================

if not AUDIT_FILE.exists():
    raise FileNotFoundError(
        f"Audit workbook not found:\n"
        f"{AUDIT_FILE}"
    )


audit_wb = load_workbook(
    AUDIT_FILE,
    read_only=True,
    data_only=True
)


if "Policy Library" not in audit_wb.sheetnames:
    raise ValueError(
        "Could not find the 'Policy Library' "
        "sheet in the audit workbook."
    )


audit_ws = audit_wb[
    "Policy Library"
]


headers = {
    clean(cell.value): cell.column
    for cell in audit_ws[1]
}


required_headers = [
    "Document",
    "Duplicate Status",
    "Duplicate / Similar To"
]


for header in required_headers:

    if header not in headers:
        raise ValueError(
            f"Required column missing: "
            f"{header}"
        )


flagged_names = set()


for row in range(
    2,
    audit_ws.max_row + 1
):

    document = clean(
        audit_ws.cell(
            row,
            headers["Document"]
        ).value
    )

    duplicate_status = clean(
        audit_ws.cell(
            row,
            headers[
                "Duplicate Status"
            ]
        ).value
    )

    duplicate_with = clean(
        audit_ws.cell(
            row,
            headers[
                "Duplicate / Similar To"
            ]
        ).value
    )

    if duplicate_status:

        if document:
            flagged_names.add(
                document
            )

        if duplicate_with:

            for name in (
                duplicate_with.split(";")
            ):

                name = clean(name)

                if name:
                    flagged_names.add(
                        name
                    )


audit_wb.close()


print(
    "Documents referenced by "
    f"duplicate flags: {len(flagged_names)}"
)

print()


# ============================================================
# BUILD NORMALISED FILE LOOKUP
# ============================================================

actual_files = list(
    POLICY_FOLDER.rglob("*.docx")
)


file_lookup = {}


for path in actual_files:

    key = filename_lookup_key(
        path.name
    )

    file_lookup.setdefault(
        key,
        []
    ).append(path)


# ============================================================
# MATCH AUDIT NAMES TO REAL FILES
# ============================================================

matched_files = {}
missing_files = []
ambiguous_files = []


for audit_name in sorted(
    flagged_names
):

    key = filename_lookup_key(
        audit_name
    )

    matches = file_lookup.get(
        key,
        []
    )

    if len(matches) == 1:

        matched_files[
            audit_name
        ] = matches[0]

    elif len(matches) == 0:

        missing_files.append(
            audit_name
        )

    else:

        ambiguous_files.append(
            (
                audit_name,
                matches
            )
        )


print(
    f"Successfully matched: "
    f"{len(matched_files)} / "
    f"{len(flagged_names)}"
)


if missing_files:

    print()
    print(
        "WARNING - files still not found:"
    )

    for name in missing_files:
        print(
            f"  {name}"
        )


if ambiguous_files:

    print()
    print(
        "WARNING - ambiguous filename matches:"
    )

    for audit_name, matches in (
        ambiguous_files
    ):

        print(
            f"  {audit_name}"
        )

        for match in matches:
            print(
                f"      -> {match.name}"
            )


print()


# ============================================================
# READ MATCHED DOCUMENTS
# ============================================================

documents = {}


for index, (
    audit_name,
    path
) in enumerate(
    sorted(
        matched_files.items()
    ),
    start=1
):

    print(
        f"[{index:02}/"
        f"{len(matched_files):02}] "
        f"Reading {path.name}"
    )

    try:

        raw_text = extract_docx(
            path
        )

        normalised_text = (
            normalise_content(
                raw_text
            )
        )

        metadata = extract_metadata(
            path
        )

        documents[audit_name] = {
            "audit_name":
                audit_name,

            "actual_filename":
                path.name,

            "path":
                path,

            "raw_file_hash":
                sha256_file(
                    path
                ),

            "content_hash":
                content_hash(
                    raw_text
                ),

            "normalised_text":
                normalised_text,

            "character_count":
                len(
                    normalised_text
                ),

            "word_count":
                len(
                    normalised_text.split()
                ),

            "metadata":
                metadata,

            "error":
                "",
        }

    except Exception as e:

        documents[audit_name] = {
            "audit_name":
                audit_name,

            "actual_filename":
                path.name,

            "path":
                path,

            "raw_file_hash":
                "",

            "content_hash":
                "",

            "normalised_text":
                "",

            "character_count":
                0,

            "word_count":
                0,

            "metadata":
                {},

            "error":
                (
                    f"{type(e).__name__}: "
                    f"{e}"
                ),
        }


# ============================================================
# BUILD COMPARISON PAIRS
# ============================================================

pairs = []


names = sorted(
    documents.keys()
)


for i in range(
    len(names)
):

    for j in range(
        i + 1,
        len(names)
    ):

        name_a = names[i]
        name_b = names[j]

        doc_a = documents[
            name_a
        ]

        doc_b = documents[
            name_b
        ]

        canonical_a = (
            canonical_policy_name(
                doc_a[
                    "actual_filename"
                ]
            )
        )

        canonical_b = (
            canonical_policy_name(
                doc_b[
                    "actual_filename"
                ]
            )
        )

        filename_similarity = (
            SequenceMatcher(
                None,
                canonical_a,
                canonical_b
            ).ratio()
        )

        same_raw_hash = (
            bool(
                doc_a[
                    "raw_file_hash"
                ]
            )
            and
            doc_a[
                "raw_file_hash"
            ]
            ==
            doc_b[
                "raw_file_hash"
            ]
        )

        same_content_hash = (
            bool(
                doc_a[
                    "content_hash"
                ]
            )
            and
            doc_a[
                "content_hash"
            ]
            ==
            doc_b[
                "content_hash"
            ]
        )

        # Candidate if filename/topic is strongly
        # related OR content itself is identical.
        candidate = (
            canonical_a
            ==
            canonical_b

            or

            filename_similarity
            >= 0.82

            or

            same_raw_hash

            or

            same_content_hash
        )


        if not candidate:
            continue


        if (
            same_raw_hash
            or
            same_content_hash
        ):

            content_similarity = 1.0

        else:

            content_similarity = (
                SequenceMatcher(
                    None,
                    doc_a[
                        "normalised_text"
                    ],
                    doc_b[
                        "normalised_text"
                    ]
                ).ratio()
            )


        metadata_a = doc_a[
            "metadata"
        ]

        metadata_b = doc_b[
            "metadata"
        ]


        metadata_differences = []


        fields_to_compare = [
            (
                "version",
                "Version"
            ),
            (
                "last_reviewed",
                "Last Reviewed"
            ),
            (
                "next_review",
                "Next Review"
            ),
            (
                "approved_date",
                "Approved Date"
            ),
            (
                "ceo_approval",
                "CEO/COO Approval"
            ),
            (
                "board_approval",
                "Board Approval"
            ),
        ]


        for key, label in (
            fields_to_compare
        ):

            value_a = clean(
                metadata_a.get(
                    key,
                    ""
                )
            )

            value_b = clean(
                metadata_b.get(
                    key,
                    ""
                )
            )

            if value_a != value_b:

                metadata_differences.append(
                    f"{label}: "
                    f"'{value_a or '[blank]'}' "
                    f"vs "
                    f"'{value_b or '[blank]'}'"
                )


        # ====================================================
        # CLASSIFICATION
        # ====================================================

        if same_raw_hash:

            classification = (
                "EXACT FILE DUPLICATE"
            )

            recommendation = (
                "Binary-identical files. "
                "Strong duplicate candidate. "
                "Confirm preferred master filename "
                "before removing anything."
            )


        elif same_content_hash:

            classification = (
                "IDENTICAL TEXT CONTENT"
            )

            recommendation = (
                "Extracted document text is identical. "
                "Check formatting, signatures and embedded "
                "objects before deciding which copy to retain."
            )


        elif content_similarity >= 0.995:

            classification = (
                "NEAR-IDENTICAL"
            )

            recommendation = (
                "Almost identical content. "
                "Review metadata differences and retain "
                "the confirmed current/approved master."
            )


        elif content_similarity >= 0.97:

            classification = (
                "VERY SIMILAR - "
                "POSSIBLE UPDATED VERSION"
            )

            recommendation = (
                "Likely version pair. "
                "Compare review/approval metadata and "
                "changed wording before selecting master."
            )


        elif content_similarity >= 0.90:

            classification = (
                "SIMILAR - REVIEW VERSIONS"
            )

            recommendation = (
                "Substantial overlap but meaningful "
                "differences exist. Do not delete "
                "automatically."
            )


        else:

            classification = (
                "DIFFERENT CONTENT / SAME TOPIC"
            )

            recommendation = (
                "Related filenames/topics but document "
                "content differs significantly. "
                "Do not treat as an automatic duplicate."
            )


        pairs.append({
            "file_a":
                doc_a[
                    "actual_filename"
                ],

            "file_b":
                doc_b[
                    "actual_filename"
                ],

            "audit_name_a":
                name_a,

            "audit_name_b":
                name_b,

            "classification":
                classification,

            "content_similarity":
                content_similarity,

            "filename_similarity":
                filename_similarity,

            "raw_hash_same":
                same_raw_hash,

            "content_hash_same":
                same_content_hash,

            "version_a":
                metadata_a.get(
                    "version",
                    ""
                ),

            "version_b":
                metadata_b.get(
                    "version",
                    ""
                ),

            "last_reviewed_a":
                metadata_a.get(
                    "last_reviewed",
                    ""
                ),

            "last_reviewed_b":
                metadata_b.get(
                    "last_reviewed",
                    ""
                ),

            "next_review_a":
                metadata_a.get(
                    "next_review",
                    ""
                ),

            "next_review_b":
                metadata_b.get(
                    "next_review",
                    ""
                ),

            "words_a":
                doc_a[
                    "word_count"
                ],

            "words_b":
                doc_b[
                    "word_count"
                ],

            "metadata_differences":
                "; ".join(
                    metadata_differences
                ),

            "recommendation":
                recommendation,
        })


# ============================================================
# SORT PAIRS
# ============================================================

pairs.sort(
    key=lambda item: (
        -item[
            "content_similarity"
        ],
        item[
            "file_a"
        ],
        item[
            "file_b"
        ]
    )
)


# ============================================================
# CREATE OUTPUT WORKBOOK
# ============================================================

wb = Workbook()


# ============================================================
# SUMMARY
# ============================================================

summary_ws = wb.active

summary_ws.title = "Summary"


summary_ws.append([
    "SAHELI DUPLICATE RECONCILIATION",
    ""
])

summary_ws.append([
    "Audit date",
    TODAY.strftime(
        "%d/%m/%Y"
    )
])

summary_ws.append([
    "Documents flagged by first audit",
    len(flagged_names)
])

summary_ws.append([
    "Successfully matched to disk",
    len(matched_files)
])

summary_ws.append([
    "Files not found",
    len(missing_files)
])

summary_ws.append([
    "Ambiguous matches",
    len(ambiguous_files)
])

summary_ws.append([
    "Comparison pairs",
    len(pairs)
])

summary_ws.append([
    "Exact file duplicates",
    sum(
        1
        for pair in pairs
        if pair[
            "classification"
        ]
        ==
        "EXACT FILE DUPLICATE"
    )
])

summary_ws.append([
    "Identical text",
    sum(
        1
        for pair in pairs
        if pair[
            "classification"
        ]
        ==
        "IDENTICAL TEXT CONTENT"
    )
])

summary_ws.append([
    "Near-identical",
    sum(
        1
        for pair in pairs
        if pair[
            "classification"
        ]
        ==
        "NEAR-IDENTICAL"
    )
])

summary_ws.append([
    "Possible updated versions",
    sum(
        1
        for pair in pairs
        if pair[
            "classification"
        ]
        ==
        "VERY SIMILAR - POSSIBLE UPDATED VERSION"
    )
])

summary_ws.append([
    "Similar / review versions",
    sum(
        1
        for pair in pairs
        if pair[
            "classification"
        ]
        ==
        "SIMILAR - REVIEW VERSIONS"
    )
])

summary_ws.append([
    "Different content / same topic",
    sum(
        1
        for pair in pairs
        if pair[
            "classification"
        ]
        ==
        "DIFFERENT CONTENT / SAME TOPIC"
    )
])

summary_ws.append([
    "Safety",
    (
        "No source policy was modified. "
        "No deletion decision is automatic."
    )
])


# ============================================================
# RECONCILIATION SHEET
# ============================================================

ws = wb.create_sheet(
    "Duplicate Reconciliation"
)


ws.append([
    "File A",
    "File B",
    "Classification",
    "Content Similarity %",
    "Filename Similarity %",
    "Same File Hash?",
    "Same Text Hash?",
    "Version A",
    "Version B",
    "Last Reviewed A",
    "Last Reviewed B",
    "Next Review A",
    "Next Review B",
    "Words A",
    "Words B",
    "Metadata Differences",
    "Recommended Action",
    "Decision",
    "Master File",
])


for pair in pairs:

    ws.append([
        pair[
            "file_a"
        ],

        pair[
            "file_b"
        ],

        pair[
            "classification"
        ],

        round(
            pair[
                "content_similarity"
            ] * 100,
            2
        ),

        round(
            pair[
                "filename_similarity"
            ] * 100,
            2
        ),

        (
            "YES"
            if pair[
                "raw_hash_same"
            ]
            else "NO"
        ),

        (
            "YES"
            if pair[
                "content_hash_same"
            ]
            else "NO"
        ),

        pair[
            "version_a"
        ],

        pair[
            "version_b"
        ],

        pair[
            "last_reviewed_a"
        ],

        pair[
            "last_reviewed_b"
        ],

        pair[
            "next_review_a"
        ],

        pair[
            "next_review_b"
        ],

        pair[
            "words_a"
        ],

        pair[
            "words_b"
        ],

        pair[
            "metadata_differences"
        ],

        pair[
            "recommendation"
        ],

        "",  # Decision

        "",  # Master File
    ])


# ============================================================
# FLAGGED DOCUMENTS
# ============================================================

docs_ws = wb.create_sheet(
    "Flagged Documents"
)


docs_ws.append([
    "Audit Name",
    "Actual Filename",
    "Version",
    "Last Reviewed",
    "Next Review",
    "Word Count",
    "File Hash",
    "Text Hash",
    "Read Error",
])


for audit_name in sorted(
    documents.keys()
):

    document = documents[
        audit_name
    ]

    metadata = document[
        "metadata"
    ]

    docs_ws.append([
        audit_name,

        document[
            "actual_filename"
        ],

        metadata.get(
            "version",
            ""
        ),

        metadata.get(
            "last_reviewed",
            ""
        ),

        metadata.get(
            "next_review",
            ""
        ),

        document[
            "word_count"
        ],

        document[
            "raw_file_hash"
        ],

        document[
            "content_hash"
        ],

        document[
            "error"
        ],
    ])


# ============================================================
# MATCH CHECK
# ============================================================

match_ws = wb.create_sheet(
    "Filename Match Check"
)


match_ws.append([
    "Audit Workbook Name",
    "Actual Filename",
    "Match Status",
])


for audit_name in sorted(
    flagged_names
):

    if audit_name in matched_files:

        match_ws.append([
            audit_name,
            matched_files[
                audit_name
            ].name,
            "MATCHED",
        ])

    elif audit_name in missing_files:

        match_ws.append([
            audit_name,
            "",
            "NOT FOUND",
        ])

    else:

        match_ws.append([
            audit_name,
            "",
            "AMBIGUOUS",
        ])


# ============================================================
# FORMATTING
# ============================================================

HEADER_FILL = PatternFill(
    "solid",
    fgColor="1F4E78"
)

HEADER_FONT = Font(
    bold=True,
    color="FFFFFF"
)

GREEN_FILL = PatternFill(
    "solid",
    fgColor="D9EAD3"
)

YELLOW_FILL = PatternFill(
    "solid",
    fgColor="FFF2CC"
)

ORANGE_FILL = PatternFill(
    "solid",
    fgColor="FCE5CD"
)

BLUE_FILL = PatternFill(
    "solid",
    fgColor="D9EAF7"
)

RED_FILL = PatternFill(
    "solid",
    fgColor="F4CCCC"
)

thin_border = Border(
    left=Side(
        style="thin",
        color="D9E1F2"
    ),
    right=Side(
        style="thin",
        color="D9E1F2"
    ),
    top=Side(
        style="thin",
        color="D9E1F2"
    ),
    bottom=Side(
        style="thin",
        color="D9E1F2"
    ),
)


def format_sheet(sheet):

    sheet.freeze_panes = "A2"

    if sheet.max_row > 1:
        sheet.auto_filter.ref = (
            sheet.dimensions
        )

    for cell in sheet[1]:

        cell.fill = (
            HEADER_FILL
        )

        cell.font = (
            HEADER_FONT
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    for row in sheet.iter_rows():

        for cell in row:

            cell.border = (
                thin_border
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    for column in sheet.columns:

        max_length = 0

        letter = get_column_letter(
            column[0].column
        )

        for cell in column:

            max_length = max(
                max_length,
                len(
                    clean(
                        cell.value
                    )
                )
            )

        sheet.column_dimensions[
            letter
        ].width = min(
            max(
                max_length + 2,
                12
            ),
            48
        )


for sheet in [
    summary_ws,
    ws,
    docs_ws,
    match_ws,
]:
    format_sheet(
        sheet
    )


# ============================================================
# COLOUR CLASSIFICATIONS
# ============================================================

for row_number in range(
    2,
    ws.max_row + 1
):

    cell = ws.cell(
        row=row_number,
        column=3
    )

    value = clean(
        cell.value
    )

    if value in {
        "EXACT FILE DUPLICATE",
        "IDENTICAL TEXT CONTENT",
    }:

        cell.fill = (
            GREEN_FILL
        )

    elif value == (
        "NEAR-IDENTICAL"
    ):

        cell.fill = (
            YELLOW_FILL
        )

    elif value in {
        "VERY SIMILAR - POSSIBLE UPDATED VERSION",
        "SIMILAR - REVIEW VERSIONS",
    }:

        cell.fill = (
            ORANGE_FILL
        )

    else:

        cell.fill = (
            BLUE_FILL
        )


# Filename matching colour

for row_number in range(
    2,
    match_ws.max_row + 1
):

    status_cell = (
        match_ws.cell(
            row=row_number,
            column=3
        )
    )

    if (
        status_cell.value
        == "MATCHED"
    ):
        status_cell.fill = (
            GREEN_FILL
        )

    else:
        status_cell.fill = (
            RED_FILL
        )


# ============================================================
# ADD TABLE
# ============================================================

if ws.max_row >= 2:

    table_ref = (
        f"A1:"
        f"{get_column_letter(ws.max_column)}"
        f"{ws.max_row}"
    )

    table = Table(
        displayName=(
            "DuplicateReconciliationTable"
        ),
        ref=table_ref
    )

    table.tableStyleInfo = (
        TableStyleInfo(
            name=(
                "TableStyleMedium2"
            ),
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
    )

    ws.add_table(
        table
    )


# ============================================================
# SAVE
# ============================================================

OUTPUT_FOLDER.mkdir(
    parents=True,
    exist_ok=True
)

wb.save(
    OUTPUT_FILE
)


# ============================================================
# TERMINAL SUMMARY
# ============================================================

print()
print("=" * 78)
print("DUPLICATE RECONCILIATION COMPLETE")
print("=" * 78)

print(
    f"Flagged documents : "
    f"{len(flagged_names)}"
)

print(
    f"Matched to disk   : "
    f"{len(matched_files)}"
)

print(
    f"Not found         : "
    f"{len(missing_files)}"
)

print(
    f"Ambiguous         : "
    f"{len(ambiguous_files)}"
)

print(
    f"Comparison pairs  : "
    f"{len(pairs)}"
)

print(
    "Exact duplicates  :",
    sum(
        1
        for pair in pairs
        if pair[
            "classification"
        ]
        ==
        "EXACT FILE DUPLICATE"
    )
)

print(
    "Identical text    :",
    sum(
        1
        for pair in pairs
        if pair[
            "classification"
        ]
        ==
        "IDENTICAL TEXT CONTENT"
    )
)

print(
    "Near-identical    :",
    sum(
        1
        for pair in pairs
        if pair[
            "classification"
        ]
        ==
        "NEAR-IDENTICAL"
    )
)

print(
    "Updated versions  :",
    sum(
        1
        for pair in pairs
        if pair[
            "classification"
        ]
        ==
        "VERY SIMILAR - POSSIBLE UPDATED VERSION"
    )
)

print(
    "Review versions   :",
    sum(
        1
        for pair in pairs
        if pair[
            "classification"
        ]
        ==
        "SIMILAR - REVIEW VERSIONS"
    )
)

print(
    "Different content :",
    sum(
        1
        for pair in pairs
        if pair[
            "classification"
        ]
        ==
        "DIFFERENT CONTENT / SAME TOPIC"
    )
)

print()
print(
    "Report created:"
)

print(
    OUTPUT_FILE
)

print()
print(
    "NO POLICY DOCUMENTS WERE "
    "MODIFIED OR DELETED."
)

print("=" * 78)