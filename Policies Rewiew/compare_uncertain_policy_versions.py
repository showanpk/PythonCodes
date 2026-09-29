from pathlib import Path
from datetime import date, datetime
from difflib import SequenceMatcher
import hashlib
import re

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# CONFIGURATION
# ============================================================

POLICY_FOLDER = Path(
    r"C:\Users\shonk\Saheli Hub\Saheli Hub - Policies"
)

OUTPUT_FOLDER = Path(
    r"C:\Users\shonk\source\PythonCodes"
)

TODAY = date.today()

OUTPUT_FILE = OUTPUT_FOLDER / (
    f"Saheli_Uncertain_Version_Comparison_{TODAY.strftime('%Y%m%d')}.xlsx"
)


# ============================================================
# SIX UNRESOLVED PAIRS ONLY
# ============================================================

PAIRS = [
    {
        "name": "Absence Management",
        "file_a": "Absence Management Policy - Copy.docx",
        "file_b": "Absence Management Policy.docx",
    },
    {
        "name": "Computers Email Internet",
        "file_a": "Computers Email  Internet Policy1 - Copy.docx",
        "file_b": "Computers Email  Internet Policy1.docx",
    },
    {
        "name": "Disciplinary",
        "file_a": "Disciplinary Policy1 - Copy.docx",
        "file_b": "Disciplinary Policy1.docx",
    },
    {
        "name": "Equal Opportunities",
        "file_a": "Equal Opportunities Policy (002)1 - Copy.docx",
        "file_b": "Equal Opportunities Policy (002)1.docx",
    },
    {
        "name": "Social Media",
        "file_a": "Social Media Policy1.docx",
        "file_b": "Social Media Policy11.docx",
    },
    {
        "name": "Data Protection",
        "file_a": "Data Protection Policy_1 - Copy.docx",
        "file_b": "Data Protection Policy_1.docx",
    },
]


# ============================================================
# SAFETY MESSAGE
# ============================================================

print("=" * 78)
print("SAHELI POLICY VERSION DIFFERENCE CHECKER")
print("=" * 78)
print()
print("LOCAL / READ-ONLY MODE")
print()
print("This script WILL NOT:")
print("  - modify policy documents")
print("  - rename policy documents")
print("  - move policy documents")
print("  - delete policy documents")
print("  - export policy body text")
print()
print("Only comparison statistics and administrative metadata")
print("will be written to the Excel report.")
print()


# ============================================================
# HELPERS
# ============================================================

def clean(value):
    if value is None:
        return ""

    value = str(value)
    value = value.replace("\xa0", " ")
    value = value.replace("\u200b", "")
    value = re.sub(r"\s+", " ", value)

    return value.strip()


def normalize_text(value):
    value = clean(value).lower()

    value = re.sub(r"[“”]", '"', value)
    value = re.sub(r"[‘’]", "'", value)

    value = re.sub(
        r"\s+",
        " ",
        value
    )

    return value.strip()


def sha256_text(value):
    return hashlib.sha256(
        value.encode(
            "utf-8",
            errors="ignore"
        )
    ).hexdigest()


# ============================================================
# METADATA LABELS
# ============================================================

METADATA_TERMS = [
    "author",
    "authorised",
    "authorized",
    "version",
    "last reviewed",
    "next review",
    "review date",
    "review due",
    "approved by",
    "approval",
    "board chair",
    "trustee",
    "effective date",
    "issue date",
]


def is_metadata_line(text):
    lower = normalize_text(text)

    return any(
        term in lower
        for term in METADATA_TERMS
    )


# ============================================================
# DOCUMENT EXTRACTION
# ============================================================

def extract_document(path):
    doc = Document(str(path))

    paragraphs = []

    for paragraph in doc.paragraphs:

        value = clean(
            paragraph.text
        )

        if value:
            paragraphs.append(
                value
            )

    table_rows = []

    for table in doc.tables:

        for row in table.rows:

            cells = [
                clean(cell.text)
                for cell in row.cells
            ]

            cells = [
                value
                for value in cells
                if value
            ]

            if cells:
                table_rows.append(
                    " | ".join(cells)
                )

    all_lines = (
        paragraphs
        +
        table_rows
    )

    return {
        "paragraphs": paragraphs,
        "tables": table_rows,
        "all_lines": all_lines,
    }


# ============================================================
# EXTRACT ADMINISTRATIVE METADATA
# ============================================================

def find_value_after_label(
    cells,
    labels
):
    for index, cell in enumerate(
        cells
    ):

        lower = cell.lower()

        for label in labels:

            if label.lower() in lower:

                # Same-cell value
                if ":" in cell:

                    after = clean(
                        cell.split(
                            ":",
                            1
                        )[1]
                    )

                    if after:
                        return after

                # Next populated cell
                for next_index in range(
                    index + 1,
                    len(cells)
                ):

                    candidate = clean(
                        cells[next_index]
                    )

                    if candidate:
                        return candidate

    return ""


def extract_metadata(path):
    doc = Document(str(path))

    result = {
        "version": "",
        "author": "",
        "last_reviewed": "",
        "next_review": "",
        "approved_date": "",
        "ceo_approval": "",
        "board_approval": "",
    }

    for table in doc.tables:

        for row in table.rows:

            cells = [
                clean(cell.text)
                for cell in row.cells
            ]

            if not result["version"]:
                result["version"] = (
                    find_value_after_label(
                        cells,
                        [
                            "version no.",
                            "version number",
                        ]
                    )
                )

            if not result["author"]:
                result["author"] = (
                    find_value_after_label(
                        cells,
                        ["author"]
                    )
                )

            if not result["last_reviewed"]:
                result["last_reviewed"] = (
                    find_value_after_label(
                        cells,
                        [
                            "last reviewed",
                            "last review",
                        ]
                    )
                )

            if not result["next_review"]:
                result["next_review"] = (
                    find_value_after_label(
                        cells,
                        [
                            "next review",
                            "review date",
                            "review due",
                        ]
                    )
                )

            if not result["approved_date"]:
                result["approved_date"] = (
                    find_value_after_label(
                        cells,
                        [
                            "approved by trustees",
                            "approved date",
                            "date approved",
                        ]
                    )
                )

            if not result["ceo_approval"]:
                result["ceo_approval"] = (
                    find_value_after_label(
                        cells,
                        [
                            "authorised by ceo/coo",
                            "authorised by ceo",
                            "authorized by ceo",
                        ]
                    )
                )

            if not result["board_approval"]:
                result["board_approval"] = (
                    find_value_after_label(
                        cells,
                        [
                            "authorised by board chairperson",
                            "authorized by board chairperson",
                            "board chairperson or vice chair",
                        ]
                    )
                )

    return result


# ============================================================
# DATE PARSING FOR VERSION RECENCY
# ============================================================

MONTHS = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}


def parse_date_for_sorting(value):
    value = clean(value)

    if not value:
        return None

    text = value.lower()
    text = text.replace("(draft)", "")
    text = clean(text)

    formats = [
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%d.%m.%Y",
        "%Y-%m-%d",
        "%d %b %Y",
        "%d %B %Y",
        "%d/%m/%y",
        "%d-%m-%y",
        "%d.%m.%y",
        "%d %b %y",
        "%d %B %y",
    ]

    for fmt in formats:

        try:
            return datetime.strptime(
                text,
                fmt
            ).date()

        except ValueError:
            pass

    match = re.fullmatch(
        r"([a-z]+)\s+(\d{4})",
        text
    )

    if match:

        month_text = (
            match.group(1)
        )

        year = int(
            match.group(2)
        )

        if month_text in MONTHS:

            return date(
                year,
                MONTHS[
                    month_text
                ],
                1
            )

    match = re.fullmatch(
        r"(20\d{2})",
        text
    )

    if match:

        return date(
            int(
                match.group(1)
            ),
            1,
            1
        )

    return None


# ============================================================
# DIFFERENCE ANALYSIS
# ============================================================

def compare_line_sets(
    lines_a,
    lines_b
):
    """
    We do NOT export changed wording.

    We only count changes and determine whether
    changed lines appear to be metadata.
    """

    normalized_a = [
        normalize_text(line)
        for line in lines_a
        if normalize_text(line)
    ]

    normalized_b = [
        normalize_text(line)
        for line in lines_b
        if normalize_text(line)
    ]

    matcher = SequenceMatcher(
        None,
        normalized_a,
        normalized_b
    )

    changed_blocks = 0
    inserted_lines = 0
    deleted_lines = 0
    replaced_lines = 0

    metadata_change_lines = 0
    substantive_change_lines = 0

    for (
        tag,
        i1,
        i2,
        j1,
        j2
    ) in matcher.get_opcodes():

        if tag == "equal":
            continue

        changed_blocks += 1

        old_lines = lines_a[
            i1:i2
        ]

        new_lines = lines_b[
            j1:j2
        ]

        if tag == "insert":

            inserted_lines += (
                j2 - j1
            )

        elif tag == "delete":

            deleted_lines += (
                i2 - i1
            )

        elif tag == "replace":

            replaced_lines += max(
                i2 - i1,
                j2 - j1
            )

        changed_lines = (
            old_lines
            +
            new_lines
        )

        for line in changed_lines:

            if is_metadata_line(
                line
            ):

                metadata_change_lines += 1

            else:

                substantive_change_lines += 1

    return {
        "changed_blocks":
            changed_blocks,

        "inserted_lines":
            inserted_lines,

        "deleted_lines":
            deleted_lines,

        "replaced_lines":
            replaced_lines,

        "metadata_change_lines":
            metadata_change_lines,

        "substantive_change_lines":
            substantive_change_lines,
    }


# ============================================================
# METADATA DIFFERENCES
# ============================================================

def compare_metadata(
    metadata_a,
    metadata_b
):
    differences = []

    fields = [
        (
            "version",
            "Version"
        ),
        (
            "author",
            "Author"
        ),
        (
            "last_reviewed",
            "Last Reviewed"
        ),
        (
            "next_review",
            "Next Review"
        ),
        (
            "approved_date",
            "Approved Date"
        ),
        (
            "ceo_approval",
            "CEO/COO Approval"
        ),
        (
            "board_approval",
            "Board Approval"
        ),
    ]

    for key, label in fields:

        a = clean(
            metadata_a.get(
                key,
                ""
            )
        )

        b = clean(
            metadata_b.get(
                key,
                ""
            )
        )

        if a != b:

            differences.append(
                label
            )

    return differences


# ============================================================
# MASTER VERSION ASSESSMENT
# ============================================================

def assess_master(
    file_a,
    file_b,
    metadata_a,
    metadata_b,
    content_similarity,
    substantive_changes,
):
    """
    This deliberately avoids selecting a master when
    substantive differences are meaningful.
    """

    last_a = parse_date_for_sorting(
        metadata_a.get(
            "last_reviewed",
            ""
        )
    )

    last_b = parse_date_for_sorting(
        metadata_b.get(
            "last_reviewed",
            ""
        )
    )

    next_a = parse_date_for_sorting(
        metadata_a.get(
            "next_review",
            ""
        )
    )

    next_b = parse_date_for_sorting(
        metadata_b.get(
            "next_review",
            ""
        )
    )

    approval_a = bool(
        clean(
            metadata_a.get(
                "ceo_approval",
                ""
            )
        )
        or
        clean(
            metadata_a.get(
                "board_approval",
                ""
            )
        )
    )

    approval_b = bool(
        clean(
            metadata_b.get(
                "ceo_approval",
                ""
            )
        )
        or
        clean(
            metadata_b.get(
                "board_approval",
                ""
            )
        )
    )

    # --------------------------------------------
    # Substantive differences = no automatic master
    # --------------------------------------------

    if (
        substantive_changes > 0
        and
        content_similarity < 0.995
    ):

        return (
            "MANUAL REVIEW REQUIRED",
            "",
            (
                "Substantive content differences detected. "
                "Do not select or delete a version automatically."
            )
        )

    # --------------------------------------------
    # Newer Last Reviewed date
    # --------------------------------------------

    if (
        last_a
        and
        last_b
        and
        last_a != last_b
    ):

        if last_a > last_b:

            return (
                "LIKELY MASTER",
                file_a,
                (
                    "File A has the more recent "
                    "Last Reviewed date."
                )
            )

        return (
            "LIKELY MASTER",
            file_b,
            (
                "File B has the more recent "
                "Last Reviewed date."
            )
        )

    # --------------------------------------------
    # Approval evidence
    # --------------------------------------------

    if (
        approval_a
        !=
        approval_b
    ):

        if approval_a:

            return (
                "LIKELY MASTER",
                file_a,
                (
                    "File A contains approval metadata "
                    "that File B does not."
                )
            )

        return (
            "LIKELY MASTER",
            file_b,
            (
                "File B contains approval metadata "
                "that File A does not."
            )
        )

    # --------------------------------------------
    # Later next-review date
    # --------------------------------------------

    if (
        next_a
        and
        next_b
        and
        next_a != next_b
        and
        content_similarity >= 0.995
    ):

        if next_a > next_b:

            return (
                "LIKELY MASTER",
                file_a,
                (
                    "Content is almost identical and "
                    "File A has the later review date."
                )
            )

        return (
            "LIKELY MASTER",
            file_b,
            (
                "Content is almost identical and "
                "File B has the later review date."
            )
        )

    return (
        "UNRESOLVED",
        "",
        (
            "No sufficiently strong administrative "
            "evidence identifies the master version."
        )
    )


# ============================================================
# RUN COMPARISON
# ============================================================

results = []


for index, pair in enumerate(
    PAIRS,
    start=1
):
    name = pair["name"]

    path_a = (
        POLICY_FOLDER
        /
        pair["file_a"]
    )

    path_b = (
        POLICY_FOLDER
        /
        pair["file_b"]
    )

    print(
        f"[{index}/"
        f"{len(PAIRS)}] "
        f"{name}"
    )

    if not path_a.exists():

        raise FileNotFoundError(
            f"File not found:\n"
            f"{path_a}"
        )

    if not path_b.exists():

        raise FileNotFoundError(
            f"File not found:\n"
            f"{path_b}"
        )

    document_a = (
        extract_document(
            path_a
        )
    )

    document_b = (
        extract_document(
            path_b
        )
    )

    metadata_a = (
        extract_metadata(
            path_a
        )
    )

    metadata_b = (
        extract_metadata(
            path_b
        )
    )

    normalized_a = (
        normalize_text(
            "\n".join(
                document_a[
                    "all_lines"
                ]
            )
        )
    )

    normalized_b = (
        normalize_text(
            "\n".join(
                document_b[
                    "all_lines"
                ]
            )
        )
    )

    similarity = (
        SequenceMatcher(
            None,
            normalized_a,
            normalized_b
        ).ratio()
    )

    difference_stats = (
        compare_line_sets(
            document_a[
                "all_lines"
            ],
            document_b[
                "all_lines"
            ]
        )
    )

    metadata_differences = (
        compare_metadata(
            metadata_a,
            metadata_b
        )
    )

    metadata_only = (
        difference_stats[
            "substantive_change_lines"
        ]
        ==
        0
        and
        difference_stats[
            "changed_blocks"
        ]
        >
        0
    )

    if similarity == 1.0:

        content_assessment = (
            "IDENTICAL EXTRACTED CONTENT"
        )

    elif similarity >= 0.995:

        content_assessment = (
            "NEAR-IDENTICAL"
        )

    elif similarity >= 0.97:

        content_assessment = (
            "VERY SIMILAR"
        )

    elif similarity >= 0.90:

        content_assessment = (
            "SIMILAR WITH MEANINGFUL DIFFERENCES"
        )

    else:

        content_assessment = (
            "SUBSTANTIALLY DIFFERENT"
        )

    (
        master_status,
        likely_master,
        master_reason
    ) = assess_master(
        pair["file_a"],
        pair["file_b"],
        metadata_a,
        metadata_b,
        similarity,
        difference_stats[
            "substantive_change_lines"
        ],
    )

    results.append({
        "policy":
            name,

        "file_a":
            pair["file_a"],

        "file_b":
            pair["file_b"],

        "similarity":
            similarity,

        "assessment":
            content_assessment,

        "changed_blocks":
            difference_stats[
                "changed_blocks"
            ],

        "inserted":
            difference_stats[
                "inserted_lines"
            ],

        "deleted":
            difference_stats[
                "deleted_lines"
            ],

        "replaced":
            difference_stats[
                "replaced_lines"
            ],

        "metadata_change_lines":
            difference_stats[
                "metadata_change_lines"
            ],

        "substantive_change_lines":
            difference_stats[
                "substantive_change_lines"
            ],

        "metadata_only":
            metadata_only,

        "metadata_differences":
            ", ".join(
                metadata_differences
            ),

        "version_a":
            metadata_a.get(
                "version",
                ""
            ),

        "version_b":
            metadata_b.get(
                "version",
                ""
            ),

        "last_reviewed_a":
            metadata_a.get(
                "last_reviewed",
                ""
            ),

        "last_reviewed_b":
            metadata_b.get(
                "last_reviewed",
                ""
            ),

        "next_review_a":
            metadata_a.get(
                "next_review",
                ""
            ),

        "next_review_b":
            metadata_b.get(
                "next_review",
                ""
            ),

        "ceo_a":
            metadata_a.get(
                "ceo_approval",
                ""
            ),

        "ceo_b":
            metadata_b.get(
                "ceo_approval",
                ""
            ),

        "board_a":
            metadata_a.get(
                "board_approval",
                ""
            ),

        "board_b":
            metadata_b.get(
                "board_approval",
                ""
            ),

        "master_status":
            master_status,

        "likely_master":
            likely_master,

        "master_reason":
            master_reason,
    })


# ============================================================
# CREATE SAFE EXCEL REPORT
# ============================================================

wb = Workbook()

ws = wb.active

ws.title = (
    "Version Comparison"
)


headers = [
    "Policy / Topic",
    "File A",
    "File B",
    "Content Similarity %",
    "Content Assessment",
    "Changed Blocks",
    "Inserted Lines",
    "Deleted Lines",
    "Replaced Lines",
    "Metadata Change Lines",
    "Substantive Change Lines",
    "Metadata Only?",
    "Metadata Fields Different",
    "Version A",
    "Version B",
    "Last Reviewed A",
    "Last Reviewed B",
    "Next Review A",
    "Next Review B",
    "Master Assessment",
    "Likely Master",
    "Reason",
    "Final Decision",
]


ws.append(
    headers
)


for result in results:

    ws.append([
        result[
            "policy"
        ],

        result[
            "file_a"
        ],

        result[
            "file_b"
        ],

        round(
            result[
                "similarity"
            ]
            *
            100,
            2
        ),

        result[
            "assessment"
        ],

        result[
            "changed_blocks"
        ],

        result[
            "inserted"
        ],

        result[
            "deleted"
        ],

        result[
            "replaced"
        ],

        result[
            "metadata_change_lines"
        ],

        result[
            "substantive_change_lines"
        ],

        (
            "YES"
            if result[
                "metadata_only"
            ]
            else "NO"
        ),

        result[
            "metadata_differences"
        ],

        result[
            "version_a"
        ],

        result[
            "version_b"
        ],

        result[
            "last_reviewed_a"
        ],

        result[
            "last_reviewed_b"
        ],

        result[
            "next_review_a"
        ],

        result[
            "next_review_b"
        ],

        result[
            "master_status"
        ],

        result[
            "likely_master"
        ],

        result[
            "master_reason"
        ],

        "",  # Final Decision
    ])


# ============================================================
# SUMMARY SHEET
# ============================================================

summary = wb.create_sheet(
    "Summary"
)


summary.append([
    "SAHELI UNCERTAIN POLICY VERSION CHECK",
    ""
])

summary.append([
    "Audit date",
    TODAY.strftime(
        "%d/%m/%Y"
    )
])

summary.append([
    "Pairs checked",
    len(results)
])

summary.append([
    "Likely masters identified",
    sum(
        1
        for result in results
        if result[
            "master_status"
        ]
        ==
        "LIKELY MASTER"
    )
])

summary.append([
    "Manual review required",
    sum(
        1
        for result in results
        if result[
            "master_status"
        ]
        ==
        "MANUAL REVIEW REQUIRED"
    )
])

summary.append([
    "Still unresolved",
    sum(
        1
        for result in results
        if result[
            "master_status"
        ]
        ==
        "UNRESOLVED"
    )
])

summary.append([
    "Privacy",
    (
        "Policy wording is not included in this workbook. "
        "Only comparison counts and administrative metadata "
        "are exported."
    )
])

summary.append([
    "Safety",
    (
        "No source document was modified, renamed, "
        "moved or deleted."
    )
])


# ============================================================
# FORMATTING
# ============================================================

HEADER_FILL = PatternFill(
    "solid",
    fgColor="1F4E78"
)

HEADER_FONT = Font(
    bold=True,
    color="FFFFFF"
)

GREEN_FILL = PatternFill(
    "solid",
    fgColor="D9EAD3"
)

YELLOW_FILL = PatternFill(
    "solid",
    fgColor="FFF2CC"
)

ORANGE_FILL = PatternFill(
    "solid",
    fgColor="FCE5CD"
)

RED_FILL = PatternFill(
    "solid",
    fgColor="F4CCCC"
)

BLUE_FILL = PatternFill(
    "solid",
    fgColor="D9EAF7"
)

thin_border = Border(
    left=Side(
        style="thin",
        color="D9E1F2"
    ),
    right=Side(
        style="thin",
        color="D9E1F2"
    ),
    top=Side(
        style="thin",
        color="D9E1F2"
    ),
    bottom=Side(
        style="thin",
        color="D9E1F2"
    ),
)


def format_sheet(sheet):

    sheet.freeze_panes = "A2"

    if sheet.max_row > 1:
        sheet.auto_filter.ref = (
            sheet.dimensions
        )

    for cell in sheet[1]:

        cell.fill = (
            HEADER_FILL
        )

        cell.font = (
            HEADER_FONT
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    for row in sheet.iter_rows():

        for cell in row:

            cell.border = (
                thin_border
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    for column in sheet.columns:

        maximum = 0

        letter = get_column_letter(
            column[0].column
        )

        for cell in column:

            maximum = max(
                maximum,
                len(
                    clean(
                        cell.value
                    )
                )
            )

        sheet.column_dimensions[
            letter
        ].width = min(
            max(
                maximum + 2,
                12
            ),
            48
        )


format_sheet(
    ws
)

format_sheet(
    summary
)


# ============================================================
# COLOUR MASTER ASSESSMENT
# ============================================================

MASTER_STATUS_COLUMN = 20


for row_number in range(
    2,
    ws.max_row + 1
):

    cell = ws.cell(
        row=row_number,
        column=MASTER_STATUS_COLUMN
    )

    value = clean(
        cell.value
    )

    if value == "LIKELY MASTER":

        cell.fill = (
            GREEN_FILL
        )

    elif value == (
        "MANUAL REVIEW REQUIRED"
    ):

        cell.fill = (
            RED_FILL
        )

    elif value == "UNRESOLVED":

        cell.fill = (
            YELLOW_FILL
        )


# ============================================================
# ADD EXCEL TABLE
# ============================================================

if ws.max_row >= 2:

    ref = (
        f"A1:"
        f"{get_column_letter(ws.max_column)}"
        f"{ws.max_row}"
    )

    table = Table(
        displayName=(
            "UncertainVersionComparison"
        ),
        ref=ref
    )

    table.tableStyleInfo = (
        TableStyleInfo(
            name=(
                "TableStyleMedium2"
            ),
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
    )

    ws.add_table(
        table
    )


# ============================================================
# SAVE
# ============================================================

OUTPUT_FOLDER.mkdir(
    parents=True,
    exist_ok=True
)

wb.save(
    OUTPUT_FILE
)


# ============================================================
# SAFE TERMINAL SUMMARY
# ============================================================

print()
print("=" * 78)
print("VERSION COMPARISON COMPLETE")
print("=" * 78)
print()


for result in results:

    print(
        result[
            "policy"
        ]
    )

    print(
        "  Similarity       : "
        f"{result['similarity'] * 100:.2f}%"
    )

    print(
        "  Changed blocks   : "
        f"{result['changed_blocks']}"
    )

    print(
        "  Metadata only    : "
        f"{'YES' if result['metadata_only'] else 'NO'}"
    )

    print(
        "  Substantive lines: "
        f"{result['substantive_change_lines']}"
    )

    print(
        "  Assessment       : "
        f"{result['master_status']}"
    )

    if result[
        "likely_master"
    ]:

        print(
            "  Likely master    : "
            f"{result['likely_master']}"
        )

    print(
        "  Reason           : "
        f"{result['master_reason']}"
    )

    print()


print("=" * 78)

print(
    "Report created:"
)

print(
    OUTPUT_FILE
)

print()

print(
    "NO POLICY TEXT WAS EXPORTED."
)

print(
    "NO POLICY DOCUMENT WAS MODIFIED OR DELETED."
)

print("=" * 78)