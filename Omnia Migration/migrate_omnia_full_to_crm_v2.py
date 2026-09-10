#!/usr/bin/env python3
"""
Saheli CRM - Omnia Medical Practice historical migration V2

Safety model
------------
* --audit-only: parses Excel only; no DB connection.
* default: DB preview inside one transaction, then ROLLBACK.
* --commit: commits only when no REVIEW_* blockers remain.
* Existing Sessions / Participants / LiteMembers / attendance are reused.
* Existing CRM profile data is never overwritten.

Expected source file (same folder as script by default):
  Register for Exercise - Omnia.xlsx

Source sheets migrated:
  2023 REGISTER  (contains history beginning 16-Nov-2022)
  2024 REGISTER
  2025 REGISTER  (continues through 28-Jan-2026)

The Breakdown sheet is reporting-only and is not migrated.
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys
import uuid
import warnings
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable, Optional

try:
    from openpyxl import load_workbook
except ImportError:
    print("ERROR: openpyxl is required. Run: py -m pip install openpyxl pyodbc", file=sys.stderr)
    raise

warnings.filterwarnings("ignore", message=r"Cell .* is marked as a date but the serial value.*")
warnings.filterwarnings("ignore", message=r"Data Validation extension is not supported.*")

SOURCE_FILE = "Register for Exercise - Omnia.xlsx"
DEFAULT_VENUE = "Omnia Medical Practice"

# Final Session.ActivityName values preserve the distinct historic source activities.
# In particular, Chair Based is NOT Omnia Chair Exercise.
ACTIVITY_ALIASES = {
    "chair": "Chair Based",
    "education": "Education",
    "crochet": "Crochet",
    "party": "Party",
}

# Session.Category is required. When the exact historic ActivityName has no existing
# session template, use only these explicit fallback activity templates for Category
# metadata. The final ActivityName remains the historic value above.
TEMPLATE_ACTIVITY_PREFERENCES = {
    "Chair Based": ["Chair Based", "Chair Based Exercise"],
    "Education": ["Education", "Workshops", "Omnia"],
    "Crochet": ["Crochet", "Crochet for Beginners"],
    "Party": ["Party", "Saheli Social", "Omnia"],
}

# Source contains start times but no end times. Keep deterministic and documented.
DURATION_MINUTES = {
    "Chair Based": 60,
    "Education": 60,
    "Crochet": 60,
    "Party": 60,
}

# Explicit chronological corrections reconstructed from the workbook. The numeric
# Excel dates in these columns have day/month reversed. Key is (sheet, Excel column).
DATE_CORRECTIONS_BY_COL = {
    # 2024 REGISTER
    ("2024 REGISTER", 7): date(2024, 1, 10),
    ("2024 REGISTER", 11): date(2024, 2, 7),
    ("2024 REGISTER", 21): date(2024, 5, 1),

    # 2023 REGISTER - every listed numeric date is a day/month reversal
    ("2023 REGISTER", 12): date(2023, 1, 11),
    ("2023 REGISTER", 13): date(2023, 1, 11),
    ("2023 REGISTER", 19): date(2023, 2, 1),
    ("2023 REGISTER", 20): date(2023, 2, 1),
    ("2023 REGISTER", 21): date(2023, 2, 1),
    ("2023 REGISTER", 22): date(2023, 2, 8),
    ("2023 REGISTER", 23): date(2023, 2, 8),
    ("2023 REGISTER", 24): date(2023, 2, 8),
    ("2023 REGISTER", 31): date(2023, 3, 1),
    ("2023 REGISTER", 32): date(2023, 3, 1),
    ("2023 REGISTER", 33): date(2023, 3, 1),
    ("2023 REGISTER", 34): date(2023, 3, 8),
    ("2023 REGISTER", 35): date(2023, 3, 8),
    ("2023 REGISTER", 36): date(2023, 3, 8),
    ("2023 REGISTER", 46): date(2023, 4, 5),
    ("2023 REGISTER", 47): date(2023, 4, 5),
    ("2023 REGISTER", 48): date(2023, 4, 5),
    ("2023 REGISTER", 56): date(2023, 5, 3),
    ("2023 REGISTER", 57): date(2023, 5, 3),
    ("2023 REGISTER", 58): date(2023, 5, 3),
    ("2023 REGISTER", 59): date(2023, 5, 10),
    ("2023 REGISTER", 60): date(2023, 5, 10),
    ("2023 REGISTER", 61): date(2023, 5, 10),
    ("2023 REGISTER", 69): date(2023, 6, 7),
    ("2023 REGISTER", 70): date(2023, 6, 7),
    ("2023 REGISTER", 75): date(2023, 7, 5),
    ("2023 REGISTER", 76): date(2023, 7, 12),
    ("2023 REGISTER", 79): date(2023, 8, 2),
    ("2023 REGISTER", 80): date(2023, 8, 9),
    ("2023 REGISTER", 84): date(2023, 10, 4),
    ("2023 REGISTER", 85): date(2023, 10, 11),
    ("2023 REGISTER", 88): date(2023, 11, 1),
    ("2023 REGISTER", 89): date(2023, 11, 8),
    ("2023 REGISTER", 93): date(2023, 12, 6),
}

INVALID_TEXT = {"", "#n/a", "#ref!", "#value!", "none", "null", "nan", "0"}

# One source-supported name variant: exact DOB continuity links the 2023 wording
# Farzand Bi to later Farzand Begum / Saheli Card 285.
SOURCE_NAME_DOB_CARD_ALIASES = {
    ("farzand bi", date(1949, 10, 3)): ("Farzand Begum", "285"),
}

# No broad fuzzy exceptions are allowed. Any card/name mismatch below the normal
# threshold remains a REVIEW_PARTICIPANT blocker until explicitly reviewed.
SAFE_REVIEWED_CARD_NAME_REUSE = {}
SAFE_EXACT_FULL_NAME_REUSE = set()

@dataclass
class PersonObservation:
    source_file: str
    sheet: str
    row: int
    name: str
    card: Optional[str] = None
    dob: Optional[date] = None
    postcode: Optional[str] = None
    emergency_name: Optional[str] = None
    emergency_phone: Optional[str] = None
    risk: Optional[str] = None

    @property
    def name_norm(self) -> str:
        return norm_name(self.name)


@dataclass
class AttendanceObservation:
    session: "SourceSession"
    person: PersonObservation
    source_ref: str


@dataclass
class SourceSession:
    source_file: str
    sheet: str
    source_ref: str
    source_date: date
    session_date: date
    source_activity: str
    activity_name: str
    start_time: time
    end_time: time
    time_quality: str
    date_correction: str = ""
    attendance: list[AttendanceObservation] = field(default_factory=list)

    @property
    def key(self) -> tuple:
        return (self.session_date, self.activity_name, self.start_time)


@dataclass
class DbParticipant:
    participant_id: int
    card: str
    full_name: str
    dob: Optional[date]
    postcode: Optional[str]
    mobile: Optional[str]


@dataclass
class DbLite:
    lite_id: str
    membership_id: str
    first_name: str
    last_name: str
    dob: Optional[date]
    phone: Optional[str]
    postcode: Optional[str]

    @property
    def full_name(self) -> str:
        return f"{self.first_name} {self.last_name}".strip()


@dataclass
class ResolvedMember:
    kind: str  # FULL / LITE
    participant_id: Optional[int]
    lite_id: Optional[str]
    display_id: str
    name: str
    card: Optional[str]
    phone: Optional[str]
    action: str
    detail: str


@dataclass
class Action:
    action: str
    severity: str = "INFO"
    source_file: str = ""
    sheet: str = ""
    source_ref: str = ""
    session_date: str = ""
    activity: str = ""
    start_time: str = ""
    source_name: str = ""
    source_card: str = ""
    member_kind: str = ""
    member_display_id: str = ""
    session_id: str = ""
    detail: str = ""


# --------------------------- normalisation helpers ---------------------------

def clean_text(v: Any) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    if s.lower() in INVALID_TEXT:
        return ""
    return re.sub(r"\s+", " ", s).strip()


def norm_name(v: Any) -> str:
    s = clean_text(v).lower()
    s = s.replace("’", "'")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def norm_simple(v: Any) -> str:
    s = clean_text(v).lower()
    return re.sub(r"[^a-z0-9]+", "", s)


def valid_card(v: Any) -> Optional[str]:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        try:
            if float(v) <= 0 or not float(v).is_integer():
                return None
            n = int(v)
            return str(n) if 0 < n <= 999999 else None
        except Exception:
            return None
    s = clean_text(v)
    if re.fullmatch(r"\d+(?:\.0+)?", s):
        n = int(float(s))
        return str(n) if 0 < n <= 999999 else None
    return None

def as_date(v: Any) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(v))).date()
        except Exception:
            return None
    s = clean_text(v)
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d-%m-%y", "%Y-%m-%d", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            pass
    return None

def as_time(v: Any) -> Optional[time]:
    if isinstance(v, datetime):
        return v.time().replace(second=0, microsecond=0)
    if isinstance(v, time):
        return v.replace(second=0, microsecond=0)
    s = clean_text(v).lower().replace(".", "")
    if not s:
        return None
    s = s.replace(" ", "")
    for fmt in ("%I%p", "%I:%M%p", "%I%M%p", "%H:%M", "%H%M"):
        try:
            return datetime.strptime(s, fmt).time()
        except Exception:
            pass
    return None


def plus_minutes(t: time, minutes: int) -> time:
    dt = datetime.combine(date(2000, 1, 1), t) + timedelta(minutes=minutes)
    return dt.time().replace(second=0, microsecond=0)


def safe_phone(v: Any) -> Optional[str]:
    # Emergency numbers only. Do not turn Excel date-formatted garbage into fake participant phones.
    if v is None or isinstance(v, (date, datetime)):
        return None
    s = clean_text(v)
    if not s:
        return None
    digits = re.sub(r"\D", "", s)
    if 9 <= len(digits) <= 15:
        if s.startswith("+"):
            return "+" + digits
        if len(digits) == 10 and not digits.startswith("0"):
            # Source frequently lost a leading zero because Excel stored the number numerically.
            return "0" + digits
        return digits
    return None


def normalize_postcode(v: Any) -> Optional[str]:
    s = clean_text(v).upper().replace(" ", "")
    return s or None


def is_attended(v: Any) -> bool:
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v) == 1.0
    return clean_text(v).lower() in {"yes", "y", "1", "x", "attended", "present", "✓", "✔"}


def map_activity(raw: Any) -> Optional[str]:
    s = norm_name(raw)
    if "chair" in s:
        return ACTIVITY_ALIASES["chair"]
    if "education" in s:
        return ACTIVITY_ALIASES["education"]
    if "crochet" in s:
        return ACTIVITY_ALIASES["crochet"]
    if "party" in s:
        return ACTIVITY_ALIASES["party"]
    return None


def source_start_time(raw_activity: str, activity: str) -> tuple[Optional[time], str]:
    s = clean_text(raw_activity).lower()
    m = re.search(r"(?<!\d)(\d{1,2})(?:[.:](\d{2}))?\s*(am|pm)?\b", s, re.I)
    if m:
        hour = int(m.group(1))
        minute = int(m.group(2) or 0)
        ap = (m.group(3) or "").lower()
        if ap == "pm" and hour < 12:
            hour += 12
        if ap == "am" and hour == 12:
            hour = 0
        if 0 <= hour < 24 and 0 <= minute < 60:
            return time(hour, minute), "SOURCE_START_END_INFERRED"
    if activity == "Crochet":
        return time(11, 0), "INFERRED_FROM_OTHER_2024_CROCHET_LABELS"
    if activity == "Party":
        return time(11, 0), "INFERRED_FROM_REGULAR_OMNIA_WEDNESDAY_SLOT"
    return None, "MISSING"

def name_similarity(a: str, b: str) -> float:
    a, b = norm_name(a), norm_name(b)
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    return SequenceMatcher(None, a, b).ratio()


def split_name(full_name: str) -> tuple[str, str]:
    parts = clean_text(full_name).split()
    if len(parts) < 2:
        return (parts[0] if parts else "Unknown", "Unknown")
    return (" ".join(parts[:-1]), parts[-1])


# ------------------------------- source parser ------------------------------

def parse_sheet_people(path: Path, ws, sheet_name: str) -> tuple[dict[int, PersonObservation], list[PersonObservation], list[Action]]:
    actions: list[Action] = []
    people_by_row: dict[int, PersonObservation] = {}
    people: list[PersonObservation] = []

    if sheet_name == "2023 REGISTER":
        card_col, name_col, dob_col = None, 2, 3
        postcode_col = emergency_name_col = emergency_phone_col = risk_col = None
    elif sheet_name == "2024 REGISTER":
        card_col, name_col, dob_col = 2, 4, 5
        postcode_col = emergency_name_col = emergency_phone_col = risk_col = None
    elif sheet_name == "2025 REGISTER":
        card_col, name_col, dob_col = 2, 4, 5
        postcode_col, emergency_name_col, emergency_phone_col, risk_col = 6, 7, 8, 9
    else:
        raise ValueError(f"Unsupported source sheet {sheet_name}")

    for row in range(4, ws.max_row + 1):
        raw_card = ws.cell(row, card_col).value if card_col else None
        card = valid_card(raw_card) if card_col else None
        name = clean_text(ws.cell(row, name_col).value)

        # Keep a valid card even if Name is blank. Only skip when both are unusable.
        if not name and not card:
            continue

        if card_col and clean_text(raw_card) and not card:
            actions.append(Action(
                action="SOURCE_NONCARD_TEXT_IGNORED", severity="INFO",
                source_file=path.name, sheet=sheet_name, source_ref=f"Row {row}",
                source_name=name, source_card="",
                detail=f"Saheli Card column contains non-card text and was not used as a card: {clean_text(raw_card)[:160]}",
            ))

        p = PersonObservation(
            source_file=path.name,
            sheet=sheet_name,
            row=row,
            name=name,
            card=card,
            dob=as_date(ws.cell(row, dob_col).value),
            postcode=normalize_postcode(ws.cell(row, postcode_col).value) if postcode_col else None,
            emergency_name=(clean_text(ws.cell(row, emergency_name_col).value) or None) if emergency_name_col else None,
            emergency_phone=safe_phone(ws.cell(row, emergency_phone_col).value) if emergency_phone_col else None,
            risk=(clean_text(ws.cell(row, risk_col).value) or None) if risk_col else None,
        )
        people_by_row[row] = p
        people.append(p)
    return people_by_row, people, actions


def parse_register_sheet(path: Path, ws, sheet_name: str) -> tuple[list[SourceSession], list[PersonObservation], list[Action], list[Action]]:
    people_by_row, people, actions = parse_sheet_people(path, ws, sheet_name)
    empty: list[Action] = []
    sessions: list[SourceSession] = []

    start_col = {"2023 REGISTER": 5, "2024 REGISTER": 7, "2025 REGISTER": 10}[sheet_name]

    for col in range(start_col, ws.max_column + 1):
        raw_activity = clean_text(ws.cell(3, col).value)
        if not raw_activity:
            continue
        activity = map_activity(raw_activity)
        if not activity:
            continue

        source_date = as_date(ws.cell(2, col).value)
        corrected_date = DATE_CORRECTIONS_BY_COL.get((sheet_name, col), source_date)
        start_time, time_quality = source_start_time(raw_activity, activity)

        attendance_rows = [row for row, p in people_by_row.items() if is_attended(ws.cell(row, col).value)]
        if not corrected_date or not start_time or not attendance_rows:
            empty.append(Action(
                action="SKIP_EMPTY_SOURCE_SESSION", severity="INFO",
                source_file=path.name, sheet=sheet_name, source_ref=f"Column {col}",
                source_name="", source_card="", activity=activity,
                detail=f"Skipped source column; date={source_date}, start={start_time}, attendance_count={len(attendance_rows)}, source_activity='{raw_activity}'.",
            ))
            continue

        correction = ""
        if source_date and corrected_date != source_date:
            correction = f"{source_date.isoformat()} -> {corrected_date.isoformat()}"
            actions.append(Action(
                action="SOURCE_DATE_CORRECTED", severity="INFO",
                source_file=path.name, sheet=sheet_name, source_ref=f"Column {col}",
                session_date=corrected_date.isoformat(), activity=activity,
                start_time=start_time.strftime("%H:%M"),
                detail=f"Day/month source correction: {correction}; source_activity='{raw_activity}'.",
            ))

        if time_quality.startswith("INFERRED"):
            actions.append(Action(
                action="SOURCE_START_TIME_INFERRED", severity="INFO",
                source_file=path.name, sheet=sheet_name, source_ref=f"Column {col}",
                session_date=corrected_date.isoformat(), activity=activity,
                start_time=start_time.strftime("%H:%M"),
                detail=f"Source label '{raw_activity}' has no explicit time; inferred {start_time.strftime('%H:%M')}.",
            ))

        s = SourceSession(
            source_file=path.name,
            sheet=sheet_name,
            source_ref=f"Column {col}",
            source_date=source_date or corrected_date,
            session_date=corrected_date,
            source_activity=raw_activity,
            activity_name=activity,
            start_time=start_time,
            end_time=plus_minutes(start_time, DURATION_MINUTES[activity]),
            time_quality=time_quality,
            date_correction=correction,
        )
        for row in attendance_rows:
            p = people_by_row[row]
            s.attendance.append(AttendanceObservation(
                session=s, person=p, source_ref=f"{sheet_name}!{ws.cell(row, col).coordinate}"
            ))
        sessions.append(s)

    return sessions, people, actions, empty


def build_source_identity_maps(people: Iterable[PersonObservation]):
    name_card_people: dict[str, list[PersonObservation]] = defaultdict(list)
    card_names: dict[str, set[str]] = defaultdict(set)
    for p in people:
        if p.card and p.name_norm:
            name_card_people[p.name_norm].append(p)
            card_names[p.card].add(p.name_norm)
    return name_card_people, card_names


def years_are_transposed(a: date, b: date) -> bool:
    if a.day != b.day or a.month != b.month:
        return False
    ya, yb = str(a.year)[-2:], str(b.year)[-2:]
    return ya == yb[::-1]


def enrich_missing_cards_profile_aware(people: Iterable[PersonObservation], name_card_people: dict[str, list[PersonObservation]]) -> list[Action]:
    actions: list[Action] = []
    for p in people:
        if p.card or not p.name_norm:
            continue

        alias = SOURCE_NAME_DOB_CARD_ALIASES.get((p.name_norm, p.dob))
        if alias:
            old_name = p.name
            p.name, p.card = alias
            actions.append(Action(
                action="SOURCE_CARD_RESOLVED_BY_NAME_DOB_ALIAS", severity="INFO",
                source_file=p.source_file, sheet=p.sheet, source_ref=f"Row {p.row}",
                source_name=p.name, source_card=p.card,
                detail=f"Source continuity: '{old_name}' with DOB {p.dob} resolves to {p.name} / Saheli Card {p.card}.",
            ))
            continue

        candidates = name_card_people.get(p.name_norm, [])
        cards = {x.card for x in candidates if x.card}
        if not candidates:
            continue
        if len(cards) > 1:
            actions.append(Action(
                action="REVIEW_SOURCE_NAME_MULTIPLE_CARDS", severity="BLOCKER",
                source_file=p.source_file, sheet=p.sheet, source_ref=f"Row {p.row}",
                source_name=p.name,
                detail=f"Exact source name appears against multiple Saheli Cards: {sorted(cards)}.",
            ))
            continue
        if len(cards) != 1:
            continue

        card = next(iter(cards))
        card_dobs = {x.dob for x in candidates if x.dob}
        if p.dob and card_dobs and p.dob not in card_dobs:
            if len(card_dobs) == 1 and years_are_transposed(p.dob, next(iter(card_dobs))):
                target_dob = next(iter(card_dobs))
                actions.append(Action(
                    action="REVIEW_SOURCE_DOB_CONFLICT", severity="BLOCKER",
                    source_file=p.source_file, sheet=p.sheet, source_ref=f"Row {p.row}",
                    source_name=p.name, source_card=card,
                    detail=f"Exact source name maps to card {card}, but no-card DOB {p.dob} conflicts with card-profile DOB {target_dob}; year digits appear transposed. Review before assigning historical attendance.",
                ))
            else:
                actions.append(Action(
                    action="SOURCE_SAME_NAME_DIFFERENT_DOB_KEPT_SEPARATE", severity="INFO",
                    source_file=p.source_file, sheet=p.sheet, source_ref=f"Row {p.row}",
                    source_name=p.name, source_card=card,
                    detail=f"Exact name also appears with card {card}, but DOB {p.dob} differs from card-profile DOB(s) {sorted(card_dobs)}. Kept as a separate no-card identity.",
                ))
            continue

        p.card = card
        actions.append(Action(
            action="SOURCE_CARD_RESOLVED_FROM_OTHER_YEAR", severity="INFO",
            source_file=p.source_file, sheet=p.sheet, source_ref=f"Row {p.row}",
            source_name=p.name, source_card=p.card,
            detail="Exact normalized name has one unique Saheli Card elsewhere in the Omnia source and no conflicting DOB evidence.",
        ))
    return actions


def source_card_conflicts(card_names: dict[str, set[str]]) -> dict[str, set[str]]:
    conflicts: dict[str, set[str]] = {}
    for card, names in card_names.items():
        ns = sorted(n for n in names if n)
        if len(ns) <= 1:
            continue
        min_sim = min(name_similarity(a, b) for i, a in enumerate(ns) for b in ns[i + 1:])
        if min_sim < 0.80:
            conflicts[card] = set(ns)
    return conflicts


def dedupe_source_attendance(sessions: list[SourceSession]) -> list[Action]:
    actions: list[Action] = []
    for s in sessions:
        seen: set[tuple] = set()
        kept: list[AttendanceObservation] = []
        for a in s.attendance:
            p = a.person
            if p.card:
                identity = ("CARD", p.card)
            elif p.name_norm:
                identity = ("NAME_DOB", p.name_norm, p.dob.isoformat() if p.dob else "")
            else:
                kept.append(a)
                continue
            if identity in seen:
                actions.append(Action(
                    action="SKIP_SOURCE_DUPLICATE_ATTENDANCE", severity="INFO",
                    source_file=s.source_file, sheet=s.sheet, source_ref=a.source_ref,
                    session_date=s.session_date.isoformat(), activity=s.activity_name,
                    start_time=s.start_time.strftime("%H:%M"), source_name=p.name,
                    source_card=p.card or "", detail=f"Duplicate source attendance identity {identity}.",
                ))
                continue
            seen.add(identity)
            kept.append(a)
        s.attendance = kept
    return actions


def build_source_distinct_name_dobs(people: Iterable[PersonObservation]) -> dict[str, set[date]]:
    d: dict[str, set[date]] = defaultdict(set)
    for p in people:
        if p.name_norm and p.dob:
            d[p.name_norm].add(p.dob)
    return {k: v for k, v in d.items() if len(v) > 1}


def parse_sources(source_dir: Path):
    path = source_dir / SOURCE_FILE
    if not path.exists():
        raise FileNotFoundError(f"Missing source file: {path}")

    wb = load_workbook(path, data_only=True)
    sessions: list[SourceSession] = []
    all_people: list[PersonObservation] = []
    source_actions: list[Action] = []
    empty: list[Action] = []

    for sheet_name in ("2023 REGISTER", "2024 REGISTER", "2025 REGISTER"):
        if sheet_name not in wb.sheetnames:
            raise RuntimeError(f"Required source sheet '{sheet_name}' is missing from {path.name}")
        ss, pp, aa, ee = parse_register_sheet(path, wb[sheet_name], sheet_name)
        sessions.extend(ss)
        all_people.extend(pp)
        source_actions.extend(aa)
        empty.extend(ee)

    name_card_people, _ = build_source_identity_maps(all_people)
    source_actions.extend(enrich_missing_cards_profile_aware(all_people, name_card_people))
    # Rebuild after enrichment.
    _, card_names = build_source_identity_maps(all_people)
    conflicts = source_card_conflicts(card_names)
    source_actions.extend(dedupe_source_attendance(sessions))

    delivered = [s for s in sessions if s.attendance]

    # Safety guard: delivered source columns must never silently collapse into one CRM key.
    key_groups: dict[tuple, list[SourceSession]] = defaultdict(list)
    for s in delivered:
        key_groups[s.key].append(s)
    for key, group in key_groups.items():
        if len(group) > 1:
            refs = "; ".join(f"{x.sheet} {x.source_ref}" for x in group)
            source_actions.append(Action(
                action="REVIEW_SOURCE_SESSION_DUPLICATE", severity="BLOCKER",
                source_file=group[0].source_file, sheet=group[0].sheet, source_ref=refs,
                session_date=group[0].session_date.isoformat(), activity=group[0].activity_name,
                start_time=group[0].start_time.strftime("%H:%M"),
                detail=f"Multiple delivered source columns share CRM session key {key}: {refs}.",
            ))

    source_distinct_name_dobs = build_source_distinct_name_dobs(all_people)
    return delivered, empty, all_people, conflicts, source_actions, source_distinct_name_dobs


# ----------------------------- database helpers -----------------------------

def get_connection(connection_string: str):
    try:
        import pyodbc
    except ImportError:
        print("ERROR: pyodbc is required for DB preview/commit. Run: py -m pip install pyodbc", file=sys.stderr)
        raise
    return pyodbc.connect(connection_string, autocommit=False)


def fetch_dicts(cursor, sql: str, params: tuple = ()) -> list[dict[str, Any]]:
    cursor.execute(sql, params)
    cols = [c[0] for c in cursor.description]
    return [dict(zip(cols, row)) for row in cursor.fetchall()]


def verify_schema(cursor):
    required = {
        "Participants": {"ParticipantID", "SaheliCardNumber", "FullName", "DateOfBirth", "Postcode", "MobileNumber", "Site", "Notes"},
        "LiteMembers": {"Id", "MembershipId", "FirstName", "LastName", "DateOfBirth", "Phone", "Postcode"},
        "Sessions": {"SessionId", "Frequency", "Category", "VenueName", "ActivityName", "Notes", "IsRecurringWeekly", "DayOfWeek", "SessionDate", "StartTime", "EndTime", "IsBookingRequired", "IsCancelled"},
        "SessionAttendance": {"AttendanceId", "SessionId", "ParticipantId", "SessionName", "SessionDate", "SessionStartTime", "SessionEndTime", "SaheliCardNumber", "Attended", "AttendanceMemberKind", "LiteMemberId", "MemberDisplayId", "MemberName"},
    }
    rows = fetch_dicts(cursor, """
        SELECT t.name AS TableName, c.name AS ColumnName
        FROM sys.tables t
        JOIN sys.columns c ON c.object_id=t.object_id
        WHERE t.name IN ('Participants','LiteMembers','Sessions','SessionAttendance')
    """)
    found: dict[str, set[str]] = defaultdict(set)
    for r in rows:
        found[r["TableName"]].add(r["ColumnName"])
    errors = []
    for table, cols in required.items():
        missing = cols - found.get(table, set())
        if missing:
            errors.append(f"{table}: missing {sorted(missing)}")
    if errors:
        raise RuntimeError("CRM schema does not match expected migration schema: " + "; ".join(errors))


def load_db_people(cursor):
    participants: list[DbParticipant] = []
    for r in fetch_dicts(cursor, """
        SELECT ParticipantID, SaheliCardNumber, FullName, DateOfBirth, Postcode, MobileNumber
        FROM Participants
    """):
        participants.append(DbParticipant(
            participant_id=int(r["ParticipantID"]),
            card=clean_text(r["SaheliCardNumber"]),
            full_name=clean_text(r["FullName"]),
            dob=as_date(r["DateOfBirth"]),
            postcode=normalize_postcode(r["Postcode"]),
            mobile=clean_text(r["MobileNumber"]) or None,
        ))

    lites: list[DbLite] = []
    for r in fetch_dicts(cursor, """
        SELECT Id, MembershipId, FirstName, LastName, DateOfBirth, Phone, Postcode
        FROM LiteMembers
    """):
        lites.append(DbLite(
            lite_id=str(r["Id"]), membership_id=clean_text(r["MembershipId"]),
            first_name=clean_text(r["FirstName"]), last_name=clean_text(r["LastName"]),
            dob=as_date(r["DateOfBirth"]), phone=clean_text(r["Phone"]) or None,
            postcode=normalize_postcode(r["Postcode"]),
        ))
    return participants, lites


def next_lite_number(lites: list[DbLite]) -> int:
    nums = []
    for l in lites:
        m = re.fullmatch(r"LITE-(\d+)", l.membership_id.strip(), re.I)
        if m:
            nums.append(int(m.group(1)))
    return (max(nums) if nums else 0) + 1


def load_session_templates(cursor, venue: str) -> dict[str, dict[str, Any]]:
    templates: dict[str, dict[str, Any]] = {}
    candidate_names = []
    for prefs in TEMPLATE_ACTIVITY_PREFERENCES.values():
        for name in prefs:
            if name not in candidate_names:
                candidate_names.append(name)

    # Prefer same-venue templates. If absent, use the same activity globally.
    for activity in candidate_names:
        rows = fetch_dicts(cursor, """
            SELECT TOP 1 Frequency, Category, SubCategory, ActivityCategory, Capacity, IsBookingRequired,
                         ActivityName, VenueName, SessionProviderId
            FROM Sessions
            WHERE LOWER(LTRIM(RTRIM(VenueName))) = LOWER(LTRIM(RTRIM(?)))
              AND LOWER(LTRIM(RTRIM(ActivityName))) = LOWER(LTRIM(RTRIM(?)))
            ORDER BY CASE WHEN SessionDate IS NULL THEN 1 ELSE 0 END,
                     ABS(DATEDIFF(day, ISNULL(SessionDate, GETDATE()), GETDATE()))
        """, (venue, activity))
        if not rows:
            rows = fetch_dicts(cursor, """
                SELECT TOP 1 Frequency, Category, SubCategory, ActivityCategory, Capacity, IsBookingRequired,
                             ActivityName, VenueName, SessionProviderId
                FROM Sessions
                WHERE LOWER(LTRIM(RTRIM(ActivityName))) = LOWER(LTRIM(RTRIM(?)))
                ORDER BY CASE WHEN SessionDate IS NULL THEN 1 ELSE 0 END,
                         ABS(DATEDIFF(day, ISNULL(SessionDate, GETDATE()), GETDATE()))
            """, (activity,))
        if rows:
            templates[activity] = rows[0]
    return templates


def choose_template_for_activity(templates: dict[str, dict[str, Any]], activity: str) -> tuple[Optional[dict[str, Any]], Optional[str]]:
    for template_name in TEMPLATE_ACTIVITY_PREFERENCES.get(activity, [activity]):
        t = templates.get(template_name)
        if t and clean_text(t.get("Category")):
            return t, template_name
    return None, None


def existing_alias_session_candidates(cursor, venue: str, s: SourceSession) -> list[dict[str, Any]]:
    aliases = [x for x in TEMPLATE_ACTIVITY_PREFERENCES.get(s.activity_name, []) if norm_name(x) != norm_name(s.activity_name)]
    if not aliases:
        return []
    rows: list[dict[str, Any]] = []
    for alias in aliases:
        rows.extend(fetch_dicts(cursor, """
            SELECT SessionId, ActivityName, SessionDate, StartTime, EndTime, VenueName, IsCancelled
            FROM Sessions
            WHERE SessionDate = ?
              AND LOWER(LTRIM(RTRIM(VenueName))) = LOWER(LTRIM(RTRIM(?)))
              AND LOWER(LTRIM(RTRIM(ActivityName))) = LOWER(LTRIM(RTRIM(?)))
            ORDER BY StartTime
        """, (s.session_date, venue, alias)))
    return rows

def existing_session_candidates(cursor, venue: str, s: SourceSession) -> list[dict[str, Any]]:
    return fetch_dicts(cursor, """
        SELECT SessionId, ActivityName, SessionDate, StartTime, EndTime, VenueName, IsCancelled
        FROM Sessions
        WHERE SessionDate = ?
          AND LOWER(LTRIM(RTRIM(VenueName))) = LOWER(LTRIM(RTRIM(?)))
          AND LOWER(LTRIM(RTRIM(ActivityName))) = LOWER(LTRIM(RTRIM(?)))
        ORDER BY StartTime
    """, (s.session_date, venue, s.activity_name))


def _session_candidate_description(candidates: list[dict[str, Any]]) -> str:
    parts = []
    for c in candidates:
        st = as_time(c.get("StartTime"))
        et = as_time(c.get("EndTime"))
        parts.append(
            f"SessionId={c.get('SessionId')} start={st.strftime('%H:%M') if st else 'NULL'} "
            f"end={et.strftime('%H:%M') if et else 'NULL'} cancelled={int(bool(c.get('IsCancelled')))}"
        )
    return "; ".join(parts)


def choose_existing_session(candidates: list[dict[str, Any]], s: SourceSession) -> tuple[Optional[dict[str, Any]], Optional[str]]:
    if not candidates:
        return None, None
    exact = []
    for c in candidates:
        ct = as_time(c["StartTime"])
        if ct == s.start_time:
            exact.append(c)
    if len(exact) == 1:
        return exact[0], None
    if len(exact) > 1:
        return None, (
            "Multiple CRM sessions match same date/activity/start time. "
            f"Source start={s.start_time.strftime('%H:%M')} time_quality={s.time_quality}; "
            f"CRM candidates: {_session_candidate_description(exact)}"
        )
    # Only inferred source times may reuse one same-day/activity candidate automatically.
    # Explicit source times remain distinct; Venue+Activity+Date+StartTime is the session identity.
    if s.time_quality.startswith("INFERRED") and len(candidates) == 1:
        return candidates[0], None
    return None, (
        "CRM has same date/activity at a different or ambiguous time; explicit source time is not auto-merged. "
        f"Source start={s.start_time.strftime('%H:%M')} end={s.end_time.strftime('%H:%M')} "
        f"time_quality={s.time_quality}; CRM candidates: {_session_candidate_description(candidates)}"
    )


def insert_session(cursor, venue: str, s: SourceSession, template: dict[str, Any], template_name: str) -> int:
    notes = (
        f"Historical Omnia migration. Source={s.source_file}; {s.sheet}; {s.source_ref}; "
        f"source_activity={s.source_activity}; time_quality={s.time_quality}; "
        f"end_time_inferred={DURATION_MINUTES[s.activity_name]}min; category_template={template_name}"
    )
    if s.date_correction:
        notes += f"; date_correction={s.date_correction}"
    cursor.execute("""
        INSERT INTO Sessions
            (Frequency, Category, SubCategory, ActivityCategory, VenueName, ActivityName, Notes,
             IsRecurringWeekly, DayOfWeek, SessionDate, ArrivalTime, StartTime, EndTime, Capacity,
             IsBookingRequired, IsCancelled, AssignedStaffId, RecurringSeriesId, SessionProviderId)
        OUTPUT INSERTED.SessionId
        VALUES (?, ?, ?, ?, ?, ?, ?, 0, NULL, ?, NULL, ?, ?, ?, ?, 0, NULL, NULL, NULL)
    """, (
        clean_text(template.get("Frequency")) or "Weekly",
        clean_text(template.get("Category")),
        clean_text(template.get("SubCategory")) or None,
        clean_text(template.get("ActivityCategory")) or None,
        venue, s.activity_name, notes,
        s.session_date, s.start_time, s.end_time,
        template.get("Capacity"), bool(template.get("IsBookingRequired") or False),
    ))
    return int(cursor.fetchone()[0])

def find_existing_attendance(cursor, session_id: int, member: ResolvedMember) -> Optional[dict[str, Any]]:
    if member.kind == "FULL":
        rows = fetch_dicts(cursor, """
            SELECT TOP 1 AttendanceId, Attended FROM SessionAttendance
            WHERE SessionId=? AND ParticipantId=?
            ORDER BY AttendanceId
        """, (session_id, member.participant_id))
    else:
        rows = fetch_dicts(cursor, """
            SELECT TOP 1 AttendanceId, Attended FROM SessionAttendance
            WHERE SessionId=? AND LiteMemberId=?
            ORDER BY AttendanceId
        """, (session_id, member.lite_id))
    return rows[0] if rows else None


def insert_attendance(cursor, session_id: int, s: SourceSession, p: PersonObservation, m: ResolvedMember):
    cursor.execute("""
        INSERT INTO SessionAttendance
            (SessionId, ParticipantId, SessionName, SessionDay, SessionDate, SessionMonth,
             SessionStartTime, SessionEndTime, SaheliCardNumber, RiskStratification, Attended,
             Notes, AttendanceMemberKind, LiteMemberId, MemberDisplayId, MemberName, Phone,
             EmergencyName, EmergencyPhone)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        session_id,
        m.participant_id if m.kind == "FULL" else None,
        s.activity_name,
        s.session_date.strftime("%A"),
        s.session_date,
        s.session_date.strftime("%B"),
        s.start_time,
        s.end_time,
        m.card if m.kind == "FULL" else None,
        p.risk,
        f"Historical Omnia migration; source={s.source_file}; {p.sheet} row {p.row}; source_activity={s.source_activity}",
        m.kind,
        m.lite_id if m.kind == "LITE" else None,
        m.display_id,
        m.name or p.name,
        m.phone,
        p.emergency_name,
        p.emergency_phone,
    ))


# ----------------------------- member resolution ----------------------------

def resolve_member(
    cursor,
    p: PersonObservation,
    participants_by_card: dict[str, DbParticipant],
    participants_by_name: dict[str, list[DbParticipant]],
    lites_by_name: dict[str, list[DbLite]],
    source_conflicting_cards: dict[str, set[str]],
    source_distinct_name_dobs: dict[str, set[date]],
    next_lite: list[int],
    venue: str,
    cache: dict[tuple, ResolvedMember],
) -> tuple[Optional[ResolvedMember], Optional[str]]:
    cache_key = ("CARD", p.card) if p.card else ("NAME", p.name_norm, p.dob, p.postcode)
    if cache_key in cache:
        return cache[cache_key], None

    if p.card:
        existing = participants_by_card.get(p.card)
        if existing:
            if p.name_norm and existing.full_name:
                sim = name_similarity(p.name, existing.full_name)
                if sim < 0.80:
                    return None, (
                        f"Saheli Card {p.card} belongs to CRM ParticipantID={existing.participant_id} "
                        f"'{existing.full_name}', but source row says '{p.name}' (name similarity {sim:.2f}). "
                        f"CRM DOB={existing.dob or 'NULL'}, postcode={existing.postcode or 'NULL'}, "
                        f"mobile={existing.mobile or 'NULL'}; source DOB={p.dob or 'NULL'}, "
                        f"postcode={p.postcode or 'NULL'}."
                    )
            m = ResolvedMember(
                kind="FULL", participant_id=existing.participant_id, lite_id=None,
                display_id=existing.card, name=existing.full_name or p.name,
                card=existing.card, phone=existing.mobile,
                action="REUSE_FULL_BY_CARD", detail=f"ParticipantID={existing.participant_id}",
            )
            cache[cache_key] = m
            return m, None

        if p.card in source_conflicting_cards:
            return None, f"Source uses card {p.card} against materially different names: {sorted(source_conflicting_cards[p.card])}"

        if not clean_text(p.name):
            return None, f"Saheli Card {p.card} is not in CRM and this source row has no usable name; refusing to create a nameless FULL participant."

        cursor.execute("""
            INSERT INTO Participants (SaheliCardNumber, FullName, DateOfBirth, Postcode, Site, Notes)
            OUTPUT INSERTED.ParticipantID
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            p.card, p.name or None, p.dob, p.postcode, venue,
            f"Created by historical Omnia migration from {p.source_file}; {p.sheet} row {p.row}. Existing CRM data is never overwritten."
        ))
        pid = int(cursor.fetchone()[0])
        dbp = DbParticipant(pid, p.card, p.name, p.dob, p.postcode, None)
        participants_by_card[p.card] = dbp
        participants_by_name[norm_name(p.name)].append(dbp)
        m = ResolvedMember(
            kind="FULL", participant_id=pid, lite_id=None, display_id=p.card,
            name=p.name, card=p.card, phone=None,
            action="CREATE_FULL_FROM_VALID_CARD", detail=f"New ParticipantID={pid}",
        )
        cache[cache_key] = m
        return m, None

    # No source card: reuse FULL only when exact normalized name plus DOB/postcode
    # provides a unique profile match. Name alone is not enough.
    full_candidates = participants_by_name.get(p.name_norm, [])
    safe_full: dict[int, DbParticipant] = {}
    for c in full_candidates:
        if p.dob and c.dob and p.dob == c.dob:
            safe_full[c.participant_id] = c
        elif p.postcode and c.postcode and p.postcode == c.postcode:
            safe_full[c.participant_id] = c
    if len(safe_full) == 1:
        c = next(iter(safe_full.values()))
        m = ResolvedMember(
            kind="FULL", participant_id=c.participant_id, lite_id=None,
            display_id=c.card, name=c.full_name or p.name, card=c.card, phone=c.mobile,
            action="REUSE_FULL_BY_NAME_AND_PROFILE", detail=f"ParticipantID={c.participant_id}",
        )
        cache[cache_key] = m
        return m, None
    if len(safe_full) > 1:
        return None, f"Multiple FULL CRM participants match name/profile for '{p.name}'."

    def source_knows_distinct_dobs(candidate_dob: Optional[date]) -> bool:
        known = source_distinct_name_dobs.get(p.name_norm, set())
        return bool(p.dob and candidate_dob and p.dob in known and candidate_dob in known and p.dob != candidate_dob)

    lite_candidates = lites_by_name.get(p.name_norm, [])
    distinct_lite_note = ""
    if len(lite_candidates) == 1:
        c = lite_candidates[0]
        if p.dob and c.dob and p.dob != c.dob:
            # A concrete conflicting DOB is positive evidence these are distinct people.
            # Do not merge them and do not block the migration merely because the name is common.
            distinct_lite_note = (
                f"Existing exact-name Lite {c.membership_id} has DOB {c.dob}, while source DOB is {p.dob}; "
                "created a separate Lite identity rather than merging conflicting DOBs. "
            )
            lite_candidates = []
        else:
            # If DOB agrees, a historical postcode difference is not enough to split
            # the person because addresses can change over a multi-year register.
            if (not p.dob or not c.dob) and p.postcode and c.postcode and p.postcode != c.postcode:
                return None, f"Exact Lite name match '{p.name}' has conflicting postcode ({p.postcode} vs {c.postcode}) and no matching DOB to confirm identity."
            m = ResolvedMember(
                kind="LITE", participant_id=None, lite_id=c.lite_id,
                display_id=c.membership_id, name=c.full_name, card=None, phone=c.phone,
                action="REUSE_LITE_BY_EXACT_NAME", detail=f"LiteMemberId={c.lite_id}",
            )
            cache[cache_key] = m
            return m, None

    if len(lite_candidates) > 1:
        narrowed: dict[str, DbLite] = {}
        if p.dob:
            for c in lite_candidates:
                if c.dob and c.dob == p.dob:
                    narrowed[c.lite_id] = c
        if not narrowed and p.postcode:
            for c in lite_candidates:
                if c.postcode and c.postcode == p.postcode:
                    narrowed[c.lite_id] = c
        if len(narrowed) == 1:
            c = next(iter(narrowed.values()))
            m = ResolvedMember(
                kind="LITE", participant_id=None, lite_id=c.lite_id,
                display_id=c.membership_id, name=c.full_name, card=None, phone=c.phone,
                action="REUSE_LITE_BY_NAME_AND_PROFILE", detail=f"LiteMemberId={c.lite_id}",
            )
            cache[cache_key] = m
            return m, None
        if len(narrowed) > 1:
            return None, f"Multiple LiteMembers share '{p.name}' with the same available source profile."
        # If Omnia source explicitly contains multiple DOBs for the same name and no
        # existing Lite matches this DOB, creating another Lite is safer than merging.
        known = source_distinct_name_dobs.get(p.name_norm, set())
        if not (p.dob and p.dob in known and len(known) > 1):
            return None, f"Multiple LiteMembers share the exact name '{p.name}' and source profile cannot resolve uniquely."

    first, last = split_name(p.name)
    lite_id = str(uuid.uuid4())
    membership_id = f"LITE-{next_lite[0]}"
    next_lite[0] += 1
    cursor.execute("""
        INSERT INTO LiteMembers
            (Id, MembershipId, FirstName, LastName, DateOfBirth, Phone, Email, Address, Postcode,
             EmergencyName, EmergencyPhone, EmergencyRelation, HealthConditions, Gender, Ethnicity, CreatedByUserId)
        VALUES (?, ?, ?, ?, ?, NULL, NULL, NULL, ?, ?, ?, NULL, NULL, NULL, NULL, NULL)
    """, (
        lite_id, membership_id, first, last, p.dob, p.postcode,
        p.emergency_name, p.emergency_phone,
    ))
    c = DbLite(lite_id, membership_id, first, last, p.dob, None, p.postcode)
    lites_by_name[p.name_norm].append(c)
    # Preserve rerun matching for one-token names stored with LastName=Unknown.
    if norm_name(last) == "unknown" and norm_name(first) != p.name_norm:
        lites_by_name[p.name_norm].append(c)
    m = ResolvedMember(
        kind="LITE", participant_id=None, lite_id=lite_id, display_id=membership_id,
        name=c.full_name, card=None, phone=None,
        action="CREATE_LITE", detail=f"{distinct_lite_note}New LiteMemberId={lite_id}",
    )
    cache[cache_key] = m
    return m, None


# --------------------------------- reports ----------------------------------

def action_from(s: SourceSession, p: Optional[PersonObservation], action: str, detail: str, severity="INFO", **kwargs) -> Action:
    return Action(
        action=action,
        severity=severity,
        source_file=s.source_file,
        sheet=s.sheet,
        source_ref=(kwargs.get("source_ref") or (f"Row {p.row}" if p else s.source_ref)),
        session_date=s.session_date.isoformat(),
        activity=s.activity_name,
        start_time=s.start_time.strftime("%H:%M"),
        source_name=p.name if p else "",
        source_card=(p.card or "") if p else "",
        member_kind=kwargs.get("member_kind", ""),
        member_display_id=kwargs.get("member_display_id", ""),
        session_id=str(kwargs.get("session_id", "") or ""),
        detail=detail,
    )


def write_actions(actions: list[Action], out_dir: Path, prefix: str) -> Path:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = out_dir / f"{prefix}_{ts}.csv"
    fields = list(Action.__dataclass_fields__.keys())
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for a in actions:
            w.writerow({k: getattr(a, k) for k in fields})
    return path


def write_review_people(actions: list[Action], out_dir: Path) -> Optional[Path]:
    rows = [a for a in actions if a.action == "REVIEW_PARTICIPANT"]
    if not rows:
        return None
    grouped: dict[tuple[str, str, str], list[Action]] = defaultdict(list)
    for a in rows:
        grouped[((a.source_card or "").strip(), norm_name(a.source_name), a.detail)].append(a)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = out_dir / f"omnia_participant_review_{ts}.csv"
    fields = ["attendance_rows", "source_card", "source_name", "reason", "example_date", "example_activity", "example_source_ref"]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for (card, _name_norm, reason), items in sorted(grouped.items(), key=lambda kv: (-len(kv[1]), kv[0][0], kv[0][1])):
            ex = items[0]
            w.writerow({
                "attendance_rows": len(items),
                "source_card": card,
                "source_name": ex.source_name,
                "reason": reason,
                "example_date": ex.session_date,
                "example_activity": ex.activity,
                "example_source_ref": ex.source_ref,
            })
    return path


def write_review_sessions(actions: list[Action], out_dir: Path) -> Optional[Path]:
    rows = [a for a in actions if a.action in {
        "REVIEW_SESSION_AMBIGUOUS",
        "REVIEW_SESSION_ACTIVITY_ALIAS_COLLISION",
        "REVIEW_EXISTING_SESSION_CANCELLED",
    }]
    if not rows:
        return None
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = out_dir / f"omnia_session_review_{ts}.csv"
    fields = ["action", "session_date", "activity", "source_start_time", "source_ref", "reason_and_crm_candidates"]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for a in sorted(rows, key=lambda x: (x.session_date, x.activity, x.start_time, x.source_ref)):
            w.writerow({
                "action": a.action,
                "session_date": a.session_date,
                "activity": a.activity,
                "source_start_time": a.start_time,
                "source_ref": a.source_ref,
                "reason_and_crm_candidates": a.detail,
            })
    return path


def write_summary(actions: list[Action], sessions: list[SourceSession], empty: list[SourceSession], out_dir: Path, mode: str) -> Path:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = out_dir / f"omnia_migration_summary_{ts}.txt"
    counts = Counter(a.action for a in actions)
    total_att = sum(len(s.attendance) for s in sessions)
    dates = [s.session_date for s in sessions]
    blockers = [a for a in actions if a.action.startswith("REVIEW_") or a.severity == "BLOCKER"]
    with path.open("w", encoding="utf-8") as f:
        f.write("Saheli CRM - Omnia Medical Practice historical migration\n")
        f.write("====================================================\n\n")
        f.write(f"Mode: {mode}\n")
        f.write(f"Delivered source sessions: {len(sessions)}\n")
        f.write(f"Empty source placeholders skipped: {len(empty)}\n")
        f.write(f"Clean source attendance rows: {total_att}\n")
        if dates:
            f.write(f"Source period: {min(dates)} to {max(dates)}\n")
        f.write(f"Review blockers: {len(blockers)}\n")
        unique_review_people = {((a.source_card or '').strip(), norm_name(a.source_name)) for a in blockers if a.action == "REVIEW_PARTICIPANT"}
        unique_review_people.discard(("", ""))
        unique_new_lites = {a.member_display_id for a in actions if a.action == "CREATE_LITE" and a.member_display_id}
        unique_new_full = {a.member_display_id for a in actions if a.action == "CREATE_FULL_FROM_VALID_CARD" and a.member_display_id}
        f.write(f"Unique participant identities needing REVIEW_PARTICIPANT: {len(unique_review_people)}\n")
        f.write(f"Unique new Lite members proposed: {len(unique_new_lites)}\n")
        f.write(f"Unique new FULL participants proposed: {len(unique_new_full)}\n\n")
        f.write("Action counts (trace rows; member-resolution actions can repeat per attendance)\n--------------------------------------------------------------------------\n")
        for k, v in sorted(counts.items()):
            f.write(f"{k}: {v}\n")
        if blockers:
            f.write("\nBLOCKERS - do not commit until resolved\n---------------------------------------\n")
            for a in blockers:
                f.write(f"{a.session_date} | {a.activity} | {a.source_name} | {a.source_card} | {a.action} | {a.detail}\n")
    return path


def print_source_summary(sessions, empty, source_actions, conflicts):
    total = sum(len(s.attendance) for s in sessions)
    dates = [s.session_date for s in sessions]
    by_activity = defaultdict(lambda: [0, 0])
    for s in sessions:
        by_activity[s.activity_name][0] += 1
        by_activity[s.activity_name][1] += len(s.attendance)
    print("\nOmnia source audit")
    print("-----------------------")
    print(f"Delivered sessions: {len(sessions)}")
    print(f"Empty placeholders skipped: {len(empty)}")
    print(f"Clean attendance rows: {total}")
    if dates:
        print(f"Period: {min(dates)} to {max(dates)}")
    for activity, (sc, ac) in sorted(by_activity.items()):
        print(f"  {activity}: {sc} sessions, {ac} attendance")
    print(f"Source card conflicts requiring CRM validation: {len(conflicts)}")
    print(f"Source resolution/dedup actions: {len(source_actions)}")


# -------------------------------- migration ---------------------------------

def run_db_preview_or_commit(args, sessions, empty, source_conflicts, source_actions, source_distinct_name_dobs):
    conn_str = args.connection_string or os.getenv("SAHELI_SQL_CONNECTION_STRING", "")
    if not conn_str:
        raise RuntimeError(
            "No SQL connection string. Set SAHELI_SQL_CONNECTION_STRING or pass --connection-string."
        )
    conn = get_connection(conn_str)
    cursor = conn.cursor()
    actions = list(source_actions)
    try:
        verify_schema(cursor)

        # Verify exact Omnia venue exists before any creation.
        venues = fetch_dicts(cursor, """
            SELECT DISTINCT VenueName FROM Sessions
            WHERE LOWER(LTRIM(RTRIM(VenueName))) = LOWER(LTRIM(RTRIM(?)))
        """, (args.venue,))
        if not venues:
            if norm_name(args.venue) == norm_name(DEFAULT_VENUE):
                actions.append(Action(
                    action="VENUE_NO_EXISTING_SESSION_YET", severity="INFO",
                    detail=f"No existing Sessions row currently uses '{args.venue}'. This is allowed for the known Omnia Medical Practice venue; new historical sessions will use the exact configured name."
                ))
            else:
                actions.append(Action(
                    action="REVIEW_VENUE_NOT_FOUND", severity="BLOCKER",
                    detail=f"No existing Sessions.VenueName exactly matches custom venue '{args.venue}'. Review before creating sessions under a new venue string."
                ))

        participants, lites = load_db_people(cursor)
        participants_by_card = {p.card: p for p in participants if p.card}
        participants_by_name: dict[str, list[DbParticipant]] = defaultdict(list)
        for p in participants:
            if norm_name(p.full_name):
                participants_by_name[norm_name(p.full_name)].append(p)
        lites_by_name: dict[str, list[DbLite]] = defaultdict(list)
        for l in lites:
            full_key = norm_name(l.full_name)
            if full_key:
                lites_by_name[full_key].append(l)

            # Historical source occasionally contains only a single given/name token
            # (e.g. "Laura" or "Muk"). split_name() must store these Lite members
            # with LastName="Unknown" because LiteMembers.LastName is required.
            # On a later migration run, index that migration placeholder back under
            # the original one-token source name as an alias. If more than one such
            # Lite exists, the normal duplicate-name safeguards below still block
            # an ambiguous reuse rather than guessing.
            if norm_name(l.last_name) == "unknown" and norm_name(l.first_name):
                alias_key = norm_name(l.first_name)
                if alias_key != full_key:
                    lites_by_name[alias_key].append(l)
        next_lite = [next_lite_number(lites)]
        member_cache: dict[tuple, ResolvedMember] = {}
        templates = load_session_templates(cursor, args.venue)

        for activity in sorted({s.activity_name for s in sessions}):
            template, template_name = choose_template_for_activity(templates, activity)
            if not template:
                actions.append(Action(
                    action="REVIEW_MISSING_SESSION_TEMPLATE", severity="BLOCKER",
                    activity=activity,
                    detail=f"Cannot safely create new '{activity}' sessions because none of the approved template activities {TEMPLATE_ACTIVITY_PREFERENCES.get(activity, [activity])} provides a non-null Category. Existing exact matching sessions can still be reused."
                ))

        session_cache: dict[tuple, Optional[int]] = {}
        # Tracks attendance inserted earlier in THIS transaction so it is not mislabeled as pre-existing CRM data.
        created_attendance_keys: set[tuple] = set()
        for s in sorted(sessions, key=lambda x: (x.session_date, x.start_time, x.activity_name)):
            session_id: Optional[int] = None
            if s.key in session_cache:
                session_id = session_cache[s.key]
            else:
                candidates = existing_session_candidates(cursor, args.venue, s)
                chosen, problem = choose_existing_session(candidates, s)
                if problem:
                    actions.append(action_from(s, None, "REVIEW_SESSION_AMBIGUOUS", problem, severity="BLOCKER"))
                    session_cache[s.key] = None
                elif chosen:
                    session_id = int(chosen["SessionId"])
                    if bool(chosen.get("IsCancelled")):
                        actions.append(action_from(s, None, "REVIEW_EXISTING_SESSION_CANCELLED", f"SessionId={session_id} is cancelled in CRM.", severity="BLOCKER", session_id=session_id))
                        session_id = None
                    else:
                        actions.append(action_from(s, None, "REUSE_SESSION", f"Existing SessionId={session_id}", session_id=session_id))
                        session_cache[s.key] = session_id
                else:
                    alias_candidates = existing_alias_session_candidates(cursor, args.venue, s)
                    alias_exact = [c for c in alias_candidates if as_time(c.get("StartTime")) == s.start_time]
                    if alias_exact:
                        desc = "; ".join(f"SessionId={c['SessionId']} ActivityName={c['ActivityName']} Start={as_time(c.get('StartTime'))}" for c in alias_exact)
                        actions.append(action_from(
                            s, None, "REVIEW_SESSION_ACTIVITY_ALIAS_COLLISION",
                            f"CRM already has same venue/date/start under an alternate related activity name: {desc}. Review before creating '{s.activity_name}'.",
                            severity="BLOCKER"
                        ))
                        session_cache[s.key] = None
                    else:
                        template, template_name = choose_template_for_activity(templates, s.activity_name)
                        if not template or not template_name:
                            actions.append(action_from(s, None, "REVIEW_CANNOT_CREATE_SESSION", "No safe session Category template is available.", severity="BLOCKER"))
                            session_cache[s.key] = None
                        else:
                            session_id = insert_session(cursor, args.venue, s, template, template_name)
                            actions.append(action_from(s, None, "CREATE_SESSION", f"New SessionId={session_id}; ActivityName='{s.activity_name}'; Category copied from approved template '{template_name}'; source start preserved; end time inferred.", session_id=session_id))
                            session_cache[s.key] = session_id

            if not session_id:
                for a in s.attendance:
                    actions.append(action_from(s, a.person, "REVIEW_ATTENDANCE_BLOCKED_BY_SESSION", "Attendance cannot be processed until its session is resolved.", severity="BLOCKER", source_ref=a.source_ref))
                continue

            for a in s.attendance:
                p = a.person
                member, problem = resolve_member(
                    cursor, p, participants_by_card, participants_by_name, lites_by_name,
                    source_conflicts, source_distinct_name_dobs, next_lite, args.venue, member_cache
                )
                if problem or not member:
                    actions.append(action_from(s, p, "REVIEW_PARTICIPANT", problem or "Unable to resolve participant", severity="BLOCKER", session_id=session_id, source_ref=a.source_ref))
                    continue

                # Log member resolution once per source attendance for traceability.
                actions.append(action_from(
                    s, p, member.action, member.detail,
                    member_kind=member.kind, member_display_id=member.display_id,
                    session_id=session_id, source_ref=a.source_ref
                ))

                resolved_key = (
                    session_id,
                    member.kind,
                    member.participant_id if member.kind == "FULL" else member.lite_id,
                )
                existing_att = find_existing_attendance(cursor, session_id, member)
                if existing_att:
                    if bool(existing_att["Attended"]):
                        if resolved_key in created_attendance_keys:
                            actions.append(action_from(
                                s, p, "SKIP_RESOLVED_DUPLICATE_ATTENDANCE",
                                f"This source row resolves to the same member/session as an attendance inserted earlier in this migration transaction (AttendanceId={existing_att['AttendanceId']}).",
                                member_kind=member.kind, member_display_id=member.display_id,
                                session_id=session_id, source_ref=a.source_ref
                            ))
                        else:
                            actions.append(action_from(
                                s, p, "SKIP_EXISTING_ATTENDANCE", f"AttendanceId={existing_att['AttendanceId']} already existed in CRM and is attended.",
                                member_kind=member.kind, member_display_id=member.display_id,
                                session_id=session_id, source_ref=a.source_ref
                            ))
                    else:
                        actions.append(action_from(
                            s, p, "REVIEW_EXISTING_ATTENDANCE_NOT_ATTENDED",
                            f"AttendanceId={existing_att['AttendanceId']} already exists with Attended=0. Script will not overwrite it.",
                            severity="BLOCKER", member_kind=member.kind, member_display_id=member.display_id,
                            session_id=session_id, source_ref=a.source_ref
                        ))
                    continue

                insert_attendance(cursor, session_id, s, p, member)
                created_attendance_keys.add(resolved_key)
                actions.append(action_from(
                    s, p, "CREATE_ATTENDANCE", "New attended SessionAttendance row.",
                    member_kind=member.kind, member_display_id=member.display_id,
                    session_id=session_id, source_ref=a.source_ref
                ))

        blockers = [a for a in actions if a.action.startswith("REVIEW_") or a.severity == "BLOCKER"]
        mode = "COMMIT" if args.commit else "PREVIEW_ROLLBACK"
        if args.commit and not blockers:
            conn.commit()
            print("\nCOMMIT COMPLETE: Omnia migration committed successfully.")
        else:
            conn.rollback()
            if args.commit and blockers:
                print(f"\nCOMMIT REFUSED: {len(blockers)} REVIEW/BLOCKER item(s) remain. Transaction rolled back.")
            else:
                print("\nPREVIEW COMPLETE: transaction rolled back; CRM was not changed.")

        out_dir = Path(args.output_dir).resolve()
        out_dir.mkdir(parents=True, exist_ok=True)
        csv_path = write_actions(actions, out_dir, "omnia_migration_preview" if not args.commit else "omnia_migration_commit")
        sum_path = write_summary(actions, sessions, empty, out_dir, mode)
        review_path = write_review_people(actions, out_dir)
        session_review_path = write_review_sessions(actions, out_dir)
        print(f"Actions CSV: {csv_path}")
        print(f"Summary:     {sum_path}")
        if review_path:
            print(f"Participant review CSV: {review_path}")
        if session_review_path:
            print(f"Session review CSV:     {session_review_path}")
        print(f"Review blockers: {len(blockers)}")
        review_rows = [a for a in blockers if a.action == "REVIEW_PARTICIPANT"]
        review_grouped: dict[tuple[str, str, str], list[Action]] = defaultdict(list)
        for a in review_rows:
            review_grouped[((a.source_card or '').strip(), norm_name(a.source_name), a.detail)].append(a)
        unique_review_people = {(card, name_norm) for (card, name_norm, _reason) in review_grouped}
        unique_review_people.discard(("", ""))
        unique_new_lites = {a.member_display_id for a in actions if a.action == "CREATE_LITE" and a.member_display_id}
        unique_new_full = {a.member_display_id for a in actions if a.action == "CREATE_FULL_FROM_VALID_CARD" and a.member_display_id}
        resolved_dupes = [a for a in actions if a.action == "SKIP_RESOLVED_DUPLICATE_ATTENDANCE"]
        print(f"Unique participant identities needing review: {len(unique_review_people)}")
        print(f"Unique new Lite members proposed: {len(unique_new_lites)}")
        print(f"Unique new FULL participants proposed: {len(unique_new_full)}")
        print(f"Resolved duplicate attendance rows skipped inside this migration: {len(resolved_dupes)}")
        if review_grouped:
            print("\nParticipants needing review (unique reason groups):")
            for (card, _name_norm, reason), items in sorted(review_grouped.items(), key=lambda kv: (-len(kv[1]), kv[0][0], kv[0][1])):
                ex = items[0]
                card_text = card or "NO CARD"
                print(f"  {len(items):>3} row(s) | {card_text} | {ex.source_name} | {reason}")
        source_blockers = [a for a in blockers if a.action != "REVIEW_PARTICIPANT"]
        if source_blockers:
            grouped_source: dict[tuple[str, str, str], list[Action]] = defaultdict(list)
            for a in source_blockers:
                grouped_source[(a.action, a.source_name, a.detail)].append(a)
            print("\nOther source/session blockers:")
            for (action_name, source_name, reason), items in sorted(grouped_source.items(), key=lambda kv: (-len(kv[1]), kv[0][0], kv[0][1])):
                who = source_name or "-"
                print(f"  {len(items):>3} item(s) | {action_name} | {who} | {reason}")
        print("\nAction summary (trace rows; member actions may repeat per attendance):")
        for k, v in sorted(Counter(a.action for a in actions).items()):
            print(f"  {k}: {v}")
        return 2 if args.commit and blockers else 0
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def main():
    parser = argparse.ArgumentParser(description="Saheli CRM Omnia Medical Practice historical migration V2")
    parser.add_argument("--source-dir", default=str(Path(__file__).resolve().parent), help="Folder containing Register for Exercise - Omnia.xlsx")
    parser.add_argument("--output-dir", default=str(Path(__file__).resolve().parent), help="Folder for audit/preview CSV and summary")
    parser.add_argument("--venue", default=DEFAULT_VENUE, help="Exact CRM Sessions.VenueName (default: Omnia Medical Practice)")
    parser.add_argument("--connection-string", default="", help="Optional ODBC connection string; safer to use SAHELI_SQL_CONNECTION_STRING env variable")
    parser.add_argument("--audit-only", action="store_true", help="Parse Excel only; no DB connection")
    parser.add_argument("--commit", action="store_true", help="Commit only if zero REVIEW/BLOCKER items remain")
    args = parser.parse_args()

    if args.audit_only and args.commit:
        parser.error("--audit-only and --commit cannot be used together")

    source_dir = Path(args.source_dir).resolve()
    out_dir = Path(args.output_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    sessions, empty, all_people, conflicts, source_actions, source_distinct_name_dobs = parse_sources(source_dir)
    print_source_summary(sessions, empty, source_actions, conflicts)

    # Explicitly report materially conflicting source cards and source-profile blockers.
    for card, names in sorted(conflicts.items()):
        print(f"  REVIEW source card {card}: {sorted(names)}")
    source_review_actions = [a for a in source_actions if a.action.startswith("REVIEW_") or a.severity == "BLOCKER"]
    if source_review_actions:
        print("Source/data review items:")
        grouped = defaultdict(list)
        for a in source_review_actions:
            grouped[(a.action, a.source_name, a.source_card, a.detail)].append(a)
        for (action_name, source_name, source_card, detail), items in grouped.items():
            print(f"  {len(items)} item(s) | {action_name} | {source_card or 'NO CARD'} | {source_name or '-'} | {detail}")

    if args.audit_only:
        audit_actions = list(source_actions)
        for card, names in sorted(conflicts.items()):
            audit_actions.append(Action(
                action="REVIEW_SOURCE_CARD_CONFLICT", severity="BLOCKER",
                source_card=card, detail=f"Card is used against materially different source names: {sorted(names)}"
            ))
        csv_path = write_actions(audit_actions, out_dir, "omnia_source_audit")
        sum_path = write_summary(audit_actions, sessions, empty, out_dir, "AUDIT_ONLY")
        print(f"\nAudit CSV: {csv_path}")
        print(f"Summary:   {sum_path}")
        print("No database connection was made.")
        return 0

    return run_db_preview_or_commit(args, sessions, empty, conflicts, source_actions, source_distinct_name_dobs)


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("Cancelled.", file=sys.stderr)
        sys.exit(130)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
